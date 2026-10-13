"""Independent rectangular-table horizontal-cell-margin regression.

Owned sources and edit references are generated with python-docx and explicit
OOXML. No application implementation is imported. Exact model/XML checks are
separate from installed-LibreOffice PDF border/text/wrapping checks. Modern
indentation remains a refusal; modern no-indent probes claim margin values,
not exact absolute border preservation. Original pinned upstream documents are
never changed and are not described as user production files.

  python scripts/check-word-table-cell-margins.py --generate-only --render
  python scripts/check-word-table-cell-margins.py --cli target/debug/rust-wps --render --corpus-dir tmp/table-corpus
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from lxml import etree as ET
from pypdf import PdfReader

_SPEC = importlib.util.spec_from_file_location("width_helpers", Path(__file__).with_name("check-word-table-widths.py"))
w = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(w)
KEY = "tableCellMarginsTwips"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
WORDS = [f"L{i:02d}" for i in range(1, 13)]
TEXTS = [[" ".join(WORDS), "RightProbe0"], ["SecondRowLeft", "RightProbe1"]]
EDIT_PAIR = {"left": 240, "right": 180}
ANCHOR_TEXT = "BetweenTablesAnchor"
PAGES = {"a4": (11906, 16838, 1440), "letter": (12240, 15840, 1440), "a5": (8391, 11906, 720)}


def check(value, message):
    if not value:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def add(parent, tag, **attrs):
    node = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    parent.append(node)
    return node


def remove(parent, tag):
    for node in list(parent.findall("w:" + tag, NS)):
        parent.remove(node)


def pair(parent, values, tag="tblCellMar"):
    node = add(parent, tag)
    for side, value in values.items():
        add(node, side, w=value, type="dxa")
    return node


def document(page="a4", mode=14):
    doc = Document()
    width, height, margin = PAGES[page]
    for section in doc.sections:
        section.page_width, section.page_height = Twips(width), Twips(height)
        section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Twips(margin)
        remove(section._sectPr, "docGrid")
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Sans", Pt(12)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for node in list(doc.styles.element.findall(".//w:contextualSpacing", NS)):
        node.getparent().remove(node)
    for node in list(doc.settings.element.findall(".//w:compatSetting", NS)):
        if node.get(qn("w:name")) == "compatibilityMode":
            if mode is None:
                node.getparent().remove(node)
            else:
                node.set(qn("w:val"), str(mode))
    return doc


def table(doc, margins=None, indent=None, align="left", placement="table", paragraph=None, right_text=False, widths=(1800, 4200)):
    tab = doc.add_table(rows=2, cols=2)
    tab.autofit = False
    props = tab._tbl.tblPr
    remove(props, "tblW")
    add(props, "tblW", w=sum(widths), type="dxa")
    if align is not None:
        add(props, "jc", val=align)
    if indent is not None:
        add(props, "tblInd", w=indent, type="dxa")
    borders = add(props, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        add(borders, edge, val="single", sz=4, color="000000")
    if margins is not None and placement == "table":
        pair(props, margins)
    for col, width in zip(tab._tbl.tblGrid, widths):
        col.set(qn("w:w"), str(width))
    for ri, row in enumerate(tab.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Twips(widths[ci])
            cell.text = TEXTS[ri][ci]
            if margins is not None and placement == "cell":
                pair(cell._tc.get_or_add_tcPr(), margins, "tcMar")
            for para in cell.paragraphs:
                para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1
                if paragraph and ri == ci == 0:
                    add(para._p.get_or_add_pPr(), "ind", **paragraph)
                if right_text and ci == 1:
                    add(para._p.get_or_add_pPr(), "jc", val="right")
                for run in para.runs:
                    run.font.name, run.font.size = "DejaVu Sans", Pt(12)
    return tab


def separator_anchor(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Twips(120)
    paragraph.paragraph_format.line_spacing = 1
    add(paragraph._p.get_or_add_pPr(), "ind", left=240, right=0, firstLine=120)
    run = paragraph.add_run(ANCHOR_TEXT)
    run.font.name, run.font.size = "DejaVu Sans", Pt(12)
    return paragraph


def verify_anchor_model(model, label):
    paragraphs = [n for n in model.get("content", []) if n["type"] == "paragraph"]
    check(len(paragraphs) == 1 and w.model_text(paragraphs[0]) == ANCHOR_TEXT, label + ": separator anchor text/order differs")
    attrs = paragraphs[0].get("attrs", {})
    for key, value in {"indentLeftTwips": 240, "indentRightTwips": 0, "firstLineIndentTwips": 120, "spaceBefore": 120, "spaceAfter": 120, "lineSpacing": 240}.items():
        check(attrs.get(key) == value, label + ": separator anchor " + key + " differs")
    check([n["type"] for n in model["content"]] == ["table", "paragraph", "table"], label + ": separated table/anchor order differs")
    for run in paragraphs[0]["content"]:
        marks = {m["type"]: m.get("attrs", {}) for m in run.get("marks", [])}
        check(marks.get("fontFamily") == {"name": "DejaVu Sans"} and marks.get("fontSize") == {"halfPoints": 24}, label + ": explicit anchor font differs")


def verify_anchor_snapshot(snapshot, label):
    paragraphs = snapshot["body_paragraphs"]
    check(len(paragraphs) == 1 and paragraphs[0]["text"] == ANCHOR_TEXT, label + ": independent separator text differs")
    para = paragraphs[0]
    check(para["indent"] == {"left": "240", "right": "0", "firstLine": "120"}, label + ": independent separator indentation differs")
    check(all(para["spacing"].get(k) == v for k, v in {"before": "120", "after": "120", "line": "240", "lineRule": "auto"}.items()), label + ": independent separator spacing differs")
    check(len(para["runs"]) == 1, label + ": separator run count differs")
    run = para["runs"][0]
    check(run["font"].get("ascii") == run["font"].get("hAnsi") == "DejaVu Sans" and run["size"] == "24", label + ": independent separator explicit font differs")


def xml_snapshot(path):
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
        settings = ET.fromstring(package.read("word/settings.xml")) if "word/settings.xml" in package.namelist() else None
        styles_root = ET.fromstring(package.read("word/styles.xml")) if "word/styles.xml" in package.namelist() else None
    doc = Document(path)
    styles = {} if styles_root is None else {n.get(qn("w:styleId")): n for n in styles_root.findall("w:style", NS)}
    def effective(raw):
        # Independent controlled-source resolver: explicitly selected unconditional
        # physical style chain -> table -> cell. It is not a general DOCX importer.
        values = {"left": 108, "right": 108}
        def apply(parent, name):
            if parent is None: return
            for node in parent.findall("w:" + name + "/*", NS):
                side = ET.QName(node).localname
                if side in values and node.get(qn("w:type")) == "dxa":
                    try: values[side] = int(node.get(qn("w:w")))
                    except (TypeError, ValueError): values[side] = None
        props = raw.find("w:tblPr", NS)
        selected = props.find("w:tblStyle", NS)
        chain, seen = [], set()
        style_id = None if selected is None else selected.get(qn("w:val"))
        while style_id is not None and style_id in styles and style_id not in seen:
            seen.add(style_id); style = styles[style_id]; chain.append(style)
            base = style.find("w:basedOn", NS)
            style_id = None if base is None else base.get(qn("w:val"))
        for style in reversed(chain): apply(style.find("w:tblPr", NS), "tblCellMar")
        apply(props, "tblCellMar")
        defaults = dict(values); rows = []
        for row in raw.findall("w:tr", NS):
            cells = []
            for cell in row.findall("w:tc", NS):
                values = dict(defaults); apply(cell.find("w:tcPr", NS), "tcMar"); cells.append(dict(values))
            rows.append(cells)
        return rows
    result = []
    for raw, py in zip(root.findall(".//w:tbl", NS), doc.tables):
        props = raw.find("w:tblPr", NS)
        decl = lambda node: None if node is None else {ET.QName(k).localname: v for k, v in node.attrib.items()}
        mar = lambda parent, name: [{"side": ET.QName(n).localname, **decl(n)} for n in parent.findall("w:" + name + "/*", NS)]
        result.append({"effective_horizontal_margins": effective(raw), "table_margins": mar(props, "tblCellMar"), "indent": decl(props.find("w:tblInd", NS)),
                       "align": decl(props.find("w:jc", NS)), "layout": decl(props.find("w:tblLayout", NS)),
                       "total": decl(props.find("w:tblW", NS)), "grid": [int(n.get(qn("w:w"))) for n in raw.findall("w:tblGrid/w:gridCol", NS)],
                       "cell_margins": [[mar(c.find("w:tcPr", NS), "tcMar") for c in r.findall("w:tc", NS)] for r in raw.findall("w:tr", NS)],
                       "paragraph_indents": [[decl(p.find("w:pPr/w:ind", NS)) for p in c.findall("w:p", NS)] for c in raw.findall("w:tr/w:tc", NS)],
                       "xml_texts": [["\n".join("".join(p.xpath(".//w:t/text()", namespaces=NS)) for p in c.findall("w:p", NS)) for c in r.findall("w:tc", NS)] for r in raw.findall("w:tr", NS)],
                       "python_texts": [[c.text for c in r.cells] for r in py.rows],
                       "python_widths": [c.width.twips for c in py.columns],
                       "python_cell_widths": [[c.width.twips for c in r.cells] for r in py.rows], "python_autofit": py.autofit})
    check(len(result) == len(doc.tables) == len(root.findall(".//w:tbl", NS)), "Independent table count mismatch")
    section = root.find("w:body/w:sectPr", NS)
    def attrs(node):
        return {} if node is None else {ET.QName(k).localname: v for k, v in node.attrib.items()}
    paragraphs = [{"text": "".join(p.xpath(".//w:t/text()", namespaces=NS)), "indent": attrs(p.find("w:pPr/w:ind", NS)),
                   "spacing": attrs(p.find("w:pPr/w:spacing", NS)),
                   "runs": [{"font": attrs(r.find("w:rPr/w:rFonts", NS)), "size": attrs(r.find("w:rPr/w:sz", NS)).get("val")} for r in p.findall("w:r", NS)]}
                  for p in root.findall("w:body/w:p", NS)]
    return {"tables": result, "body_paragraphs": paragraphs, "page_size": {ET.QName(k).localname: v for k, v in section.find("w:pgSz", NS).attrib.items()},
            "page_margins": {ET.QName(k).localname: v for k, v in section.find("w:pgMar", NS).attrib.items()},
            "compatibility_modes": [] if settings is None else [n.get(qn("w:val")) for n in settings.findall(".//w:compatSetting", NS) if n.get(qn("w:name")) == "compatibilityMode"]}


def expected(margins=None, indent=None, align="left", paragraph=None, right_text=False):
    p = {}
    for source, model in [("left", "indentLeftTwips"), ("right", "indentRightTwips"), ("firstLine", "firstLineIndentTwips")]:
        if paragraph and source in paragraph:
            p[model] = paragraph[source]
    if paragraph and "hanging" in paragraph:
        p["firstLineIndentTwips"] = -paragraph["hanging"]
    canonical = None if margins == {"left": 108, "right": 108} else margins
    return {"margins": copy.deepcopy(canonical), "effective_margins": copy.deepcopy(margins or {"left": 108, "right": 108}),
            "widths": [1800, 4200], "indent": indent, "align": align, "texts": copy.deepcopy(TEXTS),
            "paragraph": p, "source_paragraph": paragraph, "right_text": right_text}


def verify_snapshot(snapshot, targets, label, exported=False):
    check(len(snapshot["tables"]) == len(targets), label + ": table count differs")
    for actual, target in zip(snapshot["tables"], targets):
        check(actual["grid"] == actual["python_widths"] == target["widths"], label + ": grid/columns differ")
        check(all(r == target["widths"] for r in actual["python_cell_widths"]), label + ": cell widths differ")
        check(actual["layout"] == {"type": "fixed"} and actual["python_autofit"] is False, label + ": fixed layout lost")
        check(actual["xml_texts"] == actual["python_texts"] == target["texts"], label + ": text differs")
        check(all(pair == target["effective_margins"] for row in actual["effective_horizontal_margins"] for pair in row), label + ": independently resolved horizontal margins differ")
        if exported:
            margins = {n["side"]: {k: v for k, v in n.items() if k != "side"} for n in actual["table_margins"]}
            if target["margins"] is not None:
                check(margins == {s: {"w": str(v), "type": "dxa"} for s, v in target["margins"].items()}, label + ": explicit margin pair differs")
            elif target["align"] is None and target["indent"] is None and not target["paragraph"]:
                check(not margins, label + ": old table acquired new margin declarations")
            elif margins:
                check(margins == {s: {"w": "108", "type": "dxa"} for s in ["left", "right"]}, label + ": invented nondefault margins")
            check(actual["indent"] == (None if target["indent"] is None else {"w": str(target["indent"]), "type": "dxa"}), label + ": indent absent/zero differs")
            check(actual["align"] == (None if target["align"] is None else {"val": target["align"]}), label + ": alignment differs")


def fixtures(directory):
    sources = {}

    def save(name, doc, targets=None, reject=None, render=False, scope=None, separated=False):
        path = directory / (name + ".docx")
        check(not path.exists(), "Fixture destination exists")
        doc.save(path)
        snap = xml_snapshot(path)
        if separated: verify_anchor_snapshot(snap, name + " source")
        if targets is not None:
            verify_snapshot(snap, targets, name + " source")
        item = {"path": str(path), "sha256": digest(path), "bytes": path.stat().st_size, "source_independent": snap,
                "fixture_kind": "owned directed python-docx/XML fixture, not an original upstream/user file", "expected_action": "reject" if reject else "roundtrip",
                "reject_reason": reject, "expected": targets, "render": render, "scope": scope, "separated_anchor": separated}
        if targets is not None:
            edited = copy.deepcopy(targets)
            reference = document(page=next((key for key, value in PAGES.items() if int(snap["page_size"]["w"]) == value[0]), "a4"))
            for index, target in enumerate(edited):
                if separated and index: separator_anchor(reference)
                values = {"left": 180, "right": 300} if target["margins"] == EDIT_PAIR else dict(EDIT_PAIR)
                target["margins"] = target["effective_margins"] = values
                table(reference, values, target["indent"], target["align"], paragraph=target["source_paragraph"], right_text=target["right_text"])
            ref = directory / (name + "-independent-edited-reference.docx")
            reference.save(ref)
            verify_snapshot(xml_snapshot(ref), edited, name + " independent edit reference", exported=True)
            if separated: verify_anchor_snapshot(xml_snapshot(ref), name + " independent edit reference")
            item.update(edited_expected=edited, edited_reference={"path": str(ref), "sha256": digest(ref)})
        sources[name] = item

    for margins in [{"left": 108, "right": 108}, {"left": 0, "right": 0}, {"left": 180, "right": 240}, {"left": 240, "right": 180}]:
        for ind in [None, 0, 720]:
            doc = document()
            table(doc, margins, ind)
            save(f"direct-{margins['left']}-{margins['right']}-indent-{ind}", doc, [expected(margins, ind)], render=True)
    for mode in [None, 11, 12]:
        doc = document(mode=mode)
        table(doc, {"left": 180, "right": 240}, 0)
        save("legacy-mode-" + str(mode), doc, [expected({"left": 180, "right": 240}, 0)], render=True)
    for page in ["letter", "a5"]:
        doc = document(page)
        table(doc, {"left": 180, "right": 240}, 0)
        save("page-" + page, doc, [expected({"left": 180, "right": 240}, 0)], render=True)
    for placement in ["cell", "style", "derived-style"]:
        doc = document()
        margins = {"left": 180, "right": 240}
        tab = table(doc, margins if placement == "cell" else None, None, placement="cell")
        if placement != "cell":
            base = doc.styles.add_style("MarginBase", WD_STYLE_TYPE.TABLE)
            base.base_style = doc.styles["Normal Table"]
            pair(add(base.element, "tblPr"), margins)
            selected = base
            if placement == "derived-style":
                selected = doc.styles.add_style("MarginDerived", WD_STYLE_TYPE.TABLE)
                selected.base_style = base
                pair(add(selected.element, "tblPr"), {"right": 300})
                margins["right"] = 300
            add(tab._tbl.tblPr, "tblStyle", val=selected.style_id)
        save("uniform-" + placement, doc, [expected(margins, None if placement == "cell" else 0)], render=True)
    for align in ["center", "right"]:
        doc = document()
        table(doc, {"left": 180, "right": 240}, 0, align)
        save(align + "-zero-asymmetric", doc, [expected({"left": 180, "right": 240}, 0, align)], render=True)
    for name, paragraph in [("first-line", {"left": 360, "right": 180, "firstLine": 180}), ("hanging", {"left": 360, "right": 180, "hanging": 180})]:
        doc = document()
        table(doc, {"left": 180, "right": 240}, 0, paragraph=paragraph, right_text=True)
        save("paragraph-" + name, doc, [expected({"left": 180, "right": 240}, 0, paragraph=paragraph, right_text=True)], render=True)
    doc = document()
    table(doc, None, None, None)
    table(doc, {"left": 0, "right": 0}, 0)
    save("multiple-old-and-new", doc, [expected(None, None, None), expected({"left": 0, "right": 0}, 0)],
         scope="Adjacent tables may merge in this consumer. Model/XML and supplementary visual check only; not per-table measured geometry.")
    doc = document()
    table(doc, None, None, None)
    separator_anchor(doc)
    table(doc, {"left": 180, "right": 240}, 720)
    save("multiple-separated-old-and-asymmetric", doc,
         [expected(None, None, None), expected({"left": 180, "right": 240}, 720)], render=True, separated=True,
         scope="Two distinct tables separated by an explicit-font/spacing paragraph indentation anchor. Each table border/text coordinate is independently measured; unchanged and edited exports compared with independent references.")
    doc = document(mode=15)
    table(doc, {"left": 180, "right": 240}, None)
    save("modern-no-indent-values-only", doc, [expected({"left": 180, "right": 240})], render=True,
         scope="Only margin values preserved. Existing border appearance normalization permits measured 0.5pt absolute consumer shift; not modern position fidelity.")

    def bad(name, mutate, mode=14):
        doc = document(mode=mode)
        tab = table(doc, {"left": 180, "right": 240}, 0)
        mutate(doc, tab)
        save("reject-" + name, doc, reject=name)

    bad("modern-explicit-zero", lambda d, t: None, 15)
    bad("modern-explicit-positive", lambda d, t: t._tbl.tblPr.find("w:tblInd", NS).set(qn("w:w"), "720"), 15)
    bad("nonuniform-cell", lambda d, t: pair(t.cell(0, 0)._tc.get_or_add_tcPr(), {"left": 181, "right": 240}, "tcMar"))
    bad("logical-custom", lambda d, t: (remove(t._tbl.tblPr, "tblCellMar"), pair(t._tbl.tblPr, {"start": 180, "end": 240})))
    bad("partial-custom", lambda d, t: remove(t._tbl.tblPr.find("w:tblCellMar", NS), "right"))
    bad("implicit-custom-default", lambda d, t: (remove(t._tbl.tblPr, "tblCellMar"), remove(d.styles["Normal Table"].element.find("w:tblPr", NS), "tblCellMar"), pair(d.styles["Normal Table"].element.find("w:tblPr", NS), {"left": 180, "right": 240})))
    def conditional(doc, tab):
        style = doc.styles.add_style("ConditionalMargins", WD_STYLE_TYPE.TABLE)
        style.base_style = doc.styles["Normal Table"]
        pair(add(add(style.element, "tblStylePr", type="firstRow"), "tblPr"), {"left": 180, "right": 240})
        add(tab._tbl.tblPr, "tblStyle", val=style.style_id)
    bad("conditional-style", conditional)
    for kind in ["pct", "nil", "auto"]:
        bad("unit-" + kind, lambda d, t, k=kind: t._tbl.tblPr.find("w:tblCellMar/w:left", NS).set(qn("w:type"), k))
    for name, value in [("negative", "-1"), ("fractional", "1.5"), ("too-large", "31681")]:
        bad("margin-" + name, lambda d, t, v=value: t._tbl.tblPr.find("w:tblCellMar/w:left", NS).set(qn("w:w"), v))
    bad("missing-unit", lambda d, t: t._tbl.tblPr.find("w:tblCellMar/w:left", NS).attrib.pop(qn("w:type")))
    bad("duplicate-side", lambda d, t: add(t._tbl.tblPr.find("w:tblCellMar", NS), "left", w=180, type="dxa"))
    bad("duplicate-container", lambda d, t: pair(t._tbl.tblPr, {"left": 180, "right": 240}))
    bad("foreign-side", lambda d, t: setattr(t._tbl.tblPr.find("w:tblCellMar/w:left", NS), "tag", "{urn:margin-probe}left"))
    bad("alias-conflict", lambda d, t: add(t._tbl.tblPr.find("w:tblCellMar", NS), "start", w=180, type="dxa"))
    bad("unknown-attribute", lambda d, t: t._tbl.tblPr.find("w:tblCellMar/w:left", NS).set(qn("w:extra"), "1"))
    bad("zero-text-area", lambda d, t: (remove(t._tbl.tblPr, "tblCellMar"), pair(t._tbl.tblPr, {"left": 900, "right": 900})))
    for name, margins in [("plain-zero-text-area", {"left": 900, "right": 900}), ("plain-negative-text-area", {"left": 901, "right": 900})]:
        doc = document()
        table(doc, margins, None, None)
        save("reject-" + name, doc, reject="proved complete margin pair leaves non-positive fixed-column text area, despite no table indentation")
    doc = document()
    table(doc, {"left": 180, "right": 240}, None, None, paragraph={"left": 1380})
    save("reject-paragraph-zero-later-line-area", doc, reject="custom margins plus paragraph indentation leave zero later-line text area")
    manifest = directory / "expected-fixtures.json"
    manifest.write_text(json.dumps(sources, indent=2, ensure_ascii=False), encoding="utf-8")
    return sources, manifest


def verify_model(model, targets, label):
    tables = list(w.model_tables(model))
    check(len(tables) == len(targets), label + ": table count differs")
    for tab, target in zip(tables, targets):
        attrs = tab.get("attrs", {})
        for key, value in [(KEY, target["margins"]), ("columnWidthsTwips", target["widths"]), ("tableIndentTwips", target["indent"]), ("tableAlign", target["align"])]:
            check((key not in attrs) if value is None else attrs.get(key) == value, f"{label}: {key} expected {value}, got {attrs.get(key)}")
        texts = [["\n".join(w.model_text(p) for p in c.get("content", [])) for c in r["content"]] for r in tab["content"]]
        check(texts == target["texts"], label + ": model text differs")
        p = tab["content"][0]["content"][0]["content"][0]
        for key, value in target["paragraph"].items():
            check(p.get("attrs", {}).get(key) == value, label + ": cell paragraph indentation changed")


def exported(cli, evidence, directory, name, model, targets):
    jp, dp = directory / (name + ".json"), directory / (name + ".docx")
    w.write_model(jp, model)
    w.run_cli(cli, evidence["commands"], "create-docx", jp, dp)
    snap = xml_snapshot(dp)
    verify_snapshot(snap, targets, name, exported=True)
    reloaded = w.run_cli(cli, evidence["commands"], "read-docx-model", dp)
    check(reloaded == model, name + ": full trusted model differs")
    canonical = copy.deepcopy(targets)
    for target in canonical:
        if target["margins"] == {"left": 108, "right": 108}:
            target["margins"] = None
    external = w.run_cli(cli, evidence["commands"], "import-docx", dp)["document"]
    verify_model(external, canonical, name + " external reimport")
    if any(n["type"] == "paragraph" and w.model_text(n) == ANCHOR_TEXT for n in model.get("content", [])):
        verify_anchor_model(model, name + " authored model")
        verify_anchor_model(external, name + " external reimport")
        verify_anchor_snapshot(snap, name + " independent XML")
    before = digest(dp)
    error = w.run_cli(cli, evidence["commands"], "create-docx", jp, dp, success=False)
    check("Output already exists" in error and digest(dp) == before, name + ": existing output overwritten")
    return {"path": str(dp), "model_path": str(jp), "sha256": before, "independent": snap, "full_trusted_model_equal": True, "external_reimport_checked": True, "overwrite_refused": True}


def invalid_models(cli, evidence, directory, base):
    cases = []
    for i, value in enumerate([None, {}, {"left": 0}, {"right": 0}, {"left": 0, "right": 0, "top": 0}, [], 108, "108", {"left": True, "right": 0}, {"left": 0, "right": "0"}, {"left": 1.5, "right": 0}, {"left": -1, "right": 0}, {"left": 0, "right": 31681}, {"left": 900, "right": 900}]):
        model = copy.deepcopy(base)
        next(w.model_tables(model)).setdefault("attrs", {})[KEY] = value
        cases.append(("value-" + str(i), model))
    for location in ["doc", "row", "cell", "paragraph", "text"]:
        model = copy.deepcopy(base)
        tab = next(w.model_tables(model)); row = tab["content"][0]; cell = row["content"][0]; para = cell["content"][0]
        node = {"doc": model, "row": row, "cell": cell, "paragraph": para, "text": para["content"][0]}[location]
        node.setdefault("attrs", {})[KEY] = {"left": 0, "right": 0}
        cases.append(("location-" + location, model))
    for name, attrs in [("paragraph-no-later-line", {"indentLeftTwips": 1380}), ("paragraph-no-first-line", {"indentLeftTwips": 1000, "firstLineIndentTwips": 380})]:
        model = copy.deepcopy(base)
        tab = next(w.model_tables(model)); tab["attrs"][KEY] = {"left": 180, "right": 240}
        tab["content"][0]["content"][0]["content"][0].setdefault("attrs", {}).update(attrs)
        cases.append((name, model))
    nested = copy.deepcopy(base)
    outer = next(w.model_tables(nested))
    child = copy.deepcopy(outer)
    outer["content"][0]["content"][0]["content"].append(child)
    cases.append(("nested-table-new-margin", nested))
    results = []
    for name, model in cases:
        jp = directory / ("invalid-" + name + ".json"); dp = jp.with_suffix(".docx")
        w.write_model(jp, model)
        error = w.run_cli(cli, evidence["commands"], "create-docx", jp, dp, success=False)
        check(not dp.exists(), name + ": rejected model created output")
        results.append({"name": name, "error": error, "output_absent": True})
    return results


def tamper(cli, evidence, directory, good):
    with zipfile.ZipFile(good) as z:
        original = {n: z.read(n) for n in z.namelist()}
    results = []
    for mode in ["xml", "xml-forged-hash", "model"]:
        parts = copy.deepcopy(original); model = json.loads(parts["rust-wps/model-v2.json"])
        if mode.startswith("xml"):
            root = ET.fromstring(parts["word/document.xml"])
            node = root.find(".//w:tblCellMar/w:left", NS)
            node.set(qn("w:w"), str(int(node.get(qn("w:w"))) + 20))
            parts["word/document.xml"] = ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            if mode == "xml-forged-hash":
                model["hashes"]["word/document.xml"] = hashlib.sha256(parts["word/document.xml"]).hexdigest()
        else:
            next(w.model_tables(model["document"]))["attrs"][KEY]["left"] += 20
        parts["rust-wps/model-v2.json"] = json.dumps(model, separators=(",", ":")).encode()
        path = directory / ("tampered-" + mode + ".docx")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in parts.items():
                z.writestr(name, data)
        before = digest(path)
        error = w.run_cli(cli, evidence["commands"], "read-docx-model", path, success=False)
        check(digest(path) == before, "Tampered source changed")
        results.append({"mode": mode, "error": error, "source_unchanged": True})
    return results


def measure_pdf(path, table_count=1):
    reader = PdfReader(path)
    check(len(reader.pages) == 1, "Controlled probe unexpectedly spans pages: " + str(path))
    segments, pending, texts = [], [], []
    point = None
    def xy(x, y, cm):
        return float(x) * cm[0] + float(y) * cm[2] + cm[4], float(x) * cm[1] + float(y) * cm[3] + cm[5]
    def operand(op, args, cm, tm):
        nonlocal point, pending
        if op == b"m":
            point = xy(*args, cm)
        elif op == b"l":
            end = xy(*args, cm)
            if point is not None: pending.append((point, end))
            point = end
        elif op == b"re":
            x, y, width, height = map(float, args)
            vertices = [xy(a, b, cm) for a, b in [(x, y), (x + width, y), (x + width, y + height), (x, y + height), (x, y)]]
            pending.extend(zip(vertices, vertices[1:]))
        elif op in [b"S", b"s", b"B", b"B*", b"b", b"b*", b"f", b"f*", b"F"]:
            segments.extend(pending); pending = []; point = None
        elif op == b"n":
            pending = []; point = None
    def text(text, cm, tm, font, size):
        if text.strip():
            x, y = xy(tm[4], tm[5], cm)
            texts.append({"text": text.strip(), "x": x, "y": y, "font_size": size})
    reader.pages[0].extract_text(visitor_operand_before=operand, visitor_text=text)
    vertical = [{"x": round((a[0] + b[0]) / 2, 3), "low": min(a[1], b[1]), "high": max(a[1], b[1])}
                for a, b in segments if abs(a[0] - b[0]) < .1 and abs(a[1] - b[1]) > 10]
    # Separate tables by disjoint vertical intervals, rather than merging all
    # page x coordinates. A real paragraph gap must separate the mixed control.
    intervals = []
    for segment in sorted(vertical, key=lambda v: v["high"], reverse=True):
        if not intervals or segment["high"] < intervals[-1]["low"] - 1:
            intervals.append({"low": segment["low"], "high": segment["high"], "segments": [segment]})
        else:
            intervals[-1]["low"] = min(intervals[-1]["low"], segment["low"])
            intervals[-1]["segments"].append(segment)
    check(len(intervals) == table_count, "Expected distinct table vertical groups: " + str([(g["low"], g["high"]) for g in intervals]))
    tables = []
    for interval in intervals:
        xs = sorted({seg["x"] for seg in interval["segments"]})
        groups = []
        for x in xs:
            if not groups or x - groups[-1][-1] > 1: groups.append([x])
            else: groups[-1].append(x)
        borders = [sum(g) / len(g) for g in groups]
        check(len(borders) == 3, "Expected three per-table column-border coordinates: " + str(borders))
        runs = [t for t in texts if interval["low"] <= t["y"] <= interval["high"]]
        lines = [t for t in runs if any(word in t["text"] for word in WORDS)]
        check(len(lines) >= 2 and " ".join(t["text"] for t in lines).split() == WORDS, "Missing/duplicate per-table wrapping markers")
        check(all(abs((borders[i + 1] - borders[i]) - width) <= .100001 for i, width in enumerate([90, 210])), "Per-table column border distances changed")
        tables.append({"borders_x_pt": borders, "y_interval_pt": [interval["low"], interval["high"]], "text_lines": lines, "text_runs": runs,
                       "first_text_x_pt": lines[0]["x"], "subsequent_line_x_pt": lines[1]["x"]})
    body = [t for t in texts if not any(g["low"] <= t["y"] <= g["high"] for g in intervals)]
    return {"path": str(path), "sha256": digest(path), "tables": tables,
            "borders_x_pt": [x for tab in tables for x in tab["borders_x_pt"]],
            "text_lines": [line for tab in tables for line in tab["text_lines"]], "text_runs": texts, "body_text_runs": body,
            "page_width_pt": float(reader.pages[0].mediabox.width), "page_height_pt": float(reader.pages[0].mediabox.height)}


def render_all(soffice, directory, sources):
    out = directory / "pdf"; out.mkdir()
    paths = []
    for name, item in sources.items():
        if not item["render"] and name in ["multiple-old-and-new", "reject-modern-explicit-zero"]:
            paths.append(Path(item["path"]))
            for key in ["unchanged_export", "edited_export"]:
                if key in item: paths.append(Path(item[key]["path"]))
        if item["render"]:
            paths.append(Path(item["path"]))
            paths.append(Path(item["edited_reference"]["path"]))
            for key in ["unchanged_export", "edited_export"]:
                if key in item: paths.append(Path(item[key]["path"]))
    command = [soffice, f"-env:UserInstallation={(directory / 'lo-profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(out), *map(str, paths)]
    result = subprocess.run(command, text=True, capture_output=True, timeout=360, env=dict(os.environ, XDG_CACHE_HOME=str(directory / "lo-cache")))
    (directory / "libreoffice.stdout.txt").write_text(result.stdout)
    (directory / "libreoffice.stderr.txt").write_text(result.stderr)
    check(result.returncode == 0, "LibreOffice failed: " + result.stderr)
    check(all((out / (path.stem + ".pdf")).is_file() for path in paths), "LibreOffice omitted a requested PDF")
    for name, item in sources.items():
        if not item["render"]:
            if name == "reject-modern-explicit-zero":
                item["rejected_source_calibration_pdf"] = measure_pdf(out / (Path(item["path"]).stem + ".pdf"))
                item["rejected_source_calibration_pdf"]["scope"] = "Modern explicit-zero source is refused; calibration only, never a supported import"
            elif name == "multiple-old-and-new":
                item["visual_only_pdf"] = {"path": str(out / (Path(item["path"]).stem + ".pdf")), "scope": "Multiple-table contact-sheet check supplements exact model/XML checks"}
            continue
        def read(path): return measure_pdf(out / (Path(path).stem + ".pdf"), len(item["expected"]))
        item["source_pdf"] = read(item["path"])
        item["independent_edited_reference_pdf"] = read(item["edited_reference"]["path"])
        for key in ["unchanged_export", "edited_export"]:
            if key not in item: continue
            actual = read(item[key]["path"])
            reference = item["source_pdf"] if key == "unchanged_export" else item["independent_edited_reference_pdf"]
            check(actual["page_width_pt"] == reference["page_width_pt"] and actual["page_height_pt"] == reference["page_height_pt"], name + ": PDF page differs")
            check(len(actual["tables"]) == len(reference["tables"]), name + ": distinct table count differs")
            errors = [abs(a - b) for a, b in zip(actual["borders_x_pt"], reference["borders_x_pt"])]
            per_table = []
            for index, (a, b) in enumerate(zip(actual["tables"], reference["tables"])):
                check([t["text"] for t in a["text_runs"]] == [t["text"] for t in b["text_runs"]], name + ": per-table text/wrapping differs")
                delta = max([abs(x-y) for x, y in zip(a["borders_x_pt"], b["borders_x_pt"])] + [abs(x["x"]-y["x"]) for x, y in zip(a["text_runs"], b["text_runs"])])
                per_table.append({"table_index": index, "reference_borders_x_pt": b["borders_x_pt"], "actual_borders_x_pt": a["borders_x_pt"],
                                  "reference_first_text_x_pt": b["first_text_x_pt"], "actual_first_text_x_pt": a["first_text_x_pt"], "max_reference_x_delta_pt": delta})
            actual["per_table_comparison"] = per_table
            if item["separated_anchor"]:
                check([t["text"] for t in actual["body_text_runs"]] == [ANCHOR_TEXT] == [t["text"] for t in reference["body_text_runs"]], name + ": separator is not independently located between tables")
                check(abs(actual["body_text_runs"][0]["x"] - 90) <= .2, name + ": separator indentation anchor x differs")
                check(all(abs(a-b) <= .100001 for a, b in zip(actual["tables"][0]["y_interval_pt"], reference["tables"][0]["y_interval_pt"])), name + ": first table height differs")
                check(all(abs(a-b) <= .100001 for a, b in zip(actual["tables"][1]["y_interval_pt"], reference["tables"][1]["y_interval_pt"])), name + ": second table height differs")
                check(all(abs(a["y"] - b["y"]) <= .100001 for a, b in zip(actual["text_runs"], reference["text_runs"])), name + ": mixed control table/anchor vertical text differs")
                actual["separator_anchor_comparison"] = {"reference": reference["body_text_runs"][0], "actual": actual["body_text_runs"][0], "all_text_y_equal_within_0_1pt": True}
            check([t["text"] for t in actual["text_lines"]] == [t["text"] for t in reference["text_lines"]], name + ": wrapping differs")
            errors += [abs(a["x"] - b["x"]) for a, b in zip(actual["text_lines"], reference["text_lines"])]
            # Compare every actual text run, including an unbreakable second-row
            # marker split by the consumer. Requiring that marker in one PDF run
            # would falsely reject legitimate wrapping.
            check([t["text"] for t in actual["text_runs"]] == [t["text"] for t in reference["text_runs"]], name + ": non-first-cell wrapping/text runs differ")
            errors += [abs(a["x"] - b["x"]) for a, b in zip(actual["text_runs"], reference["text_runs"])]
            modern_normalization = name == "modern-no-indent-values-only" and key == "unchanged_export"
            if modern_normalization:
                check(.4 <= max(errors) <= .6, name + ": expected modern border normalization changed")
                actual["absolute_geometry_preserved"] = False
                actual["measured_normalization_delta_pt"] = max(errors)
            else:
                check(max(errors) <= .100001, name + ": source/reference border/text delta exceeds 0.1pt rounding: " + str(max(errors)))
                actual["absolute_geometry_preserved_within_0_1pt"] = True
            actual["max_reference_x_delta_pt"] = max(errors)
            actual["comparison_scope"] = item.get("scope") or "Controlled same-consumer border, first/subsequent text and wrapping comparison; no app/browser/native/Word/WPS claim"
            item[key + "_pdf"] = actual
        if name.startswith("paragraph-"):
            source = item["source_pdf"]["text_lines"]
            offset = source[0]["x"] - source[1]["x"]
            check(abs(offset - (9 if name.endswith("first-line") else -9)) < .2, "Paragraph first/subsequent-line offset changed")
            item["source_first_minus_subsequent_x_pt"] = offset


def contact_sheets(directory, sources):
    """Render actual PDF pages with Poppler; PNGs supplement measured arrays."""
    from PIL import Image, ImageDraw, ImageFont
    converter = shutil.which("pdftoppm")
    check(converter is not None, "pdftoppm needed for contact sheets")
    out = directory / "visual-qa"; out.mkdir()
    groups = {
        "legacy-source-export": [
            ("direct-0-0-indent-0", "source", "Zero margins: legacy source"),
            ("direct-0-0-indent-0", "unchanged_export", "Zero margins: unchanged export"),
            ("direct-180-240-indent-0", "source", "180/240 margins: legacy source"),
            ("direct-180-240-indent-0", "unchanged_export", "180/240 margins: unchanged export")],
        "edits-and-paragraphs": [
            ("direct-180-240-indent-0", "edited_reference", "Margin edit: independent reference"),
            ("direct-180-240-indent-0", "edited_export", "Margin edit: actual export"),
            ("paragraph-first-line", "source", "Custom margins + first-line indent: source"),
            ("paragraph-first-line", "unchanged_export", "Custom margins + first-line indent: export"),
            ("paragraph-hanging", "edited_reference", "Hanging indent + margin edit: reference"),
            ("paragraph-hanging", "edited_export", "Hanging indent + margin edit: export")],
        "mixed-edited-reference": [
            ("multiple-separated-old-and-asymmetric", "edited_reference", "Separated tables margin edit: reference"),
            ("multiple-separated-old-and-asymmetric", "edited_export", "Separated tables margin edit: actual export")],
        "mixed-and-modern-boundaries": [
            ("multiple-separated-old-and-asymmetric", "source", "Separated old/new tables + anchor: source"),
            ("multiple-separated-old-and-asymmetric", "unchanged_export", "Separated old/new tables + anchor: export"),
            ("modern-no-indent-values-only", "source", "Mode15 no indent: source (values-only claim)"),
            ("modern-no-indent-values-only", "unchanged_export", "Legacy export: 0.5 pt position normalization"),
            ("reject-modern-explicit-zero", "source", "Mode15 explicit zero: REFUSED source"),
            ("direct-180-240-indent-0", "source", "Legacy explicit zero: independent calibration")]
    }
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    font = ImageFont.truetype(font_path, 19) if Path(font_path).exists() else ImageFont.load_default()
    reports = []
    for group, selections in groups.items():
        panels, records = [], []
        for name, stage, title in selections:
            item = sources[name]
            if stage == "source": path = Path(item["path"])
            elif stage in item: path = Path(item[stage]["path"])
            else: continue
            pdf = directory / "pdf" / (path.stem + ".pdf")
            check(pdf.is_file(), "Contact-sheet PDF missing: " + str(pdf))
            prefix = out / path.stem
            png = prefix.with_suffix(".png")
            if not png.exists():
                run = subprocess.run([converter, "-f", "1", "-singlefile", "-r", "110", "-png", str(pdf), str(prefix)], text=True, capture_output=True, timeout=60)
                check(run.returncode == 0 and png.is_file(), "Poppler page render failed: " + run.stderr)
            image = Image.open(png).convert("RGB")
            # Keep page origin and horizontal margins. Only unused lower-page
            # whitespace is cropped; original full-page PNG remains beside it.
            crop = image.crop((0, 0, image.width, min(image.height, round(350 * 110 / 72))))
            crop.thumbnail((830, 495))
            panel = Image.new("RGB", (850, 535), "white")
            draw = ImageDraw.Draw(panel); draw.text((12, 8), title, fill="black", font=font)
            panel.paste(crop, (10, 36)); panels.append(panel)
            records.append({"fixture": name, "stage": stage, "pdf_path": str(pdf), "pdf_sha256": digest(pdf), "full_page_png": str(png), "png_sha256": digest(png)})
        canvas = Image.new("RGB", (1700, 535 * ((len(panels) + 1) // 2)), "#dddddd")
        for i, panel in enumerate(panels): canvas.paste(panel, ((i % 2) * 850, (i // 2) * 535))
        path = out / (group + ".png"); canvas.save(path)
        reports.append({"name": group, "path": str(path), "sha256": digest(path), "panels": records,
                        "scope": "Poppler rendering of actual LibreOffice-produced PDFs. Supplementary visual check; not browser/native/Word/WPS evidence."})
    return reports


def physical_corpus(path, cli, evidence, directory, safe_model_path):
    manifest_path = Path(__file__).with_name("table-cell-margins-corpus.json")
    manifest = json.loads(manifest_path.read_text())
    matches = [item for item in manifest["fixtures"] if item["name"] == path.name]
    check(len(matches) == 1, "Original physical-margin corpus is not pinned")
    target = matches[0]
    data = path.read_bytes()
    git_sha = hashlib.sha1(f"blob {len(data)}\0".encode() + data).hexdigest()
    check(len(data) == target["bytes"] and digest(path) == target["sha256"] and git_sha == target["git_blob_sha"], "Original physical-margin provenance mismatch")
    snap = xml_snapshot(path)
    check(len(snap["tables"]) == 1, "Original physical-margin table count changed")
    raw = snap["tables"][0]
    check(raw["total"] == {"w": "5000", "type": "pct"} and raw["indent"] == {"w": "-5", "type": "dxa"}, "Original existing unsupported geometry changed")
    horizontal = {n["side"]: {k: v for k, v in n.items() if k != "side"} for n in raw["table_margins"] if n["side"] in ["left", "right"]}
    check(horizontal == {s: {"w": "55", "type": "dxa"} for s in ["left", "right"]}, "Original physical margins changed")
    result = {**target, "path": str(path), "source_independent": snap, "manifest_sha256": digest(manifest_path),
              "fixture_kind": "original unmodified pinned public Apache POI regression fixture, not a user production file"}
    if cli:
        result["audit"] = w.run_cli(cli, evidence["commands"], "audit-docx", path)
        result["refusal"] = w.run_cli(cli, evidence["commands"], "import-docx", path, success=False)
        check("table" in result["refusal"].lower() or "page" in result["refusal"].lower(), "Original refusal unrelated to established geometry")
        error = w.run_cli(cli, evidence["commands"], "create-docx", safe_model_path, path, success=False)
        check("Output already exists" in error, "Original physical-margin source overwrite not refused")
        result["source_overwrite_refused"] = True
    check(digest(path) == target["sha256"], "Original physical-margin source changed")
    result["source_hash_unchanged"] = True
    result["custom_margin_roundtrip_claimed"] = False
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli", type=Path, default=Path("target/debug/rust-wps"))
    parser.add_argument("--root", type=Path, default=Path("tmp/word-table-cell-margins"))
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--contact-sheets", action="store_true")
    parser.add_argument("--corpus-dir", type=Path)
    parser.add_argument("--plain-corpus", type=Path)
    parser.add_argument("--physical-corpus", type=Path)
    args = parser.parse_args()
    cli = None if args.generate_only else args.cli.resolve()
    check(cli is None or cli.is_file(), "CLI missing")
    soffice = shutil.which("soffice")
    check(not args.render or soffice, "LibreOffice missing")
    check(not args.contact_sheets or args.render, "--contact-sheets requires --render")
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.root)).resolve()
    evidence = {"status": "running", "script_sha256": digest(Path(__file__)), "python": platform.python_version(),
                "python_docx": version("python-docx"), "pypdf": version("pypdf"), "cli": str(cli),
                "cli_sha256": None if cli is None else digest(cli), "commands": [], "sources": {},
                "scope": "Exact margin values and bounded rectangular geometry. PDF is installed LibreOffice only. Borders remain normalized; modern indentation is refused."}
    try:
        sources, manifest = fixtures(directory)
        evidence.update(sources=sources, expected_manifest=str(manifest), expected_manifest_sha256=digest(manifest))
        safe = None
        default_model = default_export = None
        for name, item in sources.items():
            path = Path(item["path"])
            if cli:
                item["audit"] = w.run_cli(cli, evidence["commands"], "audit-docx", path)
                if item["expected_action"] == "reject":
                    item["refusal"] = w.run_cli(cli, evidence["commands"], "import-docx", path, success=False)
                    check("table" in item["refusal"].lower() or "indent" in item["refusal"].lower(), name + ": unrelated refusal")
                else:
                    check(w.RISK not in item["audit"]["risks"], name + ": supported source marked unsafe")
                    if name == "modern-no-indent-values-only":
                        warning = "Horizontal cell margins preserved; source compatibility-dependent table position normalized"
                        check(warning in item["audit"]["risks"], name + ": compatibility-dependent position normalization warning missing")
                    imported = w.run_cli(cli, evidence["commands"], "import-docx", path)
                    item["import_risks"] = imported.get("risks")
                    model = imported["document"]
                    verify_model(model, item["expected"], name + " import")
                    if item["separated_anchor"]: verify_anchor_model(model, name + " import")
                    item["unchanged_export"] = exported(cli, evidence, directory, name + "-unchanged", model, item["expected"])
                    if name == "direct-108-108-indent-None":
                        default_model, default_export = copy.deepcopy(model), item["unchanged_export"]
                    edited = copy.deepcopy(model)
                    for tab, target in zip(w.model_tables(edited), item["edited_expected"]):
                        tab.setdefault("attrs", {})[KEY] = copy.deepcopy(target["margins"])
                    before, after = copy.deepcopy(model), copy.deepcopy(edited)
                    for doc in [before, after]:
                        for tab in w.model_tables(doc): tab.get("attrs", {}).pop(KEY, None)
                    check(before == after, name + ": margin edit changed unrelated model content")
                    item["edited_export"] = exported(cli, evidence, directory, name + "-edited", edited, item["edited_expected"])
                    if safe is None and item["expected"][0]["margins"] is not None:
                        safe = (model, item["unchanged_export"], item["expected"])
            check(digest(path) == item["sha256"], name + ": source changed")
            item["source_unchanged"] = True
        if cli:
            evidence["invalid_models"] = invalid_models(cli, evidence, directory, safe[0])
            evidence["tampered_packages"] = tamper(cli, evidence, directory, Path(safe[1]["path"]))
            explicit = copy.deepcopy(safe[0]); targets = copy.deepcopy(safe[2])
            next(w.model_tables(explicit))["attrs"][KEY] = {"left": 108, "right": 108}; targets[0]["margins"] = targets[0]["effective_margins"] = {"left": 108, "right": 108}
            evidence["authored_explicit_default_durable"] = exported(cli, evidence, directory, "authored-explicit-default", explicit, targets)
            restored = copy.deepcopy(default_model)
            next(w.model_tables(restored))["attrs"][KEY] = {"left": 180, "right": 240}
            next(w.model_tables(restored))["attrs"].pop(KEY)
            check(restored == default_model, "Margin reset changed unrelated model fields")
            reset_targets = sources["direct-108-108-indent-None"]["expected"]
            evidence["restored_absent_default"] = exported(cli, evidence, directory, "restored-absent-default", restored, reset_targets)
            check(evidence["restored_absent_default"]["sha256"] == default_export["sha256"], "Margin reset did not restore exact deterministic DOCX bytes")
            for item in sources.values():
                error = w.run_cli(cli, evidence["commands"], "create-docx", safe[1]["model_path"], item["path"], success=False)
                check("Output already exists" in error and digest(Path(item["path"])) == item["sha256"], "Source overwrite not refused")
                item["source_overwrite_refused"] = True
        if args.render:
            evidence["libreoffice_version"] = subprocess.check_output([soffice, "--version"], text=True).strip()
            render_all(soffice, directory, sources)
            evidence["pdf_file_inventory"] = [{"path": str(path), "sha256": digest(path)} for path in sorted((directory / "pdf").glob("*.pdf"))]
            if args.contact_sheets:
                evidence["visual_qa_contact_sheets"] = contact_sheets(directory, sources)
        if args.corpus_dir:
            spec = importlib.util.spec_from_file_location("position_corpus", Path(__file__).with_name("check-word-table-position.py"))
            pos = importlib.util.module_from_spec(spec); spec.loader.exec_module(pos)
            evidence["untouched_original_table_corpus"] = pos.corpus_check(args.corpus_dir.resolve(), cli, evidence, directory, None if safe is None else safe[1]["model_path"])
        if args.physical_corpus:
            evidence["untouched_original_physical_margin_refusal"] = physical_corpus(args.physical_corpus.resolve(), cli, evidence, directory, None if safe is None else safe[1]["model_path"])
        if args.plain_corpus:
            evidence["untouched_original_plain_positive"] = w.check_corpus(args.plain_corpus.resolve(), cli, evidence, directory)
        evidence["status"] = "fixtures-verified" if cli is None else "passed"
        evidence["summary"] = {"positive_sources": sum(i["expected_action"] == "roundtrip" for i in sources.values()),
                               "refusal_sources": sum(i["expected_action"] == "reject" for i in sources.values()),
                               "invalid_models": len(evidence.get("invalid_models", [])), "tamper_refusals": len(evidence.get("tampered_packages", [])),
                               "source_pdf": sum("source_pdf" in i for i in sources.values()),
                               "independent_reference_pdf": sum("independent_edited_reference_pdf" in i for i in sources.values()),
                               "actual_pdf_files": len(evidence.get("pdf_file_inventory", [])),
                               "calibration_only_pdf": sum("rejected_source_calibration_pdf" in i for i in sources.values()),
                               "adjacent_table_visual_only_pdf": sum((directory / "pdf" / (Path(path).stem + ".pdf")).is_file()
                                   for path in [sources["multiple-old-and-new"]["path"]] + [sources["multiple-old-and-new"][stage]["path"]
                                       for stage in ["unchanged_export", "edited_export"] if stage in sources["multiple-old-and-new"]]),
                               "export_pdf": sum("unchanged_export_pdf" in i for i in sources.values()) + sum("edited_export_pdf" in i for i in sources.values())}
    except Exception as error:
        evidence.update(status="failed", error=f"{type(error).__name__}: {error}")
        raise
    finally:
        evidence["source_hashes_unchanged"] = all(digest(Path(i["path"])) == i["sha256"] for i in evidence["sources"].values())
        evidence["reference_hashes_unchanged"] = all(digest(Path(i["edited_reference"]["path"])) == i["edited_reference"]["sha256"] for i in evidence["sources"].values() if "edited_reference" in i)
        evidence["script_unchanged_during_run"] = digest(Path(__file__)) == evidence["script_sha256"]
        evidence["cli_unchanged_during_run"] = cli is None or digest(cli) == evidence["cli_sha256"]
        report = directory / "report.json"; report.write_text(json.dumps(evidence, indent=2, ensure_ascii=False), encoding="utf-8")
        print(report, flush=True)
    print(json.dumps(evidence["summary"]), flush=True)


if __name__ == "__main__":
    main()
