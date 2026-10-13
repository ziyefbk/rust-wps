"""Bounded heading/body conversion from independently authored and untouched DOCX.

Uses the real browser-free application transaction through convert-word-block-type.mjs,
then the frozen Rust CLI. Source and output OOXML/python-docx reads are independent
of the application. Expected paragraph numerics are fixed in a separate manifest.
Optional LibreOffice/Poppler comparisons use separately authored target documents,
matching explicit font/size/bold/italic to isolate spacing and indentation. They do
not claim that selecting a heading retains the old body style's appearance.

python scripts/check-word-block-type.py --render
python scripts/check-word-block-type.py --generate-only
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips
from docx.text.paragraph import Paragraph
from lxml import etree as ET

ROOT = Path(__file__).resolve().parent.parent
EXPECTED = ROOT / 'tests/fixtures/word-block-type-expected.json'
HELPER = ROOT / 'scripts/convert-word-block-type.mjs'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
KEYS = ('spaceBefore', 'spaceAfter', 'lineSpacing', 'indentLeftTwips', 'indentRightTwips', 'firstLineIndentTwips', 'textAlign')


def load_script(name):
    spec = importlib.util.spec_from_file_location(name.replace('-', '_'), ROOT / 'scripts' / (name + '.py'))
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


# Reuse only independently implemented declaration/render/CLI utilities.
s = load_script('check-word-paragraph-spacing')
i = load_script('check-word-paragraph-indentation')
w = load_script('check-word-table-widths')
check, digest, run_cli = s.check, s.digest, s.run_cli


def nodes(model):
    if model['type'] in ('paragraph', 'heading'):
        yield model
    for child in model.get('content', []):
        yield from nodes(child)


def node_text(node):
    if node['type'] == 'text':
        return node['text']
    if node['type'] == 'hardBreak':
        return '\n'
    return ''.join(node_text(child) for child in node.get('content', []))


def controlled_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Twips(1440)
    # No existing source is modified or re-saved. These are new controlled probes.
    doc.styles['Normal'].font.name = 'DejaVu Sans'
    doc.styles['Normal'].font.size = Pt(12)
    return doc


def format_para(para, spacing, indent=None, align=None):
    s.set_spacing(para, *spacing)
    if indent is not None:
        i.set_indent(para, *indent)
    if align is not None:
        para.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                          'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]


def add_run(para, text, *, bold=False, italic=False, underline=False, strike=False, color='203040', size=12):
    run = para.add_run(text)
    run.font.name, run.font.size = 'DejaVu Sans', Pt(size)
    run.bold, run.italic, run.underline, run.font.strike = bold, italic, underline, strike
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def mixed_document():
    doc = controlled_document()
    p = doc.add_paragraph(); format_para(p, (240, 120, 241), (720, 360, 480), 'left')
    add_run(p, 'MixedBody '); add_run(p, 'Bold', bold=True); add_run(p, ' italic', italic=True)
    add_run(p, ' underline', underline=True); add_run(p, ' strike', strike=True)
    add_run(p, '\nMixedBodyNext', bold=True, italic=True, color='A02040', size=13.5)
    p = doc.add_paragraph(style='Heading 2'); format_para(p, (120, 240, 360), (480, 240, -120), 'right')
    add_run(p, 'MixedHeadingTwo ', bold=True); add_run(p, 'colored', italic=True, color='2050A0')
    p = doc.add_paragraph(); format_para(p, (0, 0, 240), (0, 0, 0), 'center'); add_run(p, 'ExplicitZero')
    p = doc.add_paragraph(style='Heading 3'); format_para(p, (241, 1, 241), (721, 239, 241), 'justify'); add_run(p, 'FractionalHeading', bold=True, underline=True)
    table = w.add_table(doc, [3600, 3600], rows=1)
    for cell, label, level, spacing, indent, alignment in [
        (table.cell(0, 0), 'TableBody', 0, (0, 120, 241), (180, 120, 120), 'left'),
        (table.cell(0, 1), 'TableHeading', 1, (120, 0, 300), (240, 180, -120), 'center'),
    ]:
        p = cell.paragraphs[0]; p.clear(); p.style = 'Normal' if level == 0 else f'Heading {level}'
        format_para(p, spacing, indent, alignment); add_run(p, label, bold=True, italic=True)
    p = doc.add_paragraph(); format_para(p, (0, 0, 241), (0, 0, 0))
    p = doc.add_paragraph(); format_para(p, (100, 200, 276), (360, 120, -180)); add_run(p, 'AfterTable')
    return doc


def stock_document():
    doc = Document()
    for level in range(4):
        p = doc.add_paragraph(style='Normal' if level == 0 else f'Heading {level}')
        add_run(p, 'StockBody' if level == 0 else f'StockHeading{level}')
    return doc


def render_document(source_heading, target=None):
    """Independent expected document from fixed source specifications, not model output."""
    doc = controlled_document()
    for index, (spacing, indent) in enumerate([
        ((120, 240, 241), (720, 360, 480)),
        ((0, 0, 240), (0, 0, 0)),
        ((240, 120, 360), (720, 240, -360)),
    ]):
        level = target if target is not None else (index + 1 if source_heading else 0)
        p = doc.add_paragraph(style='Normal' if level == 0 else f'Heading {level}')
        format_para(p, spacing, indent, 'left')
        # Explicit true marks neutralize inherited bold/italic style differences.
        # Hard breaks hold the line count stable; x/y still expose lost numerics.
        add_run(p, f'Render{index}First\nRender{index}Second', bold=True, italic=True, color='000000')
    return doc


def list_document():
    doc = controlled_document()
    for text, style in [('BeforeList', 'Normal'), ('RequiredListFirst', 'List Bullet'), ('AfterList', 'Normal')]:
        p = doc.add_paragraph(style=style); format_para(p, (0, 0, 241)); add_run(p, text)
    return doc


def direct_run_properties(run):
    rp = run.find('w:rPr', NS)
    if rp is None:
        return {}
    result = {}
    for tag, key in [('b', 'bold'), ('i', 'italic'), ('strike', 'strike')]:
        el = rp.find('w:' + tag, NS)
        if el is not None:
            result[key] = el.get(qn('w:val'), 'true') not in ('0', 'false', 'off')
    underline = rp.find('w:u', NS)
    if underline is not None:
        result['underline'] = underline.get(qn('w:val'), 'single') not in ('none', '0', 'false', 'off')
    for tag, key in [('sz', 'halfPoints'), ('color', 'rgb')]:
        el = rp.find('w:' + tag, NS)
        if el is not None:
            value = el.get(qn('w:val'))
            result[key] = int(value) if tag == 'sz' else value.upper()
    fonts = rp.find('w:rFonts', NS)
    if fonts is not None:
        result['fontFamily'] = fonts.get(qn('w:ascii'))
    return result


def significant(properties):
    return {key: val for key, val in properties.items() if val is not False}


def independent(path):
    """Read document.xml and python-docx, without trusting embedded rust-wps JSON."""
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read('word/document.xml'))
    spacing = s.independent_snapshot(path)['ooxml']
    with zipfile.ZipFile(path) as package:
        styles = ET.fromstring(package.read('word/styles.xml'))
    smap = {el.get(qn('w:styleId')): el for el in styles.findall('w:style', NS)}
    default_sid = next((sid for sid, el in smap.items() if el.get(qn('w:default')) in ('1', 'true', 'on') and el.get(qn('w:type')) == 'paragraph'), None)
    def ind_chain(sid, seen=()):
        check(sid not in seen, 'Independent style cycle')
        st = smap.get(sid)
        if st is None: return {}
        parent = st.find('w:basedOn', NS)
        return {**({} if parent is None else ind_chain(parent.get(qn('w:val')), seen + (sid,))), **i.direct_indent(st.find('w:pPr', NS))}
    indentation = []
    for para in root.findall('.//w:body//w:p', NS):
        pp = para.find('w:pPr', NS); st = para.find('w:pPr/w:pStyle', NS); align = para.find('w:pPr/w:jc', NS)
        direct = i.direct_indent(pp)
        indentation.append({'indentation': {**i.direct_indent(styles.find('w:docDefaults/w:pPrDefault/w:pPr', NS)), **ind_chain(default_sid), **ind_chain(None if st is None else st.get(qn('w:val'))), **direct}, 'direct': direct, 'align': None if align is None else align.get(qn('w:val')), 'in_table': any(a.tag == qn('w:tc') for a in para.iterancestors())})
    doc = Document(path)
    py_paragraphs = [Paragraph(el, doc) for el in doc.element.body.iter(qn('w:p'))]
    xml_paragraphs = root.findall('.//w:body//w:p', NS)
    check(len(xml_paragraphs) == len(py_paragraphs) == len(spacing) == len(indentation), 'Independent paragraph count differs')
    rows = []
    for para, py, sp, ind in zip(xml_paragraphs, py_paragraphs, spacing, indentation):
        fmt = py.paragraph_format
        py_indent = i.attrs3(*[None if value is None else value.twips for value in (fmt.left_indent, fmt.right_indent, fmt.first_line_indent)])
        check(py_indent == ind['direct'], f'{path.name}: independent/python-docx direct indentation differs')
        style = para.find('w:pPr/w:pStyle', NS)
        sid = None if style is None else style.get(qn('w:val'))
        atoms = []
        direct_runs = para.findall('w:r', NS)
        check(len(direct_runs) == len(py.runs), 'Independent/python-docx run count differs')
        for run, py_run in zip(direct_runs, py.runs):
            props = direct_run_properties(run)
            py_props = {key: value for key, value in [('bold', py_run.bold), ('italic', py_run.italic),
                ('strike', py_run.font.strike), ('underline', py_run.underline),
                ('halfPoints', None if py_run.font.size is None else round(py_run.font.size.pt * 2)),
                ('fontFamily', py_run.font.name),
                ('rgb', None if py_run.font.color.rgb is None else str(py_run.font.color.rgb))] if value is not None}
            check(props == py_props, f'{path.name}: independent/python-docx direct run properties differ: {props} vs {py_props}')
            run_text = ''.join((el.text or '') if el.tag == qn('w:t') else '\n' if el.tag == qn('w:br') and el.get(qn('w:type')) in (None, 'textWrapping') else '\t' if el.tag == qn('w:tab') else '' for el in run)
            # Page-break runs have no text in the bounded paragraph model.
            check(run_text == py_run.text, f'{path.name}: independent/python-docx run text differs')
            color = run.find('w:rPr/w:color', NS)
            theme_color = None if color is None else color.get(qn('w:themeColor'))
            atoms.extend({'text': char, 'properties': significant(props), 'theme_color': theme_color} for char in run_text)
        check(sp['text'] == py.text, f'{path.name}: independent/python-docx paragraph text differs')
        attrs = {**{k: (0 if k != 'lineSpacing' else 240) if v is None else v for k, v in sp['spacing'].items()}, **ind['indentation']}
        if ind.get('align') is not None:
            attrs['textAlign'] = 'justify' if ind['align'] == 'both' else ind['align']
        rows.append({'text': py.text, 'attrs': attrs, 'style': sid, 'in_table': ind['in_table'],
                     'page_break_only': sp['page_break_only'], 'atoms': atoms,
                     'direct_spacing': sp['direct_xml'], 'direct_indent': ind['direct']})
    return rows


def paragraph_rows(rows):
    return [row for row in rows if not row['page_break_only']]


def mark_properties(marks):
    result = {}
    for mark in marks:
        kind, attrs = mark['type'], mark.get('attrs', {})
        if kind in ('bold', 'italic', 'underline', 'strike'):
            result[kind] = True
        elif kind == 'fontSize': result['halfPoints'] = attrs['halfPoints']
        elif kind == 'fontFamily': result['fontFamily'] = attrs['name']
        elif kind == 'textColor': result['rgb'] = attrs['rgb']
        else: raise AssertionError('Unexpected inline mark in bounded corpus: ' + kind)
    return result


def model_atoms(node):
    result = []
    for child in node.get('content', []):
        if child['type'] == 'hardBreak':
            result.append({'text': '\n', 'properties': None})
        else:
            check(child['type'] == 'text', 'Unexpected inline content')
            result.extend({'text': char, 'properties': mark_properties(child.get('marks', []))} for char in child['text'])
    return result


def compare_to_model(rows, model, label, direct_only=False):
    rows, paragraphs = paragraph_rows(rows), list(nodes(model))
    check(len(rows) == len(paragraphs), label + ': paragraph count differs')
    for index, (row, node) in enumerate(zip(rows, paragraphs)):
        attrs = {k: node['attrs'][k] for k in KEYS if node.get('attrs', {}).get(k) is not None}
        actual_attrs = dict(row['attrs'])
        if direct_only and actual_attrs.get('textAlign') == 'left' and 'textAlign' not in attrs:
            actual_attrs.pop('textAlign')
        check(actual_attrs == attrs, f'{label} p{index}: numeric/alignment attrs differ: {row["attrs"]} != {attrs}')
        check(row['text'] == node_text(node), label + ': text differs')
        if direct_only:
            source_level = int(row['style'][-1]) if row['style'] in ('Heading1', 'Heading2', 'Heading3') else 0
            check(node['type'] == ('heading' if source_level else 'paragraph'), label + ': imported source block type differs')
            if source_level: check(node['attrs']['level'] == source_level, label + ': imported source heading level differs')
        if not direct_only:
            check(row['style'] == (f'Heading{node["attrs"]["level"]}' if node['type'] == 'heading' else None), label + ': heading/body style differs')
        expected_atoms = model_atoms(node)
        check(len(row['atoms']) == len(expected_atoms), label + ': inline length differs')
        for atom, expected in zip(row['atoms'], expected_atoms):
            check(atom['text'] == expected['text'], label + ': inline text differs')
            if expected['properties'] is not None:
                if direct_only:
                    # A theme color's w:val is only fallback; the importer resolves
                    # the theme/tint. Preserve the actual imported RGB thereafter.
                    check(all(expected['properties'].get(k) == v for k, v in atom['properties'].items() if not (k == 'rgb' and atom.get('theme_color'))), label + ': source direct inline mark differs')
                else:
                    check(atom['properties'] == expected['properties'], label + ': exported inline marks differ')
        if not direct_only:
            check(row['direct_indent'] == {k: attrs[k] for k in i.KEYS if k in attrs}, label + ': direct optional indentation/zero lost')
            for key, xml in zip(s.KEYS, ('before', 'after', 'line')):
                check(row['direct_spacing'].get(xml) == str(attrs[key]), label + ': explicit spacing/zero lost')


def compare_manifest(rows, expected, label):
    rows = paragraph_rows(rows)
    check(len(rows) == len(expected), label + ': manifest paragraph count differs')
    for row, target in zip(rows, expected):
        check(row['attrs'] == target['attrs'], label + ': fixed expected attributes differ')
        check(row['text'] == target['text'], label + ': fixed expected text differs')
        check(row['in_table'] == target.get('in_table', False), label + ': table location differs')


def comparable_model(value):
    value = copy.deepcopy(value)
    def visit(node):
        attrs = node.get('attrs', {})
        if node['type'] in ('paragraph', 'heading') and attrs.get('textAlign') is None: attrs.pop('textAlign', None)
        if node['type'] in ('tableCell', 'tableHeader'):
            for key, default in [('colspan', 1), ('rowspan', 1), ('align', None), ('colwidth', None)]:
                if attrs.get(key) == default: attrs.pop(key, None)
        if not attrs: node.pop('attrs', None)
        if 'marks' in node: node['marks'].sort(key=lambda mark: json.dumps(mark, sort_keys=True))
        for child in node.get('content', []): visit(child)
    visit(value)
    return value


def expected_conversion(source, level):
    result = copy.deepcopy(source)
    for node in nodes(result):
        node['type'] = 'heading' if level else 'paragraph'
        attrs = node.setdefault('attrs', {})
        attrs.pop('level', None)
        if level: attrs['level'] = level
        if not attrs: node.pop('attrs')
    return result


def convert(evidence, directory, name, source, level, reject=False):
    input_path, output = directory / (name + '-input.json'), directory / (name + '-converted.json')
    input_path.write_text(json.dumps(source, ensure_ascii=False, indent=2), encoding='utf-8')
    before = digest(input_path)
    command = ['node', str(HELPER), str(input_path), str(output), str(level)]
    # Node's enormous data-URL stack can be truncated when stderr is a pipe.
    # Capture to a real file so rejection assertions see the actual diagnostic.
    stderr_path = directory / (name + '-helper-stderr.txt')
    with stderr_path.open('w', encoding='utf-8') as stderr_file:
        run = subprocess.run(command, cwd=ROOT, stdout=subprocess.PIPE, stderr=stderr_file, text=True, timeout=90)
    run.stderr = stderr_path.read_text(encoding='utf-8')
    diagnostic = '\n'.join(line for line in run.stderr.splitlines() if 'data:text/javascript;base64,' not in line)
    result = {'command': command, 'returncode': run.returncode, 'stdout': run.stdout, 'stderr': diagnostic,
              'input_sha256': before, 'input_unchanged': digest(input_path) == before,
              'raw_stderr_path': str(stderr_path), 'raw_stderr_sha256': digest(stderr_path)}
    evidence['commands'].append(result)
    check(result['input_unchanged'], 'Application helper modified input JSON')
    if reject:
        check(run.returncode != 0 and not output.exists() and '列表首段须保留为正文' in run.stderr, 'Incompatible conversion was not refused with the specific list-first diagnostic and without output')
        return {'rejected_without_output': True, 'input_unchanged': True, 'diagnostic': diagnostic}
    check(run.returncode == 0 and output.is_file(), run.stdout + run.stderr)
    metadata = json.loads(run.stdout)
    changed_count = sum(node['type'] != ('heading' if level else 'paragraph') or (bool(level) and node.get('attrs', {}).get('level') != level) for node in nodes(source))
    check(metadata['changed'] == bool(changed_count) and metadata['steps'] == changed_count, name + ': changed paragraph count differs')
    check(metadata['undoDepth'] == int(bool(changed_count)) and metadata['selection'] == {'type': 'all'}, name + ': undo depth or all-selection changed')
    result['transaction_metadata'] = metadata
    result['exact_changed_count_and_single_undo_step'] = True
    model = json.loads(output.read_text())
    check(comparable_model(model) == comparable_model(expected_conversion(source, level)), name + ': application-converted model differs beyond schema defaults/mark ordering')
    result.update(output_path=str(output), output_sha256=digest(output), complete_expected_model_equal_except_declared_schema_defaults_and_mark_order=True)
    if level == 0:
        retry = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, timeout=90)
        check(retry.returncode != 0 and 'EEXIST' in retry.stderr and digest(output) == result['output_sha256'] and digest(input_path) == before, 'Application helper overwrote an existing file')
        result['helper_no_clobber'] = True
    return model, result


def export(evidence, cli, directory, name, model, source):
    path = directory / (name + '.json'); path.write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding='utf-8')
    output = path.with_suffix('.docx')
    run_cli(cli, evidence['commands'], 'create-docx', path, output)
    check(run_cli(cli, evidence['commands'], 'read-docx-model', output) == model, name + ': trusted reopen changed complete model')
    rows = independent(output); compare_to_model(rows, model, name + ' output')
    original_hash, output_hash = digest(source), digest(output)
    for target in [source, output]:
        failure = run_cli(cli, evidence['commands'], 'create-docx', path, target, success=False)
        check('Output already exists' in failure, name + ': overwrite was not refused')
    check(digest(source) == original_hash and digest(output) == output_hash, name + ': no-clobber changed existing bytes')
    return {'path': str(output), 'sha256': output_hash, 'trusted_model_exact': True,
            'independent_numeric_text_marks_equal': True, 'source_and_output_no_clobber': True, 'independent': rows}


def compare_pdf(reference, edited, label):
    check(len(reference['lines']) == len(edited['lines']) == 6, label + ': expected six lines')
    maximum = {key: 0.0 for key in ('x', 'x_end', 'y')}
    for a, b in zip(reference['lines'], edited['lines']):
        check(a['text'] == b['text'] and a['page'] == b['page'], label + ': line text/page differs')
        for key in maximum:
            maximum[key] = max(maximum[key], abs(a[key] - b[key]))
            check(abs(a[key] - b[key]) <= 0.35, f'{label}: {key} differs by {abs(a[key] - b[key]):.4f} pt')
    for group, expected_delta, expected_left in [(0, 24, 108), (1, 0, 72), (2, -18, 108)]:
        first, second = edited['lines'][2 * group:2 * group + 2]
        check(abs(first['x'] - second['x'] - expected_delta) <= 0.35, label + ': fixed first-line/hanging displacement differs')
        check(abs(second['x'] - expected_left) <= 0.35, label + ': fixed continuation left indent differs')
    return {'same_lines_pages': True, 'line_count': 6, 'max_delta_pt': maximum,
            'fixed_indent_anchors_equal': True, 'scope': 'LibreOffice PDF x/y only; explicit matching fonts/size/marks; independent target style reference'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path, default=ROOT / 'target/debug/rust-wps')
    parser.add_argument('--root', type=Path, default=ROOT / 'tmp/word-block-type')
    parser.add_argument('--generate-only', action='store_true')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--soffice')
    parser.add_argument('--corpus', type=Path, default=ROOT / 'tmp/corpus/SampleDoc.docx')
    args = parser.parse_args()
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    report = directory / 'report.json'; cli = args.cli.resolve()
    soffice = args.soffice or shutil.which('soffice') if args.render else None
    manifest = json.loads(EXPECTED.read_text())
    evidence = {'status': 'running', 'base_commit': manifest['base_commit'], 'commands': [], 'sources': {},
                'script_sha256': digest(Path(__file__)), 'manifest_sha256': digest(EXPECTED),
                'helper_sha256': digest(HELPER) if HELPER.exists() else None,
                'application_source_sha256': {str(path.relative_to(ROOT)): digest(path) for path in [ROOT / 'ui/word-block-type.ts', ROOT / 'tests/ui/word-block-type-harness.mjs'] if path.exists()},
                'cli': str(cli), 'cli_sha256': digest(cli), 'python_version': platform.python_version(),
                'python_docx_version': version('python-docx'),
                'utility_sha256': {str(Path(x.__file__).relative_to(ROOT)): digest(Path(x.__file__)) for x in (s, i, w)},
                'limitations': [
                    'Bounded independent synthetic documents and one untouched public Apache POI regression document; not user production files.',
                    'Selecting heading/body intentionally changes named heading/default style appearance. Attribute/mark retention does not promise old-style appearance.',
                    'Optional LibreOffice PDF evidence isolates explicit matching typography and x/y positions; not app print, browser, native WebView2, Word/WPS, or general layout fidelity.',
                    'The original SampleDoc uses theme/tint color declarations. Source fallback RGB is not a fixed-color oracle; actual imported RGB is preserved exactly through all conversions/exports.',
                    'Whole-model source comparison permits only declared schema defaults (null alignment and unit table-cell spans) plus mark ordering; trusted reopen is exact full JSON equality.',
                    'No browser/native execution, downloads, source mutation, commits, or publication.']}
    try:
        if args.render:
            check(soffice, 'Installed LibreOffice required for --render')
            evidence['libreoffice_version'] = subprocess.run([soffice, '--version'], capture_output=True, text=True, check=True).stdout.strip()
        factories = {'mixed-formatting': mixed_document, 'stock-defaults': stock_document,
                     'render-body': lambda: render_document(False), 'render-heading': lambda: render_document(True),
                     'list-incompatible': list_document}
        for name, factory in factories.items():
            source = directory / (name + '-source.docx'); factory().save(source)
            rows = independent(source)
            item = {'path': str(source), 'sha256': digest(source), 'kind': 'independently authored python-docx source', 'source_independent': rows}
            evidence['sources'][name] = item
            compare_manifest(rows, manifest['sources'][name], name + ' source')
        original = args.corpus.resolve(); pin = manifest['original']
        check(original.is_file() and digest(original) == pin['sha256'], 'Missing or changed untouched original corpus')
        check(hashlib.sha1(b'blob ' + str(original.stat().st_size).encode() + b'\0' + original.read_bytes()).hexdigest() == pin['git_blob_sha'], 'Original git blob differs')
        evidence['sources']['original-SampleDoc'] = {**pin, 'path': str(original), 'source_independent': independent(original)}
        if not args.generate_only:
            check(HELPER.is_file(), 'Real application conversion helper is not ready')
            for name, item in evidence['sources'].items():
                source = Path(item['path']); imported = run_cli(cli, evidence['commands'], 'import-docx', source)
                model = imported['document']; item['compatibility'] = imported['compatibility']; item['imported_model'] = model
                compare_to_model(item['source_independent'], model, name + ' import', direct_only=True)
                item['source_theme_color_normalization'] = list({(atom['theme_color'], atom['properties'].get('rgb'), expected_atom['properties'].get('rgb')) for row, node in zip(paragraph_rows(item['source_independent']), nodes(model)) for atom, expected_atom in zip(row['atoms'], model_atoms(node)) if atom.get('theme_color') and expected_atom['properties'] is not None})
                item['conversions'] = {}
                for level in [0, 1, 2, 3]:
                    label = f'{name}-level{level}'
                    if name == 'list-incompatible' and level:
                        item['conversions'][str(level)] = convert(evidence, directory, label, model, level, reject=True)
                        continue
                    converted, result = convert(evidence, directory, label, model, level)
                    result['export'] = export(evidence, cli, directory, label, converted, source)
                    item['conversions'][str(level)] = result
                    if soffice and name.startswith('render-'):
                        reference = directory / (label + '-independent-expected.docx')
                        render_document(name == 'render-heading', level).save(reference)
                        reference_rows = independent(reference)
                        compare_manifest(reference_rows, manifest['sources'][name], label + ' independent expected')
                        ref_pdf = i.render_pdf(soffice, directory, reference)
                        actual_pdf = i.render_pdf(soffice, directory, Path(result['export']['path']))
                        result['render'] = {'reference': str(reference), 'reference_sha256': digest(reference),
                            'reference_pdf': ref_pdf, 'output_pdf': actual_pdf, 'comparison': compare_pdf(ref_pdf, actual_pdf, label)}
                if name in ('mixed-formatting', 'stock-defaults', 'original-SampleDoc'):
                    heading = json.loads(Path(item['conversions']['2']['output_path']).read_text())
                    restored, result = convert(evidence, directory, name + '-heading2-to-body', heading, 0)
                    check(comparable_model(restored) == comparable_model(expected_conversion(model, 0)), name + ': two-step restore differs')
                    result['export'] = export(evidence, cli, directory, name + '-heading2-to-body', restored, source)
                    item['heading2_then_body'] = result
                check(digest(source) == item['sha256'], name + ': source modified')
                item['source_unchanged'] = True
            check(digest(cli) == evidence['cli_sha256'], 'Frozen CLI changed during verification')
            check(digest(HELPER) == evidence['helper_sha256'], 'Application helper changed during verification; rerun against final code')
            for path, sha in evidence['application_source_sha256'].items():
                check(digest(ROOT / path) == sha, 'Application source changed during verification; rerun against final code')
        evidence['status'] = 'generated' if args.generate_only else 'passed'
        evidence['counts'] = {'source_packages': len(evidence['sources']), 'synthetic_sources': len(factories), 'untouched_originals': 1,
            'successful_exports': sum(sum('export' in c for c in x.get('conversions', {}).values()) + int('heading2_then_body' in x) for x in evidence['sources'].values()),
            'atomic_refusals': sum(sum('rejected_without_output' in c for c in x.get('conversions', {}).values()) for x in evidence['sources'].values()),
            'render_comparisons': sum(sum('render' in c for c in x.get('conversions', {}).values()) for x in evidence['sources'].values())}
    except Exception as error:
        evidence.update(status='failed', error=str(error))
        raise
    finally:
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'status': evidence['status'], 'report': str(report), 'counts': evidence.get('counts'), 'error': evidence.get('error')}))


if __name__ == '__main__':
    main()
