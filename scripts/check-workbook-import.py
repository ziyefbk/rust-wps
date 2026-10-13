"""Independent openpyxl -> real Rust import -> export -> openpyxl round trip."""
from datetime import datetime
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import openpyxl

root = Path(__file__).resolve().parents[1]
results = []
for date1904 in [False, True]:
    directory = Path(tempfile.mkdtemp(prefix="import-", dir=root / "tmp/spreadsheets"))
    source, model_path, output = (directory / name for name in ["external.xlsx", "model.json", "edited.xlsx"])
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["001", 12.5, True, "=B1*2"])
    sheet["B1"].number_format = "0.00%"
    sheet["B1"].font = openpyxl.styles.Font(bold=True)
    sheet["E1"] = datetime(2024, 1, 1)
    sheet["E1"].number_format = "yyyy-mm-dd"
    sheet["F4"].number_format = "yyyy-mm-dd"
    book.create_sheet("Hidden").sheet_state = "hidden"
    book.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904 if date1904 else openpyxl.utils.datetime.CALENDAR_WINDOWS_1900
    book.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    cli = root / "target/debug/rust-wps.exe"
    def run(*arguments):
        return json.loads(subprocess.check_output([str(cli), *map(str, arguments)], encoding="utf-8"))
    report = run("import-xlsx", source)
    assert not report["blockers"] and report["model"] and report["risks"]
    model = report["model"]
    assert model["sheets"][1]["hidden"]
    assert model.get("date1904", False) == date1904
    model["sheets"][0]["cells"][1]["content"]["value"] = 24
    model_path.write_text(json.dumps(model), encoding="utf-8")
    run("create-xlsx", model_path, output)
    copy = openpyxl.load_workbook(output)
    values = [copy.active.cell(1, column).value for column in range(1, 5)]
    assert values == ["001", 24, True, "=B1*2"], values
    assert copy.epoch == book.epoch
    assert copy.active["E1"].value == datetime(2024, 1, 1)
    assert copy.active["B1"].number_format == "0.00%"
    assert copy.active["F4"].number_format == "yyyy-mm-dd"
    assert copy["Hidden"].sheet_state == "hidden"
    assert openpyxl.load_workbook(output, data_only=True).active["D1"].value is None
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    result = {"realRustImportExport": True, "independentOpenpyxl": True, "sourceUnchanged": True, "formulasUncached": True, "values": values, "date1904": date1904, "datePreserved": str(copy.active["E1"].value)}
    (directory / "report.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    results.append(result)
print(json.dumps({"systems": results}))
