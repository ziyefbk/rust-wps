"""Independent bounded contextualSpacing regression.

Every source retains the python-docx template's original contextual declarations.
Synthetic controls are authored independently of rust-wps. Fixed expectations and
explicit, newly authored numeric reference packages are separate from importer
output. --generate-only never invokes the product CLI. --render measures only
LibreOffice PDF text baselines, not browser/native printing or Word/WPS fidelity.

python scripts/check-word-contextual-spacing.py --generate-only --render
python scripts/check-word-contextual-spacing.py --cli target/debug/rust-wps --render
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docx.text.paragraph import Paragraph
from lxml import etree as ET

ROOT = Path(__file__).resolve().parent.parent
EXPECTED = ROOT / 'tests/fixtures/word-contextual-spacing-expected.json'
CORPUS = Path(__file__).with_name('contextual-spacing-corpus.json')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
KEYS = ('spaceBefore', 'spaceAfter', 'lineSpacing')
RISK = 'Unsupported paragraph spacing cannot be preserved in an editable copy'
NORMALIZATION_WARNING = 'Contextual same-style spacing resolved to fixed paragraph values; named-style and reorder dynamics are not retained'
MARKERS = ('ProbeBefore', 'ProbeOne', 'ProbeTwo', 'ProbeThree', 'ProbeAfter')


def load_spacing_helpers():
    spec = importlib.util.spec_from_file_location('spacing_helpers', Path(__file__).with_name('check-word-paragraph-spacing.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


s = load_spacing_helpers()
check, digest, run_cli = s.check, s.digest, s.run_cli


def contextual(obj, value):
    """Author a declaration, never delete a source declaration."""
    ppr = obj._element.get_or_add_pPr()
    check(ppr.find('w:contextualSpacing', NS) is None, 'Refusing duplicate contextual declaration')
    node = OxmlElement('w:contextualSpacing')
    if value is not None:
        node.set(qn('w:val'), str(value).lower() if isinstance(value, bool) else str(value))
    ppr.append(node)


def fresh():
    doc = Document()
    doc.styles['Normal'].font.name = 'DejaVu Sans'
    doc.styles['Normal'].font.size = Pt(12)
    for name in ('ProbeA', 'ProbeB', 'ProbeParent', 'ProbeChild'):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles['Normal']
    return doc


def format_paragraph(para, text, style='ProbeA', before=0, after=0, context='absent', line=240):
    para.text, para.style = text, style
    fmt = para.paragraph_format
    if before is not None:
        fmt.space_before = Twips(before)
    if after is not None:
        fmt.space_after = Twips(after)
    if line is not None:
        fmt.line_spacing = line / 240
    if context != 'absent':
        contextual(para, context)
    return para


def p(doc, text, **kwargs):
    return format_paragraph(doc.add_paragraph(), text, **kwargs)


def number(para):
    num = para._element.get_or_add_pPr().get_or_add_numPr()
    num.get_or_add_ilvl().val, num.get_or_add_numId().val = 0, 5


def controlled_table(doc, rows):
    """Isolate spacing with the editor's established half-point table borders.

    Bare-table controls are a separate geometry diagnostic: the old exporter
    already adds these borders, shifting cell baselines 0.5 pt per row edge.
    These newly authored probes must not claim arbitrary table-style fidelity.
    """
    table = doc.add_table(rows=rows, cols=1)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = OxmlElement('w:' + edge)
        node.set(qn('w:val'), 'single'); node.set(qn('w:sz'), '4')
        borders.append(node)
    table._tbl.tblPr.append(borders)
    return table


def default_context(doc, value):
    ppr = doc.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr', NS)
    check(ppr is not None, 'Missing template pPrDefault')
    node = OxmlElement('w:contextualSpacing')
    node.set(qn('w:val'), 'true' if value else 'false')
    ppr.append(node)


def template_context_signature(root):
    return {style.get(qn('w:styleId')): ET.tostring(style.find('w:pPr/w:contextualSpacing', NS), method='c14n').decode()
            for style in root.findall('w:style', NS)
            if style.find('w:pPr/w:contextualSpacing', NS) is not None}


def independent_snapshot(path):
    """Read style identity, declaration inheritance, per-paragraph numerics, and boundaries."""
    snapshot = s.independent_snapshot(path)
    with zipfile.ZipFile(path) as package:
        body = ET.fromstring(package.read('word/document.xml'))
        styles = ET.fromstring(package.read('word/styles.xml'))
    smap = {node.get(qn('w:styleId')): node for node in styles.findall('w:style', NS)}
    default_id = next((sid for sid, node in smap.items()
                       if node.get(qn('w:type')) == 'paragraph' and node.get(qn('w:default')) in ('1', 'true', 'on')), None)
    defaults = styles.find('w:docDefaults/w:pPrDefault/w:pPr', NS)

    def chain(sid, seen=()):
        check(sid not in seen, 'Independent style cycle')
        style = smap.get(sid)
        if style is None:
            return []
        based = style.find('w:basedOn', NS)
        return ([] if based is None else chain(based.get(qn('w:val')), seen + (sid,))) + [style.find('w:pPr', NS)]

    paragraphs = body.findall('.//w:body//w:p', NS)
    containers = {}
    for paragraph, row in zip(paragraphs, snapshot['ooxml']):
        ppr = paragraph.find('w:pPr', NS)
        style = paragraph.find('w:pPr/w:pStyle', NS)
        sid = default_id if style is None else style.get(qn('w:val'))
        history = []
        for declaration in [defaults, *chain(sid), ppr]:
            flag = None if declaration is None else declaration.find('w:contextualSpacing', NS)
            if flag is not None:
                history.append(flag.get(qn('w:val'), 'true'))
        parent = paragraph.getparent()
        container = containers.setdefault(parent, len(containers))
        row.update(effective_style_id=sid, contextual_history=history,
                   effective_contextual=False if not history else history[-1] not in ('0', 'false', 'off'),
                   container=container, previous_element=None if paragraph.getprevious() is None else ET.QName(paragraph.getprevious()).localname)
    snapshot['template_context_declarations'] = template_context_signature(styles)
    return snapshot


def make_fixtures(directory, manifest):
    documents = {}
    def add(name, doc):
        check(name not in documents, 'Duplicate fixture')
        documents[name] = doc

    # Exactly stock Document() plus two paragraphs. No style/default/numbering changes.
    for suffix, style in [('normal', 'Normal'), ('number', 'List Number'), ('bullet', 'List Bullet')]:
        doc = Document()
        doc.add_paragraph('ProbeOne', style=style)
        doc.add_paragraph('ProbeTwo', style=style)
        add('stock-' + suffix, doc)

    for name, left, right, after, before, second_style in [
        ('same-both', True, True, 200, 240, 'ProbeA'),
        ('same-off', False, False, 200, 240, 'ProbeA'),
        ('different-both', True, True, 200, 240, 'ProbeB'),
        ('left-own-zero', True, False, 0, 240, 'ProbeA'),
        ('left-opposite-zero', True, False, 200, 0, 'ProbeA'),
        ('right-own-zero', False, True, 200, 0, 'ProbeA'),
        ('right-opposite-zero', False, True, 0, 240, 'ProbeA'),
        ('reject-left-positive-positive', True, False, 200, 240, 'ProbeA'),
        ('reject-right-positive-positive', False, True, 200, 240, 'ProbeA'),
        ('reject-left-absent-positive-positive', True, 'absent', 200, 240, 'ProbeA'),
        ('canonical-on-off', 'on', 'off', 200, 0, 'ProbeA'),
        ('canonical-one-zero', '1', '0', 0, 240, 'ProbeA'),
        ('canonical-val-absent', None, None, 200, 240, 'ProbeA'),
    ]:
        doc = fresh()
        p(doc, 'ProbeOne', after=after, context=left)
        p(doc, 'ProbeTwo', style=second_style, before=before, context=right)
        add(name, doc)

    for name, child, direct in [('inherited-true', 'absent', 'absent'), ('inherited-false', False, 'absent'), ('direct-false', 'absent', False)]:
        doc = fresh()
        contextual(doc.styles['ProbeParent'], True)
        doc.styles['ProbeChild'].base_style = doc.styles['ProbeParent']
        if child != 'absent':
            contextual(doc.styles['ProbeChild'], child)
        p(doc, 'ProbeOne', style='ProbeChild', after=200, context=direct)
        p(doc, 'ProbeTwo', style='ProbeChild', before=240, context=direct)
        add(name, doc)

    for name, override in [('docdefaults-true', False), ('docdefaults-override-false', True)]:
        doc = fresh()
        default_context(doc, True)
        if override:
            contextual(doc.styles['ProbeA'], False)
        p(doc, 'ProbeOne', after=200)
        p(doc, 'ProbeTwo', before=240)
        add(name, doc)

    doc = fresh()
    contextual(doc.styles['ProbeParent'], True)
    doc.styles['ProbeParent'].paragraph_format.space_before = Twips(120)
    doc.styles['ProbeParent'].paragraph_format.space_after = Twips(180)
    doc.styles['ProbeChild'].base_style = doc.styles['ProbeParent']
    doc.styles['ProbeChild'].paragraph_format.space_before = Twips(240)
    p(doc, 'ProbeOne', style='ProbeChild', before=None, after=200)
    p(doc, 'ProbeTwo', style='ProbeChild', before=None, after=None)
    add('partial-spacing-inheritance', doc)

    doc = fresh()
    p(doc, 'ProbeOne', before=100, after=200, context=True)
    add('single-first-last', doc)
    doc = fresh()
    p(doc, 'ProbeOne', before=100, after=200, context=True)
    p(doc, 'ProbeTwo', before=240, after=80, context=True)
    p(doc, 'ProbeThree', style='ProbeB', before=160, after=120, context=True)
    add('three-admitted', doc)
    doc = fresh()
    p(doc, 'ProbeBefore', style='ProbeB', context=False)
    p(doc, 'ProbeOne', before=120, after=200, context=True)
    p(doc, 'ProbeTwo', before=240, after=160, context=True)
    p(doc, 'ProbeAfter', style='ProbeB', context=False)
    add('outer-margin-anchors', doc)
    doc = fresh()
    p(doc, 'ProbeOne', after=200, context=False)
    p(doc, 'ProbeTwo', before=240, after=80, context=True)
    p(doc, 'ProbeThree', before=160, context=False)
    add('reject-three-asymmetric', doc)

    for name, numbered_index in [('numbered-body-same-style', 0), ('body-numbered-same-style', 1)]:
        doc = fresh()
        paras = [p(doc, 'ProbeOne', after=200, context=True), p(doc, 'ProbeTwo', before=240, context=True)]
        number(paras[numbered_index])
        add(name, doc)

    doc = fresh()
    table = controlled_table(doc, 1)
    format_paragraph(table.cell(0, 0).paragraphs[0], 'ProbeOne', after=200, context=True)
    format_paragraph(table.cell(0, 0).add_paragraph(), 'ProbeTwo', before=240, context=True)
    add('same-cell', doc)
    doc = fresh()
    table = controlled_table(doc, 2)
    format_paragraph(table.cell(0, 0).paragraphs[0], 'ProbeOne', after=200, context=True)
    format_paragraph(table.cell(1, 0).paragraphs[0], 'ProbeTwo', before=240, context=True)
    add('separate-cells', doc)
    doc = fresh()
    p(doc, 'ProbeOne', after=200, context=True)
    table = controlled_table(doc, 1)
    format_paragraph(table.cell(0, 0).paragraphs[0], 'ProbeTwo', before=240, after=200, context=True)
    p(doc, 'ProbeThree', before=240, context=True)
    add('body-table-body', doc)

    doc = fresh()
    doc.styles['ProbeB'].base_style = doc.styles['ProbeA']
    p(doc, 'ProbeOne', after=200, context=True)
    p(doc, 'ProbeTwo', style='ProbeB', before=240, context=True)
    add('different-leaf-same-parent', doc)
    doc = fresh()
    p(doc, 'ProbeOne', after=200, context=True)
    p(doc, 'ProbeTwo', style='ProbeB', before=240, context=True)
    # Equal display names do not make different style IDs identical.
    doc.styles['ProbeB'].element.find('w:name', NS).set(qn('w:val'), 'ProbeA')
    add('different-ids-same-display-name', doc)
    doc = fresh()
    p(doc, 'ProbeOne', after=200, context=True)
    p(doc, 'ProbeTwo', style='ProbeB', before=240, context=True)
    first_name = doc.styles['ProbeA'].element.find('w:name', NS)
    first_name.getparent().remove(first_name)
    doc.styles['ProbeB'].element.find('w:name', NS).set(qn('w:val'), 'ProbeA')
    add('missing-name-collides-with-style-id', doc)
    doc = fresh()
    contextual(doc.styles['Normal'], True)
    p(doc, 'ProbeOne', style='Normal', after=200)
    last = p(doc, 'ProbeTwo', style='Normal', before=240)
    explicit = OxmlElement('w:pStyle'); explicit.set(qn('w:val'), 'Normal')
    last._element.get_or_add_pPr().insert(0, explicit)
    add('implicit-explicit-default', doc)
    doc = fresh()
    doc.styles['Normal'].element.attrib.pop(qn('w:default'), None)
    p(doc, 'ProbeOne', style='Normal', after=200, context=True)
    last = p(doc, 'ProbeTwo', style='Normal', before=240, context=True)
    explicit = OxmlElement('w:pStyle'); explicit.set(qn('w:val'), 'Normal')
    last._element.get_or_add_pPr().insert(0, explicit)
    add('missing-default-mixed-style-identity', doc)
    doc = fresh()
    # Empty paragraphs have no text run to carry imported font metadata.
    # Hold their line-box typography at the existing V2 Arial 11 pt default.
    # The independently preserved DejaVu 12 pt diagnostic documents a separate
    # preexisting empty-paragraph-font normalization, not a spacing promise.
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    p(doc, 'ProbeOne', after=200, context=True)
    p(doc, '', before=240, after=200, context=True)
    p(doc, 'ProbeTwo', before=240, context=True)
    add('empty-neighbor', doc)

    check(set(documents) == set(manifest['fixtures']), 'Fixture/expected manifest mismatch')
    sources = {}
    stock_signature = template_context_signature(Document().styles.element)
    for name, doc in documents.items():
        target = manifest['fixtures'][name]
        path = directory / (name + '.docx')
        check(not path.exists(), 'Refusing fixture overwrite')
        doc.save(path)
        snapshot = independent_snapshot(path)
        # Every original contextual template declaration survives every fixture.
        check(all(snapshot['template_context_declarations'].get(k) == v for k, v in stock_signature.items()),
              name + ': a stock contextual declaration was removed or changed')
        check([row['text'] for row in snapshot['ooxml']] == target['texts'], name + ': text mismatch')
        check([s.semantic_spacing(row['spacing']) for row in snapshot['ooxml']] ==
              [dict(zip(KEYS, values)) for values in target['source_spacing']], name + ': independent source spacing mismatch')
        item = dict(path=str(path), sha256=digest(path), fixture_kind='independent synthetic python-docx fixture',
                    expected=target, source_independent=snapshot, template_context_declarations_preserved=True)
        # Fresh explicit reference, never remove or alter declarations in the source.
        # Rebuilding a separate package from source bytes permits an explicit false
        # direct override while preserving the source's original style declarations.
        reference = Document(path)
        values = target.get('resolved_spacing', target.get('explicit_reference_spacing'))
        for para_xml, spacing in zip(reference.element.body.iter(qn('w:p')), values):
            para = Paragraph(para_xml, reference)
            ppr = para._element.get_or_add_pPr()
            flag = ppr.find('w:contextualSpacing', NS)
            if flag is None:
                contextual(para, False)
            else:
                # This is a separately named explicit reference, never an input
                # normalization presented as evidence that the source imported.
                flag.set(qn('w:val'), 'false')
            para.paragraph_format.space_before = Twips(spacing[0])
            para.paragraph_format.space_after = Twips(spacing[1])
            para.paragraph_format.line_spacing = spacing[2] / 240
        refpath = directory / (name + '-explicit-reference.docx')
        reference.save(refpath)
        item.update(reference_path=str(refpath), reference_sha256=digest(refpath), reference_independent=independent_snapshot(refpath))
        check([s.semantic_spacing(row['spacing']) for row in item['reference_independent']['ooxml']] ==
              [dict(zip(KEYS, value)) for value in values], name + ': reference XML mismatch')
        sources[name] = item
    return sources


def expected_rows(target, values=None, texts=None):
    return [dict(text=text, type='paragraph', spacing=dict(zip(KEYS, spacing)))
            for text, spacing in zip(texts or target['texts'], values or target['resolved_spacing'])]


def render_many(soffice, directory, paths):
    from pypdf import PdfReader
    out = directory / 'libreoffice-pdf'; out.mkdir(exist_ok=True)
    command = [str(soffice), f'-env:UserInstallation={(directory / "lo-profile").as_uri()}',
               '--headless', '--convert-to', 'pdf', '--outdir', str(out), *map(str, paths)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=max(120, len(paths) * 8),
                            env=dict(os.environ, XDG_CACHE_HOME=str(directory / 'lo-cache')))
    check(result.returncode == 0, result.stdout + result.stderr)
    reports = {}
    for source in paths:
        pdf = out / (source.stem + '.pdf')
        check(pdf.is_file(), f'PDF missing: {source}: {result.stdout} {result.stderr}')
        reader = PdfReader(pdf); positions = {}; texts = []
        for index, page in enumerate(reader.pages):
            def visit(text, cm, tm, font, size):
                for marker in MARKERS:
                    if marker in text:
                        check(marker not in positions, f'Duplicate PDF marker: {source} {marker}')
                        positions[marker] = dict(page=index + 1, x=tm[4] * cm[0] + tm[5] * cm[2] + cm[4],
                                                 y=tm[4] * cm[1] + tm[5] * cm[3] + cm[5], font_size=size)
            texts.append(page.extract_text(visitor_text=visit) or '')
        check('ProbeOne' in positions and len(reader.pages) == 1, f'Invalid controlled PDF: {source}')
        markers = [m for m in MARKERS if m in positions]
        reports[str(source)] = dict(path=str(pdf), sha256=digest(pdf), positions=positions, text=texts,
                                   baseline_distances_pt=[round(positions[a]['y'] - positions[b]['y'], 5) for a, b in zip(markers, markers[1:])])
    return reports


def equal_baselines(a, b):
    return len(a['baseline_distances_pt']) == len(b['baseline_distances_pt']) and all(
        abs(x - y) < 0.3 for x, y in zip(a['baseline_distances_pt'], b['baseline_distances_pt']))


def absolute_baseline_comparison(a, b):
    check(set(a['positions']) == set(b['positions']), 'PDF marker inventory differs')
    rows = {marker: dict(a_page=point['page'], b_page=b['positions'][marker]['page'],
                         a_y_pt=point['y'], b_y_pt=b['positions'][marker]['y'],
                         delta_y_pt=round(b['positions'][marker]['y'] - point['y'], 5))
            for marker, point in a['positions'].items()}
    return dict(per_marker=rows, pages_match=all(row['a_page'] == row['b_page'] for row in rows.values()),
                max_absolute_y_delta_pt=max(abs(row['delta_y_pt']) for row in rows.values()),
                comparison='Absolute PDF y coordinates; positive delta means the second PDF baseline is higher')


def run_corpus(cli, evidence):
    manifest = json.loads(CORPUS.read_text())
    rows = []
    for record in manifest['fixtures']:
        file = ROOT / record['local_default']
        check(file.is_file(), f'Pinned corpus missing: {file}')
        raw = file.read_bytes()
        check(len(raw) == record['bytes'] and digest(file) == record['sha256'], 'Pinned corpus changed')
        blob = hashlib.sha1(b'blob ' + str(len(raw)).encode() + b'\0' + raw).hexdigest()
        check(blob == record['git_blob_sha'], 'Pinned Git blob mismatch')
        row = dict(record, source_unchanged=True)
        if cli:
            audit = run_cli(cli, evidence['commands'], 'audit-docx', file)
            row['audit'] = audit
            check(all(risk in audit['risks'] for risk in record['required_remaining_risks']), 'Upstream blocker changed')
            row['import_error'] = run_cli(cli, evidence['commands'], 'import-docx', file, success=False)
        check(digest(file) == record['sha256'], 'Pinned source mutated')
        rows.append(row)
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path, default=ROOT / 'target/debug/rust-wps')
    parser.add_argument('--baseline-cli', type=Path, help='Optional frozen pre-change CLI for stock-template rejection proof')
    parser.add_argument('--root', type=Path, default=ROOT / 'tmp/word-contextual-spacing')
    parser.add_argument('--generate-only', action='store_true')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--corpus', action='store_true')
    parser.add_argument('--soffice')
    args = parser.parse_args()
    cli = None if args.generate_only else args.cli.resolve()
    check(cli is None or cli.is_file(), 'CLI missing')
    soffice = args.soffice or shutil.which('soffice')
    check(not args.render or soffice, '--render requires LibreOffice')
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    report = directory / 'report.json'
    manifest = json.loads(EXPECTED.read_text())
    evidence = dict(status='running', directory=str(directory), script_sha256=digest(Path(__file__)),
                    expected_sha256=digest(EXPECTED), cli=None if cli is None else str(cli),
                    cli_sha256=None if cli is None else digest(cli), python_version=platform.python_version(),
                    python_docx_version=version('python-docx'), commands=[], sources={}, limitations=manifest['limitations'],
                    helper_sha256=digest(Path(s.__file__)), corpus_sha256=digest(CORPUS))
    print(report, flush=True)
    try:
        evidence['sources'] = sources = make_fixtures(directory, manifest)
        if args.baseline_cli:
            baseline = args.baseline_cli.resolve()
            check(baseline.is_file(), 'Baseline CLI missing')
            evidence['baseline_cli'] = dict(path=str(baseline), sha256=digest(baseline), stock_results={})
            for name in ('stock-normal', 'stock-number', 'stock-bullet'):
                path = Path(sources[name]['path'])
                audit = run_cli(baseline, evidence['commands'], 'audit-docx', path)
                if name == 'stock-normal':
                    result = run_cli(baseline, evidence['commands'], 'import-docx', path)
                    check(result['document'] is not None and RISK not in audit['risks'], 'Baseline Normal control failed')
                    evidence['baseline_cli']['stock_results'][name] = dict(audit=audit, imported=True)
                else:
                    error = run_cli(baseline, evidence['commands'], 'import-docx', path, success=False)
                    check(RISK in error and RISK in audit['risks'], 'Baseline stock list rejection changed')
                    evidence['baseline_cli']['stock_results'][name] = dict(audit=audit, import_error=error)
        render_paths = []
        for name, item in sources.items():
            target, path = item['expected'], Path(item['path'])
            if args.render:
                render_paths.extend([path, Path(item['reference_path'])])
            if cli:
                item['audit'] = run_cli(cli, evidence['commands'], 'audit-docx', path)
                if target['action'] == 'reject':
                    check(RISK in item['audit']['risks'], name + ': missing spacing refusal')
                    item['import_error'] = run_cli(cli, evidence['commands'], 'import-docx', path, success=False)
                    check(RISK in item['import_error'], name + ': wrong import refusal')
                    continue
                check(RISK not in item['audit']['risks'], name + ': supported case rejected')
                should_warn = any(any(flag not in ('0', 'false', 'off') for flag in row['contextual_history'])
                                  for row in item['source_independent']['ooxml'])
                check((NORMALIZATION_WARNING in item['audit']['risks']) == should_warn,
                      name + ': audit normalization warning mismatch')
                imported_result = run_cli(cli, evidence['commands'], 'import-docx', path)
                check((NORMALIZATION_WARNING in imported_result['compatibility']['risks']) == should_warn,
                      name + ': import normalization warning mismatch')
                item.update(normalization_warning_expected=should_warn, audit_and_import_warning_verified=True)
                model = imported_result['document']
                targets = expected_rows(target)
                item['imported_rows'] = s.verify_model(model, targets, name + ' imported')
                modelpath = directory / (name + '-model.json'); modelpath.write_text(json.dumps(model, indent=2))
                output = directory / (name + '-export.docx')
                run_cli(cli, evidence['commands'], 'create-docx', modelpath, output)
                item.update(export_path=str(output), export_sha256=digest(output), export_independent=independent_snapshot(output))
                s.verify_independent(item['export_independent'], targets, name + ' export', require_direct=True)
                reopened = run_cli(cli, evidence['commands'], 'read-docx-model', output)
                check(reopened == model, name + ': trusted full-model mismatch')
                item['trusted_full_model_equal'] = True
                reimported = run_cli(cli, evidence['commands'], 'import-docx', output)['document']
                s.verify_model(reimported, targets, name + ' external reimport')
                item['external_full_model_equal'] = reimported == model
                for existing, sha in [(path, item['sha256']), (output, item['export_sha256'])]:
                    error = run_cli(cli, evidence['commands'], 'create-docx', modelpath, existing, success=False)
                    check('Output already exists' in error and digest(existing) == sha, 'No-clobber failure')
                item['source_and_export_overwrite_refused'] = True
                # Edit explicit spacing after import and prove this is an editable,
                # fixed-format snapshot rather than an unpreserved dynamic flag.
                edited = copy.deepcopy(model)
                def first_para(node):
                    if node['type'] in ('paragraph', 'heading'):
                        return node
                    return next((found for child in node.get('content', []) if (found := first_para(child)) is not None), None)
                first_para(edited).setdefault('attrs', {})['spaceAfter'] = 60
                edited_targets = copy.deepcopy(targets); edited_targets[0]['spacing']['spaceAfter'] = 60
                edited_path = directory / (name + '-edited.json'); edited_path.write_text(json.dumps(edited, indent=2))
                edited_output = directory / (name + '-edited-export.docx')
                run_cli(cli, evidence['commands'], 'create-docx', edited_path, edited_output)
                edited_reopened = run_cli(cli, evidence['commands'], 'read-docx-model', edited_output)
                check(edited_reopened == edited, name + ': edited trusted-model mismatch')
                edited_snapshot = independent_snapshot(edited_output)
                s.verify_independent(edited_snapshot, edited_targets, name + ' edited export', require_direct=True)
                edited_reimported = run_cli(cli, evidence['commands'], 'import-docx', edited_output)['document']
                s.verify_model(edited_reimported, edited_targets, name + ' edited external reimport')
                edited_reference = Document(item['reference_path'])
                reference_first = Paragraph(next(edited_reference.element.body.iter(qn('w:p'))), edited_reference)
                reference_first.paragraph_format.space_after = Twips(60)
                edited_reference_path = directory / (name + '-edited-explicit-reference.docx')
                edited_reference.save(edited_reference_path)
                s.verify_independent(independent_snapshot(edited_reference_path), edited_targets,
                                     name + ' independently edited reference', require_direct=True)
                item.update(edited_export_path=str(edited_output), edited_export_sha256=digest(edited_output),
                            edited_trusted_full_model_equal=True, edited_export_independent=edited_snapshot,
                            edited_reference_path=str(edited_reference_path), edited_reference_sha256=digest(edited_reference_path))
                if args.render:
                    render_paths.extend([output, edited_output, edited_reference_path])
            check(digest(path) == item['sha256'], name + ': source changed')
        if args.render:
            evidence['libreoffice_version'] = subprocess.run([soffice, '--version'], capture_output=True, text=True, check=True).stdout.strip()
            evidence['pypdf_version'] = version('pypdf')
            renders = render_many(soffice, directory, render_paths)
            for name, item in sources.items():
                item['source_pdf'] = renders[item['path']]
                item['reference_pdf'] = renders[item['reference_path']]
                target = item['expected']
                same = equal_baselines(item['source_pdf'], item['reference_pdf'])
                check(same == (target['action'] == 'roundtrip'), name + ': source/reference PDF assertion failed')
                item['source_matches_explicit_reference'] = same
                item['source_reference_absolute_baselines'] = absolute_baseline_comparison(item['source_pdf'], item['reference_pdf'])
                if target['action'] == 'roundtrip':
                    absolute = item['source_reference_absolute_baselines']
                    check(absolute['pages_match'] and absolute['max_absolute_y_delta_pt'] < 0.3,
                          name + ': source/reference absolute PDF baseline mismatch')
                if 'export_path' in item:
                    item['export_pdf'] = renders[item['export_path']]
                    check(equal_baselines(item['source_pdf'], item['export_pdf']), name + ': source/export PDF mismatch')
                    item['source_export_baselines_match'] = True
                    item['source_export_absolute_baselines'] = absolute_baseline_comparison(item['source_pdf'], item['export_pdf'])
                    absolute = item['source_export_absolute_baselines']
                    check(absolute['pages_match'] and absolute['max_absolute_y_delta_pt'] < 0.3,
                          name + ': source/export absolute PDF baseline mismatch')
                    item['edited_export_pdf'] = renders[item['edited_export_path']]
                    item['edited_reference_pdf'] = renders[item['edited_reference_path']]
                    check(equal_baselines(item['edited_export_pdf'], item['edited_reference_pdf']), name + ': edited export/reference PDF mismatch')
                    item['edited_export_reference_baselines_match'] = True
                    item['edited_reference_export_absolute_baselines'] = absolute_baseline_comparison(item['edited_reference_pdf'], item['edited_export_pdf'])
                    absolute = item['edited_reference_export_absolute_baselines']
                    check(absolute['pages_match'] and absolute['max_absolute_y_delta_pt'] < 0.3,
                          name + ': edited reference/export absolute PDF baseline mismatch')
        if args.corpus:
            evidence['upstream_corpus'] = run_corpus(cli, evidence)
        check(cli is None or digest(cli) == evidence['cli_sha256'], 'CLI changed during independent run')
        check(digest(EXPECTED) == evidence['expected_sha256'], 'Expected manifest changed during run')
        check(digest(Path(__file__)) == evidence['script_sha256'], 'Independent script changed during run')
        check(digest(Path(s.__file__)) == evidence['helper_sha256'], 'Spacing helper changed during run')
        check(digest(CORPUS) == evidence['corpus_sha256'], 'Corpus manifest changed during run')
        if args.baseline_cli:
            check(digest(args.baseline_cli.resolve()) == evidence['baseline_cli']['sha256'], 'Frozen baseline changed during run')
        evidence['summary'] = dict(supported=sum(t['action'] == 'roundtrip' for t in manifest['fixtures'].values()),
                                   rejected=sum(t['action'] == 'reject' for t in manifest['fixtures'].values()),
                                   source_pdfs=sum('source_pdf' in i for i in sources.values()),
                                   reference_pdfs=sum('reference_pdf' in i for i in sources.values()),
                                   export_pdfs=sum('export_pdf' in i for i in sources.values()),
                                   edited_export_pdfs=sum('edited_export_pdf' in i for i in sources.values()),
                                   edited_reference_pdfs=sum('edited_reference_pdf' in i for i in sources.values()))
        evidence['status'] = 'generated-and-independently-verified' if cli is None else 'passed'
    except Exception as error:
        evidence.update(status='failed', error=f'{type(error).__name__}: {error}')
        raise
    finally:
        evidence['sources_unchanged'] = all(digest(Path(i['path'])) == i['sha256'] for i in evidence['sources'].values())
        evidence['references_unchanged'] = all(digest(Path(i['reference_path'])) == i['reference_sha256'] for i in evidence['sources'].values())
        evidence['cli_unchanged'] = cli is None or digest(cli) == evidence['cli_sha256']
        evidence['script_unchanged'] = digest(Path(__file__)) == evidence['script_sha256']
        evidence['helper_unchanged'] = digest(Path(s.__file__)) == evidence['helper_sha256']
        evidence['expected_unchanged'] = digest(EXPECTED) == evidence['expected_sha256']
        evidence['corpus_unchanged'] = digest(CORPUS) == evidence['corpus_sha256']
        evidence['edited_references_unchanged'] = all(digest(Path(i['edited_reference_path'])) == i['edited_reference_sha256']
                                                      for i in evidence['sources'].values() if 'edited_reference_path' in i)
        evidence['exports_unchanged'] = all(digest(Path(i['export_path'])) == i['export_sha256']
                                           for i in evidence['sources'].values() if 'export_path' in i)
        evidence['edited_exports_unchanged'] = all(digest(Path(i['edited_export_path'])) == i['edited_export_sha256']
                                                  for i in evidence['sources'].values() if 'edited_export_path' in i)
        evidence['baseline_cli_unchanged'] = not args.baseline_cli or digest(args.baseline_cli.resolve()) == evidence['baseline_cli']['sha256']
        evidence['upstream_sources_unchanged'] = all(digest(ROOT / i['local_default']) == i['sha256']
                                                   for i in evidence.get('upstream_corpus', []))
        invariant_keys = ('sources_unchanged', 'references_unchanged', 'cli_unchanged', 'script_unchanged',
                          'helper_unchanged', 'expected_unchanged', 'corpus_unchanged', 'edited_references_unchanged',
                          'exports_unchanged', 'edited_exports_unchanged', 'baseline_cli_unchanged', 'upstream_sources_unchanged')
        changed = [key for key in invariant_keys if not evidence[key]]
        was_success = evidence['status'] in ('passed', 'generated-and-independently-verified')
        if was_success and changed:
            evidence.update(status='failed', error='Final hash invariants failed: ' + ', '.join(changed))
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
        print(report, flush=True)
        if was_success:
            check(not changed, evidence.get('error', 'Final hash invariants failed'))


if __name__ == '__main__':
    main()
