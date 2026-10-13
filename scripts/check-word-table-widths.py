"""Independent DOCX fixed-table-width regression; never modifies input files.

Directed synthetic documents are authored with python-docx, with OOXML edits
only where its public API does not express a declaration. Expected geometry is
fixed before import and checked by an independent XML reader and python-docx.
The path is source -> import -> unchanged and width-edited model -> DOCX ->
independent XML -> trusted read-docx-model exact equality. External reimport,
invalid-model refusal, source hashes and no-clobber are separate assertions.

  python scripts/check-word-table-widths.py --generate-only
  python scripts/check-word-table-widths.py --cli target/debug/rust-wps --render
  python scripts/check-word-table-widths.py --corpus tmp/corpus/SampleDoc.docx
  python scripts/check-word-table-widths.py --table-corpus tmp/table-corpus/table-indent.docx

The optional ORIGINAL Apache POI fixture is pinned in corpus.json and is kept
separate from generated fixtures. It contains no tables: it is compatibility
regression evidence, not real-world fixed-table-width evidence. --table-corpus verifies a separately pinned original
POI fixed-table document whose unsupported indentation requires refusal. --render uses
an installed LibreOffice PDF consumer for controlled border/wrapping probes;
it is not the app renderer, native WebView2, Microsoft Word or WPS validation.
"""
import argparse
import copy
import hashlib
import io
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Twips
from docx.table import Table
from lxml import etree as ET

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
KEY = "columnWidthsTwips"
# The default values are kept explicit so changing an implementation diagnostic
# cannot silently weaken a refusal assertion. CLI switches allow planned changes.
RISK = "Unsupported table geometry cannot be preserved in an editable copy"
WARNING = "Autofit table widths normalized; fixed column widths were not inferred"
WRAP_WORDS = [f"Wrap{i:02d}" for i in range(1, 13)]


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def remove(parent, tag):
    for child in list(parent.findall("w:" + tag, NS)):
        parent.remove(child)


def child(parent, tag, **attrs):
    element = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        element.set(qn("w:" + key), str(value))
    parent.append(element)
    return element


def declaration(element):
    return None if element is None else {ET.QName(k).localname: v for k, v in element.attrib.items()}


def fresh_document():
    doc = Document()
    # Isolate table geometry from unrelated template spacing and page defaults.
    for node in list(doc.styles.element.findall(".//w:contextualSpacing", NS)):
        node.getparent().remove(node)
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(25.4)
        section.left_margin = section.right_margin = Mm(25.4)
        remove(section._sectPr, "docGrid")
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Sans", Pt(12)
    normal.paragraph_format.space_before = Twips(0)
    normal.paragraph_format.space_after = Twips(0)
    normal.paragraph_format.line_spacing = 1
    return doc


def add_table(doc, widths, rows=2, header=False, layout="fixed", total="dxa", probe=False):
    table = doc.add_table(rows=rows, cols=len(widths))
    props = table._tbl.tblPr
    remove(props, "tblLayout")
    if layout is not None:
        child(props, "tblLayout", type=layout)
    remove(props, "tblW")
    if total is not None:
        child(props, "tblW", type=total, w=sum(widths) if total == "dxa" else 0)
    remove(props, "tblBorders")
    borders = child(props, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child(borders, edge, val="single", sz=4)
    for grid, width in zip(table._tbl.tblGrid, widths):
        grid.set(qn("w:w"), str(width))
    for row_index, row in enumerate(table.rows):
        if header and row_index == 0:
            child(row._tr.get_or_add_trPr(), "tblHeader")
        for column_index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Twips(width)
            cell.text = f"Row{row_index + 1}Column{column_index + 1}"
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Twips(0)
                p.paragraph_format.space_after = Twips(0)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    run.font.name, run.font.size = "DejaVu Sans", Pt(12)
    if probe:
        table.cell(0, 0).text = " ".join(WRAP_WORDS)
        table.cell(0, 1).text = "RightProbe"
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Twips(0)
                p.paragraph_format.space_after = Twips(0)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    run.font.name, run.font.size = "DejaVu Sans", Pt(12)
    return table


def xml_snapshot(path):
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    rows = []
    for table in root.findall(".//w:tbl", NS):
        props = table.find("w:tblPr", NS)
        layout = None if props is None else props.find("w:tblLayout", NS)
        total = None if props is None else props.find("w:tblW", NS)
        indent = None if props is None else props.find("w:tblInd", NS)
        alignment = None if props is None else props.find("w:jc", NS)
        cells, texts, headers = [], [], []
        for row in table.findall("w:tr", NS):
            cells.append([declaration(c.find("w:tcPr/w:tcW", NS)) for c in row.findall("w:tc", NS)])
            texts.append(["\n".join("".join(p.xpath(".//w:t/text()", namespaces=NS))
                                   for p in c.findall("w:p", NS)) for c in row.findall("w:tc", NS)])
            headers.append(row.find("w:trPr/w:tblHeader", NS) is not None)
        rows.append({"layout": declaration(layout), "total": declaration(total),
                     "indent": declaration(indent), "alignment": declaration(alignment),
                     "grid": [declaration(g) for g in table.findall("w:tblGrid/w:gridCol", NS)],
                     "cells": cells, "texts": texts, "header_rows": headers,
                     "nested": any(a.tag == qn("w:tbl") for a in table.iterancestors()),
                     "merges": len(table.findall(".//w:gridSpan", NS)) + len(table.findall(".//w:vMerge", NS)) + len(table.findall(".//w:hMerge", NS))})
    return rows


def independent_snapshot(path):
    xml = xml_snapshot(path)
    doc = Document(path)
    python = []
    for element in doc.element.body.iter(qn("w:tbl")):
        table = Table(element, doc)
        python.append({"autofit": table.autofit,
                       "column_widths_twips": [None if c.width is None else c.width.twips for c in table.columns],
                       "cell_widths_twips": [[None if c.width is None else c.width.twips for c in row.cells] for row in table.rows],
                       "texts": [[c.text for c in row.cells] for row in table.rows]})
    check(len(xml) == len(python), "Independent OOXML/python-docx table counts differ")
    return {"ooxml": xml, "python_docx": python}


def fixed_expected(widths, texts, header_rows=None):
    return {"widths": list(widths), "texts": texts,
            "header_rows": header_rows or [False] * len(texts)}


def verify_independent(snapshot, targets, label, fixed_required=True):
    check(len(snapshot["ooxml"]) == len(targets), f"{label}: table count differs")
    for index, (xml, py, target) in enumerate(zip(snapshot["ooxml"], snapshot["python_docx"], targets)):
        label_i = f"{label} table {index}"
        check(xml["texts"] == py["texts"] == target["texts"], f"{label_i}: cell text changed")
        check(xml["header_rows"] == target["header_rows"], f"{label_i}: header rows changed")
        if target["widths"] is None:
            check(xml["layout"] != {"type": "fixed"}, f"{label_i}: legacy table made a fixed-width promise")
            continue
        widths = target["widths"]
        check(len(widths) > 0 and all(type(w) is int and 1 <= w <= 31680 for w in widths)
              and sum(widths) <= 31680, f"{label_i}: invalid independent target")
        check(xml["layout"] == {"type": "fixed"} and py["autofit"] is False,
              f"{label_i}: missing explicit fixed layout")
        check(xml["grid"] == [{"w": str(w)} for w in widths], f"{label_i}: grid mismatch: {xml['grid']}")
        allowed_totals = [{"type": "dxa", "w": str(sum(widths))}]
        if not fixed_required:
            allowed_totals.extend([{"type": "auto", "w": "0"}, None])
        check(xml["total"] in allowed_totals, f"{label_i}: wrong total width: {xml['total']}")
        check(all(row == [{"type": "dxa", "w": str(w)} for w in widths] for row in xml["cells"]),
              f"{label_i}: tcW mismatch: {xml['cells']}")
        check(py["column_widths_twips"] == widths and all(row == widths for row in py["cell_widths_twips"]),
              f"{label_i}: python-docx physical width mismatch")
        check(not xml["merges"] and not xml["nested"], f"{label_i}: not a simple unmerged table")


def make_fixtures(directory):
    sources = {}

    def save(name, doc, widths=None, reject=None, normalize=False, render=False, transform=None):
        path = directory / (name + ".docx")
        check(not path.exists(), f"Fixture destination exists: {path}")
        if transform is None:
            doc.save(path)
        else:
            buffer = io.BytesIO()
            doc.save(buffer)
            with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as output:
                for entry in source.infolist():
                    data = source.read(entry)
                    output.writestr(entry, transform(data) if entry.filename == "word/document.xml" else data)
        item = {"path": str(path), "sha256": digest(path), "bytes": path.stat().st_size,
                "fixture_kind": "directed synthetic python-docx fixture; not upstream/user corpus",
                "expected_action": "reject" if reject else "normalize" if normalize else "fixed-roundtrip",
                "source_xml": xml_snapshot(path)}
        if reject:
            item["unsupported_form"] = reject
        else:
            item["source_independent"] = independent_snapshot(path)
            item["expected"] = [fixed_expected(w, x["texts"], x["header_rows"]) if w is not None else
                                {"widths": None, "texts": x["texts"], "header_rows": x["header_rows"]}
                                for w, x in zip(widths, item["source_xml"])]
            check(len(widths) == len(item["source_xml"]), name + ": fixture expectation count differs")
            verify_independent(item["source_independent"], item["expected"], name + " source", fixed_required=False)
        if render:
            item["render_probe"] = True
        sources[name] = item

    for name, widths, rows, header in (("fixed-unequal-multirow", [1800, 4200], 3, False),
                                       ("fixed-three-columns-header", [1000, 2000, 3000], 2, True),
                                       ("fixed-single-column", [5000], 2, False),
                                       ("fixed-minimum-width", [1, 2999], 1, False),
                                       ("fixed-maximum-total", [15840, 15840], 1, False),
                                       ("fixed-maximum-single", [31680], 1, False)):
        doc = fresh_document()
        add_table(doc, widths, rows=rows, header=header)
        save(name, doc, [widths])
    doc = fresh_document()
    add_table(doc, [1200, 4800])
    doc.add_paragraph("BetweenTables")
    add_table(doc, [1000, 2000, 3000])
    save("fixed-multiple-tables", doc, [[1200, 4800], [1000, 2000, 3000]])
    for name, widths in (("render-narrow-left", [1800, 7200]), ("render-wide-left", [4200, 4800])):
        doc = fresh_document()
        add_table(doc, widths, rows=1, probe=True)
        save(name, doc, [widths], render=True)
    for name, layout in (("default-layout-raw-widths", None), ("explicit-autofit-raw-widths", "autofit")):
        doc = fresh_document()
        add_table(doc, [1800, 4200], layout=layout)
        save(name, doc, [None], normalize=True)
    doc = fresh_document()
    add_table(doc, [1800, 4200], total="auto")
    save("fixed-python-docx-auto-total", doc, [[1800, 4200]])
    doc = fresh_document()
    add_table(doc, [1800, 4200], total=None)
    save("fixed-absent-total", doc, [[1800, 4200]])

    def unsupported(name, mutate, why):
        doc = fresh_document()
        table = add_table(doc, [1800, 4200])
        mutate(doc, table)
        save("reject-" + name, doc, reject=why)

    def total(table, width, kind="dxa"):
        node = table._tbl.tblPr.find("w:tblW", NS)
        node.set(qn("w:type"), kind)
        node.set(qn("w:w"), str(width))

    unsupported("percentage-total", lambda d, t: total(t, 5000, "pct"), "percentage table total")
    unsupported("mismatched-total", lambda d, t: total(t, 5999), "dxa total disagrees with grid")
    unsupported("zero-total", lambda d, t: total(t, 0), "zero dxa table total")
    unsupported("nonzero-auto-total", lambda d, t: total(t, 6000, "auto"), "nonzero automatic table total")
    unsupported("unknown-layout", lambda d, t: t._tbl.tblPr.find("w:tblLayout", NS).set(qn("w:type"), "mystery"), "unknown table layout")
    unsupported("missing-grid", lambda d, t: t._tbl.remove(t._tbl.tblGrid), "fixed layout missing grid")
    for suffix, value in (("zero", "0"), ("negative", "-1"), ("fractional", "1800.5"), ("unit", "90pt"), ("over-limit", "31681")):
        unsupported("grid-" + suffix, lambda d, t, v=value: t._tbl.tblGrid[0].set(qn("w:w"), v), "invalid grid width " + value)
    unsupported("grid-missing-width", lambda d, t: t._tbl.tblGrid[0].attrib.pop(qn("w:w")), "grid entry without width")
    unsupported("short-grid", lambda d, t: t._tbl.tblGrid.remove(t._tbl.tblGrid[-1]), "grid column count does not match rows")
    unsupported("mismatched-cell", lambda d, t: t.cell(1, 0)._tc.tcPr.find("w:tcW", NS).set(qn("w:w"), "1799"), "second row tcW disagrees with grid")
    unsupported("missing-cell-width", lambda d, t: remove(t.cell(1, 0)._tc.tcPr, "tcW"), "fixed cell without tcW")
    unsupported("percentage-cell", lambda d, t: t.cell(0, 0)._tc.tcPr.find("w:tcW", NS).set(qn("w:type"), "pct"), "percentage cell width")
    unsupported("automatic-cell", lambda d, t: t.cell(0, 0)._tc.tcPr.find("w:tcW", NS).set(qn("w:type"), "auto"), "automatic cell width under fixed layout")
    unsupported("merged", lambda d, t: t.cell(0, 0).merge(t.cell(0, 1)), "merged cells")
    unsupported("ragged", lambda d, t: t.rows[1]._tr.remove(t.rows[1]._tr.tc_lst[-1]), "ragged rows")
    unsupported("indented", lambda d, t: child(t._tbl.tblPr, "tblInd", w=240, type="dxa"), "table indentation")
    unsupported("floating", lambda d, t: child(t._tbl.tblPr, "tblpPr", vertAnchor="text", horzAnchor="margin", tblpX=0, tblpY=0), "floating table")
    unsupported("nested", lambda d, t: t.cell(0, 0).add_table(rows=1, cols=1), "nested source table")
    doc = fresh_document()
    add_table(doc, [16000, 16000])
    save("reject-total-over-limit", doc, reject="sum exceeds supported 31680 twips")
    doc = fresh_document()
    add_table(doc, [1800, 4200])
    old, new = b'<w:tblLayout w:type="fixed"', ('<w:tblLayout xmlns:alt="' + W + '" alt:type="fixed"').encode()
    def alternate(data):
        check(data.count(old) == 1, "Alternate-prefix fixture target is ambiguous")
        return data.replace(old, new)
    save("reject-alternate-prefix-layout", doc, reject="namespace-equivalent alternate-prefix fixed declaration", transform=alternate)
    manifest = directory / "expected-table-widths-manifest.json"
    manifest.write_text(json.dumps({name: {k: v for k, v in item.items() if k not in ("source_xml", "source_independent")}
                                    for name, item in sources.items()}, indent=2), encoding="utf-8")
    return sources, manifest


def run_cli(cli, log, *args, success=True):
    result = subprocess.run([str(cli), *map(str, args)], text=True, capture_output=True, encoding="utf-8", timeout=90)
    log.append({"arguments": list(map(str, args)), "returncode": result.returncode,
                "stdout": result.stdout, "stderr": result.stderr})
    check((result.returncode == 0) is success, f"CLI {args}: {result.stdout}{result.stderr}")
    return json.loads(result.stdout) if success else result.stderr


def model_tables(node):
    if node["type"] == "table":
        yield node
    for c in node.get("content", []):
        yield from model_tables(c)


def model_text(node):
    if node["type"] == "hardBreak":
        return "\n"
    return node.get("text", "") + "".join(model_text(c) for c in node.get("content", []))


def verify_model(model, targets, label):
    tables = list(model_tables(model))
    check(len(tables) == len(targets), f"{label}: model table count differs")
    result = []
    for table, target in zip(tables, targets):
        actual = table.get("attrs", {}).get(KEY)
        check(actual == target["widths"], f"{label}: expected {target['widths']}, got {actual}")
        if actual is None:
            check(KEY not in table.get("attrs", {}), f"{label}: widths must be absent, not null")
        text = [["\n".join(model_text(p) for p in cell.get("content", [])) for cell in row["content"]]
                for row in table["content"]]
        check(text == target["texts"], f"{label}: model cell text changed: {text}")
        headers = [all(cell["type"] == "tableHeader" for cell in row["content"]) for row in table["content"]]
        check(headers == target["header_rows"], f"{label}: model headers changed")
        result.append({"widths": actual, "texts": text, "header_rows": headers})
    return result


def changed_widths(widths, columns):
    if widths is None:
        return [2000] * columns
    values = list(widths)
    if len(values) == 1:
        return [max(1, values[0] - 120)] if values[0] > 1 else [2]
    # Preserve total, changing the first two columns while honoring min/max.
    change = min(400, values[1] - 1)
    if change:
        values[0] += change
        values[1] -= change
    else:
        change = min(400, values[0] - 1)
        values[0] -= change
        values[1] += change
    check(values != widths, "Width edit probe failed to change geometry")
    return values


def write_model(path, model):
    check(not path.exists(), f"Model destination exists: {path}")
    path.write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding="utf-8")


def export_verify(cli, evidence, directory, name, model, targets):
    model_path = directory / (name + ".json")
    write_model(model_path, model)
    output = directory / (name + ".docx")
    run_cli(cli, evidence["commands"], "create-docx", model_path, output)
    result = {"model_path": str(model_path), "path": str(output), "sha256": digest(output),
              "independent": independent_snapshot(output)}
    verify_independent(result["independent"], targets, name + " export")
    reopened = run_cli(cli, evidence["commands"], "read-docx-model", output)
    verify_model(reopened, targets, name + " trusted authored reopen")
    check(reopened == model, name + ": trusted authored reopen changed full model")
    result["full_model_equal_after_authored_reopen"] = True
    external = run_cli(cli, evidence["commands"], "import-docx", output)["document"]
    result["external_reimported_tables"] = verify_model(external, targets, name + " external reimport")
    result["full_model_equal_after_external_reimport"] = external == model
    failure = run_cli(cli, evidence["commands"], "create-docx", model_path, output, success=False)
    check("Output already exists" in failure and digest(output) == result["sha256"], name + ": output no-clobber failed")
    result["existing_output_overwrite_refused"] = True
    return result


def bad_models(cli, evidence, directory, model):
    cases = {"null": None, "empty": [], "short": [3000], "long": [1000, 2000, 3000],
             "string": "1800,4200", "string-entry": ["1800", 4200], "boolean-entry": [True, 4200],
             "fractional-entry": [1800.5, 4200], "zero-entry": [0, 4200], "negative-entry": [-1, 4200],
             "over-limit-entry": [31681, 1], "total-over-limit": [16000, 16000]}
    results = []
    for name, value in cases.items():
        changed = copy.deepcopy(model)
        next(model_tables(changed)).setdefault("attrs", {})[KEY] = value
        path, out = directory / ("bad-model-" + name + ".json"), directory / ("bad-model-" + name + ".docx")
        write_model(path, changed)
        error = run_cli(cli, evidence["commands"], "create-docx", path, out, success=False)
        check(not out.exists(), name + ": rejected model created an output")
        check(KEY in error or "width" in error.lower(), name + ": unrelated model refusal: " + error)
        results.append({"name": name, "value": value, "error": error, "output_absent": True})
    for location in ("doc", "paragraph", "cell", "row"):
        changed = copy.deepcopy(model)
        table = next(model_tables(changed))
        node = changed if location == "doc" else table["content"][0] if location == "row" else table["content"][0]["content"][0]
        if location == "paragraph":
            node = node["content"][0]
        node.setdefault("attrs", {})[KEY] = [1800, 4200]
        path, out = directory / ("bad-model-location-" + location + ".json"), directory / ("bad-model-location-" + location + ".docx")
        write_model(path, changed)
        error = run_cli(cli, evidence["commands"], "create-docx", path, out, success=False)
        check(not out.exists(), location + ": illegal width location created output")
        results.append({"name": "location-" + location, "error": error, "output_absent": True})
    return results


def render_pdf(soffice, directory, path, widths):
    from pypdf import PdfReader
    out = directory / "libreoffice-pdf"
    out.mkdir(exist_ok=True)
    profile = directory / "lo-profile"
    outcome = subprocess.run([str(soffice), f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "pdf",
                              "--outdir", str(out), str(path)], text=True, capture_output=True, timeout=120,
                             env=dict(os.environ, XDG_CACHE_HOME=str(directory / "lo-cache")))
    pdf = out / (path.stem + ".pdf")
    check(outcome.returncode == 0 and pdf.is_file(), outcome.stdout + outcome.stderr)
    reader = PdfReader(pdf)
    check(len(reader.pages) == 1, "Controlled table probe unexpectedly spans pages")
    segments, texts, words = [], [], {}
    pending, point = [], None

    def transform(x, y, cm):
        return (float(x) * cm[0] + float(y) * cm[2] + cm[4], float(x) * cm[1] + float(y) * cm[3] + cm[5])

    def operand(op, args, cm, tm):
        nonlocal point, pending
        if op == b"m":
            point = transform(*args, cm)
        elif op == b"l":
            end = transform(*args, cm)
            if point is not None:
                pending.append((point, end))
            point = end
        elif op == b"re":
            x, y, w, h = map(float, args)
            points = [transform(a, b, cm) for a, b in ((x, y), (x + w, y), (x + w, y + h), (x, y + h), (x, y))]
            pending.extend(zip(points, points[1:]))
        elif op in (b"S", b"s", b"B", b"B*", b"b", b"b*", b"f", b"f*", b"F"):
            segments.extend(pending)
            pending, point = [], None
        elif op == b"n":
            pending, point = [], None

    def text_visitor(text, cm, tm, font, size):
        if not text.strip():
            return
        x, y = transform(tm[4], tm[5], cm)
        texts.append({"text": text, "x": x, "y": y, "font_size": size})
        for word in WRAP_WORDS:
            if word in text:
                check(word not in words, "Duplicate PDF wrapping marker: " + word)
                words[word] = {"x": x, "y": y}

    extracted = reader.pages[0].extract_text(visitor_operand_before=operand, visitor_text=text_visitor)
    check(set(words) == set(WRAP_WORDS) and "RightProbe" in extracted, "Missing controlled PDF text markers")
    vertical = sorted({round((a[0] + b[0]) / 2, 2) for a, b in segments
                       if abs(a[0] - b[0]) < 0.1 and abs(a[1] - b[1]) > 10})
    # Thin filled rectangles can expose two near-identical edges for a border.
    clusters = []
    for x in vertical:
        if not clusters or x - clusters[-1][-1] > 1:
            clusters.append([x])
        else:
            clusters[-1].append(x)
    borders = [sum(c) / len(c) for c in clusters]
    check(len(borders) == len(widths) + 1, f"Expected table vertical borders, found {borders}")
    measured = [b - a for a, b in zip(borders, borders[1:])]
    expected = [w / 20 for w in widths]
    check(all(abs(a - b) < 0.8 for a, b in zip(measured, expected)), f"PDF column widths differ: {measured} != {expected}")
    lines = sorted({round(p["y"], 1) for p in words.values()}, reverse=True)
    check(len(lines) >= 2, "Wrapping probe did not wrap")
    return {"consumer": "LibreOffice DOCX-to-PDF export, not app renderer/print", "path": str(pdf), "sha256": digest(pdf),
            "pages": 1, "vertical_border_x_pt": borders, "column_widths_pt": measured,
            "expected_column_widths_pt": expected, "wrap_line_y_pt": lines, "wrap_line_count": len(lines),
            "word_positions": words, "text": extracted, "text_runs": texts, "stdout": outcome.stdout, "stderr": outcome.stderr}


def check_corpus(path, cli, evidence, directory):
    manifest = json.loads(Path(__file__).with_name("corpus.json").read_text(encoding="utf-8"))
    entries = [entry for entry in manifest["fixtures"] if entry["name"] == path.name]
    check(len(entries) == 1, "Corpus must be an existing pinned fixture in corpus.json")
    target = entries[0]
    check(path.stat().st_size == target["bytes"] and digest(path) == target["sha256"], "Pinned ORIGINAL corpus size/hash mismatch")
    independent = independent_snapshot(path)
    check(not independent["ooxml"], "This compatibility corpus contract expects no tables")
    result = {"fixture_kind": "original unmodified pinned Apache POI upstream regression fixture, not synthetic or user production data",
              "path": str(path), "sha256": target["sha256"], "table_count": 0,
              "scope": "No-table compatibility regression only; does not supply real-world fixed-width evidence",
              "repository": manifest["repository"], "revision": manifest["revision"],
              "source_url": f"https://github.com/{manifest['repository']}/blob/{manifest['revision']}/{manifest['directory']}/{path.name}"}
    if cli:
        imported = run_cli(cli, evidence["commands"], "import-docx", path)["document"]
        result["roundtrip"] = export_verify(cli, evidence, directory, "original-poi-export", imported, [])
        source_text = "\n".join(p.text for p in Document(path).paragraphs)
        exported_text = "\n".join(p.text for p in Document(result["roundtrip"]["path"]).paragraphs)
        check(source_text == exported_text, "Original POI paragraphs changed")
        model_path = result["roundtrip"]["model_path"]
        error = run_cli(cli, evidence["commands"], "create-docx", model_path, path, success=False)
        check("Output already exists" in error and digest(path) == target["sha256"], "Original POI no-clobber failed")
        result.update(full_body_paragraph_text_equal=True, overwrite_refused=True)
    check(digest(path) == target["sha256"], "Original POI source changed")
    result["source_unchanged"] = True
    return result



def check_table_corpus(path, cli, evidence, directory, safe_model_path):
    manifest_path = Path(__file__).with_name("table-width-corpus.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    entries = [entry for entry in manifest["fixtures"] if entry["name"] == path.name]
    check(len(entries) == 1, "Table corpus must be a pinned fixture in table-width-corpus.json")
    target = entries[0]
    data = path.read_bytes()
    actual_hash = hashlib.sha256(data).hexdigest()
    git_hash = hashlib.sha1(f"blob {len(data)}\0".encode() + data).hexdigest()
    check(len(data) == target["bytes"] and actual_hash == target["sha256"] and git_hash == target["git_blob_sha"],
          "Pinned ORIGINAL table corpus byte size, SHA-256 or Git blob SHA mismatch")
    snapshot = independent_snapshot(path)
    check(len(snapshot["ooxml"]) == len(snapshot["python_docx"]) == target["table_count"] == len(target["tables"]),
          "Pinned table corpus table count differs")
    widths, row_count = target["grid_widths_twips"], target["rows_per_table"]
    check(len(widths) == target["columns_per_table"], "Pinned manifest grid cardinality differs")
    for index, (xml, py, table) in enumerate(zip(snapshot["ooxml"], snapshot["python_docx"], target["tables"])):
        prefix = f"Original table corpus table {index}"
        expected = {"layout": target["layout"], "total": table["total"], "indent": table["indent"],
                    "alignment": target["alignment"], "grid": [{"w": str(w)} for w in widths],
                    "cells": [[{"w": str(w), "type": target["cell_width_type"]} for w in widths]] * row_count,
                    "texts": target["cell_texts"], "header_rows": [False] * row_count, "nested": False, "merges": 0}
        check(xml == expected, prefix + ": full independently read OOXML geometry differs")
        check(py == {"autofit": False, "column_widths_twips": widths, "cell_widths_twips": [widths] * row_count,
                     "texts": target["cell_texts"]}, prefix + ": python-docx geometry/text differs")
    result = {"fixture_kind": "original unmodified pinned Apache POI fixed-table regression fixture",
              "path": str(path), "sha256": actual_hash, "bytes": len(data), "git_blob_sha": git_hash,
              "repository": manifest["repository"], "revision": manifest["revision"], "source_url": target["source_url"],
              "license": manifest["license"], "scope": target["scope"], "expected_action": target["expected_action"],
              "expected_risk": target["expected_risk"], "manifest_path": str(manifest_path.resolve()),
              "manifest_sha256": digest(manifest_path), "independent": snapshot,
              "verified_fixed_tables": target["table_count"], "unsupported_indent_tables": target["table_count"] - 1}
    # Save before CLI assertions so a failing refusal still reports its provenance.
    evidence["original_fixed_table_corpus"] = result
    if cli:
        result["audit"] = run_cli(cli, evidence["commands"], "audit-docx", path)
        check(target["expected_risk"] in result["audit"]["risks"], "Original fixed-table corpus lacks precise geometry risk")
        result["import_error"] = run_cli(cli, evidence["commands"], "import-docx", path, success=False)
        check(target["expected_risk"] in result["import_error"], "Original fixed-table corpus refused for an unrelated reason")
        result["precise_editable_import_refusal_verified"] = True
        error = run_cli(cli, evidence["commands"], "create-docx", safe_model_path, path, success=False)
        check("Output already exists" in error and digest(path) == actual_hash, "Original fixed-table corpus no-clobber failed")
        result["source_overwrite_refused"] = True
    check(digest(path) == actual_hash, "Original fixed-table corpus source changed")
    result["source_unchanged"] = True
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli", type=Path, default=Path("target/debug/rust-wps"))
    parser.add_argument("--root", type=Path, default=Path("tmp/word-table-widths"))
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--soffice")
    parser.add_argument("--corpus", type=Path)
    parser.add_argument("--table-corpus", type=Path, help="Original pinned upstream fixed-table refusal file; see table-width-corpus.json")
    parser.add_argument("--risk", default=RISK)
    parser.add_argument("--normalization-warning", default=WARNING)
    args = parser.parse_args()
    cli = None if args.generate_only else args.cli.resolve()
    check(cli is None or cli.is_file(), f"CLI missing: {cli}")
    soffice = args.soffice or shutil.which("soffice")
    check(not args.render or soffice, "LibreOffice is required for --render")
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.root)).resolve()
    report = directory / "report.json"
    evidence = {"status": "running", "script_sha256": digest(Path(__file__)), "python_version": platform.python_version(),
                "python_docx_version": version("python-docx"), "lxml_version": version("lxml"),
                "cli": None if cli is None else str(cli), "cli_sha256": None if cli is None else digest(cli),
                "commands": [], "sources": {}, "risk_diagnostic": args.risk, "normalization_warning": args.normalization_warning,
                "authored_reopen_command": "read-docx-model", "external_reimport_command": "import-docx",
                "limitations": "Generated fixtures are directed synthetic python-docx documents. The separately pinned original POI document "
                "has no tables and supplies compatibility regression only. The optional separately pinned original fixed-table POI "
                "corpus supplies unsupported-indentation refusal evidence, never a supported roundtrip. LibreOffice PDF border and wrapping probes establish controlled "
                "external DOCX-consumer geometry, not app rendering/printing, native WebView2, full layout/pagination, or Word/WPS validation."}
    try:
        sources, manifest = make_fixtures(directory)
        evidence.update(sources=sources, expected_manifest=str(manifest), expected_manifest_sha256=digest(manifest))
        if args.render:
            evidence["pypdf_version"] = version("pypdf")
            evidence["libreoffice_version"] = subprocess.run([soffice, "--version"], capture_output=True, text=True, check=True, timeout=30).stdout.strip()
        safe_model_path, safe_model = None, None
        for name, item in sources.items():
            path = Path(item["path"])
            if cli:
                item["audit"] = run_cli(cli, evidence["commands"], "audit-docx", path)
                if item["expected_action"] == "reject":
                    # Merged cells already have a distinct longstanding guard.
                    diagnostic = "Merged table cells" if name == "reject-merged" else args.risk
                    check(diagnostic in item["audit"]["risks"], f"{name}: missing diagnostic {diagnostic!r}: {item['audit']}")
                    item["import_error"] = run_cli(cli, evidence["commands"], "import-docx", path, success=False)
                    check(diagnostic in item["import_error"], f"{name}: wrong import refusal")
                else:
                    check(args.risk not in item["audit"]["risks"], f"{name}: supported source rejected")
                    if item["expected_action"] == "normalize":
                        check(args.normalization_warning in item["audit"]["risks"], f"{name}: missing normalization warning")
                    imported = run_cli(cli, evidence["commands"], "import-docx", path)["document"]
                    item["imported_tables"] = verify_model(imported, item["expected"], name + " imported")
                    item["unchanged_export"] = export_verify(cli, evidence, directory, name + "-unchanged", imported, item["expected"])
                    if safe_model is None:
                        safe_model, safe_model_path = imported, item["unchanged_export"]["model_path"]
                    edited, edited_targets = copy.deepcopy(imported), copy.deepcopy(item["expected"])
                    for table, target in zip(model_tables(edited), edited_targets):
                        widths = changed_widths(target["widths"], len(table["content"][0]["content"]))
                        table.setdefault("attrs", {})[KEY], target["widths"] = widths, widths
                    item["edited_expected"] = edited_targets
                    verify_model(edited, edited_targets, name + " edited")
                    item["edited_export"] = export_verify(cli, evidence, directory, name + "-edited", edited, edited_targets)
                    if args.render and item.get("render_probe"):
                        item["unchanged_render"] = render_pdf(soffice, directory, Path(item["unchanged_export"]["path"]), item["expected"][0]["widths"])
                        item["edited_render"] = render_pdf(soffice, directory, Path(item["edited_export"]["path"]), edited_targets[0]["widths"])
            if args.render and item.get("render_probe"):
                item["source_render"] = render_pdf(soffice, directory, path, item["expected"][0]["widths"])
                if "unchanged_render" in item:
                    a, b = item["source_render"], item["unchanged_render"]
                    check(a["wrap_line_count"] == b["wrap_line_count"], f"{name}: source/unchanged wrapping differs")
                    check(all(abs(x - y) < 0.5 for x, y in zip(a["column_widths_pt"], b["column_widths_pt"])), f"{name}: source/unchanged border distances differ")
                    item["source_unchanged_export_geometry_and_wrapping_match"] = True
            check(digest(path) == item["sha256"], f"Source changed: {name}")
            item["source_unchanged"] = True
        if cli:
            evidence["bad_model_refusals"] = bad_models(cli, evidence, directory, safe_model)
            for item in sources.values():
                path = Path(item["path"])
                error = run_cli(cli, evidence["commands"], "create-docx", safe_model_path, path, success=False)
                check("Output already exists" in error and digest(path) == item["sha256"], "Source no-clobber failed")
                item["source_overwrite_refused"] = True
        if args.render:
            for stage in (["source_render", "unchanged_render"] if cli else ["source_render"]):
                narrow, wide = sources["render-narrow-left"][stage], sources["render-wide-left"][stage]
                check(narrow["wrap_line_count"] > wide["wrap_line_count"], stage + ": wider left column did not reduce wrapping")
            evidence["pdf_assertions"] = {"controlled_border_widths_verified": True, "wider_left_column_reduces_wrap_lines": True,
                                          "source_unchanged_export_match": bool(cli), "edited_output_border_widths_verified": bool(cli)}
        if args.corpus:
            evidence["original_upstream_corpus"] = check_corpus(args.corpus.resolve(), cli, evidence, directory)
        if args.table_corpus:
            check_table_corpus(args.table_corpus.resolve(), cli, evidence, directory, safe_model_path)
        evidence["status"] = "fixtures-verified" if args.generate_only else "passed"
        evidence["summary"] = {"fixed_fixtures": sum(i["expected_action"] == "fixed-roundtrip" for i in sources.values()),
                               "normalization_fixtures": sum(i["expected_action"] == "normalize" for i in sources.values()),
                               "rejection_fixtures": sum(i["expected_action"] == "reject" for i in sources.values()),
                               "bad_models_refused": len(evidence.get("bad_model_refusals", [])),
                               "rendered_sources": sum("source_render" in i for i in sources.values()),
                               "rendered_exports": sum("unchanged_render" in i for i in sources.values()) + sum("edited_render" in i for i in sources.values())}
    except Exception as error:
        evidence["status"] = "failed"
        evidence["error"] = f"{type(error).__name__}: {error}"
        raise
    finally:
        evidence["script_sha256_on_exit"] = digest(Path(__file__))
        evidence["script_unchanged_during_run"] = evidence["script_sha256_on_exit"] == evidence["script_sha256"]
        evidence["cli_sha256_on_exit"] = None if cli is None else digest(cli)
        evidence["cli_unchanged_during_run"] = evidence["cli_sha256_on_exit"] == evidence["cli_sha256"]
        evidence["source_hashes_unchanged_on_exit"] = all(Path(i["path"]).is_file() and digest(Path(i["path"])) == i["sha256"] for i in evidence["sources"].values())
        corpus = evidence.get("original_fixed_table_corpus")
        if corpus:
            corpus["source_hash_unchanged_on_exit"] = Path(corpus["path"]).is_file() and digest(Path(corpus["path"])) == corpus["sha256"]
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
        print(report)


if __name__ == "__main__":
    main()
