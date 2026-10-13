"""Independent, bounded paragraph-indentation regression.

Sources are independently authored python-docx packages. Fixed expected values
live separately in tests/fixtures/word-paragraph-indentation-expected.json.
An OOXML style reader plus python-docx inspect source/export declarations;
trusted-model reopen checks the entire model. Untouched public upstream fixtures
are separately hash-pinned in paragraph-indentation-corpus.json.
LibreOffice PDF coordinates prove that consumer only, not app/native printing,
Word/WPS behavior, all typography, or complete pagination fidelity.
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import os
import platform
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, Twips
from docx.text.paragraph import Paragraph
from lxml import etree as ET

ROOT = Path(__file__).resolve().parent.parent
EXPECTED = ROOT / 'tests/fixtures/word-paragraph-indentation-expected.json'
CORPUS = Path(__file__).with_name('paragraph-indentation-corpus.json')
KEYS = ('indentLeftTwips', 'indentRightTwips', 'firstLineIndentTwips')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
RISK = 'Unsupported paragraph indentation cannot be preserved in an editable copy'

def load_script(name):
    spec = importlib.util.spec_from_file_location(name.replace('-', '_'), Path(__file__).with_name(name + '.py'))
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module

s = load_script('check-word-paragraph-spacing')
w = load_script('check-word-table-widths')
check, digest, run_cli, qn = w.check, w.digest, w.run_cli, w.qn

def ind(ppr, **attrs):
    w.remove(ppr, 'ind')
    return w.child(ppr, 'ind', **attrs)

def set_indent(para, left=None, right=None, first=None):
    f = para.paragraph_format
    for name, value in [('left_indent', left), ('right_indent', right), ('first_line_indent', first)]:
        if value is not None: setattr(f, name, Twips(value))

def font(para):
    for run in para.runs:
        run.font.name, run.font.size = 'DejaVu Sans', Pt(12)

def fresh():
    doc = s.fresh_document()
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Twips(1440)
    return doc

def attrs3(left=None, right=None, first=None):
    return {key: value for key, value in zip(KEYS, (left, right, first)) if value is not None}

def direct_indent(ppr):
    node = None if ppr is None else ppr.find('w:ind', NS)
    if node is None: return {}
    result = {}
    for source, target in [('left', KEYS[0]), ('right', KEYS[1])]:
        if node.get(qn('w:' + source)) is not None: result[target] = int(node.get(qn('w:' + source)))
    first, hanging = node.get(qn('w:firstLine')), node.get(qn('w:hanging'))
    check(first is None or hanging is None, 'Independent supported reader refuses same-level firstLine+hanging')
    if first is not None: result[KEYS[2]] = int(first)
    elif hanging is not None: result[KEYS[2]] = -int(hanging)
    return result

def independent(path):
    with zipfile.ZipFile(path) as z:
        main = ET.fromstring(z.read('word/document.xml'))
        rels = ET.fromstring(z.read('word/_rels/document.xml.rels'))
        targets = [r.get('Target') for r in rels if r.get('Type', '').rsplit('/', 1)[-1] == 'styles']
        check(len(targets) == 1, 'Independent reader needs a unique actual styles relationship')
        target = targets[0].lstrip('/') if targets[0].startswith('/') else posixpath.normpath(posixpath.join('word', targets[0]))
        styles = ET.fromstring(z.read(target))
    style_map = {x.get(qn('w:styleId')): x for x in styles.findall('w:style', NS)}
    default_style = next((x.get(qn('w:styleId')) for x in style_map.values() if x.get(qn('w:type')) == 'paragraph' and x.get(qn('w:default')) in ('1', 'true', 'on')), None)
    defaults = direct_indent(styles.find('w:docDefaults/w:pPrDefault/w:pPr', NS))
    def chain(sid, active=()):
        check(sid not in active, 'Independent style cycle')
        style = style_map.get(sid)
        if style is None: return []
        parent = style.find('w:basedOn', NS)
        return ([] if parent is None else chain(parent.get(qn('w:val')), active + (sid,))) + [style.find('w:pPr', NS)]
    rows = []
    for para in main.findall('.//w:body//w:p', NS):
        pp = para.find('w:pPr', NS); ps = None if pp is None else pp.find('w:pStyle', NS)
        sid = None if ps is None else ps.get(qn('w:val'))
        effective = dict(defaults)
        for ppr in chain(default_style): effective.update(direct_indent(ppr))
        for ppr in chain(sid): effective.update(direct_indent(ppr))
        effective.update(direct_indent(pp))
        text = ''.join((el.text or '') if el.tag == qn('w:t') else '\n' for el in para.iter() if el.tag in (qn('w:t'), qn('w:br')))
        jc=None if pp is None else pp.find('w:jc',NS)
        rows.append({'text': text, 'indentation': effective, 'direct': direct_indent(pp), 'style': sid, 'align':None if jc is None else jc.get(qn('w:val')), 'in_table': any(a.tag == qn('w:tc') for a in para.iterancestors())})
    doc = Document(path); py = []
    for el in doc.element.body.iter(qn('w:p')):
        para = Paragraph(el, doc); f = para.paragraph_format
        py.append({'text': para.text, 'direct': attrs3(*[None if x is None else x.twips for x in (f.left_indent, f.right_indent, f.first_line_indent)])})
    check(len(rows) == len(py), 'Independent/python-docx paragraph count differs')
    for a, b in zip(rows, py):
        check(a['text'] == b['text'] and a['direct'] == b['direct'], 'Independent/python-docx direct declaration mismatch')
    return {'ooxml': rows, 'python_docx': py}

def model_rows(model, ancestors=()):
    if model['type'] in ('paragraph', 'heading'):
        def text(node):
            return node.get('text', '') if node['type'] == 'text' else ('\n' if node['type'] == 'hardBreak' else ''.join(text(c) for c in node.get('content', [])))
        yield {'text': text(model), 'type': model['type'], 'indentation': {k: model.get('attrs', {})[k] for k in KEYS if k in model.get('attrs', {})}, 'align':model.get('attrs',{}).get('textAlign'), 'in_table': 'table' in ancestors, 'ancestors':list(ancestors)}
    for child in model.get('content', []): yield from model_rows(child, ancestors + (model['type'],))

def verify_rows(rows, targets, source_rows, label):
    check(len(rows) == len(targets) == len(source_rows), label + ': paragraph count differs')
    for i, (row, target, source) in enumerate(zip(rows, targets, source_rows)):
        check(row['indentation'] == target['attrs'], f'{label} paragraph {i}: {row["indentation"]} != {target["attrs"]}')
        check(row['text'] == source['text'], f'{label} paragraph {i}: text changed')
        check(row.get('in_table', False) == target.get('in_table', False), label + ': table context differs')
        if 'align' in target:check(row.get('align')==target['align'],label+': paragraph alignment differs')
        if 'type' in row: check(row['type'] == target.get('type', 'paragraph'), label + ': node type differs')
        if 'type' in row:
            for ancestor in target.get('ancestors',[]):check(ancestor in row['ancestors'],label+': list/container structure changed')

def mutate_package(path, action):
    with zipfile.ZipFile(path) as z: parts = {n: z.read(n) for n in z.namelist()}
    doc = ET.fromstring(parts['word/document.xml']); styles = ET.fromstring(parts['word/styles.xml']); settings = ET.fromstring(parts['word/settings.xml'])
    action(doc, styles, settings, parts)
    for name, root in [('document', doc), ('styles', styles), ('settings', settings)]: parts['word/' + name + '.xml'] = ET.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items(): z.writestr(n, b)

def render_document(kind, edited=False):
    doc = fresh()
    if kind == 'table':
        table = w.add_table(doc, [3600, 3600], rows=1)
        for index, prefix in enumerate(['CellA', 'CellB']):
            para = table.cell(0, index).paragraphs[0]; para.text = ' '.join(f'{prefix}{i:02d}' for i in range(24)); font(para)
            set_indent(para, *( (360, 480, -240) if edited else (180, 240, 120)))
    elif kind=='alignment':
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for prefix,alignment,first in [('AlignedC',WD_ALIGN_PARAGRAPH.CENTER,480),('AlignedR',WD_ALIGN_PARAGRAPH.RIGHT,-360)]:
            para=doc.add_paragraph(' '.join(f'{prefix}{i:02d}' for i in range(32)));font(para);s.set_spacing(para,0,0,240);para.paragraph_format.alignment=alignment;set_indent(para,720,360,first)
    else:
        prefix = 'Head' if kind == 'heading' else 'Body'
        words = ' '.join(f'{prefix}{i:02d}' for i in range(48))
        para = doc.add_heading(words, level=2) if kind == 'heading' else doc.add_paragraph(words)
        font(para); s.set_spacing(para, 0, 0, 240)
        if kind == 'heading':
            for run in para.runs: run.bold = True
        set_indent(para, *( (1080, 1440, -360) if edited else (720, 360, 480)))
    return doc

def make_fixtures(directory, expected):
    sources = {}
    def save(name, doc, reject=False, mutation=None, render=None):
        path = directory / (name + '.docx'); doc.save(path)
        if mutation: mutate_package(path, mutation)
        item = {'path': str(path), 'sha256': digest(path), 'fixture_kind': 'directed synthetic python-docx source', 'expected_action': 'reject' if reject else 'roundtrip'}
        if reject:item['rejection_stage']='contextual-import-fit' if name in ('reject-body-width','reject-first-width','reject-first-outside','reject-sides-width','reject-cell-fit') else 'raw-source-audit'
        if render: item['render_kind'] = render
        if not reject:
            item['expected'] = expected['supported'][name]
            item['source_independent'] = independent(path)
            verify_rows(item['source_independent']['ooxml'], item['expected'], item['source_independent']['ooxml'], name + ' source')
        sources[name] = item
    doc = Document(); doc.add_paragraph('UntouchedTemplate'); save('untouched-template', doc)
    doc = fresh(); doc.add_paragraph('AbsentIndent'); save('absent-defaults', doc)
    for name, values in [('direct-left', (720,None,None)), ('direct-right', (None,360,None)), ('direct-first', (None,None,480)), ('direct-hanging', (720,360,-360)), ('explicit-zero', (0,0,0)), ('one-twip', (1,1,-1)), ('body-fit-boundary', (9025,0,0))]:
        doc = fresh(); para = doc.add_paragraph(name); set_indent(para, *values); save(name, doc)
    doc = fresh(); para = doc.add_heading('IndentedHeading', level=2); set_indent(para, 721,239,241); save('heading', doc)
    doc = fresh(); default = doc.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr', NS); ind(default,left=180,right=240,firstLine=60)
    set_indent(doc.styles['Normal'], right=360)
    parent = doc.styles.add_style('IndentParent', WD_STYLE_TYPE.PARAGRAPH); parent.base_style=doc.styles['Normal']; set_indent(parent,left=720)
    child = doc.styles.add_style('IndentChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style=parent; set_indent(child,first=-180)
    doc.add_paragraph('DefaultCascade'); doc.add_paragraph('ParentCascade',style=parent); doc.add_paragraph('ChildCascade',style=child)
    para=doc.add_paragraph('DirectPartialZero',style=child);set_indent(para,right=0,first=0);save('cascade-components',doc)
    doc=fresh()
    cases=[('FirstToHang',720,-360,None),('HangToFirst',-360,720,None),('DirectHang',720,None,-360),('DirectFirst',-360,None,720),('CancelHang',-360,None,0),('CancelFirst',720,None,0)]
    for index,(label,base,derived,direct) in enumerate(cases):
        parent=doc.styles.add_style('PairBase'+str(index),WD_STYLE_TYPE.PARAGRAPH);parent.base_style=doc.styles['Normal'];set_indent(parent,left=720,right=360,first=base)
        child=doc.styles.add_style('PairChild'+str(index),WD_STYLE_TYPE.PARAGRAPH);child.base_style=parent;set_indent(child,first=derived)
        para=doc.add_paragraph(f'Cascade{index}A {label}\nCascade{index}B continuation',style=child);font(para);set_indent(para,first=direct)
    # Attribute-specific zero is authored explicitly to test cancellation across kinds.
    ind(doc.paragraphs[4]._p.get_or_add_pPr(),firstLine=0);ind(doc.paragraphs[5]._p.get_or_add_pPr(),hanging=0)
    save('cascade-first-hanging',doc,render='cascade')
    doc=fresh();set_indent(doc.styles['Normal'],left=360,right=120,first=-180);doc.add_paragraph('EmptyStart');doc.add_paragraph();doc.add_paragraph();doc.add_paragraph('EmptyEnd');save('empty-paragraphs',doc)
    doc=fresh();table=w.add_table(doc,[3600,3600],rows=1);table.cell(0,0).text='LeftCell';table.cell(0,1).text='RightCell';set_indent(table.cell(0,0).paragraphs[0],180,120,120);set_indent(table.cell(0,1).paragraphs[0],240,180,-120);save('fixed-table',doc)
    doc=fresh();table=w.add_table(doc,[1200,1200],rows=1);set_indent(table.cell(0,0).paragraphs[0],983);save('fixed-cell-boundary',doc)
    doc=fresh();doc.add_paragraph('OrdinaryNumberedParagraph',style='List Number');save('ordinary-list',doc)
    doc=fresh();doc.add_paragraph('OrdinaryBulletParagraph',style='List Bullet');save('ordinary-bullet',doc)
    doc=fresh();para=doc.add_paragraph('NeutralFlags');set_indent(para,720,360,-180)
    for tag in ['mirrorIndents','adjustRightInd','bidi']:w.child(para._p.get_or_add_pPr(),tag,val='0')
    w.child(para._p.get_or_add_pPr(),'tabs')
    save('neutral-flags',doc)
    doc=fresh();style=doc.styles.add_style('UnusedCharacterIndent',WD_STYLE_TYPE.PARAGRAPH);ind(style.element.get_or_add_pPr(),firstLineChars=200);doc.add_paragraph('UnusedStyle');save('unused-character-style',doc)
    for kind in ['body','heading','table','alignment']:save('render-'+kind,render_document(kind),render=kind)
    for right in [0,2880]:
        doc=fresh();para=doc.add_paragraph(' '.join(f'Right{i:02d}' for i in range(48)));font(para);set_indent(para,0,right,0);save('render-right-'+str(right),doc,render='right')
    invalid=[('left-negative',{'left':'-1'}),('right-negative',{'right':'-1'}),('first-negative',{'firstLine':'-1'}),('hanging-negative',{'hanging':'-1'}),('left-large',{'left':'31681'}),('right-large',{'right':'31681'}),('first-large',{'firstLine':'31681'}),('hanging-large',{'hanging':'31681'}),('decimal',{'left':'1.5'}),('unit',{'left':'12pt'}),('boolean',{'right':'true'}),('empty-value',{'left':''}),('dual',{'firstLine':'120','hanging':'60'}),('dual-zero',{'firstLine':'0','hanging':'0'}),('first-outside',{'left':'100','hanging':'101'}),('body-width',{'left':'9026'}),('first-width',{'left':'720','firstLine':'8306'}),('sides-width',{'left':'5000','right':'4026'})]
    for key in ['leftChars','rightChars','firstLineChars','hangingChars','start','end','startChars','endChars']:invalid.append((key, {key:'0'}))
    for name,attrs in invalid:
        doc=fresh();ind(doc.add_paragraph('UnsupportedIndent')._p.get_or_add_pPr(),**attrs);save('reject-'+name,doc,True)
    for tag in ['mirrorIndents','adjustRightInd','bidi','tabs']:
        doc=fresh();para=doc.add_paragraph('UnsupportedFlag');set_indent(para,720,360,120);node=w.child(para._p.get_or_add_pPr(),tag)
        if tag=='tabs':w.child(node,'tab',val='right',pos=1440)
        save('reject-'+tag,doc,True)
    for mode in ['duplicate','foreign-attribute','foreign-element','unknown-attribute','child','wrapped']:
        doc=fresh();para=doc.add_paragraph('MalformedIndent');set_indent(para,720)
        def mutation(root,styles,settings,parts,mode=mode):
            pp=root.find('w:body/w:p/w:pPr',NS);node=pp.find('w:ind',NS)
            if mode=='duplicate':pp.append(copy.deepcopy(node))
            elif mode=='foreign-attribute':node.set('{urn:other}left','120')
            elif mode=='foreign-element':node.tag='{urn:other}ind'
            elif mode=='unknown-attribute':node.set(qn('w:mystery'),'120')
            elif mode=='child':w.child(node,'ind',left=120)
            else:pp.remove(node);w.child(pp,'ins').append(node)
        save('reject-'+mode,doc,True,mutation)
    for location in ['defaults','used-style']:
        doc=fresh();pp=doc.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr',NS) if location=='defaults' else doc.styles['Normal'].element.get_or_add_pPr();ind(pp,firstLineChars=100);doc.add_paragraph('UsedCharacterIndent');save('reject-character-'+location,doc,True)
    doc=fresh();para=doc.add_paragraph('BodyTab\tIndented');set_indent(para,720,360,120);save('reject-tab-content',doc,True)
    doc=fresh();para=doc.add_paragraph('ZeroTab\tExplicit');set_indent(para,0,0,0);save('reject-zero-tab-content',doc,True)
    for location in ['defaults','used-style']:
        doc=fresh();pp=doc.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr',NS) if location=='defaults' else doc.styles['Normal'].element.get_or_add_pPr();w.child(w.child(pp,'tabs'),'tab',val='right',pos=1440);para=doc.add_paragraph('InheritedTabStops');set_indent(para,720,0,0);save('reject-tabs-'+location,doc,True)
    doc=fresh();para=doc.add_paragraph('NumberedIndent',style='List Number');set_indent(para,720,360,120);save('reject-list-direct',doc,True)
    for kind in ['missing-abstract-reference','invalid-level-zero','invalid-abstract-zero']:
        doc=fresh();para=doc.add_paragraph('NumberingIdentity');number=w.child(para._p.get_or_add_pPr(),'numPr');w.child(number,'ilvl',val=0);w.child(number,'numId',val=1)
        def numbering_identity(root,styles,settings,parts,kind=kind):
            nums=ET.fromstring(parts['word/numbering.xml']);num=next(n for n in nums.findall('w:num',NS) if n.get(qn('w:numId'))=='1');ref=num.find('w:abstractNumId',NS)
            abstract=next(n for n in nums.findall('w:abstractNum',NS) if n.get(qn('w:abstractNumId'))=='0');level=abstract.find('w:lvl',NS);pp=level.find('w:pPr',NS)
            if pp is None:pp=w.child(level,'pPr')
            ind(pp,leftChars=100)
            if kind=='missing-abstract-reference':num.remove(ref)
            else:
                ref.set(qn('w:val'),'0')
                if kind=='invalid-level-zero':level.set(qn('w:ilvl'),'malformed')
                else:abstract.set(qn('w:abstractNumId'),'malformed')
            parts['word/numbering.xml']=ET.tostring(nums,xml_declaration=True,encoding='UTF-8',standalone=True)
        save('reject-numbering-'+kind,doc,True,numbering_identity)
    doc=fresh();table=w.add_table(doc,[1200,1200],rows=1);set_indent(table.cell(0,0).paragraphs[0],984);save('reject-cell-fit',doc,True)
    check(set(expected['supported'])=={n for n,i in sources.items() if i['expected_action']=='roundtrip'},'Expected manifest/source fixture set differs')
    return sources

def render_pdf(soffice, directory, path):
    out=directory/'libreoffice-pdf';out.mkdir(exist_ok=True);profile=directory/'lo-profile'
    run=subprocess.run([str(soffice),f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','pdf','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=120,env=dict(os.environ,XDG_CACHE_HOME=str(directory/'lo-cache')))
    pdf=out/(path.stem+'.pdf');check(run.returncode==0 and pdf.exists(),run.stdout+run.stderr)
    bbox=out/(path.stem+'.html');subprocess.run(['pdftotext','-bbox-layout',str(pdf),str(bbox)],check=True,capture_output=True,timeout=30)
    tree=ET.parse(str(bbox));lines=[]
    for page_index,page in enumerate(tree.xpath('//*[local-name()="page"]')):
        for line in page.xpath('.//*[local-name()="line"]'):
            words=line.xpath('./*[local-name()="word"]');text=' '.join(''.join(word.itertext()) for word in words)
            if not words:continue
            lines.append({'page':page_index+1,'text':text,'x':float(line.get('xMin')),'x_end':float(line.get('xMax')),'y':float(line.get('yMin')),'words':[{'text':''.join(t.itertext()),'x':float(t.get('xMin')),'x_end':float(t.get('xMax'))} for t in words]})
    return {'path':str(pdf),'sha256':digest(pdf),'lines':lines,'stdout':run.stdout,'stderr':run.stderr,'consumer':'LibreOffice PDF export, independently read by Poppler; not app printing'}

def grouped_lines(pdf):
    groups={}
    for line in pdf['lines']:
        match=re.match(r'(CellA|CellB|Body|Head|Right|AlignedC|AlignedR)\d\d|(Cascade[0-5])[AB]',line['text'])
        if match:groups.setdefault(match.group(1) or match.group(2),[]).append(line)
    return groups

def compare_pdf(a,b,label):
    check(len(a['lines'])==len(b['lines']),label+': line count differs')
    maximum=0
    for x,y in zip(a['lines'],b['lines']):
        check(x['text']==y['text'] and x['page']==y['page'],label+': text wrapping/page differs')
        delta=max(abs(x[k]-y[k]) for k in ['x','x_end'])
        check(delta<=0.35,f'{label}: horizontal positions differ by {delta:.3f} pt')
        maximum=max(maximum,delta)
    return {'same_line_breaks':True,'max_horizontal_delta_pt':maximum,'line_count':len(a['lines'])}

def assert_render_indent(pdf,expected_first,label,expected_left=0):
    result={}
    for prefix,lines in grouped_lines(pdf).items():
        check(len(lines)>=2,label+': expected wrapped continuation')
        delta=lines[0]['x']-lines[1]['x'];first=expected_first[prefix] if isinstance(expected_first,dict) else expected_first
        check(abs(delta-first/20)<=0.35,f'{label}: first/continuation x delta {delta} != {first/20}')
        check(all(abs(x['x']-lines[1]['x'])<=0.35 for x in lines[1:]),label+': continuation start moved')
        # Fixed tables have explicit 3600-twip columns and the established
        # 108-twip cell content margins. Plain text starts at the 1440-twip page
        # margin. These anchors are specified independently from PDF output.
        origin=72+(5.4+(180 if prefix=='CellB' else 0) if prefix.startswith('Cell') else 0)
        target_x=origin+expected_left/20
        check(abs(lines[1]['x']-target_x)<=0.35,f'{label}: continuation x {lines[1]["x"]} != expected {target_x}')
        result[prefix]={'first_x_pt':lines[0]['x'],'continuation_x_pt':lines[1]['x'],'expected_continuation_x_pt':target_x,'first_delta_pt':delta,'line_count':len(lines)}
    check(result,label+': no controlled marker lines')
    return result

def alignment_positions(pdf):
    groups=grouped_lines(pdf)
    check(set(groups)=={'AlignedC','AlignedR'} and all(len(lines)>=2 for lines in groups.values()),'Missing controlled centered/right-aligned multiline paragraphs')
    right_edge=11906/20-72-360/20
    check(all(abs(line['x_end']-right_edge)<=0.35 for line in groups['AlignedR']),'Right alignment failed to honor the page margin and right paragraph indent')
    return {key:[{'text':line['text'],'x':line['x'],'x_end':line['x_end']} for line in lines] for key,lines in groups.items()}

def export_model(cli,evidence,directory,name,model):
    model_path=directory/(name+'.json');model_path.write_text(json.dumps(model,ensure_ascii=False,indent=2),encoding='utf-8')
    output=directory/(name+'.docx');run_cli(cli,evidence['commands'],'create-docx',model_path,output)
    reopened=run_cli(cli,evidence['commands'],'read-docx-model',output);check(reopened==model,name+': complete trusted model differs')
    before=digest(output);error=run_cli(cli,evidence['commands'],'create-docx',model_path,output,success=False)
    check('Output already exists' in error and digest(output)==before,name+': overwrite was not refused atomically')
    return {'path':str(output),'sha256':before,'model_path':str(model_path),'complete_trusted_model_equal':True,'no_clobber':True,'independent':independent(output)}

def roundtrip(cli,evidence,directory,name,item,soffice=None):
    source=Path(item['path']);source_rows=item['source_independent']['ooxml'];targets=item['expected']
    item['audit']=run_cli(cli,evidence['commands'],'audit-docx',source)
    imported=run_cli(cli,evidence['commands'],'import-docx',source);model=imported['document'];item['compatibility']=imported['compatibility'];item['model']=model
    verify_rows(list(model_rows(model)),targets,source_rows,name+' import')
    result=export_model(cli,evidence,directory,name+'-export',model);item['export']=result
    verify_rows(result['independent']['ooxml'],targets,source_rows,name+' export')
    for row,target in zip(result['independent']['python_docx'],targets):check(row['direct']==target['attrs'],name+': export lost optional presence or zero')
    reimported=run_cli(cli,evidence['commands'],'import-docx',result['path'])['document'];verify_rows(list(model_rows(reimported)),targets,source_rows,name+' external reimport')
    item['external_reimport_complete_model_equal']=reimported==model
    error=run_cli(cli,evidence['commands'],'create-docx',result['model_path'],source,success=False)
    check('Output already exists' in error and digest(source)==item['sha256'],name+': source overwrite failed')
    item['source_no_clobber']=True
    if soffice and item.get('render_kind'):
        item['source_pdf']=render_pdf(soffice,directory,source);item['export_pdf']=render_pdf(soffice,directory,Path(result['path']))
        item['pdf_comparison']=compare_pdf(item['source_pdf'],item['export_pdf'],name)
        if item['render_kind']=='alignment':item['alignment_positions']=alignment_positions(item['export_pdf'])
        else:
            first={f'Cascade{i}':target['attrs'][KEYS[2]] for i,target in enumerate(targets)} if item['render_kind']=='cascade' else targets[0]['attrs'].get(KEYS[2],0);item['pdf_indent_assertions']=assert_render_indent(item['export_pdf'],first,name,targets[0]['attrs'].get(KEYS[0],0))
    return model

def edited_probes(cli,evidence,directory,sources,expected,soffice):
    edits={}
    for kind in ['body','heading','table']:
        item=sources['render-'+kind];model=copy.deepcopy(item['model']);targets=expected['edits'][kind]
        def paragraphs(node):
            if node['type'] in ('paragraph','heading'):yield node
            for child in node.get('content',[]):yield from paragraphs(child)
        for node,target in zip(paragraphs(model),targets):
            for key in KEYS:node.setdefault('attrs',{}).pop(key,None)
            node['attrs'].update(target['attrs'])
        output=export_model(cli,evidence,directory,'edited-'+kind,model)
        reference=directory/('reference-edited-'+kind+'.docx');render_document(kind,True).save(reference);ref=independent(reference)
        verify_rows(output['independent']['ooxml'],targets,ref['ooxml'],kind+' edited output')
        result={'output':output,'independent_reference':str(reference),'reference_sha256':digest(reference),'source_model_changed':model!=item['model']}
        check(result['source_model_changed'],'Edit did not change model')
        if soffice:
            result['reference_pdf']=render_pdf(soffice,directory,reference);result['output_pdf']=render_pdf(soffice,directory,Path(output['path']));result['pdf_comparison']=compare_pdf(result['reference_pdf'],result['output_pdf'],kind+' edited vs reference');result['pdf_indent_assertions']=assert_render_indent(result['output_pdf'],targets[0]['attrs'][KEYS[2]],kind+' edited',targets[0]['attrs'][KEYS[0]])
        edits[kind]=result
    evidence['edited_probes']=edits

def invalid_models(cli,evidence,directory,good):
    tests={}
    for key in KEYS:
        for index,value in enumerate([None,True,'20',1.5,[],{},31681,-31681]):
            m=copy.deepcopy(good);m['content'][0].setdefault('attrs',{})[key]=value;tests[f'{key}-{index}']=m
    for index,attrs in enumerate([attrs3(-1),attrs3(None,-1),attrs3(None,None,-1),attrs3(100,None,-101),attrs3(9026),attrs3(None,9026),attrs3(None,None,9026),attrs3(4500,4526),attrs3(720,None,8306)]):
        m=copy.deepcopy(good);m['content'][0]['attrs'].update(attrs);tests['fit-'+str(index)]=m
    for location in ['root','text','list','table']:
        m=copy.deepcopy(good)
        if location=='root':m['attrs'][KEYS[0]]=0
        elif location=='text':m['content'][0]['content'][0].setdefault('attrs',{})[KEYS[0]]=0
        elif location=='list':
            m['content'][0].setdefault('attrs',{})[KEYS[0]]=0;m['content']=[{'type':'bulletList','content':[{'type':'listItem','content':m['content']}]}]
        else:m['content']=[{'type':'table','attrs':{KEYS[0]:0},'content':[{'type':'tableRow','content':[{'type':'tableCell','content':m['content']}]}]}]
        tests['location-'+location]=m
    results={}
    for name,model in tests.items():
        path=directory/('invalid-'+name+'.json');path.write_text(json.dumps(model));out=path.with_suffix('.docx');before=digest(path);error=run_cli(cli,evidence['commands'],'create-docx',path,out,success=False)
        check(not out.exists() and digest(path)==before,name+': rejected model changed input or created output');results[name]={'error':error,'no_output':True}
    evidence['invalid_models']=results

def tamper_packages(cli,evidence,directory,good):
    with zipfile.ZipFile(good) as z:original={n:z.read(n) for n in z.namelist()}
    results={}
    for mode in ['xml-only','xml-forged-hash','model-only','model-forged-hash','remove-settings','settings-forged-hash']:
        parts=copy.deepcopy(original);embedded=json.loads(parts['rust-wps/model-v2.json'])
        if mode.startswith('xml'):
            root=ET.fromstring(parts['word/document.xml']);node=root.find('.//w:p/w:pPr/w:ind',NS);node.set(qn('w:left'),str(int(node.get(qn('w:left'),'0'))+20));parts['word/document.xml']=ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        elif mode.startswith('model'):embedded['document']['content'][0]['attrs'][KEYS[0]]+=20
        elif mode=='remove-settings':parts.pop('word/settings.xml')
        else:
            settings=ET.fromstring(parts['word/settings.xml']);w.child(settings,'mirrorMargins');parts['word/settings.xml']=ET.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
        if mode.endswith('forged-hash'):
            for name in embedded['hashes']:
                if name in parts:embedded['hashes'][name]=hashlib.sha256(parts[name]).hexdigest()
        parts['rust-wps/model-v2.json']=json.dumps(embedded,separators=(',',':')).encode();path=directory/('tamper-'+mode+'.docx')
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for n,b in parts.items():z.writestr(n,b)
        before=digest(path);error=run_cli(cli,evidence['commands'],'read-docx-model',path,success=False);check(digest(path)==before,'Tampered source changed');results[mode]={'error':error,'sha256':before}
    evidence['tampered_packages']=results

def corpus_checks(cli,evidence,directory,soffice):
    manifest=json.loads(CORPUS.read_text());results={}
    for fixture in manifest['fixtures']:
        path=ROOT/fixture['local_default'];check(path.exists(),'Missing pinned upstream fixture '+str(path));check(digest(path)==fixture['sha256'],'Upstream fixture SHA-256 mismatch')
        blob=hashlib.sha1(b'blob '+str(path.stat().st_size).encode()+b'\0'+path.read_bytes()).hexdigest();check(blob==fixture['git_blob_sha'],'Upstream git blob mismatch')
        item=copy.deepcopy(fixture);item['path']=str(path)
        if fixture['expected_action']=='roundtrip':
            item['source_independent']=independent(path);item['expected']=fixture['expected'];roundtrip(cli,evidence,directory,'corpus-'+path.stem,item)
            if soffice and fixture.get('render'):
                item['source_pdf']=render_pdf(soffice,directory,path);item['export_pdf']=render_pdf(soffice,directory,Path(item['export']['path']));item['pdf_comparison']=compare_pdf(item['source_pdf'],item['export_pdf'],'corpus '+path.name)
                lefts=[line['x'] for line in item['export_pdf']['lines']];check(lefts and all(abs(x-(72+567/20))<0.35 for x in lefts),'Real fixture paragraph x differs from 567-twip indent');item['left_positions_pt']=lefts
        else:
            item['audit']=run_cli(cli,evidence['commands'],'audit-docx',path);item['import_error']=run_cli(cli,evidence['commands'],'import-docx',path,success=False)
            for risk in fixture.get('required_risks',[]):check(risk in item['audit']['risks'],'Corpus diagnostic missing '+risk)
        check(digest(path)==fixture['sha256'],'Real source was changed');item['source_unchanged']=True;results[path.name]=item
    evidence['upstream_corpus']=results

def legacy_absence(cli, baseline, evidence, directory):
    expected_sha='32c4fa62845fccba00d7e29f11a11484d67f546b5ab21d3dd09597148e289e64'
    check(baseline.is_file() and digest(baseline)==expected_sha,'Missing or mismatched pre-indentation baseline CLI')
    plain={'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':'Legacy paragraph'}]}]}
    page={'widthTwips':12240,'heightTwips':15840,'marginTopTwips':1440,'marginRightTwips':1440,'marginBottomTwips':1440,'marginLeftTwips':1440}
    table={'type':'table','attrs':{'columnWidthsTwips':[1800,2400]},'content':[{'type':'tableRow','content':[{'type':'tableCell','content':copy.deepcopy(plain['content'])} for _ in range(2)]}]}
    cases={'mixed-authoring':json.loads((ROOT/'tests/fixtures/authoring.json').read_text()),'explicit-page':copy.deepcopy(plain),'fixed-table':{'type':'doc','content':[copy.deepcopy(table)]},'page-auto-table':{'type':'doc','attrs':{'pageGeometry':page},'content':[copy.deepcopy(table)]},'image':{'type':'doc','content':[copy.deepcopy(json.loads((ROOT/'tests/fixtures/word-image-legacy.json').read_text())['content'][0])]},'paragraph-spacing':copy.deepcopy(plain)}
    cases['explicit-page']['attrs']={'pageGeometry':page};cases['page-auto-table']['content'][0].pop('attrs');cases['paragraph-spacing']['content'][0]['attrs']={'spaceBefore':120,'spaceAfter':240,'lineSpacing':360}
    result={}
    for name,original in cases.items():
        for profile in ['legacy-v1','standard-v2']:
            model=copy.deepcopy(original)
            if profile=='standard-v2':model.setdefault('attrs',{})['styleProfile']=profile
            label='absence-'+name+'-'+profile;path=directory/(label+'.json');path.write_text(json.dumps(model,ensure_ascii=False));outputs=[]
            for suffix,program in [('old',baseline),('new',cli)]:
                out=directory/(label+'-'+suffix+'.docx');run_cli(program,evidence['commands'],'create-docx',path,out);outputs.append(out)
                check(run_cli(program,evidence['commands'],'read-docx-model',out)==model,label+': full trusted model changed')
            check(outputs[0].read_bytes()==outputs[1].read_bytes(),label+': absent-attribute historical bytes changed')
            with zipfile.ZipFile(outputs[1]) as package:part_hashes={n:hashlib.sha256(package.read(n)).hexdigest() for n in package.namelist()}
            result[label]={'sha256':digest(outputs[1]),'full_package_byte_equal':True,'complete_trusted_models_equal':True,'part_sha256':part_hashes}
    check(digest(baseline)==expected_sha,'Baseline CLI changed during checks')
    evidence['legacy_absence']={'baseline_cli':str(baseline),'baseline_cli_sha256':expected_sha,'cases':result}

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--cli',type=Path,default=Path('target/debug/rust-wps'));parser.add_argument('--root',type=Path,default=Path('tmp/word-paragraph-indentation'));parser.add_argument('--generate-only',action='store_true');parser.add_argument('--render',action='store_true');parser.add_argument('--soffice');parser.add_argument('--skip-corpus',action='store_true');parser.add_argument('--baseline',type=Path,default=Path('tmp/paragraph-indent-baseline/rust-wps'));args=parser.parse_args()
    args.root.mkdir(parents=True,exist_ok=True);directory=Path(tempfile.mkdtemp(prefix='run-',dir=args.root)).resolve();report=directory/'report.json'
    cli=None if args.generate_only else args.cli.resolve();soffice=(args.soffice or shutil.which('soffice')) if args.render else None;check(not args.render or soffice,'LibreOffice required');expected=json.loads(EXPECTED.read_text())
    evidence={'status':'running','script_sha256':digest(Path(__file__)),'expected_manifest_sha256':digest(EXPECTED),'corpus_manifest_sha256':digest(CORPUS),'helper_script_sha256':{str(Path(m.__file__).relative_to(ROOT)):digest(Path(m.__file__)) for m in [s,w]},'cli':str(cli) if cli else None,'cli_sha256':digest(cli) if cli else None,'python_version':platform.python_version(),'python_docx_version':version('python-docx'),'commands':[],'sources':{},'limitations':'Synthetic and untouched public upstream fixtures are separate. OOXML/model equality and LibreOffice/Poppler horizontal layout are bounded evidence, not app/browser/native Windows/WebView2/Microsoft Word/WPS or complete document fidelity.'}
    try:
        evidence['sources']=sources=make_fixtures(directory,expected)
        if soffice:evidence['libreoffice_version']=subprocess.run([soffice,'--version'],capture_output=True,text=True,check=True).stdout.strip()
        for name,item in sources.items():
            path=Path(item['path'])
            if cli:
                if item['expected_action']=='roundtrip':roundtrip(cli,evidence,directory,name,item,soffice)
                else:
                    item['audit']=run_cli(cli,evidence['commands'],'audit-docx',path);item['import_error']=run_cli(cli,evidence['commands'],'import-docx',path,success=False)
                    check(RISK in item['import_error'],name+': missing exact import indentation diagnostic')
                    if item['rejection_stage']=='raw-source-audit':check(RISK in item['audit']['risks'],name+': missing exact raw-source indentation diagnostic')
                    # Container-fit proof requires the final imported page/table
                    # model. The read-only raw audit is recorded, not misreported
                    # as a separate conversion or contextual-fit proof.
                    failure=run_cli(cli,evidence['commands'],'create-docx',sources['absent-defaults']['export']['model_path'],path,success=False)
                    check('Output already exists' in failure and digest(path)==item['sha256'],name+': rejected source overwrite was not refused');item['source_no_clobber']=True
            elif soffice and item.get('render_kind'):
                item['source_pdf']=render_pdf(soffice,directory,path)
                if item['render_kind']=='alignment':item['alignment_positions']=alignment_positions(item['source_pdf'])
                else:
                    first={f'Cascade{i}':target['attrs'][KEYS[2]] for i,target in enumerate(item['expected'])} if item['render_kind']=='cascade' else item['expected'][0]['attrs'].get(KEYS[2],0);item['pdf_indent_assertions']=assert_render_indent(item['source_pdf'],first,name,item['expected'][0]['attrs'].get(KEYS[0],0))
            check(digest(path)==item['sha256'],name+': source changed');item['source_unchanged']=True
        if cli:
            edited_probes(cli,evidence,directory,sources,expected,soffice);invalid_models(cli,evidence,directory,sources['absent-defaults']['model']);tamper_packages(cli,evidence,directory,Path(sources['render-body']['export']['path']));legacy_absence(cli,args.baseline.resolve(),evidence,directory)
            if not args.skip_corpus:corpus_checks(cli,evidence,directory,soffice)
        if soffice:
            stage='export_pdf' if cli else 'source_pdf';a=grouped_lines(sources['render-right-0'][stage])['Right'];b=grouped_lines(sources['render-right-2880'][stage])['Right'];check(len(b)>len(a),'Right indentation did not increase wrapping');evidence['right_indent_wrapping']={'unindented_lines':len(a),'right_2880_lines':len(b),'more_wrapping':True}
        evidence['status']='fixtures-verified' if args.generate_only else 'passed';evidence['summary']={'supported_sources':sum(i['expected_action']=='roundtrip' for i in sources.values()),'rejected_sources':sum(i['expected_action']=='reject' for i in sources.values()),'invalid_models':len(evidence.get('invalid_models',{})),'tampered_packages':len(evidence.get('tampered_packages',{})),'edited_probes':len(evidence.get('edited_probes',{})),'legacy_absence_byte_cases':len(evidence.get('legacy_absence',{}).get('cases',{})),'upstream_fixtures':len(evidence.get('upstream_corpus',{})),'rendered_pdfs':len(list((directory/'libreoffice-pdf').glob('*.pdf')))}
    except Exception as error:evidence['status']='failed';evidence['error']=f'{type(error).__name__}: {error}';raise
    finally:
        evidence['script_unchanged_during_run']=digest(Path(__file__))==evidence['script_sha256'];evidence['expected_manifest_unchanged']=digest(EXPECTED)==evidence['expected_manifest_sha256'];evidence['corpus_manifest_unchanged']=digest(CORPUS)==evidence['corpus_manifest_sha256'];evidence['helper_scripts_unchanged']=all(digest(ROOT/path)==sha for path,sha in evidence['helper_script_sha256'].items());evidence['cli_unchanged_during_run']=cli is None or digest(cli)==evidence['cli_sha256'];evidence['source_hashes_unchanged_on_exit']=all(digest(Path(i['path']))==i['sha256'] for i in evidence['sources'].values());report.write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8');print(report)

if __name__=='__main__':main()
