"""Self-generated PPTX fixture + independent python-pptx/real Rust inspection."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import io
import base64
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]

def make(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text = '简文 PPTX <test>'
    run = slide.shapes[0].text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.italic = Pt(24), True, True
    run.font.color.rgb = RGBColor(160, 32, 64)
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(2), Inches(3), Inches(2), Inches(1)).text = '组合文字'
    pixels = Image.new('RGB', (32, 16), (40, 120, 200))
    picture = io.BytesIO(); pixels.save(picture, format='PNG'); picture.seek(0)
    slide.shapes.add_picture(picture, Inches(7), Inches(1), Inches(3), Inches(1.5))
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second._element.set('show', '0')
    second.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'Hidden slide'
    deck.save(path)

if '--write-fixture' in sys.argv:
    make(ROOT / 'tests/fixtures/basic-slides.pptx')
else:
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / 'input.pptx'
        make(path)
        original = path.read_bytes()
        deck = Presentation(path)
        result = subprocess.run(['cargo', 'run', '--quiet', '--locked', '--', 'inspect-pptx', str(path)], cwd=ROOT, check=True, capture_output=True)
        report = json.loads(result.stdout)
        assert report['size_emu'] == [deck.slide_width, deck.slide_height]
        assert len(report['slides']) == len(deck.slides) == 2
        assert report['slides'][0]['id'] == deck.slides[0].slide_id
        shape = report['slides'][0]['shapes'][0]
        source = deck.slides[0].shapes[0]
        assert shape['bounds_emu'] == [source.left, source.top, source.width, source.height]
        assert '简文 PPTX <test>' in report['slides'][0]['text']
        assert '组合文字' in report['slides'][0]['text']
        assert shape['text'] == source.text
        run = source.text_frame.paragraphs[0].runs[0]
        assert shape['runs'][0]['font_size_hundredths_pt'] == round(run.font.size.pt * 100)
        assert shape['runs'][0]['bold'] is run.font.bold and shape['runs'][0]['italic'] is run.font.italic
        assert shape['runs'][0]['color_rgb'] == list(run.font.color.rgb)
        assert report['slides'][0]['shapes'][1]['children'][0]['text'] == deck.slides[0].shapes[1].shapes[0].text
        assert report['slides'][0]['shapes'][1]['children'][0]['kind'] == 'shape'
        image = report['slides'][0]['shapes'][2]
        assert image['image'].startswith('data:image/png;base64,')
        assert base64.b64decode(image['image'].split(',', 1)[1]) == deck.slides[0].shapes[2].image.blob
        assert report['slides'][1]['hidden'] and not report['unreadable_parts']
        assert path.read_bytes() == original
        edited = Path(temp) / 'edited.pptx'
        edit = subprocess.run(['cargo', 'run', '--quiet', '--locked', '--', 'replace-pptx', str(path), str(edited), '简文 PPTX <test>', '修改 & <内容>'], cwd=ROOT, check=True, capture_output=True)
        assert json.loads(edit.stdout)['replacements'] == 1
        reopened = Presentation(edited)
        changed = reopened.slides[0].shapes[0]
        assert changed.text == '修改 & <内容>'
        font = changed.text_frame.paragraphs[0].runs[0].font
        assert font.size.pt == 24 and font.bold and font.italic
        assert reopened.slides[0].shapes[2].image.blob == deck.slides[0].shapes[2].image.blob
        assert path.read_bytes() == original
        print(json.dumps({'independentPythonPptx': True, 'independentEditedReopen': True, 'sourceUnchanged': True, 'sha256': hashlib.sha256(original).hexdigest()}))
