"""Independent DOCX paragraph-spacing regression; never modifies input files.

The synthetic packages are authored with python-docx, with explicit XML only for
OOXML cases its public API cannot express. A fixed manifest (not the importer)
defines expected effective spacing. An independent OOXML cascade reader and
python-docx verify sources and exports. The CLI path is source -> imported JSON
-> newly authored DOCX -> trusted read-docx-model reopen (exact model equality).
External import-docx reconversion is checked separately. Every run uses a fresh
evidence directory.

  python scripts/check-word-paragraph-spacing.py --generate-only
  python scripts/check-word-paragraph-spacing.py --cli target/debug/rust-wps --render
  python scripts/check-word-paragraph-spacing.py --corpus tmp/corpus/SampleDoc.docx

--render uses LibreOffice PDF export and pypdf baseline positions. It is NOT the
app's print path, a native WebView2 test, or Microsoft Word/WPS visual validation.
The optional corpus is the separately pinned Apache POI upstream regression file
in scripts/corpus.json, never presented as synthetic or as user production data.
"""
import argparse
import hashlib
import io
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from lxml import etree as ET

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
RISK = "Unsupported paragraph spacing cannot be preserved in an editable copy"
KEYS = ("spaceBefore", "spaceAfter", "lineSpacing")
ATTRS = dict(zip(KEYS, ("before", "after", "line")))


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def xml_child(parent, tag, **attrs):
    node = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    parent.append(node)
    return node


def remove_children(parent, tag):
    for node in list(parent.findall("w:" + tag, NS)):
        parent.remove(node)


def set_xml_spacing(ppr, **attrs):
    remove_children(ppr, "spacing")
    return xml_child(ppr, "spacing", **attrs)


def set_spacing(paragraph_or_style, before=None, after=None, line=None):
    fmt = paragraph_or_style.paragraph_format
    if before is not None:
        fmt.space_before = Twips(before)
    if after is not None:
        fmt.space_after = Twips(after)
    if line is not None:
        fmt.line_spacing = line / 240


def fresh_document(defaults=(0, 0, 240)):
    doc = Document()
    # Isolate spacing: python-docx's template has contextualSpacing and unrelated spacing in
    # built-in styles. These controlled probes isolate explicit declarations.
    # A separate untouched template is a mandatory positive compatibility case.
    for node in list(doc.styles.element.findall(".//w:contextualSpacing", NS)):
        node.getparent().remove(node)
    for section in doc.sections:
        remove_children(section._sectPr, "docGrid")
    for style in doc.styles:
        ppr = style.element.find("w:pPr", NS)
        if ppr is not None:
            remove_children(ppr, "spacing")
    ppr = doc.styles.element.find("w:docDefaults/w:pPrDefault/w:pPr", NS)
    if ppr is None:
        defaults_xml = doc.styles.element.find("w:docDefaults", NS)
        if defaults_xml is None:
            defaults_xml = xml_child(doc.styles.element, "docDefaults")
        parent = defaults_xml.find("w:pPrDefault", NS)
        if parent is None:
            parent = xml_child(defaults_xml, "pPrDefault")
        ppr = xml_child(parent, "pPr")
    set_xml_spacing(ppr, before=defaults[0], after=defaults[1], line=defaults[2], lineRule="auto")
    doc.styles["Normal"].font.name = "DejaVu Sans"
    doc.styles["Normal"].font.size = Pt(12)
    return doc


def expected(text, before, after, line, node_type="paragraph", ancestors=None):
    result = {"text": text, "type": node_type,
              "spacing": dict(zip(KEYS, (before, after, line)))}
    if ancestors:
        result["ancestors"] = ancestors
    return result


def independent_snapshot(path):
    """Resolve declarations without using project code or a re-exported model."""
    with zipfile.ZipFile(path) as package:
        body = ET.fromstring(package.read("word/document.xml"))
        styles = ET.fromstring(package.read("word/styles.xml"))
    style_map = {s.get(qn("w:styleId")): s for s in styles.findall("w:style", NS)}
    base = styles.find("w:docDefaults/w:pPrDefault/w:pPr", NS)
    default_style = next((s for s in style_map.values()
                          if s.get(qn("w:type")) == "paragraph"
                          and s.get(qn("w:default")) in ("1", "true", "on")), None)

    def apply(effective, ppr):
        spacing = None if ppr is None else ppr.find("w:spacing", NS)
        if spacing is not None:
            for key, attr in ATTRS.items():
                value = spacing.get(qn("w:" + attr))
                if value is not None:
                    effective[key] = int(value)

    def chain(style_id, active=None):
        active = set() if active is None else active
        check(style_id not in active, "Independent reader found a style cycle")
        style = style_map.get(style_id)
        if style is None:
            return []
        based = style.find("w:basedOn", NS)
        parents = [] if based is None else chain(based.get(qn("w:val")), active | {style_id})
        return parents + [style.find("w:pPr", NS)]

    rows = []
    for p in body.findall(".//w:body//w:p", NS):
        text = "".join((n.text or "") if n.tag == qn("w:t") else ("\n" if n.get(qn("w:type")) in (None, "textWrapping") else "")
                       for n in p.iter() if n.tag in (qn("w:t"), qn("w:br")))
        ppr = p.find("w:pPr", NS)
        style = None if ppr is None else ppr.find("w:pStyle", NS)
        effective = {}
        apply(effective, base)
        if default_style is not None:
            for declaration in chain(default_style.get(qn("w:styleId"))):
                apply(effective, declaration)
        if style is not None:
            for declaration in chain(style.get(qn("w:val"))):
                apply(effective, declaration)
        apply(effective, ppr)
        spacing = None if ppr is None else ppr.find("w:spacing", NS)
        page_breaks = len(p.findall(".//w:br[@w:type='page']", NS))
        rows.append({"text": text, "explicit_page_breaks": page_breaks,
                     "page_break_only": not text and page_breaks > 0, "spacing": {key: effective.get(key) for key in KEYS},
                     "direct_xml": {} if spacing is None else
                     {ET.QName(k).localname: v for k, v in spacing.attrib.items()},
                     "style_id": None if style is None else style.get(qn("w:val")),
                     "in_table": any(a.tag == qn("w:tc") for a in p.iterancestors())})
    # python-docx is a second consumer for package validity, text, and direct
    # property interpretation; its API deliberately does not resolve inheritance.
    doc = Document(path)
    from docx.text.paragraph import Paragraph
    python_rows = []
    for p in doc.element.body.iter(qn("w:p")):
        paragraph = Paragraph(p, doc)
        fmt = paragraph.paragraph_format
        direct = {}
        for key, value in (("spaceBefore", fmt.space_before), ("spaceAfter", fmt.space_after)):
            if value is not None:
                direct[key] = value.twips
        if isinstance(fmt.line_spacing, float):
            direct["lineSpacing"] = round(fmt.line_spacing * 240)
        python_rows.append({"text": paragraph.text, "direct_spacing": direct,
                            "line_rule": None if fmt.line_spacing_rule is None else int(fmt.line_spacing_rule)})
    check([row["text"] for row in rows] == [row["text"] for row in python_rows],
          "Independent OOXML/python-docx paragraph text differs")
    return {"ooxml": rows, "python_docx": python_rows}


def semantic_spacing(values):
    # OOXML initial values are zero paragraph space and single line spacing.
    # Authored output may legitimately materialize absent zero declarations.
    return {key: values.get(key) if values.get(key) is not None else
            (240 if key == "lineSpacing" else 0) for key in KEYS}


def verify_independent(snapshot, expectations, label, require_direct=False):
    check(len(snapshot["ooxml"]) == len(expectations), f"{label}: paragraph count differs")
    for row, py, target in zip(snapshot["ooxml"], snapshot["python_docx"], expectations):
        check(row["text"] == target["text"], f"{label}: paragraph text differs: {row['text']!r}")
        check(semantic_spacing(row["spacing"]) == semantic_spacing(target["spacing"]),
              f"{label}: effective spacing mismatch for {target['text']!r}: {row['spacing']} != {target['spacing']}")
        for key, value in py["direct_spacing"].items():
            check(value == semantic_spacing(target["spacing"])[key], f"{label}: python-docx {key} mismatch")
        if require_direct:
            check(py["direct_spacing"] == semantic_spacing(target["spacing"]),
                  f"{label}: authored spacing was not materialized as expected: {py}")


def make_fixtures(directory):
    sources = {}

    def save(name, doc, rows=None, unsupported=None, render_probe=None, xml_transform=None):
        path = directory / (name + ".docx")
        check(not path.exists(), f"Fixture destination exists: {path}")
        if xml_transform is None:
            doc.save(path)
        else:
            buffer = io.BytesIO()
            doc.save(buffer)
            with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as output:
                for entry in source.infolist():
                    data = source.read(entry)
                    if entry.filename == "word/document.xml":
                        data = xml_transform(data)
                    output.writestr(entry, data)
        item = {"path": str(path), "sha256": digest(path), "bytes": path.stat().st_size,
                "fixture_kind": "synthetic targeted python-docx fixture",
                "expected_action": "reject" if unsupported else "roundtrip"}
        if rows is not None:
            item["expected"] = rows
            item["source_independent"] = independent_snapshot(path)
            verify_independent(item["source_independent"], rows, name + " source")
        if unsupported:
            item["unsupported_form"] = unsupported
            item["source_python_docx_paragraphs"] = [p.text for p in Document(path).paragraphs]
        if render_probe:
            item["render_probe"] = render_probe
        sources[name] = item

    doc = Document()
    doc.add_paragraph("UntouchedPythonDocxTemplate")
    save("untouched-python-docx-template", doc, [expected("UntouchedPythonDocxTemplate", None, 200, 276)])

    doc = fresh_document()
    ppr = doc.styles.element.find("w:docDefaults/w:pPrDefault/w:pPr", NS)
    remove_children(ppr, "spacing")
    doc.add_paragraph("AbsentSpacingDefaults")
    save("absent-spacing-defaults", doc, [expected("AbsentSpacingDefaults", None, None, None)])

    doc = fresh_document()
    ppr = doc.add_paragraph("ImplicitAutoLineRule")._p.get_or_add_pPr()
    set_xml_spacing(ppr, before=0, after=0, line=360)
    save("implicit-auto-line-rule", doc, [expected("ImplicitAutoLineRule", 0, 0, 360)])

    doc = fresh_document()
    doc.add_paragraph("DirectSpacing")
    set_spacing(doc.paragraphs[-1], 120, 240, 360)
    doc.add_heading("HeadingSpacing", level=2)
    set_spacing(doc.paragraphs[-1], 80, 160, 480)
    doc.add_paragraph("ExplicitZero")
    set_spacing(doc.paragraphs[-1], 0, 0, 240)
    save("direct-heading-zero", doc, [expected("DirectSpacing", 120, 240, 360),
         expected("HeadingSpacing", 80, 160, 480, "heading"), expected("ExplicitZero", 0, 0, 240)])

    doc = fresh_document((60, 100, 264))
    set_spacing(doc.styles["Normal"], after=180)
    parent = doc.styles.add_style("SpacingParent", WD_STYLE_TYPE.PARAGRAPH)
    parent.base_style = doc.styles["Normal"]
    set_spacing(parent, before=120)
    child = doc.styles.add_style("SpacingChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    set_spacing(child, line=360)
    doc.add_paragraph("NormalCascade")
    doc.add_paragraph("ParentCascade", style=parent)
    doc.add_paragraph("ChildCascade", style=child)
    p = doc.add_paragraph("DirectPartialZero", style=child)
    set_spacing(p, before=0, after=0)
    save("partial-style-cascade", doc, [expected("NormalCascade", 60, 180, 264),
         expected("ParentCascade", 120, 180, 264), expected("ChildCascade", 120, 180, 360),
         expected("DirectPartialZero", 0, 0, 360)])

    doc = fresh_document((40, 80, 300))
    p = doc.add_paragraph("BulletSpacing", style="List Bullet")
    set_spacing(p, before=100, line=360)
    p = doc.add_paragraph("NumberSpacing", style="List Number")
    set_spacing(p, after=0)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "CellSpacing"
    set_spacing(table.cell(0, 0).paragraphs[0], 0, 120, 480)
    table.cell(0, 1).text = "CellInherited"
    save("lists-table-cells", doc, [expected("BulletSpacing", 100, 80, 360, ancestors=["bulletList", "listItem"]),
         expected("NumberSpacing", 40, 0, 300, ancestors=["orderedList", "listItem"]),
         expected("CellSpacing", 0, 120, 480, ancestors=["table", "tableRow", "tableCell"]),
         expected("CellInherited", 40, 80, 300, ancestors=["table", "tableRow", "tableCell"])])

    doc = fresh_document()
    set_spacing(doc.styles["Normal"], before=120, after=180, line=360)
    doc.add_paragraph("EmptyBoundaryStart")
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("EmptyBoundaryEnd")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).add_paragraph()
    rows = [expected(text, 120, 180, 360)
            for text in ("EmptyBoundaryStart", "", "", "EmptyBoundaryEnd")]
    rows.extend(expected("", 120, 180, 360, ancestors=["table", "tableRow", "tableCell"])
                for _ in range(3))
    save("empty-body-and-cell-paragraphs", doc, rows)
    empty_source = sources["empty-body-and-cell-paragraphs"]
    with zipfile.ZipFile(empty_source["path"]) as package:
        self_closing_count = package.read("word/document.xml").count(b"<w:p/>")
    check(self_closing_count == 5, "Empty paragraph source must contain five self-closing w:p elements")
    empty_source["source_self_closing_paragraphs"] = self_closing_count

    doc = fresh_document()
    for marker, values in (("MinimumValues", (0, 0, 120)), ("MaximumValues", (31680, 31680, 2400))):
        set_spacing(doc.add_paragraph(marker), *values)
    save("supported-boundaries", doc, [expected("MinimumValues", 0, 0, 120), expected("MaximumValues", 31680, 31680, 2400)])

    doc = fresh_document()
    ppr = doc.add_paragraph("InactiveFlags")._p.get_or_add_pPr()
    set_xml_spacing(ppr, before=0, after=0, line=240, lineRule="auto", beforeAutospacing="false", afterAutospacing="off")
    xml_child(ppr, "contextualSpacing", val="0")
    xml_child(ppr, "snapToGrid", val="false")
    xml_child(doc.sections[0]._sectPr, "docGrid", type="default", linePitch=360)
    save("inactive-flags", doc, [expected("InactiveFlags", 0, 0, 240)])

    for value in (240, 360, 480):
        doc = fresh_document()
        p = doc.add_paragraph("SpacingLineA\nSpacingLineB\nSpacingLineC")
        set_spacing(p, 0, 0, value)
        save(f"render-line-{value}", doc, [expected(p.text, 0, 0, value)], "" , "line")
    for name, after, before in (("zero", 0, 0), ("before", 0, 240), ("after", 240, 0), ("both", 240, 240)):
        doc = fresh_document()
        p = doc.add_paragraph("SpacingParagraphA")
        set_spacing(p, 0, after, 240)
        p = doc.add_paragraph("SpacingParagraphB")
        set_spacing(p, before, 0, 240)
        save("render-between-" + name, doc,
             [expected("SpacingParagraphA", 0, after, 240), expected("SpacingParagraphB", before, 0, 240)],
             render_probe="between")

    invalid_spacing = {
        "line-exact": {"line": 240, "lineRule": "exact"},
        "line-at-least": {"line": 240, "lineRule": "atLeast"},
        "before-auto": {"beforeAutospacing": "1"},
        "after-auto": {"afterAutospacing": "true"},
        "before-lines-zero": {"beforeLines": "0"},
        "after-lines": {"afterLines": "100"},
        "negative-before": {"before": "-1"},
        "after-over-limit": {"after": "31681"},
        "line-below-limit": {"line": "119"},
        "line-over-limit": {"line": "2401"},
        "line-malformed": {"line": "1.5"},
        "before-malformed": {"before": "12pt"},
        "line-rule-unknown": {"line": "240", "lineRule": "mystery"},
        "line-rule-without-line": {"lineRule": "auto"},
        "auto-unknown": {"beforeAutospacing": "mystery"},
    }
    for name, attrs in invalid_spacing.items():
        doc = fresh_document()
        set_xml_spacing(doc.add_paragraph("UnsupportedProbe")._p.get_or_add_pPr(), **attrs)
        save("reject-" + name, doc, unsupported=attrs)
    for location in ("direct", "unused-style", "used-style", "defaults", "table"):
        doc = fresh_document()
        if location in ("unused-style", "used-style"):
            style = doc.styles.add_style("UnusedUnsupportedSpacing", WD_STYLE_TYPE.PARAGRAPH)
            ppr = style.element.get_or_add_pPr()
            doc.add_paragraph("UnsupportedProbe", style=style if location == "used-style" else None)
        elif location == "defaults":
            ppr = doc.styles.element.find("w:docDefaults/w:pPrDefault/w:pPr", NS)
            doc.add_paragraph("UnsupportedProbe")
        elif location == "table":
            ppr = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]._p.get_or_add_pPr()
        else:
            ppr = doc.add_paragraph("UnsupportedProbe")._p.get_or_add_pPr()
        xml_child(ppr, "contextualSpacing")
        if location == "unused-style":
            save("unused-contextual-style", doc, [expected("UnsupportedProbe", 0, 0, 240)])
        else:
            save("reject-contextual-" + location, doc, unsupported="active contextualSpacing in " + location)
    for name, attrs in (("missing-type", {"linePitch": 360}), ("lines", {"type": "lines", "linePitch": 360}),
                        ("unknown", {"type": "mystery", "linePitch": 360})):
        doc = fresh_document()
        doc.add_paragraph("UnsupportedProbe")
        xml_child(doc.sections[0]._sectPr, "docGrid", **attrs)
        if name == "missing-type":
            save("grid-missing-type-default", doc, [expected("UnsupportedProbe", 0, 0, 240)])
        else:
            save("reject-grid-" + name, doc, unsupported={"docGrid": attrs})
    for used in (False, True):
        doc = fresh_document()
        doc.add_paragraph("NumberingLevelProbe", style="List Bullet" if used else None)
        num_id = doc.styles["List Bullet"].element.find("w:pPr/w:numPr/w:numId", NS).get(qn("w:val"))
        numbering = doc.part.numbering_part.element
        num = numbering.find(f"w:num[@w:numId='{num_id}']", NS)
        abstract_id = num.find("w:abstractNumId", NS).get(qn("w:val"))
        level = numbering.find(f"w:abstractNum[@w:abstractNumId='{abstract_id}']/w:lvl", NS)
        ppr = level.find("w:pPr", NS)
        if ppr is None:
            ppr = xml_child(level, "pPr")
        set_xml_spacing(ppr, before=120)
        if used:
            save("reject-used-numbering-level-spacing", doc, unsupported="spacing inherited from an actively used numbering level")
        else:
            save("unused-numbering-level-spacing", doc, [expected("NumberingLevelProbe", 0, 0, 240)])

    doc = fresh_document()
    set_spacing(doc.add_paragraph("AlternatePrefixProbe"), before=120)
    old = b'<w:spacing w:before="120"'
    new = ('<w:spacing xmlns:alt="' + W + '" alt:before="120"').encode()
    def alternate_prefix(data):
        check(data.count(old) == 1, "Alternate-prefix fixture target is ambiguous")
        return data.replace(old, new)
    save("reject-alternate-prefix-attribute", doc, unsupported="namespace-equivalent alt:before cannot be consumed by the canonical-prefix parser",
         xml_transform=alternate_prefix)

    manifest = directory / "expected-spacing-manifest.json"
    manifest.write_text(json.dumps({name: {k: v for k, v in item.items() if k not in ("source_independent",)}
                                    for name, item in sources.items()}, indent=2), encoding="utf-8")
    return sources, manifest


def run_cli(cli, log, *args, success=True):
    result = subprocess.run([str(cli), *map(str, args)], text=True, capture_output=True,
                            encoding="utf-8", timeout=90)
    log.append({"arguments": list(map(str, args)), "returncode": result.returncode,
                "stdout": result.stdout, "stderr": result.stderr})
    check((result.returncode == 0) is success, result.stdout + result.stderr)
    return json.loads(result.stdout) if success else result.stderr


def model_paragraphs(node, ancestors=()):
    if node["type"] in ("paragraph", "heading"):
        def text_value(item):
            if item["type"] == "hardBreak":
                return "\n"
            return item.get("text", "") + "".join(text_value(n) for n in item.get("content", []))
        yield {"text": text_value(node), "type": node["type"],
               "spacing": {key: node.get("attrs", {}).get(key) for key in KEYS},
               "ancestors": list(ancestors)}
    for child in node.get("content", []):
        yield from model_paragraphs(child, ancestors + (node["type"],))


def verify_model(model, targets, label):
    rows = list(model_paragraphs(model))
    check(len(rows) == len(targets), f"{label}: model paragraph count differs: {rows}")
    for row, target in zip(rows, targets):
        check(all(value is not None for value in row["spacing"].values()),
              f"{label}: imported effective spacing was not materialized: {row}")
        row["spacing"] = semantic_spacing(row["spacing"])
        target = dict(target, spacing=semantic_spacing(target["spacing"]))
        for key in ("text", "type", "spacing"):
            check(row[key] == target[key], f"{label}: {key} differs: {row} != {target}")
        for ancestor in target.get("ancestors", []):
            check(ancestor in row["ancestors"], f"{label}: missing {ancestor}: {row}")
    return rows


def render_pdf(soffice, directory, path, probe):
    from pypdf import PdfReader
    out = directory / "libreoffice-pdf"
    out.mkdir(exist_ok=True)
    profile = directory / "lo-profile"
    result = subprocess.run([str(soffice), f"-env:UserInstallation={profile.as_uri()}",
                             "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)],
                            text=True, capture_output=True, timeout=120,
                            env=dict(os.environ, XDG_CACHE_HOME=str(directory / "lo-cache")))
    pdf = out / (path.stem + ".pdf")
    check(result.returncode == 0 and pdf.is_file(), result.stdout + result.stderr)
    reader = PdfReader(pdf)
    positions = {}
    texts = []
    markers = ("SpacingLineA", "SpacingLineB", "SpacingLineC") if probe == "line" else ("SpacingParagraphA", "SpacingParagraphB")
    for page_number, page in enumerate(reader.pages, 1):
        def visitor(text, cm, tm, font, size):
            found = [marker for marker in markers if marker in text]
            check(len(found) <= 1, f"PDF combined markers into one position: {text!r}")
            for marker in found:
                check(marker not in positions, f"PDF duplicated marker: {marker}")
                positions[marker] = {"page": page_number,
                    "x": tm[4] * cm[0] + tm[5] * cm[2] + cm[4],
                    "y": tm[4] * cm[1] + tm[5] * cm[3] + cm[5], "font_size": size}
        texts.append(page.extract_text(visitor_text=visitor) or "")
    check(len(reader.pages) == 1, "Controlled PDF probe unexpectedly spans pages")
    check(all(marker in positions for marker in markers), f"Missing PDF baseline marker: {positions}")
    distances = [abs(positions[a]["y"] - positions[b]["y"]) for a, b in zip(markers, markers[1:])]
    check(all(distance > 0 for distance in distances), "PDF baselines overlap")
    return {"consumer": "LibreOffice PDF export, not app print", "path": str(pdf), "sha256": digest(pdf),
            "pages": len(reader.pages), "text": texts, "positions": positions,
            "baseline_distances_pt": distances, "stdout": result.stdout, "stderr": result.stderr}


def verify_render_probes(sources, stage):
    values = {name: item[stage]["baseline_distances_pt"] for name, item in sources.items() if stage in item}
    single = values["render-line-240"]
    check(abs(single[0] - single[1]) < 0.15, f"{stage}: single baselines differ")
    for value, ratio in ((360, 1.5), (480, 2)):
        distances = values[f"render-line-{value}"]
        check(all(abs(distance / single[0] - ratio) < 0.04 for distance in distances),
              f"{stage}: line {value} baseline ratio mismatch: {distances} / {single}")
    base = values["render-between-zero"][0]
    for name in ("before", "after", "both"):
        distance = values["render-between-" + name][0]
        check(abs(distance - base - 12) < 0.3,
              f"{stage}: paragraph {name} delta is not 12pt: {distance} - {base}")
    return {"line_multiples_verified": [1, 1.5, 2], "paragraph_delta_pt": 12,
            "adjacent_before_after_use_max_not_sum": True}


def check_corpus(path, cli, evidence, directory):
    manifest = json.loads(Path(__file__).with_name("corpus.json").read_text(encoding="utf-8"))
    target = next(item for item in manifest["fixtures"] if item["name"] == path.name)
    check(path.stat().st_size == target["bytes"] and digest(path) == target["sha256"], "Pinned corpus hash/size mismatch")
    result = {"fixture_kind": "existing pinned Apache POI upstream regression fixture; not user production corpus",
              "path": str(path), "sha256": digest(path), "repository": manifest["repository"],
              "revision": manifest["revision"],
              "comparison_scope": "Spacing on text paragraphs plus explicit page-break count. The existing importer represents a page-break-only paragraph as a pageBreak node.",
              "source_url": f"https://github.com/{manifest['repository']}/blob/{manifest['revision']}/{manifest['directory']}/{path.name}",
              "independent": independent_snapshot(path)}
    if cli:
        audit = run_cli(cli, evidence["commands"], "audit-docx", path)
        result["audit"] = audit
        check(RISK not in audit["risks"], "Pinned POI positive corpus was rejected for spacing")
        targets = [{"text": row["text"], "type": "paragraph", "spacing": row["spacing"]}
                   for row in result["independent"]["ooxml"] if not row["page_break_only"]]
        imported = run_cli(cli, evidence["commands"], "import-docx", path)["document"]
        result["imported_spacing"] = verify_model(imported, targets, "pinned POI imported")
        model = directory / "upstream-poi-model.json"
        model.write_text(json.dumps(imported, indent=2), encoding="utf-8")
        output = directory / "upstream-poi-export.docx"
        run_cli(cli, evidence["commands"], "create-docx", model, output)
        result.update(export_path=str(output), export_sha256=digest(output))
        result["export_independent"] = independent_snapshot(output)
        filtered_export = {key: [row for index, row in enumerate(rows)
                                  if not result["export_independent"]["ooxml"][index]["page_break_only"]]
                           for key, rows in result["export_independent"].items()}
        verify_independent(filtered_export, targets, "pinned POI export")
        source_breaks = sum(row["explicit_page_breaks"] for row in result["independent"]["ooxml"])
        export_breaks = sum(row["explicit_page_breaks"] for row in result["export_independent"]["ooxml"])
        check(source_breaks == export_breaks == 1, "Pinned POI manual page break changed")
        result["source_and_export_manual_page_breaks"] = source_breaks
        reopened = run_cli(cli, evidence["commands"], "read-docx-model", output)
        result["authored_reopened_spacing"] = verify_model(reopened, targets, "pinned POI trusted authored reopen")
        result["full_model_equal_after_authored_reopen"] = imported == reopened
        check(result["full_model_equal_after_authored_reopen"], "Pinned POI trusted authored reopen changed the model")
        reimported = run_cli(cli, evidence["commands"], "import-docx", output)["document"]
        result["external_reimported_spacing"] = verify_model(reimported, targets, "pinned POI external reimport")
        result["full_model_equal_after_external_reimport"] = imported == reimported
        check(result["imported_spacing"] == result["external_reimported_spacing"], "Pinned POI external reimport changed spacing/structure")
        for existing, original_hash in ((path, target["sha256"]), (output, result["export_sha256"])):
            failure = run_cli(cli, evidence["commands"], "create-docx", model, existing, success=False)
            check("Output already exists" in failure and digest(existing) == original_hash,
                  "Pinned corpus no-clobber check failed")
        result["source_and_export_overwrite_refused"] = True
        result["verified_outcome"] = "independent source/import/export/trusted-authored-reopen roundtrip plus separate external reimport"
    check(digest(path) == target["sha256"], "Pinned upstream source changed")
    result["source_unchanged"] = True
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli", type=Path, default=Path("target/debug/rust-wps"))
    parser.add_argument("--root", type=Path, default=Path("tmp/word-paragraph-spacing"))
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--soffice")
    parser.add_argument("--corpus", type=Path)
    args = parser.parse_args()
    cli = None if args.generate_only else args.cli.resolve()
    check(cli is None or cli.is_file(), f"CLI missing: {cli}")
    soffice = args.soffice or shutil.which("soffice")
    check(not args.render or soffice, "LibreOffice is required for --render")
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.root)).resolve()
    report = directory / "report.json"
    evidence = {"status": "running", "script_sha256": digest(Path(__file__)),
                "python_version": platform.python_version(), "python_docx_version": version("python-docx"),
                "lxml_version": version("lxml"), "cli": None if cli is None else str(cli),
                "cli_sha256": None if cli is None else digest(cli), "commands": [], "sources": {},
                "authored_reopen_command": "read-docx-model",
                "external_reimport_command": "import-docx",
                "limitations": "Controlled synthetic probes isolate explicit declarations; an untouched python-docx template is also checked. "
                "The optional Apache POI file is separately pinned upstream regression evidence. "
                "LibreOffice PDF distances are controlled external-consumer probes, not app printing, "
                "native WebView2 behavior, complete layout/pagination fidelity, or Microsoft Word/WPS validation."}
    try:
        sources, manifest = make_fixtures(directory)
        evidence.update(sources=sources, expected_manifest=str(manifest), expected_manifest_sha256=digest(manifest))
        if args.render:
            evidence["pypdf_version"] = version("pypdf")
            evidence["libreoffice_version"] = subprocess.run([soffice, "--version"], capture_output=True,
                                                             text=True, check=True, timeout=30).stdout.strip()
        safe_model = None
        for name, item in sources.items():
            path = Path(item["path"])
            if cli:
                item["audit"] = run_cli(cli, evidence["commands"], "audit-docx", path)
                if item["expected_action"] == "reject":
                    check(RISK in item["audit"]["risks"], f"{name}: missing spacing diagnostic")
                    item["import_error"] = run_cli(cli, evidence["commands"], "import-docx", path, success=False)
                    check(RISK in item["import_error"], f"{name}: wrong import refusal")
                else:
                    check(RISK not in item["audit"]["risks"], f"{name}: supported spacing rejected")
                    imported = run_cli(cli, evidence["commands"], "import-docx", path)["document"]
                    item["imported_spacing"] = verify_model(imported, item["expected"], name + " imported")
                    model = directory / (name + ".json")
                    model.write_text(json.dumps(imported, ensure_ascii=False, indent=2), encoding="utf-8")
                    if safe_model is None:
                        safe_model = model
                    output = directory / (name + "-export.docx")
                    run_cli(cli, evidence["commands"], "create-docx", model, output)
                    item.update(export_path=str(output), export_sha256=digest(output))
                    item["export_independent"] = independent_snapshot(output)
                    verify_independent(item["export_independent"], item["expected"], name + " export", require_direct=True)
                    reopened = run_cli(cli, evidence["commands"], "read-docx-model", output)
                    item["authored_reopened_spacing"] = verify_model(reopened, item["expected"], name + " trusted authored reopen")
                    item["full_model_equal_after_authored_reopen"] = imported == reopened
                    check(item["full_model_equal_after_authored_reopen"], f"{name}: trusted authored reopen changed the model")
                    reimported = run_cli(cli, evidence["commands"], "import-docx", output)["document"]
                    item["external_reimported_spacing"] = verify_model(reimported, item["expected"], name + " external reimport")
                    item["full_model_equal_after_external_reimport"] = imported == reimported
                    item["spacing_structure_equal_after_external_reimport"] = item["imported_spacing"] == item["external_reimported_spacing"]
                    check(item["spacing_structure_equal_after_external_reimport"], f"{name}: external reimport changed spacing/structure")
                    failure = run_cli(cli, evidence["commands"], "create-docx", model, output, success=False)
                    check("Output already exists" in failure and digest(output) == item["export_sha256"], "Export no-clobber failed")
                    item["existing_export_overwrite_refused"] = True
                    if args.render and "render_probe" in item:
                        item["export_render"] = render_pdf(soffice, directory, output, item["render_probe"])
            if args.render and "render_probe" in item:
                item["source_render"] = render_pdf(soffice, directory, path, item["render_probe"])
                if "export_render" in item:
                    a, b = item["source_render"]["baseline_distances_pt"], item["export_render"]["baseline_distances_pt"]
                    check(all(abs(x - y) < 0.3 for x, y in zip(a, b)), f"{name}: source/export baseline mismatch: {a} != {b}")
                    item["source_export_baselines_match"] = True
            check(digest(path) == item["sha256"], f"Source changed: {name}")
            item["source_unchanged"] = True
        if cli:
            for item in sources.values():
                path = Path(item["path"])
                failure = run_cli(cli, evidence["commands"], "create-docx", safe_model, path, success=False)
                check("Output already exists" in failure and digest(path) == item["sha256"], "Source no-clobber failed")
                item["source_overwrite_refused"] = True
        if args.render:
            evidence["source_pdf_assertions"] = verify_render_probes(sources, "source_render")
            if cli:
                evidence["export_pdf_assertions"] = verify_render_probes(sources, "export_render")
        if args.corpus:
            evidence["upstream_corpus"] = check_corpus(args.corpus.resolve(), cli, evidence, directory)
        evidence["status"] = "fixtures-verified" if args.generate_only else "passed"
        evidence["summary"] = {"supported_fixtures": sum(s["expected_action"] == "roundtrip" for s in sources.values()),
                               "rejection_fixtures": sum(s["expected_action"] == "reject" for s in sources.values()),
                               "rendered_sources": sum("source_render" in s for s in sources.values()),
                               "rendered_exports": sum("export_render" in s for s in sources.values())}
    except Exception as error:
        evidence["status"] = "failed"
        evidence["error"] = f"{type(error).__name__}: {error}"
        raise
    finally:
        # Always report failures and source changes, including failed CLI checks.
        evidence["script_sha256_on_exit"] = digest(Path(__file__))
        evidence["script_unchanged_during_run"] = evidence["script_sha256_on_exit"] == evidence["script_sha256"]
        evidence["cli_sha256_on_exit"] = None if cli is None else digest(cli)
        evidence["cli_unchanged_during_run"] = evidence["cli_sha256_on_exit"] == evidence["cli_sha256"]
        evidence["source_hashes_unchanged_on_exit"] = all(Path(s["path"]).is_file() and digest(Path(s["path"])) == s["sha256"]
                                                          for s in evidence["sources"].values())
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
        print(report)


if __name__ == "__main__":
    main()
