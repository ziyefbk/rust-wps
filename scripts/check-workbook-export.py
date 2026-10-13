"""Independent openpyxl verification of the real Rust XLSX writer."""
import hashlib
import json
import subprocess
import tempfile
import warnings
import zipfile
from pathlib import Path
import openpyxl
warnings.filterwarnings("error", category=UserWarning, module=r"openpyxl\..*")

root = Path(__file__).resolve().parents[1]
(root / "tmp" / "spreadsheets").mkdir(parents=True, exist_ok=True)
directory = Path(tempfile.mkdtemp(prefix="export-", dir=root / "tmp" / "spreadsheets"))
model = {"version": 1, "sheets": [
    {"name": "中文数据", "hidden": False, "cells": [
        {"row": 0, "column": 0, "content": {"type": "text", "value": " 中文<&>😀\n第二行 "}},
        {"row": 0, "column": 1, "content": {"type": "number", "value": 12.5}, "number_format": "0.00"},
        {"row": 0, "column": 2, "content": {"type": "boolean", "value": True}},
        {"row": 0, "column": 3, "content": {"type": "formula", "value": "B1*2"}},
        {"row": 1, "column": 0, "content": {"type": "empty"}, "number_format": "0.000"},
        {"row": 1048575, "column": 16383, "content": {"type": "text", "value": "最远单元格"}},
    ]},
    {"name": "隐藏", "hidden": True, "cells": []},
]}
source = directory / "model.json"
source.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
before = source.read_bytes()
output = directory / "新工作簿.xlsx"
exe = root / "target" / "debug" / "rust-wps.exe"
command = [str(exe), "create-xlsx", str(source), str(output)]
result = subprocess.run(command, capture_output=True, check=True, encoding="utf-8")
report = json.loads(result.stdout)
book = openpyxl.load_workbook(output, data_only=False)
assert book.sheetnames == ["中文数据", "隐藏"]
sheet = book["中文数据"]
assert sheet["A1"].value == " 中文<&>😀\n第二行 "
assert sheet["B1"].value == 12.5 and sheet["B1"].number_format == "0.00"
assert sheet["C1"].value is True and sheet["C1"].data_type == "b"
assert sheet["D1"].value == "=B1*2" and sheet["D1"].data_type == "f"
assert sheet["A2"].value is None and sheet["A2"].number_format == "0.000"
assert sheet["XFD1048576"].value == "最远单元格"
assert book["隐藏"].sheet_state == "hidden"
cached = openpyxl.load_workbook(output, data_only=True)
assert cached["中文数据"]["D1"].value is None
assert not report["recalculated"]
assert source.read_bytes() == before
output_hash = hashlib.sha256(output.read_bytes()).hexdigest()
assert subprocess.run(command, capture_output=True).returncode != 0
assert hashlib.sha256(output.read_bytes()).hexdigest() == output_hash
reopened = json.loads(subprocess.run([str(exe), "read-xlsx-model", str(output)], capture_output=True, check=True, encoding="utf-8").stdout)
assert reopened["sheets"][0]["cells"][1]["content"] == {"type": "number", "value": 12.5}
reopened["sheets"][0]["cells"][1]["content"]["value"] = 24.75
edited_model = directory / "edited-model.json"
edited_model.write_text(json.dumps(reopened, ensure_ascii=False), encoding="utf-8")
edited_output = directory / "edited.xlsx"
subprocess.run([str(exe), "create-xlsx", str(edited_model), str(edited_output)], capture_output=True, check=True)
edited = openpyxl.load_workbook(edited_output, data_only=False)
assert edited["中文数据"]["B1"].value == 24.75
assert edited["中文数据"]["B1"].number_format == "0.00"
assert edited["中文数据"]["D1"].value == "=B1*2"
assert openpyxl.load_workbook(edited_output, data_only=True)["中文数据"]["D1"].value is None
assert hashlib.sha256(output.read_bytes()).hexdigest() == output_hash
variants = {}
for variant in ("stale_model", "stale_sheet", "external"):
    variant_path = directory / f"{variant}.xlsx"
    with zipfile.ZipFile(output) as original, zipfile.ZipFile(variant_path, "w", zipfile.ZIP_DEFLATED) as changed:
        for name in original.namelist():
            data = original.read(name)
            if name == "rust-wps/workbook-v1.json":
                if variant == "external":
                    continue
                if variant == "stale_model":
                    edited = json.loads(data)
                    edited["sheets"][0]["cells"][0]["content"]["value"] = "stale model"
                    data = json.dumps(edited).encode()
            if name == "xl/worksheets/sheet1.xml" and variant == "stale_sheet":
                assert b"B1*2" in data
                data = data.replace(b"B1*2", b"B1*3")
            changed.writestr(name, data)
    variants[variant] = str(variant_path)
print(json.dumps({"independentReader": "openpyxl", "typedValues": True, "formulaWithoutInventedCache": True,
                  "numberFormats": True, "sparseFarCell": True, "sourcePreserved": True,
                  "overwriteRejected": True, "verifiedModelEditAndResave": True, "output": str(output), "variants": variants}, ensure_ascii=True, indent=2))
