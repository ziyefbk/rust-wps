#!/usr/bin/env python3
"""Independent consumer calibration; no claim that cached colors are an oracle.

One python-docx/OOXML authored matrix and retained-source single-attribute
controls distinguish cached w:val consumption from theme transform evaluation.
Never imports production implementation. No original or prior evidence changed.
"""
import argparse
import colorsys
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from fractions import Fraction

from docx import Document
from docx.shared import Pt, Twips
from lxml import etree as E
import fitz
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parent.parent
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
MIN_FREE = 650 * 1024 * 1024
CLI = ROOT / 'target/debug/rust-wps'
CLI_SHA = '35d4b018e67a032c619938c62765fbeb1eaa6e09126240cbdc8448dcfa75bf47'
FROZEN = ROOT / 'tmp/paragraph-pagination-independent/resumed-u44hswfl/report.json'
FROZEN_SHA = 'e65ab364d52c77e1bd51295832d382849d25f40f2391d68f4917755c3d7ae58e'
SOFFICE = '/opt/codex/runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
CASES = [
    ('blue_shade_cached_genuine', 'accent1', 'BF', None, '365F91'),
    ('blue_shade_cached_green', 'accent1', 'BF', None, '00FF00'),
    ('blue_shade_cached_magenta', 'accent1', 'BF', None, 'FF00FF'),
    ('blue_shade_cached_absent', 'accent1', 'BF', None, None),
    ('ms_red_shade_cached_example', 'accent2', 'BF', None, '943634'),
    ('ms_red_shade_cached_green', 'accent2', 'BF', None, '00FF00'),
    ('ms_red_shade_cached_absent', 'accent2', 'BF', None, None),
    ('ms_blue_tint_cached_example', 'accent1', None, '99', '95B3D7'),
    ('ms_blue_tint_cached_green', 'accent1', None, '99', '00FF00'),
    ('ms_blue_tint_cached_absent', 'accent1', None, '99', None),
    ('blue_both_cached_green', 'accent1', 'BF', '99', '00FF00'),
    ('blue_both_cached_absent', 'accent1', 'BF', '99', None),
    ('blue_plain_cached_green', 'accent1', None, None, '00FF00'),
    ('blue_plain_cached_absent', 'accent1', None, None, None),
    ('literal_only', None, None, None, '365F91'),
    ('literal_shade_without_theme', None, 'BF', None, '365F91'),
]
APP_RGB = [
    '376092', '376092', '376092', '376092',
    '953735', '953735', '953735',
    '95B3D7', '95B3D7', '95B3D7', '95B3D7', '95B3D7',
    '4F81BD', '4F81BD', '365F91', '365F91',
]
MS_SOURCE = 'https://learn.microsoft.com/en-us/openspecs/office_standards/ms-oi29500/8229a077-7fc8-4fba-96cc-c77b6a4fc768'
CONSUMER_REV = '2c87e51eeaa2b413ff4ae097b2705eea1995d8e5'
CONSUMER_SOURCE = 'https://github.com/LibreOffice/core/blob/' + CONSUMER_REV + '/sw/source/writerfilter/dmapper/'


def exact_hls(rgb):
    """Generic rational HLS conversion independent of the product helper."""
    r, g, b = rgb
    high, low = max(rgb), min(rgb)
    light = (high + low) / 2
    if high == low:
        return Fraction(0), light, Fraction(0)
    delta = high - low
    saturation = delta / (high + low if light <= Fraction(1, 2) else 2 - high - low)
    if high == r:
        hue = (g - b) / delta
    elif high == g:
        hue = 2 + (b - r) / delta
    else:
        hue = 4 + (r - g) / delta
    return (hue / 6) % 1, light, saturation


def exact_rgb(hue, light, saturation):
    if saturation == 0:
        return (light,) * 3
    upper = light * (1 + saturation) if light <= Fraction(1, 2) else light + saturation - light * saturation
    lower = 2 * light - upper

    def channel(position):
        position %= 1
        if position < Fraction(1, 6):
            return lower + (upper - lower) * position * 6
        if position < Fraction(1, 2):
            return upper
        if position < Fraction(2, 3):
            return lower + (upper - lower) * (Fraction(2, 3) - position) * 6
        return lower

    return tuple(channel(hue + offset) for offset in (Fraction(1, 3), Fraction(0), Fraction(-1, 3)))


def arithmetic_calibration():
    vectors = []
    for base, operation, encoded, literal, provenance in [
        ('4F81BD', 'shade', 'BF', '365F91', 'Literal cached in genuine Apache POI Styles and NoHeadFoot; not independent transform evidence'),
        ('C0504D', 'shade', 'BF', '943634', 'Published MS-OI29500 section 2.1.72 example'),
        ('4F81BD', 'tint', '99', '95B3D7', 'Published MS-OI29500 section 2.1.72 example'),
    ]:
        rgb = tuple(Fraction(int(base[i:i + 2], 16), 255) for i in (0, 2, 4))
        hue, light, saturation = exact_hls(rgb)
        factor = Fraction(int(encoded, 16), 255)
        adjusted = light * factor + (1 - factor if operation == 'tint' else 0)
        raw = tuple(c * 255 for c in exact_rgb(hue, adjusted, saturation))
        float_hls = colorsys.rgb_to_hls(*(float(c) for c in rgb))
        float_light = float_hls[1] * float(factor) + (1 - float(factor) if operation == 'tint' else 0)
        float_rgb = tuple(c * 255 for c in colorsys.hls_to_rgb(float_hls[0], float_light, float_hls[2]))
        check(max(abs(float(a) - b) for a, b in zip(raw, float_rgb)) < 1e-10, 'Exact and colorsys arithmetic disagree')
        floor = ''.join(f'{c.numerator // c.denominator:02X}' for c in raw)
        nearest = ''.join(f'{(c + Fraction(1, 2)).numerator // (c + Fraction(1, 2)).denominator:02X}' for c in raw)
        vectors.append({'base_rgb': base, 'operation': operation, 'encoded': encoded,
                        'factor_exact': str(factor), 'hls_exact': [str(c) for c in (hue, light, saturation)],
                        'hls_colorsys': list(float_hls), 'adjusted_luminance_exact': str(adjusted),
                        'rgb_channels_exact': [str(c) for c in raw], 'rgb_channels_colorsys': list(float_rgb),
                        'floor_hex': floor, 'nearest_half_up_hex': nearest,
                        'literal_example': literal, 'literal_provenance': provenance,
                        'floor_matches_literal': floor == literal, 'nearest_matches_literal': nearest == literal})
    check([(v['floor_hex'], v['nearest_half_up_hex']) for v in vectors]
          == [('366091', '376092'), ('943734', '953735'), ('95B3D7', '95B3D7')], 'Arithmetic calibration changed')
    return {'method': 'Fractions exact rational HLS, independently cross-checked with Python standard-library colorsys; no intermediate quantization',
            'vectors': vectors, 'conclusion': 'Neither final-channel floor nor nearest rounding matches the published MS shade example or the genuine cached shade literal. This is not evidence for changing production quantization.',
            'ms_source': MS_SOURCE}


def check(ok, msg):
    if not ok:
        raise AssertionError(msg)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(path, data):
    with Path(path).open('x', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write('\n')


def guard(directory):
    free = shutil.disk_usage(directory).free
    check(free >= MIN_FREE + 24 * 1024 * 1024, 'Would cross retained 650 MiB floor plus 24 MiB render reservation')
    return free


def authored(path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Twips(1440)
    for name, theme, shade, tint, cached in CASES:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(name)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        color = E.SubElement(r._r.get_or_add_rPr(), W + 'color')
        for key, value in [('val', cached), ('themeColor', theme), ('themeShade', shade), ('themeTint', tint)]:
            if value is not None:
                color.set(W + key, value)
    buffer = io.BytesIO()
    doc.save(buffer)
    with zipfile.ZipFile(buffer) as z:
        theme = E.fromstring(z.read('word/theme/theme1.xml'))
        colors = {slot: theme.find('.//' + A + 'clrScheme/' + A + slot + '/' + A + 'srgbClr').get('val') for slot in ('accent1', 'accent2')}
    check(colors == {'accent1': '4F81BD', 'accent2': 'C0504D'}, 'Independent template theme differs')
    with path.open('xb') as f:
        f.write(buffer.getvalue())
    return {'kind': 'independently authored python-docx and direct OOXML matrix', 'theme': colors,
            'cases': [{'name': n, 'theme': t, 'shade': s, 'tint': i, 'cached': c} for n, t, s, i, c in CASES]}


def derivative(source, target, attribute, old, new):
    with zipfile.ZipFile(source) as z:
        parts = {n: z.read(n) for n in z.namelist()}
        styles = E.fromstring(parts['word/styles.xml'])
        colors = [c for s in styles.findall(W + 'style') if s.get(W + 'styleId') == 'berschrift1'
                  for c in s.findall(W + 'rPr/' + W + 'color')]
        check(len(colors) == 1 and colors[0].get(W + attribute) == old, 'Expected active Heading1 declaration differs')
        original = E.tostring(colors[0])
        # Change exact attribute bytes, avoiding unrelated XML reserialization.
        old_attr = ('w:' + attribute + '="' + old + '"').encode()
        new_attr = ('w:' + attribute + '="' + new + '"').encode()
        fragment = E.tostring(styles.find(".//" + W + "style[@" + W + "styleId='berschrift1']"))
        check(old_attr in fragment, 'Active style attribute bytes absent')
        raw = parts['word/styles.xml']
        start = raw.index(b'w:styleId="berschrift1"')
        end = raw.index(b'</w:style>', start)
        active = raw[start:end]
        check(active.count(old_attr) == 1, 'Active attribute not unique')
        changed = raw[:start] + active.replace(old_attr, new_attr, 1) + raw[end:]
        check(changed.replace(new_attr, old_attr, 1) == raw, 'Derivative changed more than intended attribute')
        with zipfile.ZipFile(target, 'x', compression=zipfile.ZIP_DEFLATED) as out:
            for info in z.infolist():
                out.writestr(info, changed if info.filename == 'word/styles.xml' else parts[info.filename])
    with zipfile.ZipFile(target) as z:
        differences = [n for n, data in parts.items() if data != z.read(n)]
    check(differences == ['word/styles.xml'], 'Unrelated ZIP part changed')
    return {'kind': 'controlled derivative of genuine retained fixture; not another genuine source', 'source': str(source),
            'source_sha256': sha(source), 'only_changed_part': 'word/styles.xml', 'only_changed_attribute': attribute,
            'old': old, 'new': new, 'original_color_element': original.decode(), 'unchanged_other_parts': len(parts) - 1}


def consume(directory, path, label):
    free = guard(directory)
    area = directory / label
    area.mkdir()
    (area / 'home').mkdir()
    profile = area / 'profile'
    command = [SOFFICE, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(area), str(path)]
    before = sha(path)
    proc = subprocess.run(command, capture_output=True, text=True, timeout=120,
                          env=dict(os.environ, HOME=str(area / 'home'), XDG_CACHE_HOME=str(area / 'cache')))
    save(area / 'conversion.json', {'argv': command, 'returncode': proc.returncode, 'stdout': proc.stdout, 'stderr': proc.stderr})
    pdf = area / (path.stem + '.pdf')
    check(proc.returncode == 0 and pdf.exists() and pdf.stat().st_size > 0, label + ': conversion failed')
    check(sha(path) == before, label + ': source changed')
    result = {'path': str(pdf), 'sha256': sha(pdf), 'input_sha256': before, 'command': command,
              'profile': str(profile), 'free_bytes_before': free, 'pages': [], 'spans': []}
    with fitz.open(pdf) as document:
        for i, page in enumerate(document):
            result['pages'].append({'number': i + 1, 'rect': list(page.rect), 'text': page.get_text()})
            for b in page.get_text('dict')['blocks']:
                for line in b.get('lines', []):
                    for span in line['spans']:
                        result['spans'].append({'page': i + 1, **{k: span[k] for k in ('text', 'bbox', 'origin', 'font', 'size', 'flags')},
                                                'rgb': f"{span['color']:06X}"})
            page.get_pixmap(matrix=fitz.Matrix(1, 1), alpha=False).save(str(area / f'page-{i+1}.png'))
    result['free_bytes_after'] = shutil.disk_usage(directory).free
    return result


def authored_expected_model():
    return {'type': 'doc', 'attrs': {'styleProfile': 'standard-v2', 'pageGeometry': {
        'widthTwips': 11906, 'heightTwips': 16838, 'marginTopTwips': 1440, 'marginBottomTwips': 1440,
        'marginLeftTwips': 1440, 'marginRightTwips': 1440}}, 'content': [
            {'type': 'paragraph', 'attrs': {'keepNext': False, 'keepLines': False, 'pageBreakBefore': False,
                'lineSpacing': 240, 'spaceBefore': 0, 'spaceAfter': 0}, 'content': [
                    {'type': 'text', 'text': case[0], 'marks': [
                        {'type': 'fontSize', 'attrs': {'halfPoints': 22}},
                        {'type': 'textColor', 'attrs': {'rgb': rgb}},
                        {'type': 'fontFamily', 'attrs': {'name': 'Arial'}}]}]}
            for case, rgb in zip(CASES, APP_RGB)]}


def checked_models(report, frozen):
    imports = report['baseline_imports']
    check(all(item['returncode'] == 0 for item in imports.values()), 'Baseline import rejected a control')
    responses = {name: json.loads(item['stdout']) for name, item in imports.items()}
    expected = authored_expected_model()
    check(responses['authored']['document'] == expected, 'Complete authored model differs from independent expected model')
    original = responses['genuine-Styles']
    check(original == frozen['sources']['original-Styles']['import_response'], 'Genuine full import response changed from frozen report')
    check(responses['Styles-cache-magenta'] == original, 'Cached-only derivative changed full app import response')
    altered = copy.deepcopy(original)
    heading = altered['document']['content'][0]
    check(heading['type'] == 'heading' and heading['content'][0]['text'] == 'Heading1', 'Heading identity differs')
    marks = [m for m in heading['content'][0]['marks'] if m['type'] == 'textColor']
    check(len(marks) == 1 and marks[0]['attrs']['rgb'] == '376092', 'Expected original Heading1 mark differs')
    marks[0]['attrs']['rgb'] = '254062'
    check(responses['Styles-shade-changed'] == altered, 'Shade-only derivative changed anything beyond expected Heading1 RGB')
    return responses, {'complete_authored_model_equals_independent_expectation': True,
                       'genuine_complete_response_equals_frozen_report': True,
                       'cached_only_complete_response_equals_original': True,
                       'shade_only_complete_response_differs_only_Heading1_RGB': True,
                       'expected_app_colors': dict(zip((case[0] for case in CASES), APP_RGB))}


def export_and_reopen(directory, source, model, label):
    model_path = directory / (label + '-model.json')
    target = directory / (label + '-export.docx')
    save(model_path, model)
    commands = []

    def run(args):
        result = subprocess.run([str(CLI), *map(str, args)], capture_output=True, text=True, timeout=60)
        commands.append({'argv': result.args, 'returncode': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr})
        return result

    result = run(['create-docx', model_path, target])
    check(result.returncode == 0 and target.exists(), label + ': export failed')
    result = run(['read-docx-model', target])
    check(result.returncode == 0 and json.loads(result.stdout) == model, label + ': full trusted reopened model differs')
    with zipfile.ZipFile(target) as package:
        body = E.fromstring(package.read('word/document.xml')).find(W + 'body')
    exported = []
    for paragraph in body.findall(W + 'p'):
        exported.append([{'text': ''.join(r.itertext(W + 't')),
                          'color': (r.find(W + 'rPr/' + W + 'color').get(W + 'val')
                                    if r.find(W + 'rPr/' + W + 'color') is not None else None)}
                         for r in paragraph.findall(W + 'r')])
    expected = [[{'text': r['text'], 'color': next((m['attrs']['rgb'] for m in r.get('marks', []) if m['type'] == 'textColor'), None)}
                 for r in p.get('content', [])] for p in model['content']]
    check(exported == expected, label + ': independent complete paragraph/run text and explicit color XML differs')
    for protected_target in (source, target):
        before = sha(protected_target)
        result = run(['create-docx', model_path, protected_target])
        check(result.returncode != 0 and sha(protected_target) == before, label + ': no-overwrite guard failed')
    return {'model_path': str(model_path), 'model_sha256': sha(model_path), 'path': str(target), 'sha256': sha(target),
            'complete_trusted_reopen_equal': True, 'independent_export_text_color_XML_equal': True,
            'no_overwrite_checks': 2, 'commands': commands}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--arithmetic-only', action='store_true', help='Create a standalone exact/colorsys supplement without CLI execution or rendering')
    args = parser.parse_args()
    parent = ROOT / 'tmp/docx-theme-consumer-calibration'
    parent.mkdir(exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=parent))
    report = {'status': 'running', 'directory': str(directory), 'invocation': sys.argv,
              'started_at_utc': datetime.now(timezone.utc).isoformat(), 'script_sha256': sha(__file__),
              'purpose': 'Discriminate cached literal consumption from actual theme-transform recomputation; no product repair claim',
              'inputs': {}, 'renders': {}, 'baseline_imports': {}}
    protected = {str(CLI): CLI_SHA, str(FROZEN): FROZEN_SHA, str(Path(__file__).resolve()): sha(__file__)}
    try:
        report['arithmetic'] = arithmetic_calibration()
        if args.arithmetic_only:
            report['status'] = 'completed-arithmetic-only'
            return
        check(sha(FROZEN) == FROZEN_SHA and sha(CLI) == CLI_SHA and CLI.stat().st_size == 149270000, 'Frozen baseline differs')
        frozen = json.loads(FROZEN.read_text())
        protected.update(frozen['consumers']['binaries_and_launcher_sha256'])
        retained_initial = ROOT / 'tmp/docx-theme-consumer-calibration/run-ngaiicyu'
        if retained_initial.exists():
            protected[str(retained_initial / 'report.json')] = 'a8d7912aa31d3ffedd245b464721680869bd00084c060d4e7af37a6ad59c8199'
            protected[str(retained_initial / 'initial-checker.py')] = 'a9b8f9a896427d527d389d3ba5ce2212e3d6dc0832e33195652c4b3cce7e5c21'
        for key in ('original-Styles', 'original-NoHeadFoot'):
            protected[frozen['sources'][key]['path']] = frozen['sources'][key]['sha256']
        check(all(sha(p) == h for p, h in protected.items()), 'Protected file differs before run')
        tracked = subprocess.run(['git', 'ls-files', '-z', 'src', 'vendor', 'src-tauri', 'ui', 'Cargo.toml', 'Cargo.lock', 'package.json', 'package-lock.json'],
                                 cwd=ROOT, check=True, capture_output=True).stdout.decode().split('\0')
        product_paths = [ROOT / name for name in tracked if name and (ROOT / name).is_file()]
        product_hashes = {str(p): sha(p) for p in product_paths}
        protected.update(product_hashes)
        report['protected_hashes'] = protected.copy()
        report['unrelated_product_files_protected'] = len(product_hashes)
        report['consumers'] = frozen['consumers']
        report['consumer_source_explanation'] = {
            'revision': CONSUMER_REV,
            'handler': {'url': CONSUMER_SOURCE + 'ThemeColorHandler.hxx#L32-L61', 'git_blob': 'da23f8a55d3a61994ea3ee289bb902eacd194074',
                        'finding': 'mnColor defaults to zero; only the val attribute assigns it. Theme, tint and shade are retained separately.'},
            'mapper': {'url': CONSUMER_SOURCE + 'DomainMapper.cxx#L2555-L2619',
                       'finding': 'Character color is populated directly from mnColor; complex theme metadata is retained separately.'},
            'scope': 'Verified exact source revision reported by the installed consumer; explains measured cached-value behavior, not general Word/WPS behavior.'}
        spec_zip = ROOT / 'tmp/paragraph-indentation-source/ecma-376-part1.zip'
        spec_text = ROOT / 'tmp/paragraph-indentation-source/ecma-376-part1.txt'
        check(sha(spec_zip) == '9d0bcad9cf06054785b03762fcfadbf6bab7e54a5f9d69434e34b7fd464d4129', 'Retained ECMA source ZIP differs')
        protected[str(spec_zip)] = sha(spec_zip)
        protected[str(spec_text)] = sha(spec_text)
        report['normative_sources'] = {'ecma': {
            'url': 'https://ecma-international.org/wp-content/uploads/ECMA-376-1_5th_edition_december_2016.zip',
            'zip_sha256': sha(spec_zip), 'text_sha256': sha(spec_text), 'text_lines': [13380, 13530],
            'section': '17.3.2.6 color', 'rule': 'A supplied themeColor supersedes val; tint supersedes shade.'}, 'ms_oi29500': MS_SOURCE}
        report['free_bytes_start'] = guard(directory)
        probe = directory / 'authored-color-matrix.docx'
        report['inputs']['authored'] = {'path': str(probe), **authored(probe), 'sha256': sha(probe)}
        source = Path(frozen['sources']['original-Styles']['path'])
        report['inputs']['genuine-Styles'] = {'path': str(source), 'sha256': sha(source), 'kind': 'genuine unchanged Apache POI regression fixture'}
        for name, attribute, old, new in [('Styles-cache-magenta', 'val', '365F91', 'FF00FF'), ('Styles-shade-changed', 'themeShade', 'BF', '80')]:
            path = directory / (name + '.docx')
            report['inputs'][name] = {'path': str(path), **derivative(source, path, attribute, old, new), 'sha256': sha(path)}
        for key, item in report['inputs'].items():
            command = [str(CLI), 'import-docx', item['path']]
            p = subprocess.run(command, capture_output=True, text=True, timeout=60)
            report['baseline_imports'][key] = {'argv': command, 'returncode': p.returncode, 'stdout': p.stdout, 'stderr': p.stderr}
            report['renders'][key] = consume(directory, Path(item['path']), key)
            print(key, [(s['text'], s['rgb']) for s in report['renders'][key]['spans']], flush=True)
        matrix = report['renders']['authored']
        check([s['text'] for s in matrix['spans']] == [c[0] for c in CASES], 'Authored full text order or span count differs')
        report['matrix_consumed_colors'] = {s['text']: s['rgb'] for s in matrix['spans']}
        expected_consumer = {case[0]: case[4] or '000000' for case in CASES}
        check(report['matrix_consumed_colors'] == expected_consumer, 'Pinned consumer cache behavior changed')
        report['matrix_consumer_cached_color_assertions'] = len(expected_consumer)
        responses, report['baseline_model_checks'] = checked_models(report, frozen)
        report['exports'] = {key: export_and_reopen(directory, Path(item['path']), responses[key]['document'], key)
                             for key, item in report['inputs'].items()}
        original_spans = report['renders']['genuine-Styles']['spans']
        report['genuine_derivative_comparisons'] = {}
        for key in ('Styles-cache-magenta', 'Styles-shade-changed'):
            spans = report['renders'][key]['spans']
            strip_color = lambda rows: [{k:v for k,v in s.items() if k != 'rgb'} for s in rows]
            same = strip_color(spans) == strip_color(original_spans)
            check(same, key + ': noncolor PDF text/font/geometry drift')
            check(report['renders'][key]['pages'] == report['renders']['genuine-Styles']['pages'], key + ': complete page text/dimensions differ')
            expected_rgb = 'FF00FF' if key == 'Styles-cache-magenta' else '365F91'
            check(spans[0]['rgb'] == expected_rgb, key + ': consumer no longer uses retained cached val')
            original_image = Image.open(directory / 'genuine-Styles/page-1.png').convert('RGB')
            changed_image = Image.open(directory / key / 'page-1.png').convert('RGB')
            check(original_image.size == changed_image.size, key + ': raster dimensions differ')
            difference = ImageChops.difference(original_image, changed_image)
            pixels_equal = difference.getbbox() is None
            check(pixels_equal == (key == 'Styles-shade-changed'), key + ': raster expectation changed')
            report['genuine_derivative_comparisons'][key] = {'complete_text_fonts_geometry_equal': same,
                'complete_page_text_dimensions_equal': True, 'whole_page_pixels_equal': pixels_equal,
                'color_pairs': [{'text': a['text'], 'original': a['rgb'], 'derivative': b['rgb']} for a, b in zip(original_spans, spans)]}
        report['summary'] = {'genuine_sources': 1, 'controlled_genuine_derivatives': 2, 'authored_matrix_rows': 16,
                             'fresh_consumer_PDFs': 4, 'complete_trusted_reopens': 4, 'no_overwrite_checks': 8,
                             'product_code_changed': False, 'theme_rounding_repair_proven': False}
        report['limitations'] = [
            'The two retained original fixtures are protected, but only Styles is freshly rendered here. NoHeadFoot is not counted as a newly consumed source.',
            'Consumer calibration uses LibreOffice for DOCX-to-PDF and PyMuPDF for full text/font/color/position inspection and 72-dpi rasterization. It does not independently establish Word/WPS arithmetic.',
            'Byte factors are applied exactly as /255 for the arithmetic comparison; intermediate integer quantization and Office-version differences remain unresolved.',
            'Application baseline RGB expectations pin existing behavior, not a declaration that every baseline shade result matches Word.',
            'No build, native browser test or product change is performed by this checker.',
        ]
        report['status'] = 'completed-calibration-only'
    except Exception as exc:
        report['status'] = 'failed'
        report['error'] = repr(exc)
        raise
    finally:
        report['protected_hashes'] = protected.copy()
        report['protected_hashes_unchanged'] = all(Path(p).is_file() and sha(p) == h for p, h in protected.items())
        report['inputs_unchanged'] = all(sha(item['path']) == item['sha256'] for item in report['inputs'].values())
        report['free_bytes_finish'] = shutil.disk_usage(directory).free
        report['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
        save(directory / 'report.json', report)
        print('REPORT', directory / 'report.json', sha(directory / 'report.json'), flush=True)
    check(report['protected_hashes_unchanged'] and report['inputs_unchanged'], 'Protected identity changed')


if __name__ == '__main__':
    main()
