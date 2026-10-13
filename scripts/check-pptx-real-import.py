"""Pinned Apache POI source -> Rust editable copy -> independent consumer."""
import hashlib
import json
import subprocess
import tempfile
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
CORE = ROOT / "target/debug/rust-wps.exe"
DESKTOP = ROOT / "src-tauri/target/debug/rust-wps-desktop.exe"
SOURCE = ROOT / "tests/fixtures/apache-poi/sample.pptx"


def run(*args):
    return json.loads(subprocess.check_output([str(CORE), *map(str, args)], timeout=40))


before = SOURCE.read_bytes()
assert hashlib.sha256(before).hexdigest() == "108d7219130a8460298896f45ff777dfe52b8b5527a4b4628b10731ed01d6a8e"
audit = run("import-pptx", SOURCE)
assert not audit["blockers"], audit["blockers"]
assert any("4 个对象" in risk and "固定" in risk for risk in audit["risks"])
worker = json.loads(subprocess.check_output([str(DESKTOP), "--document-worker"], input=json.dumps({"operation": "import_presentation", "path": str(SOURCE)}).encode(), timeout=40))
assert worker["result"] == "presentation_imported" and worker["report"] == audit
model = audit["model"]
original = Presentation(SOURCE)
assert [original.slide_width, original.slide_height] == model["size_emu"]
for slide, converted in zip(original.slides, model["slides"], strict=True):
    assert len(slide.shapes) == len(converted["text_boxes"])
    for shape, box in zip(slide.shapes, converted["text_boxes"], strict=True):
        assert shape.text == box["text"]
        assert [shape.left, shape.top, shape.width, shape.height] == box["bounds_emu"]
model["slides"][0]["text_boxes"][0]["text"] = "真实文稿编辑 <&>"
with tempfile.TemporaryDirectory() as directory:
    folder = Path(directory)
    data, output = folder / "model.json", folder / "edited.pptx"
    data.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
    run("create-pptx", data, output)
    assert run("read-pptx-model", output) == model
    result = Presentation(output)
    assert [result.slide_width, result.slide_height] == model["size_emu"]
    for slide, expected in zip(result.slides, model["slides"], strict=True):
        assert len(slide.shapes) == len(expected["text_boxes"])
        for shape, box in zip(slide.shapes, expected["text_boxes"], strict=True):
            # python-pptx exposes DrawingML soft breaks as VT, whereas the
            # basic model represents both paragraph and soft breaks as LF.
            # Normalize only that documented break representation.
            assert shape.text.replace("\v", "\n") == box["text"]
            assert [shape.left, shape.top, shape.width, shape.height] == box["bounds_emu"]
            run_style = shape.text_frame.paragraphs[0].runs[0].font
            assert run_style.size == box["font_size_hundredths_pt"] * 127
            assert bool(run_style.bold) == box["bold"]
            assert bool(run_style.italic) == box["italic"]
            assert list(run_style.color.rgb) == box["color_rgb"]
            if box.get("alignment"):
                assert int(shape.text_frame.paragraphs[0].alignment) == {"left": 1, "center": 2, "right": 3}[box["alignment"]]
assert SOURCE.read_bytes() == before
print("Pinned real PPTX import/edit/export passed independent text, geometry, style and source checks.")
