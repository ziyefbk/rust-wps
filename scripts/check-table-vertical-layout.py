"""Independent minimum-row-height / cell-vertical-alignment DOCX regression.

Uses an unchanged hash-pinned Apache POI document and explicitly separate owned
python-docx fixtures. Expected values come from those sources, never application
output. LibreOffice/PyMuPDF measure border-relative baselines and row growth.
No browser, native print, Word/WPS equivalence, or full-page fidelity is claimed.

  python scripts/check-table-vertical-layout.py --generate-only --render
  python scripts/check-table-vertical-layout.py --cli FROZEN_CLI --render
"""
import argparse
import copy
import hashlib
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from lxml import etree as ET

ROOT = Path(__file__).resolve().parent.parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
ORIGINAL = ROOT / 'tmp/table-corpus/TestTableCellAlign.docx'
ORIGINAL_SHA = '2968d4a5d5de21dce1cbeae3235e545cdf4bfd7ae508fe0eed3a9b032773ce17'
ORIGINAL_BLOB = 'cf40dd2213a44972cce7f1bccdfd6f9db8933386'
REVISION = 'ae62bb5116b9aee19ebd5834e3a82066132c9f7f'
HEIGHT = 'rowMinHeightTwips'
ALIGN = 'cellVerticalAlign'
TOLERANCE = 0.30


def check(value, message):
    if not value:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def add(parent, tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def remove(parent, tag):
    for node in list(parent.findall('w:' + tag, NS)):
        parent.remove(node)


def model_text(node):
    if node['type'] == 'text':
        return node['text']
    if node['type'] == 'hardBreak':
        return '\n'
    return ('\n' if node['type'] in ('tableCell', 'tableHeader') else '').join(model_text(c) for c in node.get('content', []))


def snapshot(path):
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read('word/document.xml'))
    doc = Document(path)
    raw_tables = root.findall('w:body/w:tbl', NS)
    check(len(raw_tables) == len(doc.tables), 'Independent table counts differ')
    tables = []
    for raw, py in zip(raw_tables, doc.tables):
        rows = []
        for row, pyrow in zip(raw.findall('w:tr', NS), py.rows):
            declarations = row.findall('w:trPr/w:trHeight', NS)
            height = None if not declarations else int(declarations[0].get(qn('w:val')))
            rule = None if not declarations else declarations[0].get(qn('w:hRule'))
            check(height == (None if pyrow.height is None else pyrow.height.twips), 'python-docx/raw height differs')
            cells = []
            for cell, pycell in zip(row.findall('w:tc', NS), pyrow.cells):
                values = cell.findall('w:tcPr/w:vAlign', NS)
                align = None if not values else values[0].get(qn('w:val'))
                pyalign = None if pycell.vertical_alignment is None else {0: 'top', 1: 'center', 3: 'bottom'}[int(pycell.vertical_alignment)]
                check(align == pyalign, 'python-docx/raw cell alignment differs')
                text = '\n'.join(''.join(('\n' if ET.QName(n).localname == 'br' else n.text or '') for n in p.iter() if ET.QName(n).localname in ('t', 'br')) for p in cell.findall('w:p', NS))
                check(text == pycell.text, 'python-docx/raw cell text differs')
                cells.append({'align': align, 'text': text})
            rows.append({'height': height, 'rule': rule, 'cells': cells})
        tables.append({'rows': rows, 'grid': [int(n.get(qn('w:w'))) for n in raw.findall('w:tblGrid/w:gridCol', NS)]})
    return tables


def author(heights=(1340, 1160), alignments=((None, 'bottom'), ('center', 'top')), rule=None, long=False, selected_style=False, padding=None, inactive_template=False):
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Twips(12240), Twips(15840)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Twips(1440)
        remove(section._sectPr, 'docGrid')
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'DejaVu Sans', Pt(12)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for node in list(doc.styles.element.findall('.//w:contextualSpacing', NS)):
        node.getparent().remove(node)
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    props = table._tbl.tblPr
    remove(props, 'tblW')
    add(props, 'tblW', w=6000, type='dxa')
    borders = add(props, 'tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        add(borders, edge, val='single', sz=4, color='000000')
    if selected_style or padding in ('style', 'conditional'):
        style = doc.styles.add_style('VerticalBase', WD_STYLE_TYPE.TABLE)
        style.base_style = doc.styles['Normal Table']
        derived = doc.styles.add_style('VerticalDerived', WD_STYLE_TYPE.TABLE)
        derived.base_style = style
        add(props, 'tblStyle', val=derived.style_id)
        if padding in ('style', 'conditional'):
            container = add(style.element, 'tblPr') if padding == 'style' else add(add(style.element, 'tblStylePr', type='firstRow'), 'tcPr')
            margins = add(container, 'tblCellMar' if padding == 'style' else 'tcMar')
            for edge in ('top', 'bottom'):
                add(margins, edge, w=0, type='dxa')
    if padding in ('table', 'table-top-only'):
        margins = add(props, 'tblCellMar')
        for edge in (('top',) if padding == 'table-top-only' else ('top', 'bottom')):
            add(margins, edge, w=0, type='dxa')
    if inactive_template:
        unused = doc.styles.add_style('UnusedUnsupportedVertical', WD_STYLE_TYPE.TABLE)
        add(add(unused.element, 'trPr'), 'trHeight', val='-999', hRule='exact')
        add(add(unused.element, 'tcPr'), 'vAlign', val='both')
    for col in table._tbl.tblGrid:
        col.set(qn('w:w'), '3000')
    for ri, row in enumerate(table.rows):
        if heights[ri] is not None:
            node = add(row._tr.get_or_add_trPr(), 'trHeight', val=heights[ri])
            if rule is not None:
                node.set(qn('w:hRule'), rule)
        for ci, cell in enumerate(row.cells):
            cell.width = Twips(3000)
            text = f'R{ri}C{ci}'
            if long and ri == ci == 0:
                text += '\n' + '\n'.join(f'GrowLine{i}' for i in range(1, 9))
            cell.text = text
            if alignments[ri][ci] is not None:
                add(cell._tc.get_or_add_tcPr(), 'vAlign', val=alignments[ri][ci])
            if padding == 'cell':
                margins = add(cell._tc.get_or_add_tcPr(), 'tcMar')
                for edge in ('top', 'bottom'):
                    add(margins, edge, w=0, type='dxa')
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.name, run.font.size = 'DejaVu Sans', Pt(12)
    return doc


def save_doc(doc, path):
    check(not path.exists(), 'Refusing to overwrite independent source')
    doc.save(path)
    return path


def mutate_package(source, destination, part, mutate):
    check(not destination.exists(), 'Refusing to overwrite derived fixture')
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination, 'w', zipfile.ZIP_DEFLATED) as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == part:
                root = ET.fromstring(data)
                mutate(root)
                data = ET.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
            output.writestr(info, data)


def fixtures(directory):
    check(ORIGINAL.is_file() and digest(ORIGINAL) == ORIGINAL_SHA, 'Pinned original missing or changed')
    raw = ORIGINAL.read_bytes()
    check(hashlib.sha1(f'blob {len(raw)}\0'.encode() + raw).hexdigest() == ORIGINAL_BLOB, 'Pinned Git blob differs')
    sources = {'original-TestTableCellAlign': {'path': str(ORIGINAL), 'sha256': ORIGINAL_SHA, 'kind': 'unchanged pinned public upstream fixture; not user production file', 'render': True}}
    plans = {
        'direct-omitted-rule': {},
        'direct-atleast': {'rule': 'atLeast'},
        'explicit-zero': {'heights': (0, 0), 'rule': 'atLeast'},
        'mixed-default': {'heights': (None, 1160), 'alignments': ((None, 'bottom'), ('center', None))},
        'alignment-only': {'heights': (None, None)},
        'height-only': {'alignments': ((None, None), (None, None))},
        'selected-empty-chain': {'selected_style': True},
        'long-content-growth': {'heights': (100, 1160), 'long': True},
        'zero-table-padding': {'padding': 'table'},
        'zero-single-edge-padding': {'padding': 'table-top-only'},
        'zero-cell-padding': {'padding': 'cell'},
        'zero-style-padding': {'padding': 'style'},
        'zero-conditional-padding': {'padding': 'conditional'},
        'inactive-layout-template': {'inactive_template': True},
    }
    for name, kwargs in plans.items():
        path = save_doc(author(**kwargs), directory / (name + '.docx'))
        sources[name] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned directed python-docx fixture', 'render': True}
    for name, item in sources.items():
        item['expected'] = snapshot(Path(item['path']))
        reference = Document(item['path'])
        edited_heights = (1800, 1440)
        edited_alignments = (('bottom', 'center'), ('top', 'bottom'))
        for ri, row in enumerate(reference.tables[0].rows):
            remove(row._tr.get_or_add_trPr(), 'trHeight')
            add(row._tr.get_or_add_trPr(), 'trHeight', val=edited_heights[ri], hRule='atLeast')
            for ci, cell in enumerate(row.cells):
                remove(cell._tc.get_or_add_tcPr(), 'vAlign')
                add(cell._tc.get_or_add_tcPr(), 'vAlign', val=edited_alignments[ri][ci])
        ref = save_doc(reference, directory / (name + '-independent-edited-reference.docx'))
        item['edited_reference'] = {'path': str(ref), 'sha256': digest(ref), 'expected': snapshot(ref), 'kind': 'independent python-docx derivative, original source never saved'}
    return sources


def negative_fixtures(directory, base, zero_padding_base):
    cases = {}
    def save(name, mutate, part='word/document.xml', reason='vertical', source=None):
        path = directory / ('reject-' + name + '.docx')
        mutate_package(source or base, path, part, mutate)
        cases[name] = {'path': str(path), 'sha256': digest(path), 'reason': reason}
    height = lambda root: root.find('.//w:trPr/w:trHeight', NS)
    valign = lambda root: root.find('.//w:tcPr/w:vAlign', NS)
    for value in ('exact', 'auto', 'nonsense', 'AtLeast', ''):
        save('rule-' + (value or 'empty'), lambda root, value=value: height(root).set(qn('w:hRule'), value))
    for value in ('-1', '1.5', 'NaN', '31681', '12961', '999999999999999999999999999999'):
        save('height-' + value, lambda root, value=value: height(root).set(qn('w:val'), value))
    save('height-missing', lambda root: height(root).attrib.pop(qn('w:val')))
    save('height-duplicate', lambda root: height(root).getparent().append(copy.deepcopy(height(root))))
    save('height-unknown-attribute', lambda root: height(root).set(qn('w:unknown'), '1'))
    save('height-child', lambda root: add(height(root), 'trHeight', val=200))
    save('height-wrong-namespace', lambda root: setattr(height(root), 'tag', '{urn:other}trHeight'))
    def wrap_height(root):
        node = height(root); parent = node.getparent(); parent.remove(node); add(parent, 'customXml').append(node)
    save('height-wrapped', wrap_height)
    for value in ('both', 'baseline', 'bad', 'Top', ''):
        save('align-' + (value or 'empty'), lambda root, value=value: valign(root).set(qn('w:val'), value))
    save('align-missing', lambda root: valign(root).attrib.pop(qn('w:val')))
    save('align-duplicate', lambda root: valign(root).getparent().append(copy.deepcopy(valign(root))))
    save('align-wrong-namespace', lambda root: setattr(valign(root), 'tag', '{urn:other}vAlign'))
    save('align-section-center', lambda root: add(root.find('w:body/w:sectPr', NS), 'vAlign', val='center'), reason='page')
    for tag in ('cantSplit', 'tblHeader', 'hidden'):
        save('row-' + tag, lambda root, tag=tag: add(root.find('.//w:trPr', NS), tag), reason='vertical')
    save('row-cantSplit-disabled', lambda root: add(root.find('.//w:trPr', NS), 'cantSplit', val='0'))
    save('cell-text-direction', lambda root: add(root.find('.//w:tcPr', NS), 'textDirection', val='lrTb'))
    save('cell-hidden-text', lambda root: add(add(root.find('.//w:r', NS), 'rPr'), 'vanish'), reason='hidden')
    save('cell-deleted', lambda root: add(root.find('.//w:tcPr', NS), 'cellDel'), reason='geometry')
    save('cell-merge', lambda root: add(root.find('.//w:tcPr', NS), 'vMerge', val='restart'), reason='geometry')
    save('row-revised', lambda root: add(root.find('.//w:trPr', NS), 'trPrChange'), reason='geometry')
    def nested(root):
        table = root.find('w:body/w:tbl', NS)
        table.find('w:tr/w:tc', NS).append(copy.deepcopy(table))
    save('nested-table', nested, reason='geometry')
    # Both plain and conditional style properties can be silently ignored by the
    # vendor, so the independent source deliberately keeps direct overrides too.
    for conditional in (False, True):
        for tag in ('trHeight', 'vAlign'):
            def style_mutate(root, tag=tag, conditional=conditional):
                style = root.find("w:style[@w:styleId='TableNormal']", NS)
                parent = add(style, 'tblStylePr', type='firstRow') if conditional else style
                parent = add(parent, 'trPr' if tag == 'trHeight' else 'tcPr')
                add(parent, tag, val=600 if tag == 'trHeight' else 'center')
            save(('conditional-' if conditional else 'inherited-') + tag, style_mutate, 'word/styles.xml')
    def table_padding(root, edge='top', value='20', unit='dxa', attribute=True):
        props = root.find('.//w:tbl/w:tblPr', NS)
        margins = add(props, 'tblCellMar')
        node = add(margins, edge, w=value)
        if attribute:
            node.set(qn('w:type'), unit)
        return node
    save('padding-table-positive', lambda r: table_padding(r))
    save('padding-table-negative', lambda r: table_padding(r, value='-1'))
    save('padding-table-pct', lambda r: table_padding(r, value='0', unit='pct'))
    save('padding-table-missing-unit', lambda r: table_padding(r, value='0', attribute=False))
    save('padding-table-missing-value', lambda r: table_padding(r, value='0').attrib.pop(qn('w:w')))
    save('padding-unknown-edge', lambda r: table_padding(r, edge='before', value='0'))
    save('padding-conflicting-top', lambda r: table_padding(r, value='0').getparent().append(OxmlElement('w:top')))
    def wrapped_padding(root):
        node = table_padding(root, value='0')
        parent = node.getparent(); parent.remove(node); add(parent, 'customXml').append(node)
    save('padding-wrapped-top', wrapped_padding)
    def cell_padding(root):
        margins = add(root.find('.//w:tcPr', NS), 'tcMar')
        add(margins, 'bottom', w=20, type='dxa')
    save('padding-cell-positive', cell_padding)
    def inherited_padding(root, conditional=False):
        style = root.find("w:style[@w:styleId='TableNormal']", NS)
        props = add(add(style, 'tblStylePr', type='firstRow'), 'tcPr') if conditional else style.find('w:tblPr', NS)
        if conditional:
            margins = add(props, 'tcMar')
            add(margins, 'top', w=20, type='dxa')
            add(margins, 'bottom', w=0, type='dxa')
        else:
            margins = props.find('w:tblCellMar', NS)
            check(margins is not None, 'Independent table default margin source missing')
            margins.find('w:top', NS).set(qn('w:w'), '20')
    save('padding-style-positive', inherited_padding, 'word/styles.xml')
    save('padding-conditional-positive', lambda r: inherited_padding(r, True), 'word/styles.xml')
    save('padding-style-overridden-zero', inherited_padding, 'word/styles.xml', source=zero_padding_base)
    save('padding-conditional-overridden-zero', lambda r: inherited_padding(r, True), 'word/styles.xml', source=zero_padding_base)
    return cases


def measure_pdf(path):
    document = fitz.open(path)
    check(len(document) == 1, f'{path.name}: expected one page for bounded probes')
    page = document[0]
    horizontal, vertical = [], []
    for drawing in page.get_drawings():
        for item in drawing['items']:
            if item[0] != 'l':
                continue
            a, b = item[1:]
            if abs(a.y - b.y) < .05 and abs(a.x - b.x) > 100:
                horizontal.append(float((a.y + b.y) / 2))
            if abs(a.x - b.x) < .05 and abs(a.y - b.y) > 10:
                vertical.append(float((a.x + b.x) / 2))
    def unique(values):
        result = []
        for value in sorted(values):
            if not result or value - result[-1] > .1:
                result.append(round(value, 5))
        return result
    ys, xs = unique(horizontal), unique(vertical)
    check(len(ys) == len(xs) == 3, f'{path.name}: expected 3×3 table borders, got {ys}/{xs}')
    cells = [[[], []], [[], []]]
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            for span in line['spans']:
                x, y = span['origin']
                if not span['text'].strip():
                    continue
                ri = 0 if y < ys[1] else 1
                ci = 0 if x < xs[1] else 1
                check(ys[ri] < y < ys[ri + 1], f'{path.name}: text baseline outside its row: {span}')
                cells[ri][ci].append({'text': span['text'], 'font': span['font'], 'size': span['size'],
                    'top_offset': round(y - ys[ri], 5), 'bottom_offset': round(ys[ri + 1] - y, 5),
                    'baseline_y': round(y, 5), 'x': round(x, 5)})
    return {'path': str(path), 'sha256': digest(path), 'page_count': len(document),
            'page_size': [page.rect.width, page.rect.height], 'borders_x': xs, 'borders_y': ys,
            'row_heights': [round(ys[i + 1] - ys[i], 5) for i in range(2)], 'cells': cells}


def render_all(soffice, directory, paths, commands):
    out = directory / 'pdf'; out.mkdir()
    cache = directory / 'lo-cache'; cache.mkdir()
    paths = list(dict.fromkeys(Path(p).resolve() for p in paths))
    check(len({p.stem for p in paths}) == len(paths), 'PDF names collide')
    # Each PDF gets a fresh process/profile. A retained batch investigation found
    # one zero-width PDF amid normally rendered byte-identical DOCX inputs; that
    # same input rendered correctly with an isolated profile. Never repair PDF
    # coordinates or reuse a good sibling as the failing document's evidence.
    for index, path in enumerate(paths):
        profile = directory / f'lo-profile-{index:03d}'
        command = [soffice, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(path)]
        result = subprocess.run(command, capture_output=True, text=True, timeout=60, env=dict(os.environ, XDG_CACHE_HOME=str(cache)))
        commands.append({'command': command, 'returncode': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr})
        check(result.returncode == 0, 'LibreOffice conversion failed: ' + result.stderr)
    return {str(p): measure_pdf(out / (p.stem + '.pdf')) for p in paths}


def compare_pdf(a, b, name):
    check(a['page_size'] == b['page_size'], name + ': page size differs')
    deltas = [abs(x - y) for x, y in zip(a['row_heights'], b['row_heights'])]
    absolute_y = [abs(x - y) for x, y in zip(a['borders_y'], b['borders_y'])]
    for ra, rb in zip(a['cells'], b['cells']):
        for ca, cb in zip(ra, rb):
            check([s['text'] for s in ca] == [s['text'] for s in cb], name + ': rendered text differs or wraps')
            for sa, sb in zip(ca, cb):
                check(sa['font'] == sb['font'] and sa['size'] == sb['size'], name + ': consumer font differs')
                deltas.extend(abs(sa[k] - sb[k]) for k in ('top_offset', 'bottom_offset'))
                absolute_y.append(abs(sa['baseline_y'] - sb['baseline_y']))
    maximum = max(deltas, default=0)
    check(maximum <= TOLERANCE, f'{name}: relative vertical difference {maximum:.5f}pt exceeds {TOLERANCE}')
    check(max(absolute_y, default=0) <= TOLERANCE, name + ': absolute vertical anchor differs')
    return {'max_border_relative_delta_pt': maximum, 'tolerance_pt': TOLERANCE,
            'max_absolute_y_delta_pt': max(absolute_y, default=0),
            'absolute_x_delta_record_only': [round(y - x, 5) for x, y in zip(a['borders_x'], b['borders_x'])],
            'scope': 'row border distances and each text baseline relative to own cell edges; x normalization is not a fidelity pass'}


def run_cli(cli, commands, *args, success=True):
    command = [str(cli), *map(str, args)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=45)
    commands.append({'command': command, 'returncode': result.returncode, 'stderr': result.stderr})
    check((result.returncode == 0) == success, 'Unexpected CLI result: ' + ' '.join(command) + '\n' + result.stderr)
    return json.loads(result.stdout) if success else result.stderr


def validate_model(model, expected, name):
    tables = [n for n in model['content'] if n['type'] == 'table']
    check(len(tables) == len(expected), name + ': model table count differs')
    for table, target in zip(tables, expected):
        check(len(table['content']) == len(target['rows']), name + ': model row count differs')
        for row, tr in zip(table['content'], target['rows']):
            check(row.get('attrs', {}).get(HEIGHT) == tr['height'], name + ': model minimum/omission differs')
            check(len(row['content']) == len(tr['cells']), name + ': model cell count differs')
            for cell, tc in zip(row['content'], tr['cells']):
                check(cell.get('attrs', {}).get(ALIGN) == tc['align'], name + ': model alignment/omission differs')
                check(model_text(cell) == tc['text'], name + ': model cell text differs')


def validate_export(path, expected, name):
    actual = snapshot(path)
    check(len(actual) == len(expected), name + ': output tables differ')
    for table, target in zip(actual, expected):
        check(len(table['rows']) == len(target['rows']), name + ': output rows differ')
        for row, tr in zip(table['rows'], target['rows']):
            check(row['height'] == tr['height'], name + ': independent output height differs')
            check(row['rule'] == (None if tr['height'] is None else 'atLeast'), name + ': independent output is not explicit minimum')
            check(row['cells'] == tr['cells'], name + ': independent output text/alignment differs')
    return actual


def export_model(cli, commands, directory, name, model, expected, source):
    model_path, output = directory / (name + '.json'), directory / (name + '.docx')
    with model_path.open('x') as stream:
        json.dump(model, stream, ensure_ascii=False, indent=2)
    run_cli(cli, commands, 'create-docx', model_path, output)
    check(run_cli(cli, commands, 'read-docx-model', output) == model, name + ': full trusted model differs')
    independent = validate_export(output, expected, name)
    before = {str(p): digest(p) for p in (source, output)}
    for p in (source, output):
        error = run_cli(cli, commands, 'create-docx', model_path, p, success=False)
        check('already exists' in error, name + ': no-clobber refusal differs')
    check(all(digest(p) == value for p, value in before.items()), name + ': source/output mutated')
    return {'path': str(output), 'sha256': digest(output), 'trusted_complete_model_equal': True, 'independent': independent, 'no_clobber': True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path)
    parser.add_argument('--baseline-cli', type=Path, default=ROOT / 'tmp/contextual-final2-cli')
    parser.add_argument('--root', type=Path, default=ROOT / 'tmp/table-vertical-independent')
    parser.add_argument('--generate-only', action='store_true')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--soffice', default=shutil.which('soffice'))
    parser.add_argument('--refusal-prefix', default='Unsupported table geometry')
    args = parser.parse_args()
    check(args.generate_only or args.cli, 'Pass frozen --cli or --generate-only')
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    report = directory / 'report.json'
    evidence = {'status': 'running', 'commands': [], 'sources': {}, 'cases': {}, 'rejected': {}, 'invalid_models': {}, 'tamper': {}, 'comparisons': {},
        'script_sha256': digest(__file__), 'python': platform.python_version(),
        'versions': {p: version(p) for p in ('python-docx', 'lxml', 'PyMuPDF')},
        'original_source_url': f'https://github.com/apache/poi/blob/{REVISION}/test-data/document/TestTableCellAlign.docx',
        'original_git_blob': ORIGINAL_BLOB,
        'render_strategy': 'Each source/output is converted separately with its own fresh LibreOffice profile; no batch PDF reuse or geometry correction',
        'limitations': ['Unchanged upstream regression fixture is not a user production document.',
            'Directed fixtures and edit references are independently authored and distinctly labeled.',
            'LibreOffice only; no actual browser/native printing or Word/WPS verification.',
            'Automatic column width normalization is disclosed and measured x deltas are not fidelity passes.',
            'No exact height, inheritance/conditional table layout, merged/nested tables or pagination-policy support claimed.']}
    try:
        sources = fixtures(directory)
        negatives = negative_fixtures(directory, Path(sources['selected-empty-chain']['path']), Path(sources['zero-table-padding']['path']))
        evidence['sources'] = sources
        evidence['negative_sources'] = negatives
        paths = [p for item in sources.values() for p in (item['path'], item['edited_reference']['path'])]
        if not args.generate_only:
            cli = args.cli.resolve()
            evidence['cli'] = {'path': str(cli), 'sha256': digest(cli)}
            baseline_cli = args.baseline_cli.resolve()
            check(baseline_cli.is_file(), 'Frozen previous CLI is missing')
            evidence['baseline_cli'] = {'path': str(baseline_cli), 'sha256': digest(baseline_cli)}
            imported_models = {}
            for name, item in sources.items():
                source = Path(item['path'])
                imported = run_cli(cli, evidence['commands'], 'import-docx', source)
                model = imported['document']; imported_models[name] = model
                validate_model(model, item['expected'], name)
                baseline = export_model(cli, evidence['commands'], directory, name + '-export', model, item['expected'], source)
                edited = copy.deepcopy(model)
                table = next(n for n in edited['content'] if n['type'] == 'table')
                for row, expected in zip(table['content'], item['edited_reference']['expected'][0]['rows']):
                    row.setdefault('attrs', {})[HEIGHT] = expected['height']
                    for cell, tc in zip(row['content'], expected['cells']):
                        cell.setdefault('attrs', {})[ALIGN] = tc['align']
                edit = export_model(cli, evidence['commands'], directory, name + '-edited-export', edited, item['edited_reference']['expected'], source)
                reset = copy.deepcopy(model)
                def clear(node):
                    for key in (HEIGHT, ALIGN):
                        if key in node.get('attrs', {}):
                            del node['attrs'][key]
                    if node.get('attrs') == {}:
                        del node['attrs']
                    for child in node.get('content', []):
                        clear(child)
                clear(reset)
                reset_expected = copy.deepcopy(item['expected'])
                for target in reset_expected:
                    for row in target['rows']:
                        row['height'] = row['rule'] = None
                        for cell in row['cells']:
                            cell['align'] = None
                reset_result = export_model(cli, evidence['commands'], directory, name + '-reset-export', reset, reset_expected, source)
                old_output = directory / (name + '-reset-old-cli.docx')
                run_cli(baseline_cli, evidence['commands'], 'create-docx', directory / (name + '-reset-export.json'), old_output)
                check(digest(old_output) == reset_result['sha256'], name + ': no-new-attrs whole package differs from old CLI')
                check(run_cli(baseline_cli, evidence['commands'], 'read-docx-model', old_output) == reset, name + ': old CLI reset model differs')
                reset_result['old_cli_entire_docx_bytes_equal'] = True
                evidence['cases'][name] = {'compatibility': imported['compatibility'], 'baseline': baseline, 'edited': edit, 'reset': reset_result}
                paths.extend([baseline['path'], edit['path']])
            for name, item in negatives.items():
                error = run_cli(cli, evidence['commands'], 'import-docx', item['path'], success=False)
                needle = {'vertical': args.refusal_prefix, 'page': 'Unsupported page geometry', 'hidden': 'Hidden text', 'geometry': 'table'}[item['reason']]
                check(needle.lower() in error.lower(), name + ': expected specific protection absent: ' + error)
                check(digest(item['path']) == item['sha256'], name + ': rejected source changed')
                evidence['rejected'][name] = {'stderr': error, 'source_unchanged': True}
            base = imported_models['direct-atleast']
            for attr, values, where in [(HEIGHT, [-1, 1.5, None, True, '100', 31681, 12961], 'row'), (ALIGN, [None, '', 'middle', 'Top', 1, True], 'cell')]:
                for index, value in enumerate(values):
                    name = f'{attr}-{index}'; model = copy.deepcopy(base); table = next(n for n in model['content'] if n['type'] == 'table')
                    node = table['content'][0] if where == 'row' else table['content'][0]['content'][0]
                    node.setdefault('attrs', {})[attr] = value
                    path, output = directory / ('bad-' + name + '.json'), directory / ('bad-' + name + '.docx')
                    path.write_text(json.dumps(model)); error = run_cli(cli, evidence['commands'], 'create-docx', path, output, success=False)
                    check(not output.exists(), name + ': invalid model published')
                    evidence['invalid_models'][name] = {'stderr': error, 'no_output': True}
            for name, node_path, attr, value in [('height-on-table', [], HEIGHT, 100), ('height-on-cell', [0, 0], HEIGHT, 100),
                    ('align-on-row', [0], ALIGN, 'center'), ('align-on-paragraph', [0, 0, 0], ALIGN, 'center')]:
                model = copy.deepcopy(base)
                node = next(n for n in model['content'] if n['type'] == 'table')
                for index in node_path:
                    node = node['content'][index]
                node.setdefault('attrs', {})[attr] = value
                path, output = directory / ('bad-' + name + '.json'), directory / ('bad-' + name + '.docx')
                path.write_text(json.dumps(model)); error = run_cli(cli, evidence['commands'], 'create-docx', path, output, success=False)
                check(not output.exists(), name + ': misplaced model attr published')
                evidence['invalid_models'][name] = {'stderr': error, 'no_output': True}
            good = Path(evidence['cases']['direct-atleast']['baseline']['path'])
            for name, mutation in [('height', lambda r: r.find('.//w:trHeight', NS).set(qn('w:val'), '1400')), ('alignment', lambda r: r.find('.//w:vAlign', NS).set(qn('w:val'), 'top'))]:
                bad = directory / ('tampered-' + name + '.docx')
                mutate_package(good, bad, 'word/document.xml', mutation)
                error = run_cli(cli, evidence['commands'], 'read-docx-model', bad, success=False)
                evidence['tamper'][name] = {'stderr': error, 'rejected': True}
                with zipfile.ZipFile(bad) as package:
                    parts = {n: package.read(n) for n in package.namelist()}
                stored = json.loads(parts['rust-wps/model-v2.json'])
                stored['hashes']['word/document.xml'] = hashlib.sha256(parts['word/document.xml']).hexdigest()
                parts['rust-wps/model-v2.json'] = json.dumps(stored, ensure_ascii=False, separators=(',', ':')).encode()
                forged = directory / ('tampered-' + name + '-repaired-hash.docx')
                with zipfile.ZipFile(forged, 'x', zipfile.ZIP_DEFLATED) as package:
                    for part, data in parts.items():
                        package.writestr(part, data)
                error = run_cli(cli, evidence['commands'], 'read-docx-model', forged, success=False)
                evidence['tamper'][name + '-repaired-hash'] = {'stderr': error, 'rejected': True, 'scope': 'part hash repaired, stale model still rejected by exact reencoding'}
            extra = directory / 'tampered-extra-part.docx'
            with zipfile.ZipFile(good) as original, zipfile.ZipFile(extra, 'x', zipfile.ZIP_DEFLATED) as package:
                for name in original.namelist():
                    package.writestr(name, original.read(name))
                package.writestr('word/unexpected-part.xml', b'<unexpected/>')
            error = run_cli(cli, evidence['commands'], 'read-docx-model', extra, success=False)
            evidence['tamper']['extra-part'] = {'stderr': error, 'rejected': True}
            check(digest(cli) == evidence['cli']['sha256'], 'Frozen CLI changed during validation')
            check(digest(baseline_cli) == evidence['baseline_cli']['sha256'], 'Frozen old CLI changed during validation')
        if args.render:
            check(args.soffice, 'LibreOffice unavailable')
            evidence['soffice_version'] = subprocess.check_output([args.soffice, '--version'], text=True).strip()
            rendered = render_all(args.soffice, directory, paths, evidence['commands'])
            evidence['renders'] = rendered
            for name, item in sources.items():
                source_render = rendered[str(Path(item['path']).resolve())]
                if name == 'long-content-growth':
                    check(source_render['row_heights'][0] > 100 / 20 + 50, 'Multiline source did not exceed minimum')
                    check(any(s['text'] == 'GrowLine8' for s in source_render['cells'][0][0]), 'Long source text clipped')
                if name in evidence['cases']:
                    case = evidence['cases'][name]
                    evidence['comparisons'][name + '-unchanged'] = compare_pdf(source_render, rendered[case['baseline']['path']], name)
                    evidence['comparisons'][name + '-edited'] = compare_pdf(rendered[item['edited_reference']['path']], rendered[case['edited']['path']], name + ' actual edit')
            original_pdf = Path(rendered[str(ORIGINAL.resolve())]['path'])
            preview = fitz.open(original_pdf)[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            preview.save(directory / 'original-preview.png')
        for item in sources.values():
            check(digest(item['path']) == item['sha256'], 'Source changed during suite')
            check(digest(item['edited_reference']['path']) == item['edited_reference']['sha256'], 'Independent reference changed')
        check(digest(__file__) == evidence['script_sha256'], 'Script changed during suite')
        evidence['status'] = 'generated' if args.generate_only else 'passed'
    except Exception as error:
        evidence['status'] = 'failed'; evidence['error'] = repr(error)
        raise
    finally:
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
        print(json.dumps({'report': str(report), 'status': evidence['status'], 'sources': len(evidence['sources']), 'rejections': len(evidence['rejected']), 'pdfs': len(evidence.get('renders', {}))}), flush=True)


if __name__ == '__main__':
    main()
