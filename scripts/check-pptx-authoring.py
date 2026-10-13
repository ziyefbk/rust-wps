"""Independent python-pptx validation of Rust-authored slides."""
from pathlib import Path
import json
import subprocess
import tempfile
import base64
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
model=json.loads((ROOT/'docs/examples/presentation-v1.json').read_text(encoding='utf-8'))
image_bytes=(ROOT/'tests/fixtures/slide-image.png').read_bytes()
model['slides'][0]['images']=[{'bounds_emu':[720000,360000,1828800,914400],'data_url':'data:image/png;base64,'+base64.b64encode(image_bytes).decode(),'alt':'中文 & <图片>'}]
with tempfile.TemporaryDirectory() as folder:
    source=Path(folder)/'model.json';target=Path(folder)/'output.pptx'
    source.write_text(json.dumps(model,ensure_ascii=False),encoding='utf-8')
    command=['cargo','run','--quiet','--locked','--','create-pptx',str(source),str(target)]
    subprocess.run(command,cwd=ROOT,check=True,capture_output=True)
    reopened_model=subprocess.run(['cargo','run','--quiet','--locked','--','read-pptx-model',str(target)],cwd=ROOT,check=True,capture_output=True)
    assert json.loads(reopened_model.stdout)==model
    deck=Presentation(target)
    assert [deck.slide_width,deck.slide_height]==model['size_emu']
    assert len(deck.slides)==len(model['slides'])
    for expected,slide in zip(model['slides'],deck.slides):
        assert len(slide.shapes)==len(expected['text_boxes'])+len(expected.get('images',[]))
        for wanted,shape in zip(expected['text_boxes'],slide.shapes):
            assert [shape.left,shape.top,shape.width,shape.height]==wanted['bounds_emu']
            assert shape.text.replace('\v','\n')==wanted['text']
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert round(run.font.size.pt*100)==wanted['font_size_hundredths_pt']
                    assert bool(run.font.bold)==wanted['bold'] and bool(run.font.italic)==wanted['italic']
                    assert list(run.font.color.rgb)==wanted['color_rgb']
    picture=deck.slides[0].shapes[-1]
    assert picture.image.blob==image_bytes
    assert [picture.left,picture.top,picture.width,picture.height]==model['slides'][0]['images'][0]['bounds_emu']
    assert picture._element.nvPicPr.cNvPr.get('descr')==model['slides'][0]['images'][0]['alt']
    before=target.read_bytes()
    refused=subprocess.run(command,cwd=ROOT,capture_output=True)
    assert refused.returncode and target.read_bytes()==before
    assert json.loads(source.read_text(encoding='utf-8'))==model
    print(json.dumps({'independentAuthoredPptx':True,'geometryTextStyleVerified':True,'blankSlideVerified':True,'noOverwrite':True,'pictureBytesGeometryAltVerified':True}))
