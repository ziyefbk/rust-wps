#!/usr/bin/env python3
"""Independent, bounded oracle for actual frontend preserved-list deletions.

The established geometry checker is reused read-only for OPC/style/numbering
inspection and fresh-profile PDF/raster extraction. Deletion expectations and
full paragraph/run formatting are checked here without importing application
code. References are independently authored with python-docx and OOXML.
No browser/native keyboard, Word/WPS or general list fidelity claim is made.
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import platform
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Twips
from lxml import etree as E

ROOT = Path(__file__).resolve().parent.parent
HELPER = ROOT / 'scripts/check-list-geometry-independent.py'
HELPER_SHA = 'f71c8838d06b55b27c233528b56e6e9eef526426333d1de4a905d3e3bf67c187'
CLI_SHA = '5c008894209311c84515a6432b6e973a120ad459b30c848ea019c2128808111a'
PRIOR = ROOT / 'tmp/list-geometry-final-independent/run-41zrpucu/report.json'
PRIOR_SHA = '3957dea4710f2b534a21030332f4a63578c0e174a836ef2400802b6b6cacb91b'
LEGACY = ROOT / 'tmp/list-geometry-legacy-replay/run-go2w29_o/report.json'
LEGACY_SHA = '616b8b261b4b52e12a837630a24354db015261b3b0301bc6fbdf0fce64fc5f1d'
spec = importlib.util.spec_from_file_location('independent_list_geometry', HELPER)
geometry = importlib.util.module_from_spec(spec)
spec.loader.exec_module(geometry)
check, digest, save_json = geometry.check, geometry.digest, geometry.save_json
W, NS, child, val = geometry.W, geometry.NS, geometry.child, geometry.val


def canonical(value):
    return json.dumps(value, sort_keys=True, separators=(',', ':'), ensure_ascii=False)


def model_from(path):
    wrapped = json.loads(path.read_text())
    if 'version' in wrapped:
        check(wrapped['version'] == 2, 'Snapshot version must be 2')
    return wrapped.get('document', wrapped)


def plain_blank(node):
    return (node.get('type') == 'paragraph' and not node.get('content')
            and not node.get('marks')
            and set(node.get('attrs', {})) <= {'textAlign', 'spaceBefore', 'spaceAfter', 'lineSpacing',
                                               'listGeometry', 'emptyParagraphFont', *geometry.KEYS}
            and all(v is None for v in node.get('attrs', {}).values())
            and not (set(node) - {'type', 'attrs', 'content', 'marks'}))


def deletion_expectation(operation, before, after):
    """Construct the only permitted deletion without consulting app proof code."""
    expected = copy.deepcopy(before)
    kind = operation['kind']
    if kind == 'items':
        index = operation['blockIndex']
        block = expected['content'][index]
        check(block['type'] == 'orderedList' and block['attrs'].get('markerFont'),
              'Deletion must target a preserved list')
        indices = operation['itemIndices']
        check(indices and len(indices) == len(set(indices))
              and all(type(i) is int and 0 <= i < len(block['content']) for i in indices),
              'Invalid independent item deletion indices')
        check(indices == list(range(min(indices), max(indices) + 1)),
              'Independent deletion scope must be contiguous')
        block['content'] = [n for i, n in enumerate(block['content']) if i not in indices]
        if not block['content']:
            expected['content'].pop(index)
    elif kind == 'blocks':
        indices = operation['blockIndices']
        check(indices and len(indices) == len(set(indices))
              and all(type(i) is int and 0 <= i < len(expected['content']) for i in indices),
              'Invalid independent block deletion indices')
        expected['content'] = [n for i, n in enumerate(expected['content']) if i not in indices]
    elif kind == 'document':
        expected['content'] = []
    else:
        raise AssertionError('Unknown deletion kind: ' + str(kind))
    if not expected['content']:
        check(len(after.get('content', [])) == 1 and plain_blank(after['content'][0]),
              operation['name'] + ': cleared document must contain one default blank paragraph')
        expected['content'] = copy.deepcopy(after['content'])
    check(after == expected, operation['name'] + ': survivor content, formatting, order or document/list attributes changed')
    return {'kind': kind, 'entire_expected_model_equal': True,
            'survivor_paragraphs': len(geometry.model_rows(after)),
            'blank_default_paragraph': len(after['content']) == 1 and plain_blank(after['content'][0])}


def paragraphs(model):
    for block in model.get('content', []):
        if block['type'] == 'paragraph':
            yield block
        else:
            check(block['type'] == 'orderedList' and block.get('attrs', {}).get('markerFont'),
                  'Independent deletion oracle supports only paragraphs and preserved lists')
            for item in block['content']:
                check(item['type'] == 'listItem' and len(item['content']) == 1
                      and item['content'][0]['type'] == 'paragraph', 'Unexpected item structure')
                yield item['content'][0]


def default_format():
    return {'fontFamily': 'Arial', 'fontSize': 22, 'bold': False, 'italic': False,
            'underline': False, 'strike': False, 'color': '000000', 'highlight': None}


def span_format(span):
    fmt = default_format()
    marks = {m['type']: m.get('attrs', {}) for m in span.get('marks', [])}
    check(set(marks) <= {'fontFamily', 'fontSize', 'bold', 'italic', 'underline', 'strike', 'textColor'},
          'Unsupported reference mark')
    fmt['fontFamily'] = marks.get('fontFamily', {}).get('name', 'Arial')
    fmt['fontSize'] = marks.get('fontSize', {}).get('halfPoints', 22)
    for flag in ('bold', 'italic', 'underline', 'strike'):
        fmt[flag] = flag in marks
    if 'textColor' in marks:
        fmt['color'] = marks['textColor']['rgb']
    return fmt


def compact_runs(runs):
    result = []
    for run in runs:
        if not run['text']:
            continue
        if result and result[-1]['format'] == run['format']:
            result[-1]['text'] += run['text']
        else:
            result.append(copy.deepcopy(run))
    return result


def model_formatting(model):
    check(model.get('attrs', {}).get('styleProfile') == 'standard-v2', 'Expected standard-v2 model')
    result = []
    for node in paragraphs(model):
        attrs = node.get('attrs', {})
        spacing = {key: attrs.get(key) if attrs.get(key) is not None else default
                   for key, default in [('spaceBefore', 0), ('spaceAfter', 120), ('lineSpacing', 360)]}
        runs = []
        for span in node.get('content', []):
            check(span['type'] in ('text', 'hardBreak'), 'Unsupported inline')
            runs.append({'text': span.get('text', '') if span['type'] == 'text' else '\n',
                         'format': span_format(span)})
        row = {'spacing': spacing, 'alignment': attrs.get('textAlign') or 'left',
               'runs': compact_runs(runs)}
        if not runs:
            row['emptyFont'] = attrs.get('emptyParagraphFont') or {'family': 'Arial', 'halfPoints': 22}
        result.append(row)
    return result


def inspect_formatting(path):
    """Resolve document/style/direct paragraph and Latin run properties ourselves."""
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    info = geometry.inspect_docx(path)
    document = E.fromstring(parts[info['parts']['document']])
    styles = E.fromstring(parts[info['parts']['styles']])
    defaults = styles.find('w:docDefaults', NS)
    style_map = geometry.unique_map(styles.findall('w:style', NS), 'styleId', 'style identity')
    default_ids = [sid for sid, n in style_map.items() if n.get(W+'type') == 'paragraph'
                   and n.get(W+'default') in ('1', 'true', 'on')]
    check(len(default_ids) <= 1, 'Ambiguous default style')
    theme = E.fromstring(parts[info['parts']['theme']]) if info['parts']['theme'] else None
    themes = {}
    if theme is not None:
        for name in ('major', 'minor'):
            node = theme.find('.//a:fontScheme/a:'+name+'Font/a:latin', NS)
            if node is not None:
                themes[name+'HAnsi'] = node.get('typeface')

    def lineage(sid, seen=()):
        if sid is None:
            return []
        check(sid in style_map and sid not in seen, 'Invalid style lineage')
        node = style_map[sid]
        return lineage(val(node, 'w:basedOn'), seen+(sid,)) + [node]

    def apply_rpr(fmt, rpr):
        if rpr is None:
            return
        font = rpr.find('w:rFonts', NS)
        if font is not None:
            family = font.get(W+'ascii') or themes.get(font.get(W+'asciiTheme'))
            if family:
                fmt['fontFamily'] = family
        size = val(rpr, 'w:sz')
        if size is not None:
            fmt['fontSize'] = int(size)
        for tag, key in [('b', 'bold'), ('i', 'italic'), ('u', 'underline'), ('strike', 'strike')]:
            element = rpr.find('w:'+tag, NS)
            if element is not None:
                fmt[key] = element.get(W+'val', '1') not in ('0', 'false', 'off', 'none')
        color = val(rpr, 'w:color')
        if color is not None:
            fmt['color'] = '000000' if color == 'auto' else color.upper()
        highlight = val(rpr, 'w:highlight')
        if highlight is not None:
            fmt['highlight'] = None if highlight == 'none' else highlight

    rows = []
    for para in document.findall('w:body/w:p', NS):
        ppr = para.find('w:pPr', NS)
        styles_here = lineage(val(ppr, 'w:pStyle', default_ids[0] if default_ids else None))
        pprs = ([defaults.find('w:pPrDefault/w:pPr', NS)] if defaults is not None else [])
        pprs += [n.find('w:pPr', NS) for n in styles_here] + [ppr]
        spacing = {'spaceBefore': 0, 'spaceAfter': 0, 'lineSpacing': 240}
        alignment = 'left'
        for pp in pprs:
            if pp is None:
                continue
            alignment = val(pp, 'w:jc', alignment)
            sp = pp.find('w:spacing', NS)
            if sp is not None:
                check(sp.get(W+'lineRule', 'auto') == 'auto', 'Unexpected reference line rule')
                for xmlkey, key in [('before', 'spaceBefore'), ('after', 'spaceAfter'), ('line', 'lineSpacing')]:
                    if sp.get(W+xmlkey) is not None:
                        spacing[key] = int(sp.get(W+xmlkey))
        base = default_format()
        apply_rpr(base, defaults.find('w:rPrDefault/w:rPr', NS) if defaults is not None else None)
        for style in styles_here:
            apply_rpr(base, style.find('w:rPr', NS))
        runs = []
        for run in para.findall('w:r', NS):
            fmt = copy.deepcopy(base)
            apply_rpr(fmt, run.find('w:rPr', NS))
            text = geometry.paragraph_text(run)
            runs.append({'text': text, 'format': fmt})
        row = {'spacing': spacing, 'alignment': alignment, 'runs': compact_runs(runs)}
        if not row['runs']:
            empty = copy.deepcopy(base)
            apply_rpr(empty, ppr.find('w:rPr', NS) if ppr is not None else None)
            row['emptyFont'] = {'family': empty['fontFamily'], 'halfPoints': empty['fontSize']}
        rows.append(row)
    return rows


def reference_from_model(original, model, path):
    """Build new paragraphs/numbering, retaining only source settings/theme context."""
    doc = Document(original)
    for node in list(doc.element.body):
        if node.tag != W+'sectPr':
            doc.element.body.remove(node)
    page = model['attrs']['pageGeometry']
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(page['widthTwips']), Twips(page['heightTwips'])
    for side in ('top', 'right', 'bottom', 'left'):
        setattr(section, side+'_margin', Twips(page['margin'+side.title()+'Twips']))
    numbering = doc.part.numbering_part.element
    for node in list(numbering):
        numbering.remove(node)

    def paragraph(node, num_id=None):
        attrs = node.get('attrs', {})
        p = doc.add_paragraph()
        pp = p._p.get_or_add_pPr()
        if num_id is not None:
            num = child(pp, 'numPr'); child(num, 'ilvl', val=0); child(num, 'numId', val=num_id)
            geo = attrs['listGeometry']
            p.paragraph_format.left_indent = Twips(geo['leftTwips'])
            p.paragraph_format.right_indent = Twips(geo['rightTwips'])
            p.paragraph_format.first_line_indent = Twips(-geo['hangingTwips'])
        else:
            for key, property_name in zip(geometry.KEYS, ('left_indent', 'right_indent', 'first_line_indent')):
                if attrs.get(key) is not None:
                    setattr(p.paragraph_format, property_name, Twips(attrs[key]))
        fmt = model_formatting({'type': 'doc', 'attrs': {'styleProfile': 'standard-v2'}, 'content': [node]})[0]
        p.paragraph_format.space_before = Twips(fmt['spacing']['spaceBefore'])
        p.paragraph_format.space_after = Twips(fmt['spacing']['spaceAfter'])
        p.paragraph_format.line_spacing = fmt['spacing']['lineSpacing']/240
        child(pp, 'jc', val=fmt['alignment'])
        if not node.get('content'):
            empty = fmt['emptyFont']; rp = child(pp, 'rPr')
            child(rp, 'rFonts', ascii=empty['family'], hAnsi=empty['family'])
            child(rp, 'sz', val=empty['halfPoints'])
        for span in node.get('content', []):
            run = p.add_run(span.get('text', ''))
            if span['type'] == 'hardBreak':
                run.add_break()
            rf = span_format(span)
            run.font.name = rf['fontFamily']; run.font.size = Pt(rf['fontSize']/2)
            for flag in ('bold', 'italic', 'underline', 'strike'):
                setattr(run.font, flag, rf[flag])
            run.font.color.rgb = RGBColor.from_string(rf['color'])
            if rf['highlight'] is not None:
                child(run._r.get_or_add_rPr(), 'highlight', val=rf['highlight'])

    for i, block in enumerate(model['content']):
        if block['type'] == 'paragraph':
            paragraph(block)
            continue
        check(block['type'] == 'orderedList', 'Reference block outside deletion subset')
        attrs = block['attrs']; font = attrs['markerFont']
        abstract = child(numbering, 'abstractNum', abstractNumId=i)
        child(abstract, 'multiLevelType', val='singleLevel'); level = child(abstract, 'lvl', ilvl=0)
        for tag, value in [('start', attrs.get('start', 1)), ('numFmt', 'decimal'), ('lvlText', '%1.'),
                           ('lvlJc', 'left'), ('suff', 'tab')]:
            child(level, tag, val=value)
        rp = child(level, 'rPr')
        child(rp, 'rFonts', ascii=font['fontFamily'], hAnsi=font['fontFamily'])
        child(rp, 'sz', val=font['fontSize']); child(rp, 'b', val=0)
        child(rp, 'i', val=0); child(rp, 'color', val='000000')
        instance = child(numbering, 'num', numId=i+1); child(instance, 'abstractNumId', val=i)
        for item in block['content']:
            paragraph(item['content'][0], i+1)
    check(not path.exists(), 'Independent reference must be fresh')
    doc.save(path)
    return geometry.inspect_docx(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path, default=ROOT/'tmp/list-geometry-final-cli')
    parser.add_argument('--original', type=Path, default=ROOT/'tmp/paragraph-indentation-source/testParaListRightIndent.docx')
    parser.add_argument('--editor-models', type=Path, required=True)
    parser.add_argument('--root', type=Path, default=ROOT/'tmp/list-deletion-independent')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--soffice')
    args = parser.parse_args()
    check(digest(HELPER) == HELPER_SHA, 'Established read-only helper identity differs')
    check(digest(args.cli) == CLI_SHA, 'Frozen backend CLI identity differs')
    check(digest(PRIOR) == PRIOR_SHA, 'Historical reference report identity differs')
    check(digest(LEGACY) == LEGACY_SHA, 'Historical legacy report identity differs')
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    original, cli, inputs = args.original.resolve(), args.cli.resolve(), args.editor_models.resolve()
    manifest_path = inputs/'manifest.json'
    evidence = {'status': 'running', 'started_at_utc': datetime.now(timezone.utc).isoformat(),
                'invocation': [sys.executable, *sys.argv], 'python_version': platform.python_version(),
                'script_sha256': digest(__file__), 'helper_sha256': digest(HELPER),
                'cli': str(cli), 'cli_sha256': digest(cli), 'commands': [],
                'limitations': geometry.LIMITATIONS, 'input_hashes': {}, 'snapshots': {},
                'operations': {}, 'render_comparisons': {}}
    try:
        evidence['original'] = {'path': str(original), 'sha256': digest(original),
                                'independent': geometry.original_preflight(original)}
        manifest = json.loads(manifest_path.read_text())
        evidence['manifest'] = {'path': str(manifest_path), 'sha256': digest(manifest_path), 'contents': manifest}
        evidence['generation_input'] = {'path': manifest['input'], 'sha256': manifest['inputSha256']}
        check(digest(manifest['input']) == manifest['inputSha256'], 'Frontend generation input identity differs')
        evidence['product_files'] = {str((ROOT/path).resolve()): sha for path, sha in manifest['productFiles'].items()}
        check(all(digest(path) == sha for path, sha in evidence['product_files'].items()),
              'Actual frontend snapshots are not bound to the current product/generator files')
        models = {}
        for filename in manifest['files']:
            path = (inputs/filename).resolve()
            check(path.parent == inputs and filename not in models, 'Snapshot path must be unique and directly inside input folder')
            evidence['input_hashes'][str(path)] = digest(path)
            models[filename] = model_from(path)
        check(models and manifest.get('operations'), 'Missing actual frontend operation evidence')
        for operation in manifest['operations']:
            name = operation['name']
            before, after = models[operation['before']], models[operation['after']]
            result = deletion_expectation(operation, before, after)
            for stage, expected in [('undo', before), ('redo', after), ('draft', after)]:
                check(stage in operation, name+': missing actual '+stage+' snapshot')
                check(models[operation[stage]] == expected, name+': exact '+stage+' identity differs')
                result[stage+'_entire_model_equal'] = True
            evidence['operations'][name] = result
        # A bounded no-op golden control, not a broad historical replay.
        prior = json.loads(PRIOR.read_text())
        imported = geometry.run_cli(cli, evidence, 'import-docx', original)['document']
        check(model_from(Path(manifest['input'])) == imported, 'Frontend generation input differs from actual unchanged original import')
        retained_initial = prior['frontend_snapshots']['initial']
        check(digest(retained_initial['model_source_path']) == retained_initial['model_source_sha256'],
              'Retained initial frontend identity differs')
        check(models['initial.json'] == model_from(Path(retained_initial['model_source_path'])),
              'Actual initial frontend snapshot differs from the retained original model')
        evidence['initial_full_retained_model_equal'] = True
        noop = geometry.create(cli, evidence, directory, 'original-noop', imported)
        historical = prior['sources']['unchanged-original']['export']
        check(digest(historical['path']) == historical['sha256'] == noop['sha256'],
              'Original no-op export no longer matches the retained whole-package golden')
        evidence['original_noop'] = {**noop, 'retained_golden': historical['path'],
                                     'whole_package_bytes_equal': True}
        legacy = json.loads(LEGACY.read_text())
        old = next(g for g in legacy['goldens'] if g['source'].endswith('/ordinary-list-export.docx'))
        check(digest(old['source']) == old['source_sha256'] and digest(old['model_input']) == old['model_input_sha256'],
              'Retained ordinary-list input identities differ')
        old_model = model_from(Path(old['model_input']))
        check('markerFont' not in canonical(old_model) and 'listGeometry' not in canonical(old_model),
              'Old writer control unexpectedly carries new list attributes')
        old_export = geometry.create(cli, evidence, directory, 'ordinary-list-legacy', old_model)
        check(old_export['sha256'] == old['source_sha256'], 'Legacy ordinary-list whole package changed')
        evidence['ordinary_list_legacy'] = {**old_export, 'retained_golden': old['source'],
                                             'retained_model': old['model_input'],
                                             'whole_package_bytes_equal': True}
        # Every actual snapshot is exported/reopened, including undo/redo/drafts.
        # Render only distinct full models, identifying reuse explicitly.
        rendered = {}
        if args.render:
            soffice = args.soffice or shutil.which('soffice')
            check(soffice, 'LibreOffice is required')
            evidence['libreoffice'] = {'path': soffice, 'sha256': digest(soffice),
                                      'version': subprocess.run([soffice, '-env:UserInstallation='+(directory/'version-profile').as_uri(), '--version'],
                                                                text=True, capture_output=True, check=True).stdout.strip()}
        for filename, model in models.items():
            label = Path(filename).stem
            result = evidence['snapshots'][filename] = {'source': str(inputs/filename),
                                                        'source_sha256': digest(inputs/filename)}
            output = geometry.create(cli, evidence, directory, label+'-output', model)
            result['export'] = output
            geometry.compare_rows(output['independent']['rows'], geometry.model_rows(model), label+' XML/model')
            check(output['independent']['pageGeometry'] == model['attrs']['pageGeometry'], label+': page geometry changed')
            expected_format = model_formatting(model)
            actual_format = inspect_formatting(Path(output['path']))
            check(actual_format == expected_format, label+': output run/paragraph/empty-font formatting differs: '+canonical({'actual': actual_format, 'expected': expected_format}))
            result['all_paragraph_and_run_formatting_equal'] = True
            key = canonical(model)
            if key in rendered:
                previous = evidence['snapshots'][rendered[key]]
                check(output['sha256'] == previous['export']['sha256'], label+': equal models produced different package bytes')
                result['identical_model_and_bytes_to'] = rendered[key]
                continue
            rendered[key] = filename
            reference = directory/(label+'-reference.docx')
            independent = reference_from_model(original, model, reference)
            result['reference'] = {'path': str(reference), 'sha256': digest(reference), 'independent': independent,
                                   'method': 'Fresh python-docx paragraphs and independently constructed OOXML numbering; no application writer/resolver'}
            geometry.compare_rows(independent['rows'], geometry.model_rows(model), label+' reference/model')
            check(independent['pageGeometry'] == model['attrs']['pageGeometry'], label+': reference page geometry differs')
            check(inspect_formatting(reference) == expected_format, label+': independent reference formatting differs')
            if args.render:
                check(shutil.disk_usage(directory).free > 300*1024*1024, 'Less than 300 MiB free; pause consumer rendering')
                result['reference_pdf'] = geometry.render(soffice, directory, reference, label+'-reference')
                result['output_pdf'] = geometry.render(soffice, directory, Path(output['path']), label+'-output')
                comparison = geometry.compare_render(result['reference_pdf'], result['output_pdf'])
                evidence['render_comparisons'][filename] = comparison
                check(not comparison['has_layout_drift'], label+': exact PDF coordinates/line breaks differ: '+canonical(comparison))
                check(comparison['raster_comparison'] and all(p.get('pixel_equal') for p in comparison['raster_comparison']), label+': full-page pixels differ')
        for target in (original, Path(historical['path']), Path(old['source'])):
            before = digest(target)
            error = geometry.run_cli(cli, evidence, 'create-docx', noop['model_path'], target, success=False)
            check('Output already exists' in error and digest(target) == before, 'Protected source/golden target overwritten')
        evidence['summary'] = {'operations': len(evidence['operations']), 'actual_snapshot_files': len(models),
                               'distinct_models': len(rendered), 'trusted_reopens': len(models)+2,
                               'rendered_pdfs': 2*len(evidence['render_comparisons']),
                               'exact_coordinate_and_pixel_pairs': len(evidence['render_comparisons']),
                               'nonempty_render_pairs': sum(bool(evidence['snapshots'][name]['reference_pdf']['lines'])
                                                            for name in evidence['render_comparisons']),
                               'blank_render_pairs': sum(not evidence['snapshots'][name]['reference_pdf']['lines']
                                                         for name in evidence['render_comparisons']),
                               'existing_targets_protected': len(models)+5,
                               'whole_package_noop_goldens': 2}
        evidence['status'] = 'passed'
    except Exception as error:
        evidence['status'] = 'failed'; evidence['error'] = type(error).__name__+': '+str(error)
        raise
    finally:
        evidence['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
        evidence['script_unchanged'] = digest(__file__) == evidence['script_sha256']
        evidence['helper_unchanged'] = digest(HELPER) == HELPER_SHA
        evidence['cli_unchanged'] = digest(cli) == CLI_SHA
        evidence['genuine_original_unchanged'] = digest(original) == geometry.ORIGINAL_SHA
        evidence['historical_report_unchanged'] = digest(PRIOR) == PRIOR_SHA
        evidence['legacy_report_unchanged'] = digest(LEGACY) == LEGACY_SHA
        evidence['frontend_inputs_unchanged'] = all(digest(p) == h for p, h in evidence['input_hashes'].items())
        evidence['product_files_unchanged'] = all(digest(p) == h for p, h in evidence.get('product_files', {}).items())
        evidence['generation_input_unchanged'] = ('generation_input' not in evidence or
                                                 digest(evidence['generation_input']['path']) == evidence['generation_input']['sha256'])
        evidence['manifest_unchanged'] = 'manifest' not in evidence or digest(manifest_path) == evidence['manifest']['sha256']
        report = directory/'report.json'; save_json(report, evidence); print(report)


if __name__ == '__main__':
    main()
