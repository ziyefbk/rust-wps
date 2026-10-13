#!/usr/bin/env python3
"""Independent signed paragraph margin-overhang checks.

The application is used only through its frozen CLI. Expectations, python-docx
sources/references, physical bounds, OOXML reader, PDF coordinates and pixels
are independent of application code. Fresh LibreOffice profile for every file.
Original indents.docx is hash-pinned and never edited/re-saved in place.
No browser/native Word/WPS or arbitrary-glyph/full-document fidelity claim.
"""
import argparse
import copy
import hashlib
import json
import os
import platform
import shutil
import subprocess
import sys
from datetime import datetime, timezone
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree as ET
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parent.parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
KEYS = ('indentLeftTwips', 'indentRightTwips', 'firstLineIndentTwips')
ORIGINAL_SHA = 'd62a3e3162a737cbe45e21fe544f5dc845e8cdc621bd8a2e3b7e8dcfed52386d'
PAGE = dict(widthTwips=11906, heightTwips=16838, marginTopTwips=1440,
            marginRightTwips=1440, marginBottomTwips=1440, marginLeftTwips=1440)
LIMITATIONS = ('Independent synthetic controls and one unchanged upstream regression fixture; '
               'not a user production document. PDF/pixel evidence is LibreOfficeDev only. '
               'No real app/browser, WebView2, native printing, Microsoft Word or WPS acceptance. '
               'Physical line-box bounds do not prove all glyphs or unbreakable words fit.')


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def add(parent, tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def remove(parent, tag):
    for node in list(parent.findall('w:' + tag, NS)):
        parent.remove(node)


def set_indent(para, values):
    fmt = para.paragraph_format
    for name, value in zip(('left_indent', 'right_indent', 'first_line_indent'), values):
        setattr(fmt, name, None if value is None else Twips(value))


def attrs(values):
    return {key: value for key, value in zip(KEYS, values) if value is not None}


def raw_indent(parent):
    node = None if parent is None else parent.find('w:ind', NS)
    if node is None:
        return {}
    check(not (node.get(qn('w:firstLine')) is not None and node.get(qn('w:hanging')) is not None),
          'Independent supported reader refuses competing firstLine/hanging')
    out = {}
    for source, target in zip(('left', 'right'), KEYS[:2]):
        if node.get(qn('w:' + source)) is not None:
            out[target] = int(node.get(qn('w:' + source)))
    if node.get(qn('w:firstLine')) is not None:
        out[KEYS[2]] = int(node.get(qn('w:firstLine')))
    elif node.get(qn('w:hanging')) is not None:
        out[KEYS[2]] = -int(node.get(qn('w:hanging')))
    return out


def inspect_docx(path):
    """Independent unconditional paragraph-style resolver; raw and python-docx cross-check."""
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read('word/document.xml'))
        styles = ET.fromstring(package.read('word/styles.xml'))
    by_id = {s.get(qn('w:styleId')): s for s in styles.findall('w:style', NS)}
    default = next((s.get(qn('w:styleId')) for s in by_id.values()
                    if s.get(qn('w:type')) == 'paragraph' and s.get(qn('w:default')) in ('1', 'true', 'on')), None)
    defaults = raw_indent(styles.find('w:docDefaults/w:pPrDefault/w:pPr', NS))
    def chain(style_id, seen=()):
        check(style_id not in seen, 'Style cycle in independent controlled reader')
        style = by_id.get(style_id)
        if style is None:
            return []
        base = style.find('w:basedOn', NS)
        return ([] if base is None else chain(base.get(qn('w:val')), seen + (style_id,))) + [style]
    rows = []
    for para in root.findall('.//w:body//w:p', NS):
        ppr = para.find('w:pPr', NS)
        selected = None if ppr is None else ppr.find('w:pStyle', NS)
        sid = default if selected is None else selected.get(qn('w:val'))
        effective = dict(defaults)
        for style in chain(sid):
            effective.update(raw_indent(style.find('w:pPr', NS)))
        effective.update(raw_indent(ppr))
        text = ''.join((el.text or '') if el.tag == qn('w:t') else '\n'
                       for el in para.iter() if el.tag in (qn('w:t'), qn('w:br')))
        rows.append({'text': text, 'indent': effective, 'direct': raw_indent(ppr),
                     'in_table': any(a.tag == qn('w:tc') for a in para.iterancestors()), 'style': sid})
    python_rows = []
    doc = Document(path)
    for element in doc.element.body.iter(qn('w:p')):
        para = Paragraph(element, doc)
        f = para.paragraph_format
        python_rows.append({'text': para.text, 'direct': attrs([None if v is None else v.twips
                            for v in (f.left_indent, f.right_indent, f.first_line_indent)])})
    check(len(rows) == len(python_rows), 'OOXML/python-docx paragraph count differs')
    for xml, py in zip(rows, python_rows):
        check(xml['text'] == py['text'] and xml['direct'] == py['direct'], 'OOXML/python-docx disagreement')
    sec = root.find('w:body/w:sectPr', NS)
    geometry = None
    if sec is not None:
        size, margins = sec.find('w:pgSz', NS), sec.find('w:pgMar', NS)
        if size is not None and margins is not None:
            geometry = {'widthTwips': int(size.get(qn('w:w'))), 'heightTwips': int(size.get(qn('w:h')))}
            geometry.update({'margin' + edge.title() + 'Twips': int(margins.get(qn('w:' + edge)))
                             for edge in ('top', 'right', 'bottom', 'left')})
    return {'ooxml': rows, 'python_docx': python_rows, 'pageGeometry': geometry}


def model_rows(model, parents=()):
    def text(node):
        if node['type'] == 'hardBreak':
            return '\n'
        return node.get('text', '') + ''.join(text(c) for c in node.get('content', []))
    if model['type'] in ('paragraph', 'heading'):
        yield {'text': text(model), 'indent': {k: model.get('attrs', {})[k] for k in KEYS if k in model.get('attrs', {})},
               'in_table': 'table' in parents, 'type': model['type']}
    for child in model.get('content', []):
        yield from model_rows(child, parents + (model['type'],))


def verify_rows(actual, expected, label):
    check(len(actual) == len(expected), label + ': paragraph count differs')
    for i, (left, right) in enumerate(zip(actual, expected)):
        for key in ('text', 'indent', 'in_table'):
            check(left[key] == right[key], f'{label} paragraph {i} {key}: {left[key]!r} != {right[key]!r}')
        if 'type' in left and 'type' in right:
            check(left['type'] == right['type'], label + ': paragraph/heading type differs')


def physical_bounds(values, geometry):
    l, r, f = [0 if v is None else v for v in values]
    left, right, width = geometry['marginLeftTwips'], geometry['marginRightTwips'], geometry['widthTwips']
    start, first, end = left + l, left + l + f, width - right - r
    check(min(start, first) >= 0 and max(start, first) < end <= width, 'Independent physical line-box proof failed')
    return {'continuation_start_twips': start, 'first_start_twips': first, 'end_twips': end,
            'continuation_width_twips': end - start, 'first_width_twips': end - first}


def document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(PAGE['widthTwips']), Twips(PAGE['heightTwips'])
    sec.top_margin = sec.right_margin = sec.bottom_margin = sec.left_margin = Twips(1440)
    remove(sec._sectPr, 'docGrid')
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'DejaVu Sans', Pt(12)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for node in list(doc.styles.element.findall('.//w:contextualSpacing', NS)):
        node.getparent().remove(node)
    return doc


def paragraph(doc, values, prefix='Probe', heading=False, align=None, long=True):
    text = ' '.join(f'{prefix}{i:02}' for i in range(48)) if long else prefix
    para = doc.add_heading(text, level=2) if heading else doc.add_paragraph(text)
    set_indent(para, values)
    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1
    if align is not None:
        para.alignment = align
    for run in para.runs:
        run.font.name, run.font.size = 'DejaVu Sans', Pt(12)
        run.bold = heading
    return para


def mutate(path, callback):
    with zipfile.ZipFile(path) as package:
        parts = {n: package.read(n) for n in package.namelist()}
    main = ET.fromstring(parts['word/document.xml'])
    styles = ET.fromstring(parts['word/styles.xml'])
    settings = ET.fromstring(parts['word/settings.xml'])
    callback(main, styles, settings, parts)
    for name, root in [('document', main), ('styles', styles), ('settings', settings)]:
        parts['word/' + name + '.xml'] = ET.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def fixtures(directory):
    sources = {}
    def save(name, doc, supported=True, render=False, mutation=None, values=None, expected=None):
        path = directory / (name + '.docx')
        check(not path.exists(), 'Fixture must not overwrite')
        doc.save(path)
        if mutation:
            mutate(path, mutation)
        item = {'path': str(path), 'sha256': digest(path), 'fixture_kind': 'synthetic directed python-docx/OOXML',
                'expected_action': 'roundtrip' if supported else 'reject', 'render': render}
        if supported:
            snapshot = inspect_docx(path)
            item['independent'] = snapshot
            item['expected'] = expected if expected is not None else [dict(r) for r in snapshot['ooxml']]
            if values is not None:
                check(len(snapshot['ooxml']) == 1 and snapshot['ooxml'][0]['indent'] == attrs(values), name + ': authored declarations differ')
                item['physical_bounds'] = physical_bounds(values, PAGE)
                item['values'] = values
            verify_rows(snapshot['ooxml'], item['expected'], name + ' source')
        sources[name] = item
    probes = {
        'negative-left-positive-first': (-720, -360, 240),
        'negative-left-hanging': (-360, 0, -720),
        'positive-left-hanging': (360, -480, -720),
        'negative-right-positive-first': (360, -720, 240),
        'heading-negative-sides': (-720, -360, -360),
        'left-physical-edge': (-1440, 0, 0),
        'first-physical-edge': (-720, 0, -720),
        'right-physical-edge': (0, -1440, 0),
        'both-physical-edges': (-1440, -1440, 240),
        'one-twip-inside-edge': (-1439, -1439, 1),
        'ordinary-positive-control': (720, 360, -360),
        'explicit-zero-control': (0, 0, 0),
    }
    for name, values in probes.items():
        doc = document()
        paragraph(doc, values, heading=name.startswith('heading'))
        save(name, doc, render=True, values=values)
    for name, align in [('signed-right-aligned', WD_ALIGN_PARAGRAPH.RIGHT), ('signed-centered', WD_ALIGN_PARAGRAPH.CENTER)]:
        doc = document(); values = (-360, -720, 240)
        paragraph(doc, values, align=align); save(name, doc, render=True, values=values)
    doc = document(); paragraph(doc, (10465, -1440, 0), long=False)
    save('one-twip-line-width', doc, values=(10465, -1440, 0))
    doc = document(); paragraph(doc, (0, 0, None), long=False)
    save('absent-first-control', doc, values=(0, 0, None))
    for name in ['inherit-partial', 'inherit-opposite-slot', 'inherit-explicit-zero', 'docdefaults-signed']:
        doc = document()
        para = paragraph(doc, (None, None, None), prefix='Cascade')
        parent = doc.styles.add_style('OverhangParent', WD_STYLE_TYPE.PARAGRAPH)
        parent.base_style = doc.styles['Normal']; set_indent(parent, (-720, -480, 360))
        child = doc.styles.add_style('OverhangChild', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = parent; para.style = child
        if name == 'inherit-partial':
            set_indent(para, (-360, None, None)); expected = (-360, -480, 360)
        elif name == 'inherit-opposite-slot':
            set_indent(child, (None, None, -360)); expected = (-720, -480, -360)
        elif name == 'inherit-explicit-zero':
            set_indent(child, (0, 0, 0)); expected = (0, 0, 0)
        else:
            para.style = doc.styles['Normal']
            default = doc.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr', NS)
            add(default, 'ind', left=-720, right=-480, hanging=360); expected = (-720, -480, -360)
        save(name, doc, render=True, values=expected)
    # Existing ordinary table and list support must remain independent.
    doc = document(); tab = doc.add_table(rows=1, cols=1); tab.autofit = False
    tab.columns[0].width = Twips(3600); tab.cell(0, 0).width = Twips(3600)
    tab.cell(0, 0).text = 'TableControl'; set_indent(tab.cell(0, 0).paragraphs[0], (180, 120, -120))
    save('table-positive-control', doc)
    doc = document(); doc.add_paragraph('Ordinary list', style='List Number'); save('ordinary-list-control', doc)
    invalid_values = {
        'left-off-paper': (-1441, 0, 0), 'right-off-paper': (0, -1441, 0),
        'first-off-paper': (-720, 0, -721), 'positive-left-hang-off-paper': (360, 0, -1801),
        'zero-width': (9026, 0, 0), 'negative-width': (9027, 0, 0),
        'first-zero-width': (0, 0, 9026), 'first-negative-width': (0, 0, 9027),
        'negative-right-zero-width': (10466, -1440, 0),
        'negative-left-too-large': (-31681, 0, 0), 'negative-right-too-large': (0, -31681, 0),
        'positive-left-too-large': (31681, 0, 0), 'first-too-large': (0, 0, 31681),
    }
    for name, values in invalid_values.items():
        doc = document(); paragraph(doc, values, long=False)
        save('reject-' + name, doc, False)
    for name in ['missing-section', 'missing-size', 'missing-margins', 'two-columns', 'gutter', 'mirror-margins']:
        doc = document(); paragraph(doc, (-720, -360, -240), long=False)
        def change(root, styles, settings, parts, name=name):
            sec = root.find('w:body/w:sectPr', NS)
            if name == 'missing-section': sec.getparent().remove(sec)
            elif name == 'missing-size': remove(sec, 'pgSz')
            elif name == 'missing-margins': remove(sec, 'pgMar')
            elif name == 'two-columns': remove(sec, 'cols'); add(sec, 'cols', num=2, space=720)
            elif name == 'gutter': sec.find('w:pgMar', NS).set(qn('w:gutter'), '120')
            else: add(settings, 'mirrorMargins', val=1)
        save('reject-' + name, doc, False, mutation=change)
    for name, raw in {
        'firstLine-negative': {'left': '-720', 'firstLine': '-1'},
        'hanging-negative': {'left': '-720', 'hanging': '-1'},
        'logical-start': {'start': '-720'}, 'logical-end': {'end': '-720'},
        'character-left': {'leftChars': '-100'}, 'fractional-left': {'left': '-720.5'},
        'dual-slot': {'left': '-720', 'firstLine': '120', 'hanging': '60'},
        'dual-zero-slot': {'left': '-720', 'firstLine': '0', 'hanging': '0'},
    }.items():
        doc = document(); para = paragraph(doc, (None, None, None), long=False)
        add(para._p.get_or_add_pPr(), 'ind', **raw); save('reject-' + name, doc, False)
    for name in ['tab-content', 'tab-stop', 'bidi', 'mirror-indents', 'list-negative', 'table-negative-left',
                 'table-negative-right', 'table-hanging', 'table-negative-left-page-margin']:
        doc = document()
        if name.startswith('table'):
            tab = doc.add_table(rows=1, cols=1); tab.autofit = False
            tab.columns[0].width = Twips(3600); tab.cell(0, 0).width = Twips(3600)
            para = tab.cell(0, 0).paragraphs[0]; para.text = 'Cell refusal'
            values = (0, -1, 0) if name.endswith('right') else ((0, 0, -1) if name.endswith('hanging') else (-1, 0, 0))
        else:
            para = doc.add_paragraph('Tab\tContent' if name == 'tab-content' else 'Overhang refusal',
                                     style='List Number' if name == 'list-negative' else None)
            values = (-360, 0, -360)
        set_indent(para, values)
        if name == 'tab-stop': add(add(para._p.get_or_add_pPr(), 'tabs'), 'tab', val='right', pos=1440)
        if name == 'bidi': add(para._p.get_or_add_pPr(), 'bidi', val=1)
        if name == 'mirror-indents': add(para._p.get_or_add_pPr(), 'mirrorIndents', val=1)
        save('reject-' + name, doc, False)
    for name, values in [('missing-page-hanging-only', (360, 0, -720)), ('missing-page-right-only', (0, -360, 0)),
                         ('missing-page-left-only', (-360, 0, 0))]:
        doc = document(); paragraph(doc, values, long=False)
        def missing_page(root, styles, settings, parts):
            sec = root.find('w:body/w:sectPr', NS); sec.getparent().remove(sec)
        save('reject-' + name, doc, False, mutation=missing_page)
    return sources


def run_cli(cli, evidence, *args, success=True):
    result = subprocess.run([str(cli), *map(str, args)], text=True, capture_output=True, timeout=90)
    entry = {'program': str(cli), 'arguments': list(map(str, args)), 'returncode': result.returncode,
             'stdout': result.stdout, 'stderr': result.stderr}
    evidence['commands'].append(entry)
    check((result.returncode == 0) == success, 'Unexpected CLI result: ' + json.dumps(entry))
    return json.loads(result.stdout) if success else result.stderr


def create(cli, evidence, directory, name, model):
    model_path, output = directory / (name + '.json'), directory / (name + '.docx')
    save_json(model_path, model)
    run_cli(cli, evidence, 'create-docx', model_path, output)
    reopened = run_cli(cli, evidence, 'read-docx-model', output)
    check(reopened == model, name + ': entire trusted model changed')
    before = digest(output)
    error = run_cli(cli, evidence, 'create-docx', model_path, output, success=False)
    check('Output already exists' in error and digest(output) == before, name + ': no-clobber failed')
    return {'path': str(output), 'sha256': before, 'model_path': str(model_path),
            'complete_trusted_model_equal': True, 'no_clobber': True, 'independent': inspect_docx(output)}


def roundtrip(cli, evidence, directory, name, item):
    path = Path(item['path'])
    item['audit'] = run_cli(cli, evidence, 'audit-docx', path)
    imported = run_cli(cli, evidence, 'import-docx', path)
    item['import_response'] = imported
    model = imported['document']; expected = item['expected']
    # Lists have an explicitly normalized numbering path. Their ordinary
    # paragraph default indent is not misrepresented as preserved metadata.
    if name != 'ordinary-list-control':
        verify_rows(list(model_rows(model)), expected, name + ' import')
    if item['independent']['pageGeometry'] is not None:
        check(model.get('attrs', {}).get('pageGeometry') == item['independent']['pageGeometry'], name + ': explicit page proof missing/changed')
    out = create(cli, evidence, directory, name + '-export', model)
    item['export'] = out
    if name != 'ordinary-list-control':
        verify_rows(out['independent']['ooxml'], expected, name + ' independent export')
    reimport = run_cli(cli, evidence, 'import-docx', out['path'])['document']
    verify_rows(list(model_rows(reimport)), list(model_rows(model)), name + ' external reimport')
    item['external_reimport_complete_model_equal'] = reimport == model
    error = run_cli(cli, evidence, 'create-docx', out['model_path'], path, success=False)
    check('Output already exists' in error and digest(path) == item['sha256'], name + ': original/source no-clobber failed')
    item['source_no_clobber'] = True
    return model


def edited(cli, evidence, directory, sources):
    results = {}
    for name in ['negative-left-positive-first', 'negative-left-hanging', 'negative-right-positive-first', 'heading-negative-sides', 'original-indents']:
        item = sources[name]
        model = copy.deepcopy(item['import_response']['document'])
        edits = {0: (-900, -600, 300)} if name != 'original-indents' else {1: (-1000, -300, 200), 2: (-400, -600, -500)}
        reference = Document(item['path'])
        for index, values in edits.items():
            model['content'][index].setdefault('attrs', {}).update(attrs(values))
            set_indent(reference.paragraphs[index], values)
            physical_bounds(values, model['attrs']['pageGeometry'])
        ref_path = directory / (name + '-edited-reference.docx'); reference.save(ref_path)
        ref = inspect_docx(ref_path)
        output = create(cli, evidence, directory, name + '-edited-output', model)
        verify_rows(output['independent']['ooxml'], ref['ooxml'], name + ' edited vs independently authored reference')
        results[name] = {'reference_path': str(ref_path), 'reference_sha256': digest(ref_path), 'export': output,
                         'operation': 'actual imported full model attr edits followed by create-docx and trusted reopen',
                         'edited_paragraph_values': edits}
    evidence['edited'] = results


def invalid_models(cli, evidence, directory, good):
    cases = {}
    def add_case(name, mutate_model):
        model = copy.deepcopy(good); mutate_model(model); cases[name] = model
    def para(m): return m['content'][0].setdefault('attrs', {})
    for key in KEYS:
        for value in [None, '12', 1.5, True, -31681, 31681]:
            add_case(key + '-' + str(value), lambda m, k=key, v=value: para(m).__setitem__(k, v))
    for name, values in [('left-off-paper', (-1441, 0, 0)), ('right-off-paper', (0, -1441, 0)),
                         ('first-off-paper', (-720, 0, -721)), ('zero-width', (9026, 0, 0)),
                         ('first-zero-width', (0, 0, 9026)), ('negative-right-zero-width', (10466, -1440, 0))]:
        add_case(name, lambda m, v=values: para(m).update(attrs(v)))
    add_case('missing-page', lambda m: m['attrs'].pop('pageGeometry'))
    add_case('null-page', lambda m: m['attrs'].__setitem__('pageGeometry', None))
    add_case('shrink-left-margin', lambda m: m['attrs']['pageGeometry'].__setitem__('marginLeftTwips', 719))
    add_case('shrink-right-margin', lambda m: m['attrs']['pageGeometry'].__setitem__('marginRightTwips', 359))
    for kind in ['table', 'bulletList', 'orderedList']:
        def wrap(m, kind=kind):
            p = m['content'][0]
            if kind == 'table': m['content'] = [{'type': 'table', 'attrs': {'columnWidthsTwips': [5000]}, 'content': [{'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [p]}]}]}]
            else: m['content'] = [{'type': kind, 'content': [{'type': 'listItem', 'content': [p]}]}]
        add_case('container-' + kind, wrap)
    results = {}
    for name, model in cases.items():
        path, output = directory / ('invalid-' + name + '.json'), directory / ('invalid-' + name + '.docx')
        save_json(path, model)
        error = run_cli(cli, evidence, 'create-docx', path, output, success=False)
        check(not output.exists(), name + ': rejected model produced output')
        sentinel = directory / ('invalid-' + name + '-existing.docx'); sentinel.write_bytes(b'unchanged existing destination')
        before = digest(sentinel)
        run_cli(cli, evidence, 'create-docx', path, sentinel, success=False)
        check(digest(sentinel) == before, name + ': invalid model clobbered existing destination')
        results[name] = {'error': error, 'no_partial_output': True, 'existing_destination_unchanged': True}
    evidence['invalid_models'] = results


def tamper(cli, evidence, directory, exported):
    with zipfile.ZipFile(exported) as package:
        original = {name: package.read(name) for name in package.namelist()}
    member = 'rust-wps/model-v2.json'
    check(member in original, 'Expected complete V2 model member')
    results = {}
    for mode in ['model-signed-value', 'model-missing-page', 'xml-signed-value', 'xml-page-margin', 'xml-heading-type',
                 'xml-signed-value-forged-hash', 'xml-page-margin-forged-hash', 'xml-heading-type-forged-hash']:
        parts = dict(original)
        base_mode = mode.removesuffix('-forged-hash')
        if mode.startswith('model'):
            model = json.loads(parts[member])
            # Embedded model envelope is preserved; locate the document node.
            doc = model.get('document', model)
            if mode == 'model-signed-value': doc['content'][0]['attrs'][KEYS[0]] = -719
            else: doc['attrs'].pop('pageGeometry')
            parts[member] = json.dumps(model, ensure_ascii=False, separators=(',', ':')).encode()
        else:
            root = ET.fromstring(parts['word/document.xml'])
            if base_mode == 'xml-signed-value': root.find('w:body/w:p/w:pPr/w:ind', NS).set(qn('w:left'), '-719')
            elif base_mode == 'xml-page-margin': root.find('w:body/w:sectPr/w:pgMar', NS).set(qn('w:left'), '1439')
            else:
                ppr = root.find('w:body/w:p/w:pPr', NS); remove(ppr, 'pStyle'); add(ppr, 'pStyle', val='Heading2')
            parts['word/document.xml'] = ET.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        if mode.endswith('forged-hash'):
            envelope = json.loads(parts[member])
            for name in envelope['hashes']:
                envelope['hashes'][name] = hashlib.sha256(parts[name]).hexdigest()
            parts[member] = json.dumps(envelope, ensure_ascii=False, separators=(',', ':')).encode()
        path = directory / ('tampered-' + mode + '.docx')
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as package:
            for name, data in parts.items(): package.writestr(name, data)
        results[mode] = {'path': str(path), 'sha256': digest(path), 'error': run_cli(cli, evidence, 'read-docx-model', path, success=False)}
    evidence['tampered_packages'] = results


def render(soffice, directory, path, label):
    area = directory / 'renders' / label
    area.mkdir(parents=True)
    profile, home = area / 'lo-profile', area / 'home'
    home.mkdir()
    before = digest(path)
    command = [soffice, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(area), str(path)]
    run = subprocess.run(command, capture_output=True, text=True, timeout=120,
                         env=dict(os.environ, HOME=str(home), XDG_CACHE_HOME=str(area / 'cache')))
    (area / 'stdout.txt').write_text(run.stdout); (area / 'stderr.txt').write_text(run.stderr)
    pdf = area / (Path(path).stem + '.pdf')
    check(run.returncode == 0 and pdf.is_file() and pdf.stat().st_size, label + ': fresh isolated LibreOffice render failed')
    check(digest(path) == before, label + ': consumer mutated input')
    bbox = area / 'bbox.html'
    subprocess.run(['pdftotext', '-bbox-layout', str(pdf), str(bbox)], capture_output=True, check=True, timeout=30)
    tree = ET.parse(str(bbox)); lines = []; pages = []
    for pi, page in enumerate(tree.xpath('//*[local-name()="page"]')):
        pages.append({'width': float(page.get('width')), 'height': float(page.get('height'))})
        for line in page.xpath('.//*[local-name()="line"]'):
            words = line.xpath('./*[local-name()="word"]')
            if not words: continue
            lines.append({'page': pi + 1, 'text': ' '.join(''.join(w.itertext()) for w in words),
                          **{k: float(line.get(k)) for k in ['xMin', 'yMin', 'xMax', 'yMax']}})
    subprocess.run(['pdftoppm', '-r', '90', '-png', str(pdf), str(area / 'page')], capture_output=True, check=True, timeout=60)
    pngs = sorted(area.glob('page-*.png'))
    check(len(pngs) == len(pages), label + ': raster page count differs')
    return {'path': str(pdf), 'sha256': digest(pdf), 'input_sha256': before, 'profile': str(profile), 'command': command,
            'fresh_profile': True, 'pages': pages, 'lines': lines, 'pngs': [str(p) for p in pngs],
            'png_sha256': {p.name: digest(p) for p in pngs}}


def compare_render(a, b):
    same_text = [(r['page'], r['text']) for r in a['lines']] == [(r['page'], r['text']) for r in b['lines']]
    result = {'same_pages': a['pages'] == b['pages'], 'same_line_breaks': same_text,
              'source_line_count': len(a['lines']), 'output_line_count': len(b['lines'])}
    if same_text:
        result['max_horizontal_delta_pt'] = max((abs(x[k] - y[k]) for x, y in zip(a['lines'], b['lines']) for k in ['xMin', 'xMax']), default=0)
        result['max_vertical_delta_pt'] = max((abs(x[k] - y[k]) for x, y in zip(a['lines'], b['lines']) for k in ['yMin', 'yMax']), default=0)
    pixels = []
    for left, right in zip(a['pngs'], b['pngs']):
        x, y = Image.open(left).convert('RGB'), Image.open(right).convert('RGB')
        if x.size != y.size:
            pixels.append({'same_size': False}); continue
        diff = ImageChops.difference(x, y)
        pixels.append({'same_size': True, 'pixel_equal': diff.getbbox() is None,
                       'changed_pixel_count': x.width * x.height - ImageChops.lighter(ImageChops.lighter(diff.getchannel('R'), diff.getchannel('G')), diff.getchannel('B')).histogram()[0], 'pixels': x.width * x.height})
    result['raster_comparison'] = pixels
    result['has_layout_drift'] = (not result['same_pages'] or not same_text or result.get('max_horizontal_delta_pt', 0) > .100001 or result.get('max_vertical_delta_pt', 0) > .100001)
    return result


def assert_positions(item, pdf, name):
    if name == 'original-indents':
        expected_x = [70.85, 14.15, 42.5, 14.15, 99.2]
        expected_text = ['Some indents', 'Negative 2cm, no firstline', 'Negative 2cm, first line 1cm',
                         'Negative 1cm, first line -1cm', 'Positive indent 2cm, first line -1cm']
        check([r['text'] for r in pdf['lines']] == expected_text, 'Original five visible paragraphs/text order changed')
        check(all(abs(r['xMin'] - x) <= .11 for r, x in zip(pdf['lines'], expected_x)),
              'Original physical paragraph starts differ from independently fixed twip geometry')
        return {'expected_first_starts_x_pt': expected_x, 'actual_first_starts_x_pt': [r['xMin'] for r in pdf['lines']],
                'tolerance_pt': .11, 'five_original_texts_equal': True}
    if name.startswith('signed-') or 'physical_bounds' not in item:
        return None
    rows = pdf['lines']; check(len(rows) >= 2, name + ': expected wrapped continuation')
    bounds = item['physical_bounds']
    expected_first, expected_next = bounds['first_start_twips'] / 20, bounds['continuation_start_twips'] / 20
    check(abs(rows[0]['xMin'] - expected_first) <= .11, name + ': first baseline does not match independent physical origin')
    check(all(abs(row['xMin'] - expected_next) <= .11 for row in rows[1:]), name + ': continuation origin mismatch')
    return {'expected_first_x_pt': expected_first, 'expected_continuation_x_pt': expected_next,
            'actual_first_x_pt': rows[0]['xMin'], 'actual_continuation_x_pt': [r['xMin'] for r in rows[1:]]}


def render_all(soffice, directory, evidence, sources, generated_only):
    comparisons = {}
    for name, item in sources.items():
        if not item.get('render'): continue
        item['source_pdf'] = render(soffice, directory, Path(item['path']), name + '-source')
        item['positions'] = assert_positions(item, item['source_pdf'], name)
        if not generated_only:
            item['output_pdf'] = render(soffice, directory, Path(item['export']['path']), name + '-output')
            comparisons[name] = compare_render(item['source_pdf'], item['output_pdf'])
            item['output_positions'] = assert_positions(item, item['output_pdf'], name)
            if name != 'original-indents':
                check(not comparisons[name]['has_layout_drift'], name + ': controlled source/output layout drift: ' + json.dumps(comparisons[name]))
    for name, item in evidence.get('edited', {}).items():
        item['reference_pdf'] = render(soffice, directory, Path(item['reference_path']), name + '-edited-reference')
        item['output_pdf'] = render(soffice, directory, Path(item['export']['path']), name + '-edited-output')
        comparisons[name + '-edited'] = compare_render(item['reference_pdf'], item['output_pdf'])
        if name != 'original-indents':
            check(not comparisons[name + '-edited']['has_layout_drift'], name + ': controlled edited/reference layout drift')
    for name, item in evidence.get('frontend_snapshots', {}).items():
        item['reference_pdf'] = render(soffice, directory, Path(item['reference_path']), 'frontend-' + name + '-reference')
        item['output_pdf'] = render(soffice, directory, Path(item['export']['path']), 'frontend-' + name + '-output')
        comparisons['frontend-' + name] = compare_render(item['reference_pdf'], item['output_pdf'])
        check(not comparisons['frontend-' + name]['has_layout_drift'], 'Frontend ' + name + ': independently authored reference layout drift')
    evidence['render_comparisons'] = comparisons
    evidence['layout_drift'] = {name: value for name, value in comparisons.items() if value['has_layout_drift']}
    # Explicit signed right alignment ensures the physical right bound is used,
    # and source/output wrapping comparison covers right indentation semantics.
    item = sources['signed-right-aligned']
    pdf = item['source_pdf'] if generated_only else item['output_pdf']
    expected_right = item['physical_bounds']['end_twips'] / 20
    check(all(abs(row['xMax'] - expected_right) <= .35 for row in pdf['lines']), 'Signed right alignment did not reach independent physical right bound')
    evidence['signed_right_edge'] = {'expected_x_pt': expected_right, 'glyph_endpoint_tolerance_pt': .35, 'max_glyph_endpoint_residual_pt': max(abs(row['xMax'] - expected_right) for row in pdf['lines']), 'scope': 'Glyph-endpoint approximation only; the paragraph frame itself is independently proved from XML/model twips', 'actual_x_pt': [r['xMax'] for r in pdf['lines']]}


def legacy(cli, baseline, directory, evidence, sources):
    plain = {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Historic plain body'}]}]}
    cases = {'plain': plain, 'page': copy.deepcopy(plain), 'positive-indentation': copy.deepcopy(plain),
             'heading': copy.deepcopy(plain), 'spacing': copy.deepcopy(plain)}
    cases['page']['attrs'] = {'pageGeometry': dict(PAGE)}
    cases['positive-indentation']['content'][0]['attrs'] = attrs((720, 360, -360))
    cases['heading']['content'][0]['type'] = 'heading'; cases['heading']['content'][0]['attrs'] = {'level': 2}
    cases['spacing']['content'][0]['attrs'] = {'spaceBefore': 120, 'spaceAfter': 240, 'lineSpacing': 360}
    results = {}
    for name, original in cases.items():
        for profile in ['legacy-v1', 'standard-v2']:
            model = copy.deepcopy(original)
            if profile == 'standard-v2': model.setdefault('attrs', {})['styleProfile'] = profile
            label = 'legacy-' + name + '-' + profile
            model_path = directory / (label + '.json'); save_json(model_path, model)
            outputs = []
            for suffix, program in [('baseline', baseline), ('candidate', cli)]:
                output = directory / (label + '-' + suffix + '.docx')
                run_cli(program, evidence, 'create-docx', model_path, output); outputs.append(output)
                check(run_cli(program, evidence, 'read-docx-model', output) == model, label + ': historic complete model changed')
            check(outputs[0].read_bytes() == outputs[1].read_bytes(), label + ': historic full DOCX bytes changed')
            results[label] = {'sha256': digest(outputs[0]), 'complete_package_byte_equal': True, 'complete_trusted_model_equal': True}
    responses = {}
    for name in ['ordinary-positive-control', 'explicit-zero-control', 'absent-first-control', 'table-positive-control', 'ordinary-list-control',
                 'reject-logical-start', 'reject-table-negative-left', 'original-indents']:
        path = sources[name]['path']
        comparison = {}
        for command in ['audit-docx', 'import-docx']:
            out = []
            for program in [baseline, cli]:
                result = subprocess.run([str(program), command, path], capture_output=True, text=True, timeout=90)
                record = {'program': str(program), 'arguments': [command, path], 'returncode': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr}
                evidence['commands'].append(record); out.append({k: record[k] for k in ['returncode', 'stdout', 'stderr']})
            comparison[command] = {'exact_response_equal': out[0] == out[1], 'baseline': out[0], 'candidate': out[1]}
            if name != 'original-indents': check(out[0] == out[1], name + ': unrelated historic CLI response changed')
        responses[name] = comparison
    evidence['legacy_bytes'] = results; evidence['legacy_responses'] = responses


def frontend_snapshots(cli, evidence, directory, input_directory):
    manifest_path = input_directory / 'manifest.json'
    manifest = json.loads(manifest_path.read_text())
    results = {}
    for filename in manifest['files']:
        source = input_directory / filename
        wrapped = json.loads(source.read_text())
        check(wrapped['version'] == 2, 'Frontend emitted unexpected model version')
        model = wrapped['document']
        doc = document()
        geometry = model['attrs']['pageGeometry']; sec = doc.sections[0]
        sec.page_width, sec.page_height = Twips(geometry['widthTwips']), Twips(geometry['heightTwips'])
        for side in ['top', 'right', 'bottom', 'left']:
            setattr(sec, side + '_margin', Twips(geometry['margin' + side.title() + 'Twips']))
        for node in model['content']:
            check(node['type'] in ['paragraph', 'heading'], 'Frontend reference needs a plain paragraph/heading model')
            a = node.get('attrs', {})
            para = doc.add_heading('', level=a['level']) if node['type'] == 'heading' else doc.add_paragraph()
            set_indent(para, [a.get(k) for k in KEYS])
            para.paragraph_format.space_before = Twips(a.get('spaceBefore', 0))
            para.paragraph_format.space_after = Twips(a.get('spaceAfter', 0))
            para.paragraph_format.line_spacing = a.get('lineSpacing', 240) / 240
            if a.get('textAlign') in ['left', 'center', 'right', 'justify']:
                para.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                                  'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[a['textAlign']]
            for span in node.get('content', []):
                if span['type'] == 'hardBreak': para.add_run().add_break(); continue
                check(span['type'] == 'text', 'Frontend reference needs text-only content')
                run = para.add_run(span['text'])
                marks = {m['type']: m.get('attrs', {}) for m in span.get('marks', [])}
                check(marks.get('fontFamily') == {'name': 'DejaVu Sans'} and marks.get('fontSize') == {'halfPoints': 24},
                      'Frontend evidence must carry independently specified controlled font')
                run.font.name, run.font.size = 'DejaVu Sans', Pt(12)
                # The actual snapshot has no color mark: the standard-v2 model
                # uses black. python-docx's stock heading blue theme is unrelated
                # to this independent model reference and must not leak into it.
                run.font.color.rgb = RGBColor.from_string('000000')
                if 'bold' in marks: run.bold = True
                if 'italic' in marks: run.italic = True
        label = 'frontend-' + source.stem
        reference = directory / (label + '-reference.docx'); doc.save(reference)
        independent = inspect_docx(reference)
        output = create(cli, evidence, directory, label + '-output', model)
        verify_rows(output['independent']['ooxml'], independent['ooxml'], label + ' independently authored reference')
        verify_rows(list(model_rows(model)), independent['ooxml'], label + ' actual frontend model')
        results[source.stem] = {'model_source_path': str(source), 'model_source_sha256': digest(source),
                               'reference_path': str(reference), 'reference_sha256': digest(reference), 'export': output,
                               'reference_control': 'Explicit DejaVu Sans 12pt and black model-default text; no stock python-docx heading theme color'}
    check(json.loads((input_directory / 'undo.json').read_text()) == json.loads((input_directory / 'initial.json').read_text()), 'Frontend undo identity changed')
    check(json.loads((input_directory / 'redo.json').read_text()) == json.loads((input_directory / 'edited.json').read_text()), 'Frontend redo identity changed')
    check(json.loads((input_directory / 'heading-undo.json').read_text()) == json.loads((input_directory / 'edited.json').read_text()), 'Frontend heading undo identity changed')
    evidence['frontend_manifest'] = {'path': str(manifest_path), 'sha256': digest(manifest_path), 'contents': manifest}
    evidence['frontend_snapshots'] = results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path)
    parser.add_argument('--baseline', type=Path)
    parser.add_argument('--root', type=Path, default=ROOT / 'tmp/paragraph-overhang-independent')
    parser.add_argument('--original', type=Path, default=ROOT / 'tmp/paragraph-indentation-source/indents.docx')
    parser.add_argument('--generate-only', action='store_true')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--soffice')
    parser.add_argument('--editor-models', type=Path)
    args = parser.parse_args()
    check(args.generate_only or args.cli is not None, 'Pass a frozen --cli, or --generate-only')
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    cli, baseline = (args.cli.resolve() if args.cli else None), (args.baseline.resolve() if args.baseline else None)
    soffice = (args.soffice or shutil.which('soffice')) if args.render else None
    check(not args.render or soffice is not None, 'LibreOffice is required for --render')
    evidence = {'status': 'running', 'script_sha256': digest(__file__), 'cli': str(cli), 'cli_sha256': digest(cli) if cli else None,
                'baseline': str(baseline), 'baseline_sha256': digest(baseline) if baseline else None,
                'python_version': platform.python_version(), 'python_docx_version': version('python-docx'),
                'corpus_manifest_sha256': digest(ROOT / 'scripts/paragraph-indentation-corpus.json'),
                'commands': [], 'sources': {}, 'limitations': LIMITATIONS,
                'invocation': [sys.executable, *sys.argv], 'started_at_utc': datetime.now(timezone.utc).isoformat(),
                'baseline_repository_commit': '5f812c2612d1d1de1363b61786c3e435d0163616'}
    try:
        original = args.original.resolve()
        check(digest(original) == ORIGINAL_SHA and original.stat().st_size == 11553, 'Genuine original hash/length mismatch')
        corpus = json.loads((ROOT / 'scripts/paragraph-indentation-corpus.json').read_text())
        provenance = next(i for i in corpus['fixtures'] if i['name'] == 'indents.docx')
        check(provenance['sha256'] == ORIGINAL_SHA, 'Historical corpus identity changed')
        evidence['original_provenance'] = {'manifest_path': str(ROOT / 'scripts/paragraph-indentation-corpus.json'),
                                          'manifest_sha256': evidence['corpus_manifest_sha256'],
                                          'source_url': provenance['source_url'], 'raw_url': provenance['raw_url'],
                                          'repository': provenance['repository'], 'revision': provenance['revision'],
                                          'git_blob_sha': provenance['git_blob_sha'], 'sha256': provenance['sha256'],
                                          'historical_manifest_expected_action': provenance['expected_action'],
                                          'current_scoped_expected_action': 'roundtrip',
                                          'note': 'Historical corpus expected_action and prior reports are unchanged. This checker independently asserts newly supported bounded top-level overhang.',
                                          'licensing_caveat': next(i for i in corpus['licenses'] if i['repository'] == provenance['repository'])['caveat']}
        original_snapshot = inspect_docx(original)
        expected_original = [attrs(v) for v in [(None, None, None), (-1134, None, None), (-1134, None, 567), (-567, None, -567), (1134, None, -567)]]
        check([r['indent'] for r in original_snapshot['ooxml']] == expected_original, 'Genuine original declarations differ')
        evidence['original_preflight'] = {'sha256': ORIGINAL_SHA, 'bytes': 11553, 'independent': original_snapshot,
                                          'physical_bounds': [physical_bounds([r['indent'].get(k) for k in KEYS], original_snapshot['pageGeometry']) for r in original_snapshot['ooxml']]}
        sources = evidence['sources'] = fixtures(directory)
        sources['original-indents'] = {'path': str(original), 'sha256': ORIGINAL_SHA,
                                      'fixture_kind': 'genuine unchanged public upstream regression fixture; not user production',
                                      'expected_action': 'roundtrip', 'render': True, 'independent': original_snapshot, 'expected': original_snapshot['ooxml']}
        if cli:
            for name, item in sources.items():
                if item['expected_action'] == 'roundtrip': roundtrip(cli, evidence, directory, name, item)
                else:
                    item['audit'] = run_cli(cli, evidence, 'audit-docx', item['path'])
                    item['import_error'] = run_cli(cli, evidence, 'import-docx', item['path'], success=False)
                    check('Unsupported' in item['import_error'] or 'Invalid' in item['import_error'], name + ': unclear refusal')
                    before = digest(item['path'])
                    error = run_cli(cli, evidence, 'create-docx', sources['negative-left-positive-first']['export']['model_path'], item['path'], success=False)
                    check('Output already exists' in error and digest(item['path']) == before, name + ': rejected source was clobbered')
                    item['source_no_clobber'] = True
                check(digest(item['path']) == item['sha256'], name + ': source mutated')
            edited(cli, evidence, directory, sources)
            invalid_models(cli, evidence, directory, sources['negative-left-positive-first']['import_response']['document'])
            tamper(cli, evidence, directory, sources['negative-left-positive-first']['export']['path'])
            if baseline: legacy(cli, baseline, directory, evidence, sources)
            if args.editor_models: frontend_snapshots(cli, evidence, directory, args.editor_models.resolve())
        if soffice:
            evidence['libreoffice_executable'] = {'path': str(Path(soffice).resolve()), 'sha256': digest(soffice)}
            underlying = Path(soffice).parent.parent.parent / 'native/libreoffice-headless/libreoffice/program/soffice'
            if underlying.is_file():
                evidence['libreoffice_underlying_launcher'] = {'path': str(underlying.resolve()), 'sha256': digest(underlying)}
                binary = underlying.with_name('soffice.bin')
                if binary.is_file(): evidence['libreoffice_underlying_binary'] = {'path': str(binary.resolve()), 'sha256': digest(binary)}
            evidence['libreoffice_version'] = subprocess.run([soffice, '-env:UserInstallation=' + (directory / 'lo-version-profile').as_uri(), '--version'], capture_output=True, text=True, check=True).stdout.strip()
            render_all(soffice, directory, evidence, sources, args.generate_only)
        evidence['summary'] = {'supported_sources': sum(v['expected_action'] == 'roundtrip' for v in sources.values()),
                               'rejected_sources': sum(v['expected_action'] == 'reject' for v in sources.values()),
                               'invalid_models': len(evidence.get('invalid_models', {})), 'tampered_packages': len(evidence.get('tampered_packages', {})),
                               'edited_files': len(evidence.get('edited', {})), 'frontend_snapshot_files': len(evidence.get('frontend_snapshots', {})), 'legacy_byte_cases': len(evidence.get('legacy_bytes', {})),
                               'rendered_pdfs': len(list((directory / 'renders').glob('*/*.pdf'))),
                               'layout_drift_pairs': len(evidence.get('layout_drift', {})),
                               'render_comparison_pairs': len(evidence.get('render_comparisons', {})),
                               'pixel_identical_pairs': sum(all(p.get('pixel_equal') for p in c['raster_comparison']) for c in evidence.get('render_comparisons', {}).values()),
                               'source_no_clobber_cases': sum(i.get('source_no_clobber', False) for i in sources.values()),
                               'successful_trusted_model_reopens': sum(c['arguments'][0] == 'read-docx-model' and c['returncode'] == 0 for c in evidence['commands']),
                               'historic_exact_response_pairs': sum(c['exact_response_equal'] for r in evidence.get('legacy_responses', {}).values() for c in r.values())}
        evidence['status'] = 'fixtures-verified' if args.generate_only else 'passed-with-disclosed-drift' if evidence.get('layout_drift') else 'passed'
    except Exception as error:
        evidence['status'] = 'failed'; evidence['error'] = type(error).__name__ + ': ' + str(error)
        raise
    finally:
        evidence['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
        evidence['corpus_manifest_unchanged'] = digest(ROOT / 'scripts/paragraph-indentation-corpus.json') == evidence['corpus_manifest_sha256']
        evidence['script_unchanged_during_run'] = digest(__file__) == evidence['script_sha256']
        evidence['cli_unchanged_during_run'] = cli is None or digest(cli) == evidence['cli_sha256']
        evidence['baseline_unchanged_during_run'] = baseline is None or digest(baseline) == evidence['baseline_sha256']
        evidence['source_hashes_unchanged'] = all(digest(i['path']) == i['sha256'] for i in evidence['sources'].values())
        evidence['frontend_inputs_unchanged'] = all(digest(i['model_source_path']) == i['model_source_sha256'] for i in evidence.get('frontend_snapshots', {}).values())
        evidence['genuine_original_unchanged'] = args.original.exists() and digest(args.original) == ORIGINAL_SHA
        report = directory / 'report.json'; save_json(report, evidence); print(report)


if __name__ == '__main__':
    main()
