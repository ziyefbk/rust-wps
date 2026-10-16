#!/usr/bin/env python3
"""Independent file reader for the bounded real-browser Word workflow.

Reads standard DOCX/PDF content with python-docx, ZIP/XML, pypdf and Poppler.
It never imports rust-wps or trusts its embedded model as the content oracle.
"""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from pypdf import PdfReader

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def docx_report(file):
    doc = Document(file)
    with zipfile.ZipFile(file) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
        breaks = sum(n.attrib.get(W + 'type') == 'page' for n in root.iter(W + 'br'))
        blocks = []
        for node in root.find(W + 'body'):
            if node.tag == W + 'p':
                text = ''.join(n.text or '' for n in node.iter(W + 't'))
                if text:
                    blocks.append(text)
            elif node.tag == W + 'tbl':
                blocks.extend(''.join(n.text or '' for n in cell.iter(W + 't'))
                              for row in node.findall(W + 'tr') for cell in row.findall(W + 'tc'))
    return {
        'paragraphs': [p.text for p in doc.paragraphs],
        'nonempty_paragraphs': [p.text for p in doc.paragraphs if p.text],
        'paragraph_styles': [p.style.style_id for p in doc.paragraphs if p.text],
        'bold_run_text': [r.text for p in doc.paragraphs for r in p.runs if r.bold is True],
        'tables': [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables],
        'page_breaks': breaks, 'text_in_body_order': blocks,
        'sections': [{'width_points': s.page_width.pt, 'height_points': s.page_height.pt,
                      'margins_points': [s.top_margin.pt, s.right_margin.pt, s.bottom_margin.pt, s.left_margin.pt]}
                     for s in doc.sections],
    }


def without_whitespace(text):
    # PDF extractors introduce layout line breaks. No NFKC, punctuation folding,
    # case folding, or character substitution is allowed.
    return ''.join(text.split())


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['docx', 'pdf'])
    parser.add_argument('file', type=Path)
    parser.add_argument('--expected', type=Path, required=True)
    args = parser.parse_args()
    expected = json.loads(args.expected.read_text())
    result = {'file': str(args.file), 'sha256': hashlib.sha256(args.file.read_bytes()).hexdigest(),
              'reader': 'python-docx + independent ZIP/XML' if args.mode == 'docx' else 'pypdf + Poppler',
              'errors': []}
    errors = result['errors']
    if args.mode == 'docx':
        result.update(docx_report(args.file))
        for key in ['nonempty_paragraphs', 'tables', 'page_breaks', 'paragraph_styles', 'bold_run_text']:
            if key in expected and result[key] != expected[key]:
                errors.append({'field': key, 'expected': expected[key], 'actual': result[key]})
        if expected.get('margins_points'):
            for section in result['sections']:
                if section['margins_points'] != expected['margins_points']:
                    errors.append({'field': 'margins_points', 'actual': section['margins_points']})
        if expected.get('page_points'):
            for section in result['sections']:
                actual = [section['width_points'], section['height_points']]
                if any(abs(a-b) > 1 for a,b in zip(actual, expected['page_points'])):
                    errors.append({'field': 'page_points', 'actual': actual})
    else:
        reader = PdfReader(args.file)
        result['pages'] = [{'width_points': float(p.mediabox.width), 'height_points': float(p.mediabox.height),
                            'text': p.extract_text() or ''} for p in reader.pages]
        poppler = subprocess.run(['pdftotext', '-layout', str(args.file), '-'], check=True,
                                 capture_output=True, text=True, timeout=30)
        result['poppler_text'] = poppler.stdout
        result['comparison'] = 'Exact Unicode characters after layout whitespace removal; raw text retained'
        if 'pages' in expected and len(result['pages']) != expected['pages']:
            errors.append({'field': 'pages', 'expected': expected['pages'], 'actual': len(result['pages'])})
        for page in result['pages']:
            if any(abs(a-b) > 1 for a,b in zip([page['width_points'],page['height_points']], expected['page_points'])):
                errors.append({'field': 'page_points', 'actual': [page['width_points'], page['height_points']]})
        for name,text in [('pypdf', '\n'.join(p['text'] for p in result['pages'])), ('poppler', poppler.stdout)]:
            canonical = without_whitespace(text)
            if 'full_text' in expected and canonical != without_whitespace(expected['full_text']):
                errors.append({'reader': name, 'field': 'full_text_order_and_multiplicity', 'expected': without_whitespace(expected['full_text']), 'actual': canonical})
            for phrase in expected['contains']:
                if without_whitespace(phrase) not in canonical:
                    errors.append({'reader': name, 'missing': phrase})
            for phrase in expected.get('excludes', []):
                if without_whitespace(phrase) in canonical:
                    errors.append({'reader': name, 'unexpected': phrase})
        for index, marker in enumerate(expected.get('page_markers', [])):
            if index >= len(result['pages']) or without_whitespace(marker) not in without_whitespace(result['pages'][index]['text']):
                errors.append({'field': 'page_marker', 'page': index+1, 'expected': marker})
            for other_index, other in enumerate(expected.get('page_markers', [])):
                if other_index != index and index < len(result['pages']) and without_whitespace(other) in without_whitespace(result['pages'][index]['text']):
                    errors.append({'field': 'page_marker_leak', 'page': index+1, 'unexpected': other})
    result['passed'] = not errors
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result['passed'] else 1)


if __name__ == '__main__':
    main()
