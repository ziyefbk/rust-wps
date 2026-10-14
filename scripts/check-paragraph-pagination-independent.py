#!/usr/bin/env python3
"""Independent, strict DOCX paragraph-pagination source and consumer checker.

No product Python, Rust or frontend implementation is imported. Sources are
resolved through OPC relationships and a field-by-field defaults/style/direct
cascade. References are authored with python-docx/OOXML, never the app writer.
Each LibreOffice conversion has its own fresh process/profile and HOME. Complete
Poppler page text, line positions and page pixels are compared with zero drift.
Original upstream fixtures and historical evidence are hash-checked, never edited.
"""
import argparse
import copy
import hashlib
import json
import os
import platform
import posixpath
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.shared import Pt, Twips, RGBColor
from lxml import etree as E
from PIL import Image, ImageChops
import fitz

ROOT = Path(__file__).resolve().parent.parent
GAP = ROOT / 'tmp/next-docx-gap-20261014'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'
R = '{http://schemas.openxmlformats.org/package/2006/relationships}'
FLAGS = ('keepNext', 'keepLines', 'pageBreakBefore')
MIN_FREE = 650 * 1024 * 1024
ORIGINALS = {
    'Styles': {'sha256': '63f3fe4c3b4fbdc211eca8fa7fd31a799ff4438afbf4ca6e771b3aef107c732e',
               'git_blob': '27f66f48bff955494065233a3980fb65793b2e2f', 'bytes': 10905},
    'NoHeadFoot': {'sha256': 'dc6a67b1c118658acadc87ba41200f0b90ead05ada95e7ce6fa72a97700430c5',
                   'git_blob': '7af2a14583d09c337a7993fb744c33e31e28d687', 'bytes': 10380},
}
REVISION = 'ae62bb5116b9aee19ebd5834e3a82066132c9f7f'
COLOR_LIMITATION_SOURCES = {
    ORIGINALS['Styles']['sha256'], ORIGINALS['NoHeadFoot']['sha256'],
    'b599b8fd12147bb7fbef041dea7d962953b061c7705140abf62bacf14a71fa98',
}
OLD_CLI_SHA = '5c008894209311c84515a6432b6e973a120ad459b30c848ea019c2128808111a'
LIMITATIONS = [
    'Genuine sources are unchanged public Apache POI regression fixtures, not user production files; local validation only, no newly reviewed redistribution rights.',
    'Synthetic source controls and independently authored frontend reference DOCX files are labeled separately from genuine sources.',
    'Actual frontend snapshots are produced by executable editor code under a DOM transport substitute; they do not establish real browser keyboard, IME, clipboard, native WebView2, or print acceptance.',
    'Consumer evidence is LibreOffice plus Poppler only, not Microsoft Word/WPS or a general automatic pagination engine.',
    'Historical byte goldens use frozen old packages/models; they are labeled historical and are not fresh real-world source or fresh PDF evidence.',
]


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def disk_guard(path, reserve=8 * 1024 * 1024):
    free = shutil.disk_usage(path).free
    check(free - reserve >= MIN_FREE,
          f'Disk stop: {free / 1024**2:.1f} MiB free; reserve would cross 650 MiB safety threshold. Notify parent; do not clean up.')
    return free


def child(parent, tag, **attrs):
    node = E.SubElement(parent, W + tag)
    for key, value in attrs.items():
        node.set(W + key, str(value))
    return node


def singleton(parent, tag):
    nodes = [] if parent is None else parent.findall(W + tag)
    check(len(nodes) <= 1, 'Independent XML resolver found duplicate direct ' + tag)
    return nodes[0] if nodes else None


def value(parent, tag, default=None):
    node = singleton(parent, tag)
    return default if node is None else node.get(W + 'val', default)


def boolean(raw):
    check(raw in ('1', 'true', 'on', '0', 'false', 'off'), 'Invalid independently read OOXML boolean: ' + repr(raw))
    return raw in ('1', 'true', 'on')


def flags(ppr):
    result = {}
    if ppr is None:
        return result
    for name in FLAGS:
        node = singleton(ppr, name)
        if node is not None:
            check(set(node.attrib) <= {W + 'val'} and len(node) == 0,
                  'Invalid independently read pagination declaration: ' + name)
            result[name] = boolean(node.get(W + 'val', 'true'))
    return result


def package_parts(path):
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        check(len(names) == len(set(names)), 'Duplicate ZIP member')
        return {name: package.read(name) for name in names}


def related(parts, base, kind, required=True):
    relpath = '_rels/.rels' if base is None else posixpath.join(posixpath.dirname(base), '_rels', posixpath.basename(base) + '.rels')
    check(relpath in parts, 'Missing relationship part ' + relpath)
    rels = E.fromstring(parts[relpath])
    check(rels.tag == R + 'Relationships', 'Wrong package relationship namespace')
    matches = [r for r in rels if r.tag == R + 'Relationship' and r.get('Type', '').rsplit('/', 1)[-1] == kind]
    check(len(matches) == 1 or (not required and not matches), 'Missing/duplicate actual ' + kind + ' relationship')
    if not matches:
        return None
    rel = matches[0]
    check(rel.get('TargetMode') != 'External', 'External related source is not independently trusted')
    target = rel.get('Target')
    check(target and '#' not in target and '?' not in target, 'Invalid actual related part target')
    target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(base or ''), target))
    check(not target.startswith('../') and target in parts, 'Missing/escaping actual related part')
    return target


def paragraph_text(para):
    result = []
    for node in para.iter():
        if node.tag == W + 't':
            result.append(node.text or '')
        elif node.tag == W + 'tab':
            result.append('\t')
        elif node.tag == W + 'br':
            result.append('\f' if node.get(W + 'type') == 'page' else '\n')
    return ''.join(result)


def inspect_docx(path):
    """Source oracle: docDefaults -> explicit/default active chain -> direct pPr."""
    parts = package_parts(path)
    main_name = related(parts, None, 'officeDocument')
    styles_name = related(parts, main_name, 'styles')
    main, styles = E.fromstring(parts[main_name]), E.fromstring(parts[styles_name])
    check(main.tag == W + 'document' and styles.tag == W + 'styles', 'Wrong active main/styles namespace')
    body = singleton(main, 'body')
    check(body is not None, 'Missing body')
    style_map = {}
    for style in styles.findall(W + 'style'):
        sid = style.get(W + 'styleId')
        check(sid is not None and sid not in style_map, 'Missing/duplicate style identity')
        style_map[sid] = style
    default_ids = [sid for sid, style in style_map.items()
                   if style.get(W + 'type') == 'paragraph' and boolean(style.get(W + 'default', 'false'))]
    check(len(default_ids) <= 1, 'Ambiguous default paragraph style')
    defaults = singleton(styles, 'docDefaults')
    default_ppr = singleton(singleton(defaults, 'pPrDefault'), 'pPr')
    default_flags = flags(default_ppr)

    def chain(sid, seen=()):
        if sid is None:
            return []
        check(sid in style_map and sid not in seen and len(seen) < 64,
              'Missing/cyclic/excessive active style chain: ' + str(sid))
        style = style_map[sid]
        check(style.get(W + 'type') == 'paragraph', 'Active paragraph chain has nonparagraph style')
        based = singleton(style, 'basedOn')
        if based is not None:
            check(set(based.attrib) == {W + 'val'} and based.get(W + 'val'), 'Invalid basedOn identity')
        return chain(None if based is None else based.get(W + 'val'), seen + (sid,)) + [style]

    rows = []
    for index, para in enumerate(body.findall(W + 'p')):
        ppr = singleton(para, 'pPr')
        sid = value(ppr, 'pStyle', default_ids[0] if default_ids else None)
        lineage = chain(sid)
        layers = [{'kind': 'docDefaults', 'values': default_flags}]
        layers.extend({'kind': 'style', 'style': s.get(W + 'styleId'), 'values': flags(singleton(s, 'pPr'))} for s in lineage)
        layers.append({'kind': 'direct', 'values': flags(ppr)})
        effective = dict.fromkeys(FLAGS, False)
        for layer in layers:
            effective.update(layer['values'])
        text = paragraph_text(para)
        sole_page_break = text == '\f' and len(para.findall('.//' + W + 'br')) == 1
        outline = None
        for pp in [default_ppr, *[singleton(s, 'pPr') for s in lineage], ppr]:
            candidate = value(pp, 'outlineLvl')
            if candidate is not None:
                outline = int(candidate)
        kind = 'pageBreak' if sole_page_break else 'heading' if outline is not None and 0 <= outline <= 2 else 'paragraph'
        rows.append({'index': index, 'type': kind, 'text': '' if sole_page_break else text,
                     'level': outline + 1 if kind == 'heading' else None, 'style': sid,
                     'style_chain': [s.get(W + 'styleId') for s in lineage],
                     'layers': layers, 'direct': flags(ppr), 'effective': effective,
                     'contains_page_break': '\f' in text,
                     'declarations_present': sorted({k for layer in layers for k in layer['values']})})
    check(len(body.findall('.//' + W + 'p')) == len(rows), 'Independent supported source oracle requires body-level paragraphs')
    check(all(n.tag in (W + 'p', W + 'sectPr') for n in body), 'Unsupported body wrapper in source oracle')
    section = singleton(body, 'sectPr')
    size, margins = singleton(section, 'pgSz'), singleton(section, 'pgMar')
    page = None
    if size is not None and margins is not None:
        page = {'widthTwips': int(size.get(W + 'w')), 'heightTwips': int(size.get(W + 'h'))}
        page.update({'margin' + s.title() + 'Twips': int(margins.get(W + s)) for s in ('top', 'right', 'bottom', 'left')})
    # A second library independently checks direct on/off values and text.
    pyrows = []
    for para in Document(path).paragraphs:
        pf = para.paragraph_format
        direct = {key: val for key, val in zip(FLAGS, (pf.keep_with_next, pf.keep_together, pf.page_break_before)) if val is not None}
        pyrows.append({'text': para.text, 'direct': direct})
    check(len(pyrows) == len(rows), 'python-docx/OOXML body paragraph count disagrees')
    for row, py in zip(rows, pyrows):
        check(row['direct'] == py['direct'], 'python-docx/OOXML direct flags disagree')
        check(row['text'].replace('\f', '') == py['text'], 'python-docx/OOXML text disagrees')
    return {'parts': {'main': main_name, 'styles': styles_name}, 'rows': rows,
            'pageGeometry': page, 'python_docx_direct': pyrows,
            'oracle': 'Strict independent OPC/defaults/active basedOn root-leaf/direct cascade; never inferred from app output'}


def model_rows(model, preserve_absence=False):
    rows = []
    for index, node in enumerate(model.get('content', [])):
        check(node['type'] in ('paragraph', 'heading', 'pageBreak'), 'Nonbody context in independent supported model comparison')
        attrs = node.get('attrs', {})
        direct = {k: attrs[k] for k in FLAGS if k in attrs}
        check(all(type(v) is bool for v in direct.values()), 'Model flags are not strictly boolean')
        text = ''.join(n.get('text', '') if n['type'] == 'text' else '\n' if n['type'] == 'hardBreak' else 'UNSUPPORTED' for n in node.get('content', []))
        historical = {'keepNext': node['type'] == 'heading', 'keepLines': False, 'pageBreakBefore': False}
        rows.append({'index': index, 'type': node['type'], 'level': attrs.get('level'), 'text': text,
                     'direct': direct, 'effective': historical | direct})
    return rows


def compare_model_source(model, source, label):
    actual, expected = model_rows(model), source['rows']
    check(len(actual) == len(expected), label + ': complete block order/count differs')
    for got, want in zip(actual, expected):
        for key in ('text', 'type', 'level'):
            check(got[key] == want[key], f'{label}: paragraph {want["index"]} {key}: {got[key]!r} != {want[key]!r}')
        if want['type'] != 'pageBreak':
            check(got['direct'] == want['effective'], f'{label}: explicit source flags lost/changed: {got} != {want}')
    check(model.get('attrs', {}).get('pageGeometry') == source['pageGeometry'], label + ': page geometry differs')


def compare_export_model(source, model, label):
    expected = model_rows(model)
    check(len(source['rows']) == len(expected), label + ': output paragraph count/order changed')
    for got, want in zip(source['rows'], expected):
        for key in ('text', 'type', 'level'):
            check(got[key] == want[key], f'{label}: exported {key} changed')
        check(got['direct'] == want['direct'], f'{label}: exact direct XML declarations differ: {got["direct"]} != {want["direct"]}')
        if want['type'] != 'pageBreak':
            check(got['effective'] == want['effective'], label + ': actual output cascade differs from model semantics')


def run_cli(cli, evidence, *args, success=True):
    disk_guard(evidence['directory'])
    result = subprocess.run([str(cli), *map(str, args)], text=True, capture_output=True, timeout=90)
    entry = {'program': str(cli), 'arguments': list(map(str, args)), 'returncode': result.returncode,
             'stdout': result.stdout, 'stderr': result.stderr}
    evidence['commands'].append(entry)
    check((result.returncode == 0) == success, 'Unexpected CLI result: ' + json.dumps(entry))
    return json.loads(result.stdout) if success else result.stderr


def create(cli, evidence, directory, name, model, inspect=True):
    model_path, output = directory / (name + '.json'), directory / (name + '.docx')
    check(not model_path.exists() and not output.exists(), 'New artifact would overwrite existing evidence')
    save_json(model_path, model)
    run_cli(cli, evidence, 'create-docx', model_path, output)
    reopened = run_cli(cli, evidence, 'read-docx-model', output)
    check(reopened == model, name + ': entire trusted re-opened model differs')
    before = digest(output)
    err = run_cli(cli, evidence, 'create-docx', model_path, output, success=False)
    check('Output already exists' in err and digest(output) == before, name + ': output overwrite protection failed')
    independent = inspect_docx(output) if inspect else None
    if independent is not None:
        compare_export_model(independent, model, name)
    return {'path': str(output), 'sha256': before, 'model_path': str(model_path),
            'model_sha256': digest(model_path),
            'complete_trusted_reopen_equal': True, 'no_overwrite': True, 'independent': independent}


def roundtrip(cli, evidence, directory, name, item):
    path = Path(item['path'])
    item['audit'] = run_cli(cli, evidence, 'audit-docx', path)
    item['import_response'] = run_cli(cli, evidence, 'import-docx', path)
    model = item['import_response']['document']
    compare_model_source(model, item['independent'], name)
    item['export'] = create(cli, evidence, directory, name + '-export', model)
    before = digest(path)
    err = run_cli(cli, evidence, 'create-docx', item['export']['model_path'], path, success=False)
    check('Output already exists' in err and digest(path) == before, name + ': original overwrite protection failed')
    item['source_no_overwrite'] = True
    return model


def write_package(parts, path):
    check(not path.exists(), 'Derived package must be fresh')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def mutate(source, target, action):
    parts = package_parts(source)
    main = E.fromstring(parts['word/document.xml'])
    styles = E.fromstring(parts['word/styles.xml'])
    action(main, styles, parts)
    parts['word/document.xml'] = E.tostring(main, xml_declaration=True, encoding='UTF-8', standalone=True)
    parts['word/styles.xml'] = E.tostring(styles, xml_declaration=True, encoding='UTF-8', standalone=True)
    write_package(parts, target)


def set_flags(ppr, values):
    for name, val in values.items():
        for old in ppr.findall(W + name):
            ppr.remove(old)
        node = child(ppr, name)
        if val is not None:
            node.set(W + 'val', str(int(val)) if type(val) is bool else str(val))


def base_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Twips(1440)
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(12)
        style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.widow_control = False
    return doc


def para(doc, text, values=None, style=None):
    paragraph = doc.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    for run in paragraph.runs:
        run.font.name, run.font.size = 'Arial', Pt(12)
    if values:
        set_flags(paragraph._p.get_or_add_pPr(), values)
    return paragraph


def fixtures(directory):
    sources = {}

    def add(name, path, kind='independently authored synthetic source control', render=True, expected=None):
        independent = inspect_docx(path)
        if expected is not None:
            check([r['effective'] for r in independent['rows']] == expected, name + ': manually specified source cascade mismatch')
        sources[name] = {'path': str(path), 'sha256': digest(path), 'classification': kind,
                         'expected_action': 'roundtrip', 'render': render, 'independent': independent}

    for name, spec in ORIGINALS.items():
        path = ROOT / 'tmp/paragraph-indentation-source/poi' / (name + '.docx')
        raw = path.read_bytes()
        check(len(raw) == spec['bytes'] and digest(path) == spec['sha256'], name + ': genuine source identity mismatch')
        check(hashlib.sha1(b'blob ' + str(len(raw)).encode() + b'\0' + raw).hexdigest() == spec['git_blob'], name + ': Git blob mismatch')
        add('original-' + name, path, 'genuine unchanged Apache POI regression fixture, not a production file')
        sources['original-' + name]['provenance'] = spec | {'revision': REVISION,
            'url': 'https://github.com/apache/poi/blob/' + REVISION + '/test-data/document/' + name + '.docx'}
    for feature in FLAGS:
        for active in (True, False):
            name = f'control-{feature}-{str(active).lower()}'
            add(name, GAP / (name + '.docx'), 'retained independently authored python-docx synthetic control; freshly re-rendered this run')
    add('control-precedence', GAP / 'control-precedence.docx',
        'retained independently authored precedence control; expected vectors manually authored before implementation',
        expected=[dict(zip(FLAGS, v)) for v in ((True, True, True), (True, False, False), (False, False, False), (False, False, False))])

    def authored(name, build, expected=None, render=True):
        doc = base_document()
        build(doc)
        path = directory / (name + '.docx')
        doc.save(path)
        add(name, path, render=render, expected=expected)

    for enabled in (True, False):
        def chained(doc, enabled=enabled):
            for i in range(48):
                para(doc, f'CHAIN FILL {i+1:02}')
            for label in ('ANCHOR A', 'ANCHOR B'):
                para(doc, label, {'keepNext': enabled})
            para(doc, 'CHAIN FOLLOWER')
        authored('chain-keepNext-' + str(enabled).lower(), chained)

        def long(doc, enabled=enabled):
            para(doc, 'LONG BEFORE')
            para(doc, '\n'.join(f'LONG LINE {i+1:03}' for i in range(111)), {'keepLines': enabled})
            para(doc, 'LONG AFTER')
        authored('long-keepLines-' + str(enabled).lower(), long)

        def first(doc, enabled=enabled):
            para(doc, 'FIRST PARAGRAPH', {'pageBreakBefore': enabled})
            para(doc, 'SECOND PARAGRAPH')
        authored('document-start-' + str(enabled).lower(), first)

        def adjacent(doc, enabled=enabled):
            para(doc, 'BEFORE MANUAL BREAK')
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            para(doc, 'AFTER MANUAL BREAK', {'pageBreakBefore': enabled})
            para(doc, 'AFTER FOLLOWER')
        authored('adjacent-manual-break-' + str(enabled).lower(), adjacent)

    def override(doc):
        style = doc.styles.add_style('PaginationKeep', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles['Normal']
        set_flags(style.element.get_or_add_pPr(), dict.fromkeys(FLAGS, True))
        for i in range(49):
            para(doc, f'FALSE FILL {i+1:02}')
        para(doc, 'DIRECT FALSE ANCHOR', dict.fromkeys(FLAGS, False), 'PaginationKeep')
        para(doc, 'DIRECT FALSE FOLLOWER')
    authored('direct-false-over-style-true', override)

    def style_false(doc):
        defaults = singleton(singleton(doc.styles.element, 'docDefaults'), 'pPrDefault')
        default_ppr = singleton(defaults, 'pPr')
        if default_ppr is None:
            default_ppr = child(defaults, 'pPr')
        set_flags(default_ppr, dict.fromkeys(FLAGS, True))
        style = doc.styles.add_style('PaginationFalse', WD_STYLE_TYPE.PARAGRAPH)
        set_flags(style.element.get_or_add_pPr(), dict.fromkeys(FLAGS, False))
        para(doc, 'STYLE FALSE FIRST', style='PaginationFalse')
        para(doc, 'STYLE FALSE SECOND', style='PaginationFalse')
    authored('style-false-over-default-true', style_false,
             expected=[dict.fromkeys(FLAGS, False)] * 2)

    def explicit_no_default_overlay(doc):
        set_flags(doc.styles['Normal'].element.get_or_add_pPr(), dict.fromkeys(FLAGS, True))
        style = doc.styles.add_style('DetachedParagraph', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = None
        para(doc, 'DETACHED EXPLICIT STYLE', style='DetachedParagraph')
        para(doc, 'DETACHED SECOND', style='DetachedParagraph')
    authored('explicit-style-does-not-overlay-default', explicit_no_default_overlay,
             expected=[dict.fromkeys(FLAGS, False)] * 2)

    # Every accepted on/off spelling is exercised against all three properties.
    def spellings(doc):
        for spelling in ('1', 'true', 'on', '0', 'false', 'off', None):
            para(doc, 'BOOLEAN ' + str(spelling), dict.fromkeys(FLAGS, spelling))
    authored('valid-boolean-spellings', spellings, render=False)

    base = Path(sources['control-precedence']['path'])
    rejects = {}

    def directed(name, action, reject=True):
        path = directory / (name + '.docx')
        mutate(base, path, action)
        if reject:
            rejects[name] = {'path': str(path), 'sha256': digest(path),
                             'classification': 'directed synthetic invalid/ambiguous source control',
                             'expected_action': 'reject', 'render': False}
        else:
            add(name, path, render=False)

    def direct(main):
        paragraph = main.find('w:body/w:p', NS)
        ppr = singleton(paragraph, 'pPr')
        if ppr is None:
            ppr = E.Element(W + 'pPr')
            paragraph.insert(0, ppr)
        return ppr

    for name in FLAGS:
        for invalid in ('yes', 'TRUE', '2', '', ' false '):
            directed('reject-' + name + '-' + repr(invalid).replace(' ', '_').replace("'", 'q'),
                     lambda m, s, p, key=name, raw=invalid: set_flags(direct(m), {key: raw}))
        directed('reject-duplicate-' + name,
                 lambda m, s, p, key=name: (set_flags(direct(m), {key: True}), child(direct(m), key, val='0')))
        directed('reject-foreign-attribute-' + name,
                 lambda m, s, p, key=name: (set_flags(direct(m), {key: True}), singleton(direct(m), key).set('val', 'false')))
        directed('reject-nested-declaration-' + name,
                 lambda m, s, p, key=name: child(child(direct(m), 'pPrChange'), key, val='1'))
    for mode in ('missing', 'cycle', 'wrong-type', 'duplicate-id'):
        def invalid_style(main, styles, parts, mode=mode):
            leaf = next(s for s in styles.findall(W + 'style') if s.get(W + 'styleId') == 'PaginationLeaf')
            if mode == 'missing':
                singleton(leaf, 'basedOn').set(W + 'val', 'MissingPaginationBase')
            elif mode == 'cycle':
                singleton(leaf, 'basedOn').set(W + 'val', 'PaginationLeaf')
            elif mode == 'wrong-type':
                leaf.set(W + 'type', 'character')
            else:
                styles.append(copy.deepcopy(leaf))
        directed('reject-active-style-' + mode, invalid_style)

    def unused_bad(main, styles, parts):
        style = child(styles, 'style', type='paragraph', styleId='UnusedBrokenPagination')
        child(style, 'basedOn', val='MissingUnusedStyle')
        set_flags(child(style, 'pPr'), dict.fromkeys(FLAGS, 'unrecognized'))
    directed('unused-malformed-style-ignored', unused_bad, reject=False)

    def relocated(main, styles, parts):
        rels = E.fromstring(parts['word/_rels/document.xml.rels'])
        for rel in rels:
            if rel.get('Type', '').endswith('/styles'):
                rel.set('Target', 'renamed/pagination-styles.xml')
        parts['word/_rels/document.xml.rels'] = E.tostring(rels, xml_declaration=True, encoding='UTF-8')
        parts['word/renamed/pagination-styles.xml'] = E.tostring(styles, xml_declaration=True, encoding='UTF-8')
        types = E.fromstring(parts['[Content_Types].xml'])
        old = next(n for n in types if n.get('PartName') == '/word/styles.xml')
        new = copy.deepcopy(old)
        new.set('PartName', '/word/renamed/pagination-styles.xml')
        types.append(new)
        parts['[Content_Types].xml'] = E.tostring(types, xml_declaration=True, encoding='UTF-8')
        # Conventional styles.xml remains as an inactive decoy.
        set_flags(singleton(singleton(singleton(styles, 'docDefaults'), 'pPrDefault'), 'pPr'), dict.fromkeys(FLAGS, False))
    directed('actual-relocated-styles-not-decoy', relocated, reject=False)
    sources.update(rejects)
    return sources


def invalid_models(cli, evidence, directory, good):
    results = {}
    for key in FLAGS:
        for number, val in enumerate((None, 0, 1, 'true', 'false', [], {}, 1.5)):
            model = copy.deepcopy(good)
            model['content'][0].setdefault('attrs', {})[key] = val
            name = key + '-' + str(number)
            path, output = directory / ('invalid-' + name + '.json'), directory / ('invalid-' + name + '.docx')
            save_json(path, model)
            error = run_cli(cli, evidence, 'create-docx', path, output, success=False)
            check(not output.exists(), 'Invalid model left partial output')
            results[name] = {'error': error, 'no_partial_output': True}
        for context in ('doc', 'pageBreak', 'list', 'cell'):
            model = copy.deepcopy(good)
            paragraph = copy.deepcopy(model['content'][0])
            paragraph.setdefault('attrs', {})[key] = False
            if context == 'doc':
                model.setdefault('attrs', {})[key] = False
            elif context == 'pageBreak':
                model['content'] = [{'type': 'pageBreak', 'attrs': {key: False}}]
            elif context == 'list':
                model['content'] = [{'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [paragraph]}]}]
            else:
                model['content'] = [{'type': 'table', 'content': [{'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [paragraph]}]}]}]
            name = key + '-context-' + context
            path, output = directory / ('invalid-' + name + '.json'), directory / ('invalid-' + name + '.docx')
            save_json(path, model)
            error = run_cli(cli, evidence, 'create-docx', path, output, success=False)
            check(not output.exists(), 'Invalid context left partial output')
            results[name] = {'error': error, 'no_partial_output': True}
    evidence['invalid_models'] = results


def tamper(cli, evidence, directory, exported):
    original = package_parts(exported)
    member = 'rust-wps/model-v2.json'
    check(member in original, 'Expected V2 envelope')
    results = {}
    for key in FLAGS:
        for mode in ('model-flip', 'xml-flip', 'xml-remove', 'xml-duplicate', 'xml-invalid'):
            for forged in ((False,) if mode.startswith('model') else (False, True)):
                name = key + '-' + mode + ('-forged-hashes' if forged else '')
                parts = dict(original)
                envelope = json.loads(parts[member])
                if mode == 'model-flip':
                    attrs = envelope['document']['content'][0]['attrs']
                    attrs[key] = not attrs[key]
                else:
                    main = E.fromstring(parts['word/document.xml'])
                    pp = main.find('w:body/w:p/w:pPr', NS)
                    node = singleton(pp, key)
                    check(node is not None, 'Tamper requires explicit exported flag')
                    if mode == 'xml-flip':
                        node.set(W + 'val', '0' if boolean(node.get(W + 'val', 'true')) else '1')
                    elif mode == 'xml-remove':
                        pp.remove(node)
                    elif mode == 'xml-duplicate':
                        pp.append(copy.deepcopy(node))
                    else:
                        node.set(W + 'val', 'not-a-boolean')
                    parts['word/document.xml'] = E.tostring(main, xml_declaration=True, encoding='UTF-8', standalone=True)
                if forged:
                    for part in envelope['hashes']:
                        envelope['hashes'][part] = hashlib.sha256(parts[part]).hexdigest()
                parts[member] = json.dumps(envelope, ensure_ascii=False, separators=(',', ':')).encode()
                path = directory / ('tampered-' + name + '.docx')
                write_package(parts, path)
                results[name] = {'path': str(path), 'sha256': digest(path),
                                 'error': run_cli(cli, evidence, 'read-docx-model', path, success=False)}
    evidence['tampered_packages'] = results


def reference_from_model(model, path):
    """Fresh independently authored source package representing actual editor state."""
    doc = base_document()
    check(model.get('attrs', {}).get('styleProfile') == 'standard-v2', 'Frontend reference requires explicit standard-v2 profile')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Twips(120)
    normal.paragraph_format.line_spacing = 1.5
    page = model.get('attrs', {}).get('pageGeometry')
    if page:
        section = doc.sections[0]
        section.page_width, section.page_height = Twips(page['widthTwips']), Twips(page['heightTwips'])
        for side in ('top', 'right', 'bottom', 'left'):
            setattr(section, side + '_margin', Twips(page['margin' + side.title() + 'Twips']))
    for node in model.get('content', []):
        if node['type'] == 'pageBreak':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        check(node['type'] in ('paragraph', 'heading'), 'Unsupported frontend reference context')
        attrs = node.get('attrs', {})
        p = doc.add_paragraph(style='Heading ' + str(attrs['level']) if node['type'] == 'heading' else None)
        pp = p._p.get_or_add_pPr()
        # Model absence is app heading keepNext=true, body flags=false.
        effective = dict.fromkeys(FLAGS, False)
        effective['keepNext'] = node['type'] == 'heading'
        effective.update({key: attrs[key] for key in FLAGS if key in attrs})
        set_flags(pp, effective)
        pf = p.paragraph_format
        heading = node['type'] == 'heading'
        pf.space_before, pf.space_after = Twips(attrs.get('spaceBefore', 240 if heading else 0)), Twips(attrs.get('spaceAfter', 120))
        pf.line_spacing = attrs.get('lineSpacing', 360) / 240
        for key, property_name in [('indentLeftTwips', 'left_indent'), ('indentRightTwips', 'right_indent'), ('firstLineIndentTwips', 'first_line_indent')]:
            if key in attrs:
                setattr(pf, property_name, Twips(attrs[key]))
        if attrs.get('textAlign'):
            child(pp, 'jc', val=attrs['textAlign'])
        if not node.get('content'):
            font = attrs.get('emptyParagraphFont', {'family': 'Arial', 'halfPoints': 22})
            rp = child(pp, 'rPr')
            child(rp, 'rFonts', ascii=font['family'], hAnsi=font['family'])
            child(rp, 'sz', val=font['halfPoints'])
            for key in ('bold', 'italic'):
                child(rp, 'b' if key == 'bold' else 'i', val='1' if font.get(key, False) else '0')
        for span in node.get('content', []):
            if span['type'] == 'hardBreak':
                p.add_run().add_break()
                continue
            check(span['type'] == 'text', 'Unsupported frontend reference inline')
            run = p.add_run(span['text'])
            marks = {m['type']: m.get('attrs', {}) for m in span.get('marks', [])}
            check(set(marks) <= {'fontFamily', 'fontSize', 'bold', 'italic', 'underline', 'strike', 'textColor'}, 'Unsupported independent reference marks')
            run.font.name = marks.get('fontFamily', {}).get('name', 'Arial')
            default_size = (36, 30, 26)[attrs['level'] - 1] if heading else 22
            run.font.size = Pt(marks.get('fontSize', {}).get('halfPoints', default_size) / 2)
            for key in ('bold', 'italic', 'underline', 'strike'):
                setattr(run.font, key, key in marks or (key == 'bold' and heading))
            run.font.color.rgb = RGBColor.from_string(marks.get('textColor', {}).get('rgb', '000000'))
    check(not path.exists(), 'Reference must be fresh')
    doc.save(path)
    independent = inspect_docx(path)
    expected = model_rows(model)
    check(len(independent['rows']) == len(expected), 'Reference/model block count differs')
    for got, want in zip(independent['rows'], expected):
        check(all(got[k] == want[k] for k in ('type', 'text', 'level', 'effective')), 'Independently authored reference disagrees with editor model')
    return independent


def frontend_snapshots(cli, evidence, directory, input_directory):
    manifest_path = input_directory / 'manifest.json'
    manifest = json.loads(manifest_path.read_text())
    evidence['frontend_manifest'] = {'path': str(manifest_path), 'sha256': digest(manifest_path), 'contents': manifest}
    evidence['frontend_operation_checks'] = frontend_operation_expectations(input_directory, manifest)
    results = evidence['frontend_snapshots'] = {}
    for filename in manifest['files']:
        source = input_directory / filename
        wrapped = json.loads(source.read_text())
        model = wrapped.get('document', wrapped)
        label = 'frontend-' + source.stem
        reference = directory / (label + '-reference.docx')
        independent = reference_from_model(model, reference)
        output = create(cli, evidence, directory, label + '-output', model)
        results[source.stem] = {'model_path': str(source), 'model_sha256': digest(source),
                               'reference_path': str(reference), 'reference_sha256': digest(reference),
                               'reference_independent': independent, 'export': output,
                               'reference_method': 'New python-docx/OOXML reference from actual frontend model; no application-generated expected DOCX'}
    for left, right in manifest.get('equalities', []):
        check(json.loads((input_directory / left).read_text()) == json.loads((input_directory / right).read_text()), 'Frontend exact identity failed: ' + left + '/' + right)


def frontend_operation_expectations(directory, manifest):
    """Derive complete expected edits without importing any frontend code."""
    input_path = Path(manifest['input'])
    check(digest(input_path) == manifest['inputSha256'], 'Frontend original input identity changed')
    raw = json.loads(input_path.read_text())
    raw = raw.get('document', raw)
    snapshots = {Path(name).stem: json.loads((directory / name).read_text())['document'] for name in manifest['files']}
    initial = snapshots['initial']

    def without_schema_defaults(node):
        node = copy.deepcopy(node)
        if node.get('attrs', {}).get('textAlign', 'absent') is None:
            del node['attrs']['textAlign']
        if 'attrs' in node and not node['attrs']:
            del node['attrs']
        if 'content' in node:
            node['content'] = [without_schema_defaults(n) for n in node['content']]
        return node

    check(without_schema_defaults(initial) == without_schema_defaults(raw), 'Frontend initial changed source beyond known null textAlign schema default')
    index = manifest['target']['blockIndex']
    check(type(index) is int and 0 <= index < len(initial['content']), 'Invalid frontend target index')
    check(initial['content'][index]['type'] == 'paragraph', 'Independent edit oracle currently requires paragraph target')
    expected = {'initial': copy.deepcopy(initial)}

    def changed(base='initial'):
        return copy.deepcopy(expected[base])

    edited = changed()
    edited['content'][index]['attrs'].update(keepNext=False, keepLines=True, pageBreakBefore=False)
    expected['edited'] = edited
    reset = changed('edited')
    for key in FLAGS:
        reset['content'][index]['attrs'].pop(key, None)
    expected['reset'] = reset
    text_edited = changed()
    check(text_edited['content'][index]['content'][0]['type'] == 'text', 'Independent prefix edit requires leading text')
    text_edited['content'][index]['content'][0]['text'] = 'Edited ' + text_edited['content'][index]['content'][0]['text']
    expected['text-edited'] = text_edited
    converted = changed()
    converted['content'][index]['type'] = 'heading'
    converted['content'][index]['attrs']['level'] = 2
    expected['heading-converted'] = converted
    prepared = changed()
    prepared['content'][index]['attrs'].update(keepNext=True, keepLines=False, pageBreakBefore=True)
    expected['split-ready'] = prepared
    split = changed('split-ready')
    block = split['content'][index]
    offset = manifest['target']['splitContentOffset']
    left, right, remaining = [], [], offset
    for span in block.get('content', []):
        if span['type'] == 'text':
            encoded = span['text'].encode('utf-16-le')
            size = len(encoded) // 2
            if remaining >= size:
                left.append(copy.deepcopy(span))
                remaining -= size
            elif remaining <= 0:
                right.append(copy.deepcopy(span))
            else:
                l, r = copy.deepcopy(span), copy.deepcopy(span)
                l['text'], r['text'] = encoded[:remaining * 2].decode('utf-16-le'), encoded[remaining * 2:].decode('utf-16-le')
                left.append(l)
                right.append(r)
                remaining = 0
        else:
            check(span['type'] == 'hardBreak', 'Unsupported inline for independent split oracle')
            (left if remaining > 0 else right).append(copy.deepcopy(span))
            remaining = max(0, remaining - 1)
    check(remaining == 0 and left and right, 'Independent split must be bounded within content')
    leading, trailing = copy.deepcopy(block), copy.deepcopy(block)
    leading['content'], trailing['content'] = left, right
    trailing['attrs'].pop('pageBreakBefore', None)
    split['content'][index:index + 1] = [leading, trailing]
    expected['enter-split'] = split
    expected['merge-refused'] = copy.deepcopy(split)
    for name in ('edited', 'reset', 'text-edited', 'heading-converted', 'split-ready', 'enter-split'):
        base = 'edited' if name == 'reset' else 'split-ready' if name == 'enter-split' else 'initial'
        expected[name + '-undo'] = copy.deepcopy(expected[base])
        expected[name + '-redo'] = copy.deepcopy(expected[name])
        expected[name + '-draft'] = copy.deepcopy(expected[name])
    # The full registered chain includes the established automatic body tail.
    def with_body_tail(model):
        model = copy.deepcopy(model)
        if model['content'][-1]['type'] != 'paragraph':
            model['content'].append({'type': 'paragraph', 'attrs': {'textAlign': None}})
        return model
    expected['registered-enter'] = with_body_tail(expected['enter-split'])
    expected['registered-enter-redo'] = copy.deepcopy(expected['registered-enter'])
    expected['registered-enter-before'] = with_body_tail(expected['split-ready'])
    expected['registered-enter-undo'] = with_body_tail(expected['split-ready'])
    compatible = changed()
    compatible['content'][index]['attrs'].update(keepNext=False, keepLines=False, pageBreakBefore=False)
    expected['merge-ready'] = compatible
    merge_split = copy.deepcopy(expected['enter-split'])
    merge_split['content'][index]['attrs'].update(keepNext=False, keepLines=False, pageBreakBefore=False)
    merge_split['content'][index + 1]['attrs'].update(keepNext=False, keepLines=False)
    expected['merge-split'] = merge_split
    expected['merge-compatible'] = copy.deepcopy(compatible)
    clipboard = changed()
    donor_index = next((i for i in range(index + 1, len(initial['content']))
                        if initial['content'][i]['type'] in ('paragraph', 'heading')), index)
    recipient = clipboard['content'][index]
    donor = initial['content'][donor_index]
    check(len(recipient['content']) == 1 and recipient['content'][0]['type'] == 'text',
          'Independent inline-paste oracle requires one recipient text run')
    first = recipient['content'][0]
    encoded = first['text'].encode('utf-16-le')
    left, right = copy.deepcopy(first), copy.deepcopy(first)
    left['text'], right['text'] = encoded[:2].decode('utf-16-le'), encoded[2:].decode('utf-16-le')
    pieces = [left, *copy.deepcopy(donor['content']), right]
    merged = []
    for piece in pieces:
        if piece['type'] == 'text' and not piece['text']:
            continue
        if (merged and piece['type'] == 'text' and merged[-1]['type'] == 'text'
                and {k: v for k, v in piece.items() if k != 'text'} == {k: v for k, v in merged[-1].items() if k != 'text'}):
            merged[-1]['text'] += piece['text']
        else:
            merged.append(piece)
    recipient['content'] = merged
    expected['clipboard-compatible'] = clipboard
    for name, base in (('merge-ready', 'initial'), ('merge-split', 'merge-ready'),
                       ('merge-compatible', 'merge-split'), ('clipboard-compatible', 'initial')):
        expected[name + '-undo'] = copy.deepcopy(expected[base])
        expected[name + '-redo'] = copy.deepcopy(expected[name])
        expected[name + '-draft'] = copy.deepcopy(expected[name])
    check(set(snapshots) == set(expected), 'Actual snapshot set differs from complete independent operation oracle')
    for name, model in expected.items():
        check(snapshots[name] == model, 'Independent complete expected frontend operation differs: ' + name)
    return {'input_path': str(input_path), 'input_sha256': digest(input_path), 'target_block_index': index,
            'split_utf16_offset': offset, 'complete_expected_snapshot_count': len(expected),
            'verified': sorted(expected), 'method': 'Source identity plus independent Python full-model apply/reset/prefix/Heading2/split/history/draft/refusal/body-tail/equal-effective-merge/inline-clipboard expectations, without invoking frontend implementation'}


def consumer_identity(soffice, directory):
    launcher = Path(soffice).resolve()
    candidates = [launcher, launcher.parent / 'soffice.bin',
                  launcher.parent / '../../native/libreoffice-headless/libreoffice/program/soffice.bin',
                  launcher.parent / '../../native/libreoffice-headless/libreoffice/program/libswlo.so']
    files = {}
    for path in candidates:
        path = path.resolve()
        if path.is_file():
            files[str(path)] = digest(path)
    check(any(Path(path).name == 'soffice.bin' for path in files), 'Need underlying LibreOffice ELF identity; supply a locatable soffice launcher')
    version = subprocess.run([soffice, '-env:UserInstallation=' + (directory / 'lo-version-profile').as_uri(), '--version'],
                             capture_output=True, text=True, check=True, timeout=30)
    versions = {'libreoffice': version.stdout.strip() + version.stderr.strip()}
    for name in ('pdftotext', 'pdftoppm'):
        path = Path(shutil.which(name)).resolve()
        files[str(path)] = digest(path)
        result = subprocess.run([str(path), '-v'], capture_output=True, text=True, check=True, timeout=30)
        versions[name] = result.stdout.strip() + result.stderr.strip()
    return {'binaries_and_launcher_sha256': files, 'versions': versions}


def render(soffice, directory, path, label):
    disk_guard(directory, 24 * 1024 * 1024)
    area = directory / 'renders' / label
    area.mkdir(parents=True)
    profile, home = area / 'profile', area / 'home'
    home.mkdir()
    check(not profile.exists(), 'LibreOffice profile must be new')
    before = digest(path)
    command = [soffice, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(area), str(path)]
    result = subprocess.run(command, text=True, capture_output=True, timeout=120,
                            env=dict(os.environ, HOME=str(home), XDG_CACHE_HOME=str(area / 'cache')))
    (area / 'stdout.txt').write_text(result.stdout)
    (area / 'stderr.txt').write_text(result.stderr)
    pdf = area / (path.stem + '.pdf')
    check(result.returncode == 0 and pdf.exists() and pdf.stat().st_size, label + ': fresh-profile conversion failed')
    check(digest(path) == before, label + ': consumer modified DOCX')
    bbox = area / 'bbox.html'
    subprocess.run(['pdftotext', '-bbox-layout', str(pdf), str(bbox)], capture_output=True, check=True, timeout=30)
    plain = subprocess.run(['pdftotext', '-layout', str(pdf), '-'], capture_output=True, text=True, check=True, timeout=30).stdout
    (area / 'complete-page-text.txt').write_text(plain)
    tree = E.parse(str(bbox))
    pages, lines = [], []
    for pi, page in enumerate(tree.xpath('//*[local-name()="page"]')):
        page_lines = []
        for line in page.xpath('.//*[local-name()="line"]'):
            words = line.xpath('./*[local-name()="word"]')
            if words:
                row = {'page': pi + 1, 'text': ' '.join(''.join(w.itertext()) for w in words),
                       **{k: float(line.get(k)) for k in ('xMin', 'yMin', 'xMax', 'yMax')}}
                lines.append(row)
                page_lines.append(row['text'])
        pages.append({'width': float(page.get('width')), 'height': float(page.get('height')), 'text_lines': page_lines})
    subprocess.run(['pdftoppm', '-r', '90', '-png', str(pdf), str(area / 'page')], capture_output=True, check=True, timeout=60)
    pngs = sorted(area.glob('page-*.png'))
    check(len(pngs) == len(pages), 'Consumer raster/page count differs')
    return {'path': str(pdf), 'sha256': digest(pdf), 'input_sha256': before, 'command': command,
            'fresh_process_profile_home': True, 'profile': str(profile), 'pages': pages,
            'complete_page_text': plain.split('\f')[:-1], 'lines': lines,
            'bbox_path': str(bbox), 'bbox_sha256': digest(bbox),
            'pngs': [{'path': str(p), 'sha256': digest(p)} for p in pngs]}


def compare_render(left, right):
    same_lines = [(r['page'], r['text']) for r in left['lines']] == [(r['page'], r['text']) for r in right['lines']]
    result = {'same_pages_and_complete_text': left['pages'] == right['pages'],
              'same_poppler_layout_text': left['complete_page_text'] == right['complete_page_text'],
              'same_line_breaks': same_lines, 'source_line_count': len(left['lines']),
              'output_line_count': len(right['lines'])}
    if same_lines:
        result['max_horizontal_delta_pt'] = max((abs(x[k] - y[k]) for x, y in zip(left['lines'], right['lines']) for k in ('xMin', 'xMax')), default=0)
        result['max_vertical_delta_pt'] = max((abs(x[k] - y[k]) for x, y in zip(left['lines'], right['lines']) for k in ('yMin', 'yMax')), default=0)
    pixels = []
    for a, b in zip(left['pngs'], right['pngs']):
        x, y = Image.open(a['path']).convert('RGB'), Image.open(b['path']).convert('RGB')
        if x.size != y.size:
            pixels.append({'pixel_equal': False, 'same_size': False})
            continue
        diff = ImageChops.difference(x, y)
        maximum = ImageChops.lighter(ImageChops.lighter(diff.getchannel('R'), diff.getchannel('G')), diff.getchannel('B'))
        pixels.append({'same_size': True, 'pixel_equal': diff.getbbox() is None,
                       'changed_pixel_count': x.width * x.height - maximum.histogram()[0]})
    result['raster_comparison'] = pixels
    result['strictly_equal'] = (result['same_pages_and_complete_text'] and result['same_poppler_layout_text'] and same_lines
                               and result.get('max_horizontal_delta_pt') == 0.0 and result.get('max_vertical_delta_pt') == 0.0
                               and len(left['pngs']) == len(right['pngs']) and all(p['pixel_equal'] for p in pixels))
    return result


def historical_color_witness(old_cli, cli, evidence, directory, name, item, soffice):
    """Exact historical witness for three identity-bound theme-shade limitations.

    The source raster comparison stays FAILED. This proves an existing color
    normalization independently; it does not turn those pixels into a pass or
    widen either pixel or coordinate thresholds for any other comparison.
    """
    check(digest(old_cli) == OLD_CLI_SHA, 'Wrong frozen historical theme-color CLI')
    check(item['sha256'] in COLOR_LIMITATION_SOURCES, 'Unapproved source identity for historical color witness')
    warning = 'Theme and automatic text colors normalized to fixed RGB or editor defaults'
    check(warning in item['import_response']['compatibility']['risks'], 'Missing documented theme-color normalization warning')
    comparison = compare_render(item['source_pdf'], item['output_pdf'])
    check(comparison['same_pages_and_complete_text'] and comparison['same_poppler_layout_text']
          and comparison['same_line_breaks'] and comparison['max_horizontal_delta_pt'] == 0.0
          and comparison['max_vertical_delta_pt'] == 0.0, 'Theme-color limitation may not cover any geometry/text drift')
    parts = package_parts(item['path'])
    main_name = related(parts, None, 'officeDocument')
    styles_name, theme_name = related(parts, main_name, 'styles'), related(parts, main_name, 'theme')
    styles, theme = E.fromstring(parts[styles_name]), E.fromstring(parts[theme_name])
    active_ids = {sid for row in item['independent']['rows'] for sid in row['style_chain']}
    declarations = [dict(color.attrib) for style in styles.findall(W + 'style') if style.get(W + 'styleId') in active_ids
                    for color in style.findall('w:rPr/w:color', NS)
                    if color.get(W + 'val') == '365F91' and color.get(W + 'themeColor') == 'accent1' and color.get(W + 'themeShade') == 'BF']
    check(declarations, 'Expected actual active theme-shade source XML absent')
    a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    accent = theme.find('.//' + a + 'clrScheme/' + a + 'accent1/' + a + 'srgbClr')
    check(accent is not None and accent.get('val') == '4F81BD', 'Actual source theme accent differs')

    def spans(pdf):
        out = []
        with fitz.open(pdf['path']) as doc:
            for index, page in enumerate(doc):
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines', []):
                        for span in line['spans']:
                            out.append({'page': index + 1, **span})
        return out

    source_spans, final_spans = spans(item['source_pdf']), spans(item['output_pdf'])
    check(len(source_spans) == len(final_spans), 'Theme witness cannot hide changed PDF text span count')
    color_changes = []
    for source, final in zip(source_spans, final_spans):
        check({k: v for k, v in source.items() if k != 'color'} == {k: v for k, v in final.items() if k != 'color'},
              'Theme witness may not hide any consumed text/font/size/origin/bbox/flags changes')
        if source['color'] != final['color']:
            check((source['color'], final['color']) == (0x365F91, 0x376092), 'Unexpected consumed color difference')
            color_changes.append({'page': source['page'], 'text': source['text'], 'source_rgb': '365F91', 'final_rgb': '376092'})
    check(color_changes, 'Source raster failure has no proved theme color difference')
    imported = run_cli(old_cli, evidence, 'import-docx', item['path'])
    def without_flags(node):
        node = copy.deepcopy(node)
        for key in FLAGS:
            node.get('attrs', {}).pop(key, None)
        if 'content' in node:
            node['content'] = [without_flags(n) for n in node['content']]
        return node
    check(imported['document'] == without_flags(item['import_response']['document']), 'Historical import differs beyond new source flags')
    old_output = create(old_cli, evidence, directory, name + '-historical-color-witness', imported['document'])
    old_pdf = render(soffice, directory, Path(old_output['path']), name + '-historical-color-witness')
    old_final = compare_render(old_pdf, item['output_pdf'])
    if name.startswith('original-'):
        check(old_final['strictly_equal'], 'Fresh genuine historical/final exact pixels or positions differ')
    old_spans = spans(old_pdf)
    nonlayout = lambda rows: [{k: v for k, v in row.items() if k not in ('page', 'origin', 'bbox')} for row in rows]
    check(nonlayout(old_spans) == nonlayout(final_spans), 'Historical/final consumed text/fonts/colors differ')
    def run_properties(path):
        p = package_parts(path)
        main = E.fromstring(p[related(p, None, 'officeDocument')])
        return [E.tostring(n, method='c14n').decode() for n in main.findall('.//' + W + 'rPr')]
    old_rpr, final_rpr = run_properties(old_output['path']), run_properties(item['export']['path'])
    check(old_rpr == final_rpr, 'Historical/final direct XML run formatting differs')
    check(digest(old_cli) == OLD_CLI_SHA and digest(cli) == evidence['cli_sha256'], 'CLI identity changed during historical witness')
    return {'classification': 'FAILED genuine/retained-source raster equality; known preexisting theme-shade normalization only',
            'source_sha256': item['sha256'], 'source_strict_comparison': comparison,
            'documented_warning': warning, 'source_styles_part': styles_name, 'source_styles_sha256': hashlib.sha256(parts[styles_name]).hexdigest(),
            'source_theme_part': theme_name, 'source_theme_sha256': hashlib.sha256(parts[theme_name]).hexdigest(),
            'active_source_color_declarations': declarations, 'source_accent1': '4F81BD', 'consumed_color_changes': color_changes,
            'historical_cli': str(old_cli), 'historical_cli_sha256': OLD_CLI_SHA, 'historical_export': old_output,
            'fresh_historical_pdf': old_pdf, 'historical_final_strict_comparison': old_final,
            'historical_final_consumed_text_fonts_colors_equal': True, 'historical_final_xml_run_properties_equal': True,
            'historical_and_final_run_properties': final_rpr,
            'scope': ('No source-raster pass claimed. Full source text/geometry stays exact; genuine originals additionally require fresh frozen-old output exactly equal final pixels/positions. '
                      'The retained precedence control uses historical imported model and run-color/font identity only, since old pagination flags were lost; a separately authored neutral-color precedence control must pass full raster equality.')}


def neutral_precedence(cli, evidence, directory):
    """Literal-black derivative is a separate strict pixel-positive source control."""
    original = evidence['sources']['control-precedence']
    check(original['sha256'] == 'b599b8fd12147bb7fbef041dea7d962953b061c7705140abf62bacf14a71fa98', 'Retained precedence source changed')
    active_ids = {sid for row in original['independent']['rows'] for sid in row['style_chain']}
    changed = []
    def neutral(main, styles, parts):
        for style in styles.findall(W + 'style'):
            if style.get(W + 'styleId') not in active_ids:
                continue
            for color in style.findall('w:rPr/w:color', NS):
                changed.append({'style': style.get(W + 'styleId'), 'before': dict(color.attrib)})
                color.attrib.clear()
                color.set(W + 'val', '000000')
    path = directory / 'control-precedence-neutral-color.docx'
    mutate(original['path'], path, neutral)
    check(len(changed) == 1 and changed[0]['style'] == 'Heading1', 'Neutral derivative must change only active Heading1 color')
    independent = inspect_docx(path)
    check(independent['rows'] == original['independent']['rows'] and independent['pageGeometry'] == original['independent']['pageGeometry'],
          'Neutral derivative changed source flags, order, type, text, style chain, or page geometry')
    name = 'control-precedence-neutral-color'
    item = {'path': str(path), 'sha256': digest(path), 'classification': 'independently authored literal-black derivative of retained precedence fixture; not a genuine original',
            'expected_action': 'roundtrip', 'render': True, 'independent': independent,
            'neutral_color_derivative': {'original_path': original['path'], 'original_sha256': original['sha256'], 'exact_changes': changed,
                                        'unchanged_pagination_text_type_style_page_oracle': True}}
    evidence['sources'][name] = item
    roundtrip(cli, evidence, directory, name, item)


def boundary_assertions(evidence):
    """Predetermined page ownership, never a result inferred from app output."""
    specifications = {
        'control-keepLines-true': {**{'TARGET ' + str(i): 2 for i in range(1, 6)}},
        'control-keepLines-false': {**{'TARGET ' + str(i): 1 if i <= 3 else 2 for i in range(1, 6)}},
        'control-keepNext-true': {'ANCHOR': 2, 'FOLLOWER': 2},
        'control-keepNext-false': {'ANCHOR': 1, 'FOLLOWER': 2},
        'control-pageBreakBefore-true': {'BEFORE': 1, 'AFTER': 2},
        'control-pageBreakBefore-false': {'BEFORE': 1, 'AFTER': 1},
        'chain-keepNext-true': {'ANCHOR A': 2, 'ANCHOR B': 2, 'CHAIN FOLLOWER': 2},
        'chain-keepNext-false': {'ANCHOR A': 1, 'ANCHOR B': 1, 'CHAIN FOLLOWER': 2},
        'document-start-true': {'FIRST PARAGRAPH': 1, 'SECOND PARAGRAPH': 1},
        'document-start-false': {'FIRST PARAGRAPH': 1, 'SECOND PARAGRAPH': 1},
        'adjacent-manual-break-true': {'BEFORE MANUAL BREAK': 1, 'AFTER MANUAL BREAK': 2, 'AFTER FOLLOWER': 2},
        'adjacent-manual-break-false': {'BEFORE MANUAL BREAK': 1, 'AFTER MANUAL BREAK': 2, 'AFTER FOLLOWER': 2},
        'direct-false-over-style-true': {'DIRECT FALSE ANCHOR': 1, 'DIRECT FALSE FOLLOWER': 2},
        'style-false-over-default-true': {'STYLE FALSE FIRST': 1, 'STYLE FALSE SECOND': 1},
        'explicit-style-does-not-overlay-default': {'DETACHED EXPLICIT STYLE': 1, 'DETACHED SECOND': 1},
    }
    results = {}
    for name, expected in specifications.items():
        if name not in evidence['sources']:
            continue
        result = {}
        for consumer in ('source_pdf', 'output_pdf'):
            rows = evidence['sources'][name][consumer]['lines']
            actual = {}
            for text, page in expected.items():
                found = [r for r in rows if r['text'] == text]
                check(len(found) == 1 and found[0]['page'] == page,
                      f'{name}/{consumer}: predetermined paragraph boundary {text!r} should be page {page}; got {found}')
                actual[text] = found[0]
            result[consumer] = actual
        results[name] = {'expected_pages': expected, 'actual_boundaries': result}
    for enabled in ('true', 'false'):
        name = 'long-keepLines-' + enabled
        if name not in evidence['sources']:
            continue
        for consumer in ('source_pdf', 'output_pdf'):
            rows = [r for r in evidence['sources'][name][consumer]['lines'] if r['text'].startswith('LONG LINE ')]
            check([r['text'] for r in rows] == ['LONG LINE ' + str(i).zfill(3) for i in range(1, 112)], 'Long paragraph content lost/duplicated')
            check(len({r['page'] for r in rows}) >= 3, 'Oversize keepLines control must exceed a page; keep is best effort')
    evidence['predetermined_boundary_checks'] = results


def historical_goldens(cli, old_cli, evidence, directory, historical_report):
    """Replay frozen no-pagination old models with byte-for-byte package assertions."""
    prior = json.loads(historical_report.read_text())
    records = evidence['historical_goldens'] = []
    evidence['historical_provenance'] = {'report': str(historical_report), 'report_sha256': digest(historical_report),
        'old_cli': str(old_cli), 'old_cli_sha256': digest(old_cli),
        'classification': 'Historical complete-package byte goldens and trusted models, not fresh source/render evidence'}
    unique = set()
    for i, item in enumerate(prior['goldens']):
        path = Path(item['source'])
        expected = item.get('source_sha256', item.get('sha256'))
        check(digest(path) == expected, 'Historical golden identity changed')
        check(expected not in unique, 'Duplicate historical package identity')
        unique.add(expected)
        parts = package_parts(path)
        names = [n for n in parts if n.startswith('rust-wps/model-v') and n.endswith('.json')]
        check(len(names) == 1, 'Expected one historical embedded model')
        model = json.loads(parts[names[0]])['document']
        def absent(node):
            return not any(k in node.get('attrs', {}) for k in FLAGS) and all(absent(n) for n in node.get('content', []))
        check(absent(model), 'Historical model has new pagination metadata')
        model_path = directory / ('historical-' + str(i) + '.json')
        save_json(model_path, model)
        record = {'source': str(path), 'source_sha256': expected, 'embedded_model_part': names[0], 'executions': {}}
        for label, program in (('frozen_old', old_cli), ('candidate', cli)):
            check(run_cli(program, evidence, 'read-docx-model', path) == model, 'Historical full model re-open changed')
            output = directory / ('historical-' + str(i) + '-' + label + '.docx')
            run_cli(program, evidence, 'create-docx', model_path, output)
            check(digest(output) == expected, 'Historical whole-package bytes changed: ' + str(path))
            check(run_cli(program, evidence, 'read-docx-model', output) == model, 'Historical generated full reopen differs')
            for target in (path, output):
                err = run_cli(program, evidence, 'create-docx', model_path, target, success=False)
                check('Output already exists' in err and digest(target) == expected, 'Historical no-overwrite failed')
            record['executions'][label] = {'output': str(output), 'sha256': expected, 'complete_package_equal': True,
                                            'complete_trusted_reopens': 2, 'no_overwrite_checks': 2}
        records.append(record)
        if (i + 1) % 50 == 0:
            print('historical byte goldens', i + 1, flush=True)


def resume_consumers(args):
    """Resume only missing consumers; preserve and hash-bind the blocked report."""
    previous_path = args.resume_render_report.resolve()
    previous = json.loads(previous_path.read_text())
    check(previous['status'] == 'failed' and 'Disk stop:' in previous.get('error', ''),
          'Resume mode only accepts an explicitly disk-blocked consumer report')
    cli = args.cli.resolve()
    check(digest(cli) == args.expected_cli_sha256 == previous['cli_sha256'], 'Resume CLI identity mismatch')
    args.root.mkdir(parents=True, exist_ok=True)
    disk_guard(args.root)
    directory = Path(tempfile.mkdtemp(prefix='resumed-', dir=args.root)).resolve()
    evidence = copy.deepcopy(previous)
    evidence.update(status='running', directory=str(directory), script_sha256=digest(__file__),
                    started_at_utc=datetime.now(timezone.utc).isoformat(), invocation=[sys.executable, *sys.argv],
                    free_disk_bytes_at_start=shutil.disk_usage(directory).free, commands=[])
    evidence.pop('error', None)
    evidence['reused_partial_consumer_report'] = {'path': str(previous_path), 'sha256': digest(previous_path),
        'script_sha256': previous['script_sha256'], 'original_status': previous['status'], 'original_error': previous['error'],
        'method': 'Every retained DOCX/model/PDF/bbox/raster and consumer binary hash is checked; all completed comparisons are recomputed without rendering. Only missing consumer PDFs use new per-file profiles in this new directory.',
        'retained_complete_trusted_reopens': sum(c['arguments'][0] == 'read-docx-model' and c['returncode'] == 0 for c in previous['commands'])}
    protected = {str(previous_path): digest(previous_path)}

    def protect(path, expected):
        check(digest(path) == expected, 'Retained resume artifact changed: ' + str(path))
        protected[str(path)] = expected

    def protect_export(item):
        protect(item['path'], item['sha256'])
        protect(item['model_path'], item['model_sha256'])

    def protect_render(item):
        protect(item['path'], item['sha256'])
        protect(item['bbox_path'], item['bbox_sha256'])
        for raster in item['pngs']:
            protect(raster['path'], raster['sha256'])

    try:
        for item in evidence['sources'].values():
            protect(item['path'], item['sha256'])
            if 'export' in item:
                protect_export(item['export'])
            for key in ('source_pdf', 'output_pdf'):
                if key in item:
                    protect_render(item[key])
        for item in evidence['frontend_snapshots'].values():
            protect(item['model_path'], item['model_sha256'])
            protect(item['reference_path'], item['reference_sha256'])
            protect_export(item['export'])
            for key in ('reference_pdf', 'output_pdf'):
                if key in item:
                    protect_render(item[key])
        manifest = evidence['frontend_manifest']
        protect(manifest['path'], manifest['sha256'])
        evidence['frontend_operation_checks'] = frontend_operation_expectations(Path(manifest['path']).parent, manifest['contents'])
        for witness in evidence.get('known_source_raster_failures', {}).values():
            protect(witness['historical_cli'], witness['historical_cli_sha256'])
            protect_export(witness['historical_export'])
            protect_render(witness['fresh_historical_pdf'])
        for path, sha in evidence['consumers']['binaries_and_launcher_sha256'].items():
            protect(path, sha)
        original_check = evidence['reused_cli_check_report']
        protect(original_check['path'], original_check['sha256'])
        retained_pdfs = {p for p in protected if p.endswith('.pdf')}
        evidence['reused_partial_consumer_report']['retained_pdf_count'] = len(retained_pdfs)
        comparisons = evidence['render_comparisons']
        for name, item in evidence['sources'].items():
            if not item['render']:
                continue
            check('source_pdf' in item and 'output_pdf' in item, 'Resume requires previously completed source pairs')
            actual = compare_render(item['source_pdf'], item['output_pdf'])
            check(actual == comparisons[name], 'Retained source comparison changed')
            if not actual['strictly_equal']:
                witness = evidence['known_source_raster_failures'][name]
                check(actual == witness['source_strict_comparison'], 'Retained source color failure changed')
                old_final = compare_render(witness['fresh_historical_pdf'], item['output_pdf'])
                check(old_final == witness['historical_final_strict_comparison'], 'Historical witness changed')
                if name.startswith('original-'):
                    check(old_final['strictly_equal'], 'Genuine old/final pixel witness no longer exact')
        soffice = args.soffice or shutil.which('soffice')
        check(soffice is not None, 'LibreOffice needed to resume')
        check(str(Path(soffice).resolve()) in evidence['consumers']['binaries_and_launcher_sha256'], 'Resume must use same LibreOffice launcher')
        for name, item in evidence['frontend_snapshots'].items():
            if 'reference_pdf' not in item:
                item['reference_pdf'] = render(soffice, directory, Path(item['reference_path']), 'frontend-' + name + '-reference')
            if 'output_pdf' not in item:
                item['output_pdf'] = render(soffice, directory, Path(item['export']['path']), 'frontend-' + name + '-output')
            comparison = comparisons['frontend-' + name] = compare_render(item['reference_pdf'], item['output_pdf'])
            check(comparison['strictly_equal'], 'Frontend resume ' + name + ': strict zero-drift comparison failed: ' + json.dumps(comparison))
            print('strict frontend consumer equal', name, flush=True)
        boundary_assertions(evidence)
        fresh_pdfs = list((directory / 'renders').glob('*/*.pdf'))
        evidence['summary'] = {'supported_sources': sum(i['expected_action'] == 'roundtrip' for i in evidence['sources'].values()),
            'rejected_sources': sum(i['expected_action'] == 'reject' for i in evidence['sources'].values()),
            'frontend_snapshots': len(evidence['frontend_snapshots']), 'newly_rendered_pdfs_this_resume': len(fresh_pdfs),
            'hash_verified_fresh_pdfs_from_prior_partial_run': len(retained_pdfs), 'total_current_feature_consumer_pdfs': len(retained_pdfs) + len(fresh_pdfs),
            'render_comparison_pairs': len(comparisons), 'strict_consumer_pairs': sum(c['strictly_equal'] for c in comparisons.values()),
            'known_source_raster_failures': len(evidence.get('known_source_raster_failures', {})),
            'historical_byte_goldens_reused_from_final_source_check': original_check['summary']['historical_byte_goldens']}
        evidence['status'] = 'passed-with-known-source-raster-failures'
    except BaseException as error:
        evidence['status'], evidence['error'] = 'failed', type(error).__name__ + ': ' + str(error)
        raise
    finally:
        evidence['resume_protected_hashes'] = protected
        evidence['resume_protected_hashes_unchanged'] = all(digest(p) == h for p, h in protected.items())
        evidence['checker_unchanged'] = digest(__file__) == evidence['script_sha256']
        evidence['cli_unchanged'] = digest(cli) == args.expected_cli_sha256
        evidence['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
        evidence['free_disk_bytes_at_finish'] = shutil.disk_usage(directory).free
        okay = evidence['resume_protected_hashes_unchanged'] and evidence['checker_unchanged'] and evidence['cli_unchanged']
        if not okay:
            evidence['status'], evidence['identity_error'] = 'failed', 'Protected resume identities changed'
        report = directory / 'report.json'
        save_json(report, evidence)
        print('REPORT', report, flush=True)
        check(okay, 'Protected resume identities changed')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path, required=True)
    parser.add_argument('--expected-cli-sha256', required=True)
    parser.add_argument('--root', type=Path, default=ROOT / 'tmp/paragraph-pagination-independent')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--soffice')
    parser.add_argument('--editor-models', type=Path)
    parser.add_argument('--frontend-only', action='store_true')
    parser.add_argument('--render-report', type=Path,
                        help='Freshly render hash-verified source/export pairs from a passed frozen CLI check report; do not repeat old CLI/golden calls')
    parser.add_argument('--resume-render-report', type=Path,
                        help='Resume only missing frontend consumers from a preserved disk-blocked render report, after exact artifact hash verification')
    parser.add_argument('--historical-cli', type=Path)
    parser.add_argument('--historical-render-cli', type=Path,
                        help='Fresh old-CLI pixel witness for exact identity-bound known theme-shade source raster limitations; never relax geometry/pixel tolerances')
    parser.add_argument('--historical-report', type=Path, default=ROOT / 'tmp/contextual-legacy/run-oegw_f99/report.json')
    args = parser.parse_args()
    if args.resume_render_report:
        check(not args.render_report and not args.editor_models and not args.historical_cli, 'Resume uses only its hash-bound retained inputs')
        resume_consumers(args)
        return
    check(not args.frontend_only or args.editor_models, '--frontend-only needs --editor-models')
    check(not args.render_report or (args.render and not args.frontend_only and not args.historical_cli),
          '--render-report requires --render and excludes --frontend-only/--historical-cli')
    cli = args.cli.resolve()
    check(digest(cli) == args.expected_cli_sha256, 'CLI is not the authorized frozen build')
    args.root.mkdir(parents=True, exist_ok=True)
    disk_guard(args.root)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    evidence = {'status': 'running', 'directory': str(directory), 'script_sha256': digest(__file__),
                'cli': str(cli), 'cli_sha256': digest(cli), 'commands': [], 'sources': {},
                'limitations': LIMITATIONS, 'invocation': [sys.executable, *sys.argv],
                'started_at_utc': datetime.now(timezone.utc).isoformat(), 'python_version': platform.python_version(),
                'free_disk_bytes_at_start': shutil.disk_usage(directory).free}
    try:
        if args.render_report:
            previous_path = args.render_report.resolve()
            previous = json.loads(previous_path.read_text())
            check(previous['status'] == 'passed', 'Only passed CLI evidence may feed fresh consumer checks')
            check(previous['cli_sha256'] == evidence['cli_sha256'], 'Reused CLI evidence belongs to another binary')
            for item in previous['sources'].values():
                check(digest(item['path']) == item['sha256'], 'Reused source identity changed')
                if 'export' in item:
                    output = item['export']
                    check(digest(output['path']) == output['sha256'] and digest(output['model_path']) == output['model_sha256'],
                          'Reused model/export identity changed')
            evidence['sources'] = copy.deepcopy(previous['sources'])
            evidence['reused_cli_check_report'] = {'path': str(previous_path), 'sha256': digest(previous_path),
                'summary': previous['summary'], 'script_sha256': previous['script_sha256'],
                'method': 'Prior passed final-CLI source/model/tamper/302-golden checks explicitly reused after source/model/export hash verification. Source PDFs in this report are newly rendered with separate profiles.'}
        elif not args.frontend_only:
            sources = evidence['sources'] = fixtures(directory)
            for name, item in sources.items():
                if item['expected_action'] == 'roundtrip':
                    roundtrip(cli, evidence, directory, name, item)
                else:
                    if name == 'reject-active-style-duplicate-id':
                        # Package parsing itself correctly refuses this identity
                        # ambiguity before an audit JSON can be produced.
                        item['audit_error'] = run_cli(cli, evidence, 'audit-docx', item['path'], success=False)
                        check('duplicate DOCX style identifier' in item['audit_error'], 'Unexpected duplicate-style audit refusal')
                    else:
                        item['audit'] = run_cli(cli, evidence, 'audit-docx', item['path'])
                    item['error'] = run_cli(cli, evidence, 'import-docx', item['path'], success=False)
                    check(any(token in item['error'] for token in ('Unsupported', 'Invalid', 'malformed XML')), name + ': unclear refusal')
                check(digest(item['path']) == item['sha256'], name + ': source changed')
                print('verified source', name, flush=True)
            good = sources['control-precedence']['import_response']['document']
            save_json(directory / 'editor-input.json', good)
            invalid_models(cli, evidence, directory, good)
            tamper(cli, evidence, directory, sources['control-precedence']['export']['path'])
        if args.historical_render_cli and 'control-precedence' in evidence['sources']:
            neutral_precedence(cli, evidence, directory)
        if args.editor_models:
            frontend_snapshots(cli, evidence, directory, args.editor_models.resolve())
        if args.historical_cli:
            historical_goldens(cli, args.historical_cli.resolve(), evidence, directory, args.historical_report.resolve())
        if args.render:
            soffice = args.soffice or shutil.which('soffice')
            check(soffice is not None, 'LibreOffice required')
            evidence['consumers'] = consumer_identity(soffice, directory)
            comparisons = evidence['render_comparisons'] = {}
            for name, item in evidence['sources'].items():
                if not item['render']:
                    continue
                item['source_pdf'] = render(soffice, directory, Path(item['path']), name + '-source')
                item['output_pdf'] = render(soffice, directory, Path(item['export']['path']), name + '-output')
                comparison = comparisons[name] = compare_render(item['source_pdf'], item['output_pdf'])
                if not comparison['strictly_equal'] and args.historical_render_cli:
                    witness = historical_color_witness(args.historical_render_cli.resolve(), cli, evidence, directory, name, item, soffice)
                    evidence.setdefault('known_source_raster_failures', {})[name] = witness
                    print('source raster FAILED; exact historical color witness verified', name, flush=True)
                else:
                    check(comparison['strictly_equal'], name + ': consumer zero-drift comparison failed: ' + json.dumps(comparison))
                    print('strict consumer equal', name, flush=True)
            for name, item in evidence.get('frontend_snapshots', {}).items():
                item['reference_pdf'] = render(soffice, directory, Path(item['reference_path']), 'frontend-' + name + '-reference')
                item['output_pdf'] = render(soffice, directory, Path(item['export']['path']), 'frontend-' + name + '-output')
                comparison = comparisons['frontend-' + name] = compare_render(item['reference_pdf'], item['output_pdf'])
                check(comparison['strictly_equal'], 'Frontend ' + name + ': consumer zero-drift comparison failed: ' + json.dumps(comparison))
                print('strict frontend consumer equal', name, flush=True)
            boundary_assertions(evidence)
        evidence['summary'] = {'supported_sources': sum(i['expected_action'] == 'roundtrip' for i in evidence['sources'].values()),
                               'rejected_sources': sum(i['expected_action'] == 'reject' for i in evidence['sources'].values()),
                               'invalid_models': len(evidence.get('invalid_models', {})),
                               'tampered_packages': len(evidence.get('tampered_packages', {})),
                               'frontend_snapshots': len(evidence.get('frontend_snapshots', {})),
                               'historical_byte_goldens': len(evidence.get('historical_goldens', [])),
                               'newly_rendered_pdfs': len(list((directory / 'renders').glob('*/*.pdf'))),
                               'render_comparison_pairs': len(evidence.get('render_comparisons', {})),
                               'strict_consumer_pairs': sum(c['strictly_equal'] for c in evidence.get('render_comparisons', {}).values()),
                               'known_source_raster_failures': len(evidence.get('known_source_raster_failures', {})),
                               'complete_trusted_reopens': sum(c['arguments'][0] == 'read-docx-model' and c['returncode'] == 0 for c in evidence['commands'])}
        evidence['status'] = ('passed-with-known-source-raster-failures'
                              if evidence.get('known_source_raster_failures') else 'passed')
    except BaseException as error:
        evidence['status'], evidence['error'] = 'failed', type(error).__name__ + ': ' + str(error)
        raise
    finally:
        evidence['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
        evidence['cli_unchanged'] = digest(cli) == evidence['cli_sha256']
        evidence['checker_unchanged'] = digest(__file__) == evidence['script_sha256']
        evidence['sources_unchanged'] = all(digest(i['path']) == i['sha256'] for i in evidence['sources'].values())
        evidence['genuine_originals_unchanged'] = all(digest(ROOT / 'tmp/paragraph-indentation-source/poi' / (name + '.docx')) == spec['sha256'] for name, spec in ORIGINALS.items())
        evidence['frontend_inputs_unchanged'] = all(digest(i['model_path']) == i['model_sha256'] for i in evidence.get('frontend_snapshots', {}).values())
        evidence['consumer_binaries_unchanged'] = all(digest(p) == h for p, h in evidence.get('consumers', {}).get('binaries_and_launcher_sha256', {}).items())
        evidence['historical_render_cli_unchanged'] = all(digest(i['historical_cli']) == i['historical_cli_sha256'] for i in evidence.get('known_source_raster_failures', {}).values())
        evidence['historical_identities_unchanged'] = all(digest(i['source']) == i['source_sha256'] for i in evidence.get('historical_goldens', []))
        if 'historical_provenance' in evidence:
            historic = evidence['historical_provenance']
            evidence['historical_identities_unchanged'] &= digest(historic['report']) == historic['report_sha256'] and digest(historic['old_cli']) == historic['old_cli_sha256']
        evidence['reused_cli_evidence_unchanged'] = True
        if 'reused_cli_check_report' in evidence:
            previous = evidence['reused_cli_check_report']
            evidence['reused_cli_evidence_unchanged'] = digest(previous['path']) == previous['sha256']
            evidence['reused_cli_evidence_unchanged'] &= all(
                digest(i['export']['path']) == i['export']['sha256'] and digest(i['export']['model_path']) == i['export']['model_sha256']
                for i in evidence['sources'].values() if 'export' in i)
        identity_keys = ('cli_unchanged', 'checker_unchanged', 'sources_unchanged', 'genuine_originals_unchanged',
                         'frontend_inputs_unchanged', 'consumer_binaries_unchanged', 'historical_identities_unchanged', 'reused_cli_evidence_unchanged', 'historical_render_cli_unchanged')
        identity_failures = [key for key in identity_keys if not evidence[key]]
        if identity_failures:
            evidence['status'] = 'failed'
            evidence['identity_error'] = 'Protected identities changed: ' + ', '.join(identity_failures)
        evidence['free_disk_bytes_at_finish'] = shutil.disk_usage(directory).free
        report = directory / 'report.json'
        save_json(report, evidence)
        print('REPORT', report, flush=True)
        check(not identity_failures, evidence.get('identity_error', ''))


if __name__ == '__main__':
    main()
