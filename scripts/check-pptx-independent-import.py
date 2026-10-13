"""Independent producer/consumer check for the supported basic PPTX subset."""
import hashlib
import json
import subprocess
import tempfile
import zipfile
from pathlib import Path
from lxml import etree

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
CORE = ROOT / "target/debug/rust-wps.exe"
DESKTOP = ROOT / "src-tauri/target/debug/rust-wps-desktop.exe"


def run(*args):
    return json.loads(subprocess.check_output([str(CORE), *map(str, args)], timeout=40))


def produce(path):
    deck = Presentation()
    # Retain the standard template and its placeholder authoring prompts.
    # Slide content below has direct geometry; template style fidelity is
    # still outside the supported import subset.
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
    # Exercise the standard lexical boolean spelling and integer zero.
    transform = box._element.spPr.xfrm
    transform.set("flipH", "false")
    transform.set("flipV", "0")
    transform.set("rot", "-0")
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "独立工具文稿 <&>"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    slide.notes_slide.notes_text_frame.text = "独立备注\n第二行"
    # This fixture deliberately opts out of Office extension metadata and
    # printer settings. They remain blocked by the Rust importer.
    for part in list(deck.part.package.iter_parts()):
        element = getattr(part, "_element", None)
        generic_xml = element is None and str(part.partname).endswith(".xml")
        if generic_xml:
            element = etree.fromstring(part.blob)
        if element is not None:
            for extension in list(element.iter()):
                if extension.tag.endswith("}extLst"):
                    extension.getparent().remove(extension)
        if generic_xml:
            part._blob = etree.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)
        for relationship in list(part.rels.values()):
            if relationship.reltype.endswith("/printerSettings"):
                part.drop_rel(relationship.rId)
    deck.save(path)


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        source = folder / "source.pptx"
        produce(source)
        with zipfile.ZipFile(source) as package:
            assert b"Click to edit Master title style" in package.read("ppt/slideMasters/slideMaster1.xml")
        before = hashlib.sha256(source.read_bytes()).digest()
        report = run("import-pptx", source)
        assert not report["blockers"], report["blockers"]
        worker = json.loads(subprocess.check_output(
            [str(DESKTOP), "--document-worker"],
            input=json.dumps({"operation": "import_presentation", "path": str(source)}).encode(),
            timeout=40,
        ))
        assert worker["result"] == "presentation_imported"
        assert worker["report"] == report
        model = report["model"]
        box = model["slides"][0]["text_boxes"][0]
        assert box["text"] == "独立工具文稿 <&>"
        assert box["font_size_hundredths_pt"] == 2800
        assert box["bold"] and box["color_rgb"] == [18, 52, 86]
        assert model["slides"][0]["notes"].rstrip("\n") == "独立备注\n第二行"
        box["text"] = "编辑后中文 <&>"
        model_path = folder / "model.json"
        model_path.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
        output = folder / "edited.pptx"
        run("create-pptx", model_path, output)
        actual = Presentation(output)
        shape = actual.slides[0].shapes[0]
        assert shape.text == box["text"]
        assert shape.text_frame.paragraphs[0].runs[0].font.size == Pt(28)
        assert shape.text_frame.paragraphs[0].runs[0].font.bold
        assert tuple(shape.text_frame.paragraphs[0].runs[0].font.color.rgb) == (18, 52, 86)
        assert actual.slides[0].notes_slide.notes_text_frame.text.rstrip("\n") == "独立备注\n第二行"
        assert [actual.slide_width, actual.slide_height] == model["size_emu"]
        assert [shape.left, shape.top, shape.width, shape.height] == box["bounds_emu"]
        assert run("read-pptx-model", output) == model
        assert hashlib.sha256(source.read_bytes()).digest() == before
        print("Independent python-pptx author -> Rust import/edit/export -> python-pptx checks passed.")
