#!/usr/bin/env python3
"""Independent bounded decimal-list geometry verification.

Uses only the frozen CLI, independently resolved source OOXML, python-docx,
LibreOffice with a new profile per conversion, PDF coordinates and raster pixels.
The genuine upstream source is hash-pinned and never edited or resaved in place.
Derived ZIP mutations are labelled directed synthetic controls, not originals.
No browser, native Word/WPS or arbitrary-list/full-document fidelity claim.
"""
import argparse
import copy
import hashlib
import json
import os
import platform
import posixpath
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt, Twips
from docx.text.paragraph import Paragraph
from lxml import etree as E
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parent.parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
W = '{' + NS['w'] + '}'
KEYS = ('indentLeftTwips', 'indentRightTwips', 'firstLineIndentTwips')
ORIGINAL_SHA = '48209f95d7a2ddfbf0fb3291658b1463e7b8121dac982f26501d8775e5450707'
ORIGINAL_BLOB = '975b598f476e97e08366b102f5b60b8211e384b1'
ORIGINAL_REVISION = '6e5f11edfc5f080cf99e1807916df5cedec785b7'
ORIGINAL_URL = 'https://github.com/LibreOffice/core/blob/' + ORIGINAL_REVISION + '/sw/qa/extras/ooxmlexport/data/testParaListRightIndent.docx'
LIMITATIONS = ('Genuine unchanged public upstream regression fixture, not a user production file. '
               'Directed synthetic controls are separate derived files. Render evidence covers '
               'LibreOfficeDev only; no real application/browser, WebView2, native printing, '
               'Microsoft Word/WPS, arbitrary-list, or whole-document fidelity acceptance.')


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def child(parent, tag, **attrs):
    node = E.SubElement(parent, W + tag)
    for key, value in attrs.items():
        node.set(W + key, str(value))
    return node


def val(parent, path, default=None):
    node = None if parent is None else parent.find(path, NS)
    return default if node is None else node.get(W + 'val', default)


def indentation(ppr):
    node = None if ppr is None else ppr.find('w:ind', NS)
    if node is None:
        return {}
    check(not (node.get(W + 'firstLine') is not None and node.get(W + 'hanging') is not None),
          'Independent resolver refuses competing firstLine/hanging')
    result = {target: int(node.get(W + source)) for source, target in zip(('left', 'right'), KEYS[:2])
              if node.get(W + source) is not None}
    if node.get(W + 'firstLine') is not None:
        result[KEYS[2]] = int(node.get(W + 'firstLine'))
    elif node.get(W + 'hanging') is not None:
        result[KEYS[2]] = -int(node.get(W + 'hanging'))
    return result


def unique_map(nodes, attribute, label):
    result = {}
    for node in nodes:
        key = node.get(W + attribute)
        check(key is not None and key not in result, 'Independent resolver ambiguous ' + label)
        result[key] = node
    return result


def relationship(parts, base, kind, required=True):
    relpath = posixpath.join(posixpath.dirname(base), '_rels', posixpath.basename(base) + '.rels')
    rels = E.fromstring(parts[relpath])
    matches = [r for r in rels if r.get('Type', '').rsplit('/', 1)[-1] == kind]
    check(len(matches) == 1 or (not required and not matches), 'Independent resolver needs unique ' + kind + ' relationship')
    if not matches:
        return None
    rel = matches[0]
    check(rel.get('TargetMode') != 'External', 'External ' + kind + ' is outside independent scope')
    target = rel.get('Target')
    result = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(base), target))
    check(result in parts, 'Missing actual relationship part: ' + result)
    return result


def paragraph_text(para):
    return ''.join((n.text or '') if n.tag == W + 't' else '\n' if n.tag == W + 'br' else '\t'
                   for n in para.iter() if n.tag in (W + 't', W + 'br', W + 'tab'))


def inspect_docx(path):
    """Resolve identities, style ancestry and level geometry without application code."""
    with zipfile.ZipFile(path) as package:
        check(len(package.namelist()) == len(set(package.namelist())), 'Duplicate ZIP member')
        parts = {n: package.read(n) for n in package.namelist()}
    rels = E.fromstring(parts['_rels/.rels'])
    main_rels = [n for n in rels if n.get('Type', '').endswith('/officeDocument')]
    check(len(main_rels) == 1, 'Independent reader needs unique officeDocument')
    main_name = main_rels[0].get('Target').lstrip('/')
    main = E.fromstring(parts[main_name])
    styles_name, numbering_name = [relationship(parts, main_name, k) for k in ('styles', 'numbering')]
    styles, numbering = E.fromstring(parts[styles_name]), E.fromstring(parts[numbering_name])
    style_map = unique_map(styles.findall('w:style', NS), 'styleId', 'style identity')
    num_map = unique_map(numbering.findall('w:num', NS), 'numId', 'num identity')
    abstract_map = unique_map(numbering.findall('w:abstractNum', NS), 'abstractNumId', 'abstract identity')
    defaults = styles.find('w:docDefaults', NS)
    default_styles = [n for n in style_map.values() if n.get(W+'type') == 'paragraph' and n.get(W+'default') in ('1','true','on')]
    check(len(default_styles) <= 1, 'Ambiguous paragraph default')
    default_id = default_styles[0].get(W+'styleId') if default_styles else None
    theme_name = relationship(parts, main_name, 'theme', False)
    theme = E.fromstring(parts[theme_name]) if theme_name else None
    themes = {}
    if theme is not None:
        for name in ('major', 'minor'):
            latin = theme.find('.//a:fontScheme/a:' + name + 'Font/a:latin', NS)
            if latin is not None:
                themes[name + 'HAnsi'] = latin.get('typeface')
    def chain(sid, active=()):
        if sid is None:
            return []
        check(sid not in active and sid in style_map, 'Independent resolver missing/cyclic style ' + str(sid))
        style = style_map[sid]
        return chain(val(style, 'w:basedOn'), active + (sid,)) + [style]
    def run_style(rpr, effective):
        if rpr is None:
            return
        fonts = rpr.find('w:rFonts', NS)
        if fonts is not None:
            face = fonts.get(W+'ascii') or themes.get(fonts.get(W+'asciiTheme'))
            if face:
                effective['fontFamily'] = face
        size = val(rpr, 'w:sz')
        if size is not None:
            effective['fontSizeHalfPoints'] = int(size)
    rows, counters = [], {}
    for index, para in enumerate(main.findall('.//w:body//w:p', NS)):
        ppr = para.find('w:pPr', NS)
        sid = val(ppr, 'w:pStyle', default_id)
        lineage = chain(sid)
        pprs = ([defaults.find('w:pPrDefault/w:pPr', NS)] if defaults is not None else []) + [s.find('w:pPr', NS) for s in lineage] + [ppr]
        identity = {}
        for pp in pprs:
            np = None if pp is None else pp.find('w:numPr', NS)
            for key in ('numId', 'ilvl'):
                value = val(np, 'w:' + key)
                if value is not None:
                    identity[key] = int(value)
        effective = {}
        marker = None
        if identity.get('numId', 0):
            num_id = str(identity['numId']); level_id = str(identity.get('ilvl', 0))
            check(num_id in num_map, 'Missing independently referenced numId')
            instance = num_map[num_id]
            abstract_id = val(instance, 'w:abstractNumId')
            check(abstract_id in abstract_map, 'Missing independently referenced abstractNum')
            abstract = abstract_map[abstract_id]
            levels = unique_map(abstract.findall('w:lvl', NS), 'ilvl', 'level identity')
            check(level_id in levels, 'Missing independently referenced level')
            level = levels[level_id]
            overrides = [n for n in instance.findall('w:lvlOverride', NS) if n.get(W+'ilvl') == level_id]
            check(len(overrides) <= 1, 'Ambiguous level override')
            start = int(val(level, 'w:start', '1'))
            if overrides:
                check(overrides[0].find('w:lvl', NS) is None, 'Full level override outside reference resolver')
                start = int(val(overrides[0], 'w:startOverride', str(start)))
            effective.update(indentation(level.find('w:pPr', NS)))
            marker = {'numId': int(num_id), 'abstractNumId': int(abstract_id), 'level': int(level_id),
                      'start': start, 'format': val(level,'w:numFmt'), 'text': val(level,'w:lvlText'),
                      'suffix': val(level,'w:suff','tab'), 'alignment': val(level,'w:lvlJc','left')}
            key = (num_id, level_id)
            counters[key] = counters.get(key, start-1) + 1
            marker['ordinal'] = counters[key]
            marker['visible'] = marker['text'].replace('%'+str(int(level_id)+1), str(counters[key]))
            font = {}
            run_style(None if defaults is None else defaults.find('w:rPrDefault/w:rPr', NS), font)
            for st in lineage:
                run_style(st.find('w:rPr', NS), font)
            run_style(None if ppr is None else ppr.find('w:rPr', NS), font)
            run_style(level.find('w:rPr', NS), font)
            marker.update(font)
        for pp in pprs:
            effective.update(indentation(pp))
        rows.append({'index': index, 'text': paragraph_text(para), 'style': sid,
                     'style_chain': [s.get(W+'styleId') for s in lineage], 'indent': effective,
                     'direct': indentation(ppr), 'marker': marker,
                     'in_table': any(p.tag == W+'tc' for p in para.iterancestors())})
    doc = Document(path)
    pyrows = []
    for element in doc.element.body.iter(W+'p'):
        p = Paragraph(element, doc); f = p.paragraph_format
        pyrows.append({'text': p.text, 'direct': {k:v.twips for k,v in zip(KEYS,(f.left_indent,f.right_indent,f.first_line_indent)) if v is not None}})
    check(len(pyrows) == len(rows), 'Independent OOXML/python-docx paragraph count differs')
    for xml, py in zip(rows,pyrows):
        check(xml['text'] == py['text'] and xml['direct'] == py['direct'], 'Independent OOXML/python-docx paragraph text/direct indentation disagreement')
    section = main.find('w:body/w:sectPr', NS)
    size, margins = section.find('w:pgSz',NS), section.find('w:pgMar',NS)
    page = {'widthTwips': int(size.get(W+'w')), 'heightTwips': int(size.get(W+'h'))}
    page.update({'margin'+side.title()+'Twips': int(margins.get(W+side)) for side in ('top','right','bottom','left')})
    return {'rows': rows, 'pageGeometry': page, 'parts': {'document':main_name,'styles':styles_name,'numbering':numbering_name,'theme':theme_name}, 'python_docx':pyrows}


def compare_rows(actual, expected, label):
    check(len(actual) == len(expected), label + ': paragraph count changed')
    for i,(a,b) in enumerate(zip(actual,expected)):
        for key in ('text','indent','in_table'):
            check(a[key] == b[key], f'{label} paragraph {i} {key}: {a[key]!r} != {b[key]!r}')
        check((a['marker'] is None) == (b['marker'] is None), label + ': marker presence changed')
        if a['marker'] is not None:
            for key in ('level','ordinal','format','text','suffix','alignment','fontFamily','fontSizeHalfPoints','visible'):
                check(a['marker'].get(key) == b['marker'].get(key), f'{label} marker {i} {key}: {a["marker"].get(key)!r} != {b["marker"].get(key)!r}')


def mutate(source, target, action):
    check(not target.exists(), 'Derived control must be fresh')
    with zipfile.ZipFile(source) as package:
        parts = {n: package.read(n) for n in package.namelist()}
    trees = {k:E.fromstring(parts['word/'+k+'.xml']) for k in ('document','styles','numbering','settings')}
    action(trees,parts)
    for key, tree in trees.items():
        parts['word/'+key+'.xml'] = E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as package:
        for name,data in parts.items():
            package.writestr(name,data)


def run_cli(cli, evidence, *args, success=True):
    result = subprocess.run([str(cli),*map(str,args)],text=True,capture_output=True,timeout=90)
    entry = {'program':str(cli),'arguments':list(map(str,args)),'returncode':result.returncode,'stdout':result.stdout,'stderr':result.stderr}
    evidence['commands'].append(entry)
    check((result.returncode == 0) == success, 'Unexpected CLI outcome: '+json.dumps(entry))
    return json.loads(result.stdout) if success else result.stderr


def render(soffice, directory, path, label):
    area = directory/'renders'/label; area.mkdir(parents=True)
    profile, home = area/'lo-profile', area/'home'; home.mkdir()
    check(not profile.exists(), 'LibreOffice profile must be new')
    before = digest(path)
    command = [soffice,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(area),str(path)]
    result = subprocess.run(command,text=True,capture_output=True,timeout=120,env=dict(os.environ,HOME=str(home),XDG_CACHE_HOME=str(area/'cache')))
    (area/'stdout.txt').write_text(result.stdout);(area/'stderr.txt').write_text(result.stderr)
    pdf = area/(Path(path).stem+'.pdf')
    check(result.returncode == 0 and pdf.is_file() and pdf.stat().st_size, label+': fresh-profile LibreOffice conversion failed')
    check(digest(path) == before, label+': consumer mutated input')
    bbox=area/'bbox.html'
    subprocess.run(['pdftotext','-bbox-layout',str(pdf),str(bbox)],capture_output=True,check=True,timeout=30)
    tree=E.parse(str(bbox));pages=[];lines=[]
    for pi,page in enumerate(tree.xpath('//*[local-name()="page"]')):
        pages.append({'width':float(page.get('width')),'height':float(page.get('height'))})
        for line in page.xpath('.//*[local-name()="line"]'):
            words=line.xpath('./*[local-name()="word"]')
            if words:
                lines.append({'page':pi+1,'text':' '.join(''.join(w.itertext()) for w in words),**{k:float(line.get(k)) for k in ('xMin','yMin','xMax','yMax')}})
    subprocess.run(['pdftoppm','-r','90','-png',str(pdf),str(area/'page')],capture_output=True,check=True,timeout=60)
    pngs=sorted(area.glob('page-*.png'));check(len(pngs)==len(pages),label+': raster count mismatch')
    return {'path':str(pdf),'sha256':digest(pdf),'input_sha256':before,'command':command,'fresh_profile':True,'profile':str(profile),'pages':pages,'lines':lines,'pngs':[str(p) for p in pngs]}


def compare_render(a,b):
    same_lines=[(r['page'],r['text']) for r in a['lines']]==[(r['page'],r['text']) for r in b['lines']]
    result={'same_pages':a['pages']==b['pages'],'same_line_breaks':same_lines,'source_line_count':len(a['lines']),'output_line_count':len(b['lines'])}
    if same_lines:
        result['max_horizontal_delta_pt']=max((abs(x[k]-y[k]) for x,y in zip(a['lines'],b['lines']) for k in ('xMin','xMax')),default=0)
        result['max_vertical_delta_pt']=max((abs(x[k]-y[k]) for x,y in zip(a['lines'],b['lines']) for k in ('yMin','yMax')),default=0)
    pixels=[]
    for left,right in zip(a['pngs'],b['pngs']):
        x,y=Image.open(left).convert('RGB'),Image.open(right).convert('RGB')
        if x.size!=y.size:
            pixels.append({'same_size':False});continue
        diff=ImageChops.difference(x,y); maxdiff=ImageChops.lighter(ImageChops.lighter(diff.getchannel('R'),diff.getchannel('G')),diff.getchannel('B'))
        pixels.append({'same_size':True,'pixel_equal':diff.getbbox() is None,'changed_pixel_count':x.width*x.height-maxdiff.histogram()[0],'pixels':x.width*x.height})
    result['raster_comparison']=pixels
    result['has_layout_drift']=not result['same_pages'] or not same_lines or result.get('max_horizontal_delta_pt',0)!=0.0 or result.get('max_vertical_delta_pt',0)!=0.0
    return result


def original_preflight(original):
    check(digest(original)==ORIGINAL_SHA,'Original SHA-256 mismatch')
    raw=original.read_bytes()
    check(hashlib.sha1(b'blob '+str(len(raw)).encode()+b'\0'+raw).hexdigest()==ORIGINAL_BLOB,'Original upstream Git blob mismatch')
    source=inspect_docx(original)
    check([r['style_chain'] for r in source['rows']]==[['Normal','Style1'],['Normal','Style1','Style2']],'Original style identities changed')
    check([r['indent'] for r in source['rows']]==[dict(zip(KEYS,(1134,1134,-567))),dict(zip(KEYS,(3969,2835,-2268)))],'Original effective geometry mismatch')
    check([r['marker']['visible'] for r in source['rows']]==['1.','2.'],'Original numbering sequence changed')
    check(all(r['marker']['fontFamily']=='Calibri' and r['marker']['fontSizeHalfPoints']==22 for r in source['rows']),'Original marker typeface/size mismatch')
    check(len(source['rows'])==2 and source['rows'][0]['text']==source['rows'][1]['text'] and len(source['rows'][0]['text'])==525 and hashlib.sha256(source['rows'][0]['text'].encode()).hexdigest()=='49c5f4adf3173685cea46572098c6be8a09b9a3c6d0b5cd6eb816dd1321d95a9','Original exact paragraph text/count mismatch')
    return source


def assert_original_positions(pdf):
    lines=pdf['lines'];markers=[r for r in lines if r['text'] in ('1.','2.')]
    check(len(markers)==2 and [r['text'] for r in markers]==['1.','2.'],'Original visible marker count/order changed')
    groups=[]
    for i,m in enumerate(markers):
        stop=markers[i+1]['yMin'] if i+1<len(markers) else float('inf')
        body=[r for r in lines if r not in markers and m['yMin']-.01<=r['yMin']<stop]
        groups.append(body)
    check([len(g) for g in groups]==[7,18],'Original wrapped body line counts changed')
    check(all(abs(m['xMin']-x)<=.11 for m,x in zip(markers,(113.50,170.20))),'Original marker physical x positions differ')
    check(all(abs(r['xMin']-x)<=.11 for rows,x in zip(groups,(141.85,283.60)) for r in rows),'Original body continuation positions differ')
    return {'markers_x_pt':[r['xMin'] for r in markers],'body_x_pt':[[r['xMin'] for r in g] for g in groups],'body_line_counts':[len(g) for g in groups],'tolerance_pt':.11}


def fixtures(directory, original):
    sources = {}
    def save(name, action, supported, render=False):
        path=directory/(name+'.docx');mutate(original,path,action)
        item={'path':str(path),'sha256':digest(path),'fixture_kind':'directed synthetic derivative of original; not an unchanged original',
              'expected_action':'roundtrip' if supported else 'reject','render':render}
        if supported:
            item['independent']=inspect_docx(path)
        sources[name]=item
    def style(t,sid):
        return t['styles'].find("w:style[@w:styleId='"+sid+"']",NS)
    def level(t):
        return t['numbering'].find("w:abstractNum/w:lvl[@w:ilvl='0']",NS)
    def ppr(t,index=0):
        return t['document'].findall('w:body/w:p',NS)[index].find('w:pPr',NS)
    def direct(t,parts):
        for index,values in enumerate(((1440,720,360),(3600,1440,1800))):
            child(ppr(t,index),'ind',**dict(zip(('left','right','hanging'),values)))
    save('complete-direct-override',direct,True,True)
    def zero_right(t,parts):
        for sid in ('Style1','Style2'):
            style(t,sid).find('w:pPr/w:ind',NS).set(W+'right','0')
    save('explicit-zero-right',zero_right,True,True)
    def explicit_level(t,parts):
        child(style(t,'Style1').find('w:pPr/w:numPr',NS),'ilvl',val=0)
    save('explicit-active-level-zero',explicit_level,True)
    save('start-seven',lambda t,p:level(t).find('w:start',NS).set(W+'val','7'),True,True)
    def tiny_hanging(t,parts,value=0):
        level(t).find('w:start',NS).set(W+'val','99')
        for sid in ('Style1','Style2'):style(t,sid).find('w:pPr/w:ind',NS).set(W+'hanging',str(value))
    save('zero-hanging-start-99',tiny_hanging,True,True)
    save('one-twip-hanging-start-99',lambda t,p:tiny_hanging(t,p,1),True,True)
    def detached_unused(t,parts):
        unused=copy.deepcopy(level(t));unused.set(W+'ilvl','8')
        # Unused template levels already include unsupported Roman/letter formats.
        # Add another inactive abstract definition to prove active identity scoping.
        abstract=child(t['numbering'],'abstractNum',abstractNumId=77)
        abstract.append(unused)
    save('unused-unsupported-definition',detached_unused,True)
    def relocated(t,parts):
        rels=E.fromstring(parts['word/_rels/document.xml.rels'])
        types=E.fromstring(parts['[Content_Types].xml'])
        for kind,target in [('styles','renamed/styles-main.xml'),('numbering','renamed/numbers-main.xml')]:
            for rel in rels:
                if rel.get('Type','').endswith('/'+kind):rel.set('Target',target)
            parts['word/'+target]=E.tostring(t[kind],xml_declaration=True,encoding='UTF-8',standalone=True)
            typ=next(n for n in types if n.get('PartName')=='/word/'+kind+'.xml')
            duplicate=copy.deepcopy(typ);duplicate.set('PartName','/word/'+target);types.append(duplicate)
        parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
        parts['[Content_Types].xml']=E.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
    save('actual-renamed-related-parts',relocated,True)

    for name,tag,value in [('active-roman','numFmt','upperRoman'),('active-letter','numFmt','lowerLetter'),
                           ('custom-pattern','lvlText','Chapter %1:'),('multilevel-pattern','lvlText','%1.%2.'),
                           ('right-marker','lvlJc','right'),('center-marker','lvlJc','center')]:
        def action(t,parts,tag=tag,value=value):level(t).find('w:'+tag,NS).set(W+'val',value)
        save('reject-'+name,action,False)
    for suffix in ('space','nothing'):
        save('reject-suffix-'+suffix,lambda t,p,suffix=suffix:child(level(t),'suff',val=suffix),False)
    save('reject-active-level-one',lambda t,p:child(style(t,'Style1').find('w:pPr/w:numPr',NS),'ilvl',val=1),False)
    save('reject-negative-left',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'left','-1'),False)
    save('reject-negative-right',lambda t,p:style(t,'Style2').find('w:pPr/w:ind',NS).set(W+'right','-1'),False)
    save('reject-hanging-beyond-left',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'hanging','1135'),False)
    save('reject-zero-body-width',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'left','8555'),False)
    save('reject-partial-effective-triple',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).attrib.pop(W+'right'),False)
    save('reject-dual-first-hanging',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'firstLine','0'),False)
    save('reject-logical-start',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'start','1134'),False)
    save('reject-character-left',lambda t,p:style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'leftChars','100'),False)
    save('reject-paragraph-tabs',lambda t,p:child(child(ppr(t),'tabs'),'tab',val='num',pos=1440),False)
    save('reject-level-tabs',lambda t,p:child(child(level(t).find('w:pPr',NS),'tabs'),'tab',val='num',pos=1440),False)
    save('reject-bidi',lambda t,p:child(ppr(t),'bidi',val=1),False)
    save('reject-mirror-indents',lambda t,p:child(ppr(t),'mirrorIndents',val=1),False)
    save('reject-marker-bold',lambda t,p:child(child(level(t),'rPr'),'b'),False)
    save('reject-level-restart',lambda t,p:child(level(t),'lvlRestart',val=0),False)
    save('reject-legal-numbering',lambda t,p:child(level(t),'isLgl'),False)
    save('reject-numbering-style-link',lambda t,p:child(t['numbering'].find('w:abstractNum',NS),'numStyleLink',val='NoList'),False)
    save('reject-duplicate-num',lambda t,p:t['numbering'].append(copy.deepcopy(t['numbering'].find('w:num',NS))),False)
    save('reject-duplicate-abstract',lambda t,p:t['numbering'].append(copy.deepcopy(t['numbering'].find('w:abstractNum',NS))),False)
    save('reject-duplicate-active-level',lambda t,p:t['numbering'].find('w:abstractNum',NS).append(copy.deepcopy(level(t))),False)
    def ambiguous(t,parts):
        n=t['numbering'].find("w:abstractNum/w:lvl[@w:ilvl='1']",NS);child(n,'pStyle',val='Style1')
    save('reject-ambiguous-pstyle-level',ambiguous,False)
    save('reject-style-cycle',lambda t,p:style(t,'Style1').find('w:basedOn',NS).set(W+'val','Style2'),False)
    def full_override(t,parts):
        child(t['numbering'].find('w:num',NS),'lvlOverride',ilvl=0).append(copy.deepcopy(level(t)))
    save('reject-full-level-override',full_override,False)
    save('reject-simple-start-override',lambda t,p:child(child(t['numbering'].find('w:num',NS),'lvlOverride',ilvl=0),'startOverride',val=3),False)
    save('reject-paragraph-mark-format',lambda t,p:child(child(ppr(t),'rPr'),'b'),False)
    for tabstop in ('360','1440'):
        def nondefault_tab(t,parts,tabstop=tabstop):
            node=t['settings'].find('w:defaultTabStop',NS)
            if node is None:child(t['settings'],'defaultTabStop',val=tabstop)
            else:node.set(W+'val',tabstop)
            tiny_hanging(t,parts,0)
        save('reject-default-tab-stop-'+tabstop,nondefault_tab,False)
    def noncontiguous(t,parts):
        body=t['document'].find('w:body',NS);para=E.Element(W+'p');child(child(para,'r'),'t').text='Interrupting body paragraph';body.insert(1,para)
    save('reject-noncontiguous-instance',noncontiguous,False)
    def table_list(t,parts):
        body=t['document'].find('w:body',NS);para=body.find('w:p',NS);body.remove(para)
        table=E.Element(W+'tbl');props=child(table,'tblPr');child(props,'tblW',w=8000,type='dxa');grid=child(table,'tblGrid');child(grid,'gridCol',w=8000)
        cell=child(child(table,'tr'),'tc');child(child(cell,'tcPr'),'tcW',w=8000,type='dxa');cell.append(para);body.insert(0,table)
    save('reject-table-list',table_list,False)
    def negative_canceled(t,parts):
        style(t,'Style1').find('w:pPr/w:ind',NS).set(W+'right','-1')
        for i in range(2):child(ppr(t,i),'ind',left=1440,right=0,hanging=360)
    save('reject-negative-provenance-canceled',negative_canceled,False)
    save('reject-missing-page',lambda t,p:t['document'].find('w:body',NS).remove(t['document'].find('w:body/w:sectPr',NS)),False)
    def empty_source(t,parts,kind):
        para=t['document'].find('w:body/w:p',NS)
        for n in list(para):
            if n.tag!=W+'pPr':para.remove(n)
        if kind!='no-runs':child(child(para,'r'),'t').text='   ' if kind=='spaces' else ''
    for kind in ('no-runs','empty-text','spaces'):
        save('reject-empty-source-'+kind,lambda t,p,kind=kind:empty_source(t,p,kind),False)
    def leading_break(t,parts):t['document'].find('w:body/w:p/w:r',NS).insert(0,E.Element(W+'br'))
    save('reject-source-hard-break',leading_break,False)
    def empty_formatted_run(t,parts):
        para=t['document'].find('w:body/w:p',NS);run=E.Element(W+'r');child(child(run,'rPr'),'sz',val=48);para.insert(1,run)
    save('reject-empty-formatted-run',empty_formatted_run,False)
    return sources


def model_rows(model):
    rows=[]
    def text(node):
        return '\n' if node['type']=='hardBreak' else node.get('text','')+''.join(text(n) for n in node.get('content',[]))
    def visit(node, ancestors=(), marker=None):
        if node['type']=='orderedList' and node.get('attrs',{}).get('markerFont') is not None:
            check(not any(a in ('table','orderedList','bulletList','listItem') for a in ancestors),'Preservation list is not top-level')
            attrs=node['attrs']; font=attrs['markerFont']; start=attrs.get('start',1)
            for index,item in enumerate(node.get('content',[])):
                check(item['type']=='listItem' and len(item.get('content',[]))==1 and item['content'][0]['type']=='paragraph','Preserved list item structure differs')
                visit(item,ancestors+('orderedList',),{'level':0,'ordinal':start+index,'start':start,'format':'decimal','text':'%1.',
                      'suffix':'tab','alignment':'left','fontFamily':font['fontFamily'],'fontSizeHalfPoints':font['fontSize'],'visible':str(start+index)+'.'})
            return
        if node['type'] in ('paragraph','heading'):
            attrs=node.get('attrs',{});geometry=attrs.get('listGeometry')
            if marker is not None:
                check(isinstance(geometry,dict) and set(geometry)=={'leftTwips','rightTwips','hangingTwips'},'Preserved list geometry incomplete')
                check(not any(k in attrs and attrs[k] is not None for k in KEYS),'Preserved list unexpectedly has body indentation metadata')
                indent=dict(zip(KEYS,(geometry['leftTwips'],geometry['rightTwips'],-geometry['hangingTwips'])))
            else:
                check(geometry is None,'Orphan preserved geometry')
                indent={k:attrs[k] for k in KEYS if k in attrs and attrs[k] is not None}
            rows.append({'text':text(node),'indent':indent,'marker':marker,'in_table':'table' in ancestors})
        for n in node.get('content',[]):visit(n,ancestors+(node['type'],),marker)
    visit(model)
    return rows


def create(cli,evidence,directory,name,model):
    model_path,output=directory/(name+'.json'),directory/(name+'.docx');save_json(model_path,model)
    run_cli(cli,evidence,'create-docx',model_path,output)
    reopened=run_cli(cli,evidence,'read-docx-model',output)
    check(reopened==model,name+': entire trusted model differs')
    before=digest(output);err=run_cli(cli,evidence,'create-docx',model_path,output,success=False)
    check('Output already exists' in err and digest(output)==before,name+': existing output clobbered')
    return {'path':str(output),'sha256':before,'model_path':str(model_path),'complete_trusted_model_equal':True,'no_clobber':True,'independent':inspect_docx(output)}


def roundtrip(cli,evidence,directory,name,item):
    path=Path(item['path']);item['audit']=run_cli(cli,evidence,'audit-docx',path)
    imported=run_cli(cli,evidence,'import-docx',path);item['import_response']=imported;model=imported['document']
    check(model.get('attrs',{}).get('pageGeometry')==item['independent']['pageGeometry'],name+': explicit source page proof changed')
    compare_rows(model_rows(model),item['independent']['rows'],name+' source/import')
    output=item['export']=create(cli,evidence,directory,name+'-export',model)
    compare_rows(output['independent']['rows'],item['independent']['rows'],name+' independent source/export')
    external_error=run_cli(cli,evidence,'import-docx',output['path'],success=False)
    check('Unsupported' in external_error,name+': external source reimport must refuse unsupported canonical marker declarations')
    item['external_reimport']={'accepted':False,'error':external_error,'scope':'External-source proof excludes explicit marker run properties and numbering overrides; trusted read-docx-model is the exact supported own-export reopen path'}
    before=digest(path);error=run_cli(cli,evidence,'create-docx',output['model_path'],path,success=False)
    check('Output already exists' in error and digest(path)==before,name+': source target clobbered')
    item['source_no_clobber']=True
    return model


def invalid_models(cli,evidence,directory,good):
    cases={}
    def add(name,action):
        m=copy.deepcopy(good);action(m);cases[name]=m
    def ol(m):return m['content'][0]
    def para(m,index=0):return ol(m)['content'][index]['content'][0]
    def geom(m):return para(m)['attrs']['listGeometry']
    for key in ('leftTwips','rightTwips','hangingTwips'):
        for value in (None,True,'12',1.5,-1,31681):
            add('geometry-'+key+'-'+str(value),lambda m,k=key,v=value:geom(m).__setitem__(k,v))
        add('geometry-missing-'+key,lambda m,k=key:geom(m).pop(k))
    add('geometry-extra',lambda m:geom(m).__setitem__('firstLineTwips',0))
    add('geometry-null',lambda m:para(m)['attrs'].__setitem__('listGeometry',None))
    add('geometry-nonobject',lambda m:para(m)['attrs'].__setitem__('listGeometry',[]))
    add('geometry-hanging-outside',lambda m:geom(m).__setitem__('hangingTwips',geom(m)['leftTwips']+1))
    add('geometry-zero-width',lambda m:geom(m).update(leftTwips=8555,rightTwips=1134))
    add('geometry-partial-list',lambda m:para(m,1)['attrs'].pop('listGeometry'))
    add('body-indent-conflict',lambda m:para(m)['attrs'].__setitem__('indentLeftTwips',0))
    add('missing-marker-font',lambda m:ol(m)['attrs'].pop('markerFont'))
    for value in (None,True,[],{},'Calibri'):
        add('marker-font-object-'+str(value),lambda m,v=value:ol(m)['attrs'].__setitem__('markerFont',v))
    for value in (None,True,'22',1.5,1,3277):
        add('marker-font-size-'+str(value),lambda m,v=value:ol(m)['attrs']['markerFont'].__setitem__('fontSize',v))
    for value in (None,True,22,'',' Calibri ','A\nB'):
        add('marker-font-family-'+str(value).replace('\n','newline'),lambda m,v=value:ol(m)['attrs']['markerFont'].__setitem__('fontFamily',v))
    add('marker-font-extra',lambda m:ol(m)['attrs']['markerFont'].__setitem__('bold',True))
    add('custom-marker-type',lambda m:ol(m)['attrs'].__setitem__('type','A'))
    add('missing-page',lambda m:m['attrs'].pop('pageGeometry'))
    add('page-too-narrow',lambda m:m['attrs']['pageGeometry'].__setitem__('widthTwips',5000))
    add('lift-preserved-paragraph',lambda m:m.__setitem__('content',[para(m)]))
    add('bullet-container',lambda m:ol(m).__setitem__('type','bulletList'))
    add('second-paragraph',lambda m:ol(m)['content'][0]['content'].append({'type':'paragraph'}))
    add('list-heading',lambda m:para(m).__setitem__('type','heading'))
    add('list-hard-page-break',lambda m:para(m)['content'].append({'type':'pageBreak'}))
    add('list-carriage-return',lambda m:para(m)['content'][0].__setitem__('text','Before\rAfter'))
    add('nested-preserved-list',lambda m:ol(m)['content'][0]['content'].append(copy.deepcopy(ol(m))))
    results={}
    for name,model in cases.items():
        path,output=directory/('invalid-'+name+'.json'),directory/('invalid-'+name+'.docx');save_json(path,model)
        error=run_cli(cli,evidence,'create-docx',path,output,success=False)
        check(not output.exists(),name+': invalid model produced partial output')
        sentinel=directory/('invalid-'+name+'-existing.docx');sentinel.write_bytes(b'unchanged target sentinel');before=digest(sentinel)
        run_cli(cli,evidence,'create-docx',path,sentinel,success=False);check(digest(sentinel)==before,name+': invalid model changed existing target')
        results[name]={'error':error,'no_partial_output':True,'existing_destination_unchanged':True}
    evidence['invalid_models']=results


def tamper(cli,evidence,directory,exported):
    with zipfile.ZipFile(exported) as package:original={n:package.read(n) for n in package.namelist()}
    member='rust-wps/model-v2.json';check(member in original,'Expected V2 model envelope')
    results={}
    for mode in ('model-geometry','model-marker-font','xml-geometry','xml-pattern','xml-marker-font','xml-numbering-start','xml-numbering-id'):
        for forged in (False,True) if mode.startswith('xml-') else (False,):
            name=mode+('-forged-hashes' if forged else '');parts=dict(original)
            if mode.startswith('model-'):
                envelope=json.loads(parts[member]);model=envelope.get('document',envelope)
                if mode=='model-geometry':model['content'][0]['content'][0]['content'][0]['attrs']['listGeometry']['leftTwips']+=1
                else:model['content'][0]['attrs']['markerFont']['fontSize']+=1
                parts[member]=json.dumps(envelope,ensure_ascii=False,separators=(',',':')).encode()
            else:
                main=E.fromstring(parts['word/document.xml']);numbering=E.fromstring(parts['word/numbering.xml']);para=main.find('w:body/w:p',NS)
                num_id=val(para,'w:pPr/w:numPr/w:numId');instance=numbering.find("w:num[@w:numId='"+num_id+"']",NS)
                abstract_id=val(instance,'w:abstractNumId');level=numbering.find("w:abstractNum[@w:abstractNumId='"+abstract_id+"']/w:lvl[@w:ilvl='0']",NS)
                if mode=='xml-geometry':para.find('w:pPr/w:ind',NS).set(W+'left','1135')
                elif mode=='xml-pattern':level.find('w:lvlText',NS).set(W+'val','(%1)')
                elif mode=='xml-marker-font':level.find('w:rPr/w:sz',NS).set(W+'val','24')
                elif mode=='xml-numbering-start':level.find('w:start',NS).set(W+'val','2')
                else:para.find('w:pPr/w:numPr/w:numId',NS).set(W+'val','9999')
                parts['word/document.xml']=E.tostring(main,xml_declaration=True,encoding='UTF-8',standalone=True)
                parts['word/numbering.xml']=E.tostring(numbering,xml_declaration=True,encoding='UTF-8',standalone=True)
            if forged:
                envelope=json.loads(parts[member])
                for part in envelope['hashes']:envelope['hashes'][part]=hashlib.sha256(parts[part]).hexdigest()
                parts[member]=json.dumps(envelope,ensure_ascii=False,separators=(',',':')).encode()
            path=directory/('tampered-'+name+'.docx')
            with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as package:
                for part,data in parts.items():package.writestr(part,data)
            results[name]={'path':str(path),'sha256':digest(path),'error':run_cli(cli,evidence,'read-docx-model',path,success=False)}
    evidence['tampered_packages']=results


def reference_from_model(original,model,path):
    """Independently authored python-docx/OOXML reference, never an app export."""
    doc=Document(original)
    for paragraph in list(doc.element.body.findall(W+'p')):doc.element.body.remove(paragraph)
    page=model['attrs']['pageGeometry'];section=doc.sections[0]
    section.page_width,section.page_height=Twips(page['widthTwips']),Twips(page['heightTwips'])
    for side in ('top','right','bottom','left'):setattr(section,side+'_margin',Twips(page['margin'+side.title()+'Twips']))
    numbering=doc.part.numbering_part.element
    # Rebuild the active reference numbering explicitly; unused source template
    # definitions are removed from this derivative, never from the original.
    for item in list(numbering):numbering.remove(item)
    for block_index,block in enumerate(model.get('content',[])):
        check(block['type']=='orderedList','Frontend reference currently requires direct ordered lists')
        a=block['attrs'];font=a['markerFont'];abstract=child(numbering,'abstractNum',abstractNumId=block_index)
        child(abstract,'multiLevelType',val='singleLevel');level=child(abstract,'lvl',ilvl=0)
        child(level,'start',val=a.get('start',1));child(level,'numFmt',val='decimal');child(level,'lvlText',val='%1.');child(level,'lvlJc',val='left');child(level,'suff',val='tab')
        level_rpr=child(level,'rPr');child(level_rpr,'rFonts',ascii=font['fontFamily'],hAnsi=font['fontFamily']);child(level_rpr,'sz',val=font['fontSize'])
        instance=child(numbering,'num',numId=block_index+1);child(instance,'abstractNumId',val=block_index)
        for item in block['content']:
            check(item['type']=='listItem' and len(item['content'])==1 and item['content'][0]['type']=='paragraph','Unsupported reference item')
            node=item['content'][0];attrs=node['attrs'];g=attrs['listGeometry'];p=doc.add_paragraph();ppr=p._p.get_or_add_pPr()
            np=child(ppr,'numPr');child(np,'ilvl',val=0);child(np,'numId',val=block_index+1)
            p.paragraph_format.left_indent=Twips(g['leftTwips']);p.paragraph_format.right_indent=Twips(g['rightTwips']);p.paragraph_format.first_line_indent=Twips(-g['hangingTwips'])
            p.paragraph_format.space_before=Twips(attrs.get('spaceBefore',0));p.paragraph_format.space_after=Twips(attrs.get('spaceAfter',0));p.paragraph_format.line_spacing=attrs.get('lineSpacing',240)/240
            if attrs.get('textAlign') is not None:child(ppr,'jc',val=attrs['textAlign'])
            if not node.get('content'):
                # Empty source-derived paragraphs must use the MODEL default,
                # not the unrelated original DOCX's Calibri11 docDefaults.
                empty=attrs.get('emptyParagraphFont') or {'family':'Arial','halfPoints':22}
                rp=child(ppr,'rPr');child(rp,'rFonts',ascii=empty['family'],hAnsi=empty['family']);child(rp,'sz',val=empty['halfPoints'])
            for span in node.get('content',[]):
                if span['type']=='hardBreak':p.add_run().add_break();continue
                check(span['type']=='text','Independent reference inline must be text/hardBreak')
                run=p.add_run(span['text']);marks={m['type']:m.get('attrs',{}) for m in span.get('marks',[])}
                check(set(marks)<={'fontFamily','fontSize','bold','italic','underline','strike','color','highlight'},'Independent reference unsupported marks')
                run.font.name=marks.get('fontFamily',{}).get('name','Arial');run.font.size=Pt(marks.get('fontSize',{}).get('halfPoints',22)/2)
                for mark,property_name in [('bold','bold'),('italic','italic'),('underline','underline'),('strike','strike')]:
                    if mark in marks:setattr(run.font,property_name,True)
                if 'color' in marks:
                    from docx.shared import RGBColor
                    color=marks['color'].get('color',marks['color'].get('value','')).lstrip('#');run.font.color.rgb=RGBColor.from_string(color)
                if 'highlight' in marks:child(run._r.get_or_add_rPr(),'highlight',val=marks['highlight'].get('color'))
    check(not path.exists(),'Reference must be fresh');doc.save(path)
    return inspect_docx(path)


def frontend_snapshots(cli,evidence,directory,input_directory,original):
    manifest_path=input_directory/'manifest.json';manifest=json.loads(manifest_path.read_text());results={}
    for filename in manifest['files']:
        source=input_directory/filename;wrapped=json.loads(source.read_text());model=wrapped.get('document',wrapped)
        if 'version' in wrapped:check(wrapped['version']==2,'Unexpected frontend model version')
        label='frontend-'+source.stem;reference=directory/(label+'-reference.docx')
        independent=reference_from_model(original,model,reference)
        compare_rows(model_rows(model),independent['rows'],label+' independent model/reference')
        output=create(cli,evidence,directory,label+'-output',model)
        compare_rows(output['independent']['rows'],independent['rows'],label+' independent output/reference')
        results[source.stem]={'model_source_path':str(source),'model_source_sha256':digest(source),'reference_path':str(reference),
                              'reference_sha256':digest(reference),'independent_reference':independent,'export':output,
                              'reference_method':'Fresh python-docx paragraphs and real OOXML decimal numbering independently constructed from actual frontend snapshot in a separate derivative of original styles/theme/settings'}
    for restored,reference in (('undo.json','initial.json'),('redo.json','edited.json')):
        if (input_directory/restored).exists() and (input_directory/reference).exists():
            check(json.loads((input_directory/restored).read_text())==json.loads((input_directory/reference).read_text()),'Frontend exact '+restored+' identity differs')
    for filename in manifest['files']:
        if filename.endswith('-undo.json'):
            check(json.loads((input_directory/filename).read_text())==json.loads((input_directory/'initial.json').read_text()),'Frontend exact undo differs: '+filename)
        elif filename.endswith('-redo.json'):
            base=filename.removesuffix('-redo.json')+'.json'
            check(json.loads((input_directory/filename).read_text())==json.loads((input_directory/base).read_text()),'Frontend exact redo differs: '+filename)
    evidence['frontend_manifest']={'path':str(manifest_path),'sha256':digest(manifest_path),'contents':manifest}
    evidence['frontend_snapshots']=results



def recheck_report(report_path):
    """Apply current strict comparison to hash-verified existing consumer artifacts.

    This never converts, overwrites, or replaces the original run report. It is
    useful when the comparator is strengthened after a finished consumer run.
    """
    prior=json.loads(report_path.read_text());check(prior['status']=='passed','Only a successful frozen run may be rechecked')
    check(digest(prior['cli'])==prior['cli_sha256'],'Frozen CLI changed since consumer run')
    pairs={}
    for name,item in prior['sources'].items():
        check(digest(item['path'])==item['sha256'],'Original/source control changed')
        if 'source_pdf' in item:
            check(digest(item['export']['path'])==item['export']['sha256'],'Exported source consumer input changed')
            pairs[name]=(item['source_pdf'],item['output_pdf'])
    for name,item in prior.get('frontend_snapshots',{}).items():
        check(digest(item['model_source_path'])==item['model_source_sha256'],'Actual editor snapshot changed')
        check(digest(item['reference_path'])==item['reference_sha256'],'Independent reference changed')
        check(digest(item['export']['path'])==item['export']['sha256'],'Exported frontend consumer input changed')
        pairs['frontend-'+name]=(item['reference_pdf'],item['output_pdf'])
    checked={};artifacts=[]
    for name,(left,right) in pairs.items():
        for render_data in (left,right):
            pdf=Path(render_data['path']);check(digest(pdf)==render_data['sha256'],'Consumer PDF changed')
            actual_bbox=subprocess.run(['pdftotext','-bbox-layout',str(pdf),'-'],capture_output=True,check=True,timeout=30).stdout
            tree=E.fromstring(actual_bbox);pages=[];lines=[]
            artifacts.append({'pdf':str(pdf),'pdf_sha256':digest(pdf),'bbox_sha256':digest(pdf.parent/'bbox.html'),
                              'rasters':[{"path":p,"sha256":digest(p)} for p in render_data['pngs']]})
            for pi,page in enumerate(tree.xpath('//*[local-name()="page"]')):
                pages.append({'width':float(page.get('width')),'height':float(page.get('height'))})
                for line in page.xpath('.//*[local-name()="line"]'):
                    words=line.xpath('./*[local-name()="word"]')
                    if words:lines.append({'page':pi+1,'text':' '.join(''.join(w.itertext()) for w in words),**{k:float(line.get(k)) for k in ('xMin','yMin','xMax','yMax')}})
            check(pages==render_data['pages'] and lines==render_data['lines'],'Stored coordinate evidence disagrees with actual bbox file')
        result=compare_render(left,right)
        check(not result['has_layout_drift'] and all(p.get('pixel_equal') for p in result['raster_comparison']),name+': strict zero-delta/exact-pixel comparison failed')
        checked[name]=result
    out={'status':'passed','original_report':str(report_path),'original_report_sha256':digest(report_path),
         'original_run_script_sha256':prior['script_sha256'],'strict_checker_sha256':digest(__file__),
         'cli_sha256':prior['cli_sha256'],'comparisons':checked,'artifacts':artifacts,'strict_coordinate_delta_pt':0.0,
         'method':'Re-extract all38 hash-verified final-run PDFs with pdftotext and compare exact coordinates/text plus actual raster pixels; source, exported DOCX, frontend snapshot and CLI hashes unchanged; no new LibreOffice conversion or overwritten original report',
         'checked_at_utc':datetime.now(timezone.utc).isoformat()}
    target=report_path.with_name('strict-coordinate-recheck-'+out['strict_checker_sha256'][:12]+'.json');check(not target.exists(),'Strict recheck must not clobber existing evidence');save_json(target,out);print(target)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--recheck-report',type=Path,help='Strictly recheck frozen existing PDF/bbox/raster evidence without new conversions')
    parser.add_argument('--cli',type=Path);parser.add_argument('--original',type=Path,default=ROOT/'tmp/paragraph-indentation-source/testParaListRightIndent.docx')
    parser.add_argument('--root',type=Path,default=ROOT/'tmp/list-geometry-independent');parser.add_argument('--generate-only',action='store_true')
    parser.add_argument('--render',action='store_true');parser.add_argument('--soffice');parser.add_argument('--editor-models',type=Path)
    parser.add_argument('--frontend-only',action='store_true',help='Only verify actual editor snapshots, with source identity preflight')
    args=parser.parse_args()
    if args.recheck_report:
        recheck_report(args.recheck_report.resolve());return
    check(args.generate_only or args.cli is not None,'Pass --cli or --generate-only')
    check(not args.frontend_only or args.editor_models is not None,'--frontend-only requires --editor-models')
    args.root.mkdir(parents=True,exist_ok=True);directory=Path(tempfile.mkdtemp(prefix='run-',dir=args.root)).resolve()
    cli=args.cli.resolve() if args.cli else None;original=args.original.resolve();soffice=args.soffice or shutil.which('soffice')
    evidence={'status':'running','script_sha256':digest(__file__),'cli':str(cli),'cli_sha256':digest(cli) if cli else None,
              'commands':[],'sources':{},'limitations':LIMITATIONS,'python_version':platform.python_version(),
              'started_at_utc':datetime.now(timezone.utc).isoformat(),'invocation':[sys.executable,*sys.argv],
              'original_provenance':{'path':str(original),'sha256':ORIGINAL_SHA,'git_blob_sha':ORIGINAL_BLOB,'revision':ORIGINAL_REVISION,'source_url':ORIGINAL_URL,
              'bytes':12774,'fixture_kind':'genuine unchanged public upstream regression fixture; not user production','licensing_caveat':'Local validation only; original binary not committed or redistributed'}}
    try:
        source=original_preflight(original);evidence['original_preflight']=source
        if not args.frontend_only:
            sources=evidence['sources']=fixtures(directory,original)
            sources['unchanged-original']={'path':str(original),'sha256':ORIGINAL_SHA,'fixture_kind':'genuine unchanged public upstream regression fixture',
                                          'expected_action':'roundtrip','render':True,'independent':source}
            if cli:
                # Original first so the primary result is available before controls.
                original_item=sources['unchanged-original'];good=roundtrip(cli,evidence,directory,'unchanged-original',original_item)
                save_json(directory/'original-imported-model.json',{'version':2,'document':good})
                for name,item in sources.items():
                    if name=='unchanged-original':continue
                    if item['expected_action']=='roundtrip':roundtrip(cli,evidence,directory,name,item)
                    else:
                        item['audit']=run_cli(cli,evidence,'audit-docx',item['path'])
                        item['import_error']=run_cli(cli,evidence,'import-docx',item['path'],success=False)
                        check('Unsupported' in item['import_error'] or 'Invalid' in item['import_error'],name+': unclear refusal')
                        before=digest(item['path']);err=run_cli(cli,evidence,'create-docx',original_item['export']['model_path'],item['path'],success=False)
                        check('Output already exists' in err and digest(item['path'])==before,name+': refused source changed');item['source_no_clobber']=True
                    check(digest(item['path'])==item['sha256'],name+': source mutated')
                invalid_models(cli,evidence,directory,good);tamper(cli,evidence,directory,original_item['export']['path'])
        if cli and args.editor_models:frontend_snapshots(cli,evidence,directory,args.editor_models.resolve(),original)
        if args.render:
            check(soffice is not None,'LibreOffice required');evidence['libreoffice_executable']={'path':str(Path(soffice).resolve()),'sha256':digest(soffice)}
            evidence['libreoffice_version']=subprocess.run([soffice,'-env:UserInstallation='+(directory/'lo-version-profile').as_uri(),'--version'],text=True,capture_output=True,check=True).stdout.strip()
            comparisons=evidence['render_comparisons']={}
            for name,item in evidence['sources'].items():
                if not item.get('render'):continue
                item['source_pdf']=render(soffice,directory,Path(item['path']),name+'-source')
                if name=='unchanged-original':item['source_positions']=assert_original_positions(item['source_pdf'])
                if cli:
                    item['output_pdf']=render(soffice,directory,Path(item['export']['path']),name+'-output');comparison=compare_render(item['source_pdf'],item['output_pdf']);comparisons[name]=comparison
                    if name=='unchanged-original':item['output_positions']=assert_original_positions(item['output_pdf'])
                    check(not comparison['has_layout_drift'],name+': independent source/output layout drift: '+json.dumps(comparison))
                    check(all(p.get('pixel_equal') for p in comparison['raster_comparison']),name+': independent source/output pixels differ')
            for name,item in evidence.get('frontend_snapshots',{}).items():
                item['reference_pdf']=render(soffice,directory,Path(item['reference_path']),'frontend-'+name+'-reference')
                item['output_pdf']=render(soffice,directory,Path(item['export']['path']),'frontend-'+name+'-output');comparison=compare_render(item['reference_pdf'],item['output_pdf']);comparisons['frontend-'+name]=comparison
                check(not comparison['has_layout_drift'],'Frontend '+name+': independent reference/output layout drift: '+json.dumps(comparison))
                check(all(p.get('pixel_equal') for p in comparison['raster_comparison']),'Frontend '+name+': reference/output pixels differ')
        evidence['summary']={'supported_sources':sum(i['expected_action']=='roundtrip' for i in evidence['sources'].values()),'rejected_sources':sum(i['expected_action']=='reject' for i in evidence['sources'].values()),
                             'invalid_models':len(evidence.get('invalid_models',{})),'tampered_packages':len(evidence.get('tampered_packages',{})),
                             'frontend_snapshot_files':len(evidence.get('frontend_snapshots',{})),'rendered_pdfs':len(list((directory/'renders').glob('*/*.pdf'))),
                             'render_comparison_pairs':len(evidence.get('render_comparisons',{})),
                             'pixel_identical_pairs':sum(all(p.get('pixel_equal') for p in c['raster_comparison']) for c in evidence.get('render_comparisons',{}).values()),
                             'successful_trusted_model_reopens':sum(c['arguments'][0]=='read-docx-model' and c['returncode']==0 for c in evidence['commands'])}
        evidence['status']='fixtures-verified' if args.generate_only else 'passed'
    except Exception as error:
        evidence['status']='failed';evidence['error']=type(error).__name__+': '+str(error);raise
    finally:
        evidence['finished_at_utc']=datetime.now(timezone.utc).isoformat();evidence['script_unchanged_during_run']=digest(__file__)==evidence['script_sha256']
        evidence['cli_unchanged_during_run']=cli is None or digest(cli)==evidence['cli_sha256'];evidence['source_hashes_unchanged']=all(digest(i['path'])==i['sha256'] for i in evidence['sources'].values())
        evidence['genuine_original_unchanged']=digest(original)==ORIGINAL_SHA
        evidence['frontend_inputs_unchanged']=all(digest(i['model_source_path'])==i['model_source_sha256'] for i in evidence.get('frontend_snapshots',{}).values())
        report=directory/'report.json';save_json(report,evidence);print(report)


if __name__=='__main__':main()
