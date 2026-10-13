"""Independent bounded empty-paragraph-font DOCX regression and provenance.

Fixed expectations are authored before calling the product. Source XML cascade,
python-docx typed properties, output XML, trusted reopen, rejection and consumer
PDF checks are distinct. Originals are hash-pinned, never rewritten. Every PDF
uses one new soffice PROCESS and UNIQUE profile (never a shared-profile batch).

  python scripts/check-empty-paragraph-fonts.py --generate-only --render
  python scripts/check-empty-paragraph-fonts.py --cli FROZEN_CLI --render

No browser, native printing, Word/WPS equivalence, font availability, or complete
document fidelity is claimed. Only controlled Latin baseline geometry is gated.
"""
import argparse
import copy
import hashlib
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docx.text.paragraph import Paragraph
from lxml import etree as ET

ROOT = Path(__file__).resolve().parent.parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
ATTR = 'emptyParagraphFont'
OLD_SHA = '0c67ef3110ae794d371123313fe01ccfa53bef452f2316dac8a5e11c999da67e'
REVISION = 'ae62bb5116b9aee19ebd5834e3a82066132c9f7f'
PRESERVED = 'Empty paragraph Latin baseline fonts preserved as fixed names'
UNPROVED = 'Some empty paragraph baseline fonts are unproved'
TOLERANCE = 0.30


def check(value, message):
    if not value:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def add(parent, tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def remove(parent, tag):
    for node in list(parent.findall('w:' + tag, NS)):
        parent.remove(node)


def font(family='DejaVu Sans', halfPoints=24, bold=False, italic=False):
    return dict(family=family, halfPoints=halfPoints, bold=bold, italic=italic)


def font_xml(parent, values):
    """Write only supplied fields, so partial-cascade tests are genuine."""
    if 'family' in values:
        remove(parent, 'rFonts')
        add(parent, 'rFonts', ascii=values['family'], hAnsi=values['family'])
    if 'halfPoints' in values:
        remove(parent, 'sz')
        add(parent, 'sz', val=values['halfPoints'])
    for key, tag in [('bold', 'b'), ('italic', 'i')]:
        if key in values:
            remove(parent, tag)
            add(parent, tag, val='1' if values[key] else '0')


def save(doc, path):
    check(not path.exists(), 'Refusing to overwrite independent source')
    doc.save(path)
    return path


def mutate_package(source, output, part, mutate):
    check(not output.exists(), 'Refusing to overwrite derivative')
    with zipfile.ZipFile(source) as old, zipfile.ZipFile(output, 'x', zipfile.ZIP_DEFLATED) as new:
        for info in old.infolist():
            data = old.read(info.filename)
            if info.filename == part:
                root = ET.fromstring(data)
                mutate(root)
                data = ET.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
            new.writestr(info, data)


def true_empty(p):
    return all(ET.QName(n).localname in ('pPr', 'bookmarkStart', 'bookmarkEnd', 'proofErr') or
               (ET.QName(n).localname == 'r' and len(n) == 0 and not n.attrib and not n.text) for n in p)


def paragraph_text(p):
    return ''.join(n.text or '' if n.tag == qn('w:t') else '\t' if n.tag == qn('w:tab') else '\n'
                   for n in p.iter() if n.tag in (qn('w:t'), qn('w:tab'), qn('w:br')))


def snapshot(path):
    """Independent limited Latin cascade, plus separately parsed typed fields.

    The fixed fixture manifest is the oracle. This reader corroborates it; it
    does not call rust-wps or learn expectations from its model/export.
    """
    with zipfile.ZipFile(path) as package:
        main = ET.fromstring(package.read('word/document.xml'))
        styles = ET.fromstring(package.read('word/styles.xml'))
        theme = ET.fromstring(package.read('word/theme/theme1.xml')) if 'word/theme/theme1.xml' in package.namelist() else None
    index = {s.get(qn('w:styleId')): s for s in styles.findall('w:style', NS)}
    default = next((sid for sid, s in index.items() if s.get(qn('w:type')) == 'paragraph' and s.get(qn('w:default')) in ('1', 'true', 'on')), None)

    def chain(sid, seen=()):
        check(sid not in seen, 'Independent reader found cyclic style chain')
        style = index.get(sid)
        if style is None:
            return []
        based = style.find('w:basedOn', NS)
        return ([] if based is None else chain(based.get(qn('w:val')), seen + (sid,))) + [style]

    def overlay(state, props, toggle):
        if props is None:
            return
        fonts = props.find('w:rFonts', NS)
        for slot in ('ascii', 'hAnsi'):
            if fonts is not None:
                themed, literal = fonts.get(qn('w:' + slot + 'Theme')), fonts.get(qn('w:' + slot))
                if themed is not None:
                    node = None if theme is None else theme.find('a:themeElements/a:fontScheme/a:' + ('majorFont' if themed.startswith('major') else 'minorFont') + '/a:latin', NS)
                    state[slot] = None if node is None else node.get('typeface')
                elif literal is not None:
                    state[slot] = literal
        size = props.find('w:sz', NS)
        if size is not None:
            state['halfPoints'] = int(size.get(qn('w:val')))
        for field, tag in [('bold', 'b'), ('italic', 'i')]:
            node = props.find('w:' + tag, NS)
            if node is not None:
                flag = node.get(qn('w:val'), '1') not in ('0', 'false', 'off')
                state[field] = (state[field] != flag) if toggle else flag

    doc = Document(path)
    python_paragraphs = list(doc.element.body.iter(qn('w:p')))
    raw_paragraphs = list(main.find('w:body', NS).iter(qn('w:p')))
    check(len(python_paragraphs) == len(raw_paragraphs), 'Independent paragraph count differs')
    rows = []
    body = main.find('w:body', NS)
    for p, python_p in zip(raw_paragraphs, python_paragraphs):
        ppr = p.find('w:pPr', NS)
        sid = p.find('w:pPr/w:pStyle', NS)
        sid = default if sid is None else sid.get(qn('w:val'))
        state = dict(bold=False, italic=False)
        overlay(state, styles.find('w:docDefaults/w:rPrDefault/w:rPr', NS), False)
        for style in chain(sid):
            overlay(state, style.find('w:rPr', NS), True)
        mark = p.find('w:pPr/w:rPr', NS)
        overlay(state, mark, False)
        resolved = None if state.get('ascii') != state.get('hAnsi') or not state.get('ascii') or 'halfPoints' not in state else font(state['ascii'], state['halfPoints'], state['bold'], state['italic'])
        raw_direct = None
        py_direct = None
        if mark is not None:
            fonts, sz, b, i = [mark.find('w:' + tag, NS) for tag in ('rFonts', 'sz', 'b', 'i')]
            raw_direct = {'ascii': None if fonts is None else fonts.get(qn('w:ascii')), 'hAnsi': None if fonts is None else fonts.get(qn('w:hAnsi')),
                          'halfPoints': None if sz is None else int(sz.get(qn('w:val'))),
                          'bold': None if b is None else b.get(qn('w:val'), '1') not in ('0', 'false', 'off'),
                          'italic': None if i is None else i.get(qn('w:val'), '1') not in ('0', 'false', 'off')}
            typed = python_p.find(qn('w:pPr')).find(qn('w:rPr'))
            py_direct = {'ascii': None if typed.rFonts is None else typed.rFonts.ascii, 'hAnsi': None if typed.rFonts is None else typed.rFonts.hAnsi,
                         'halfPoints': None if typed.sz is None else int(typed.sz.val.pt * 2),
                         'bold': None if typed.b is None else typed.b.val, 'italic': None if typed.i is None else typed.i.val}
            check(raw_direct == py_direct, 'python-docx typed paragraph-mark fields differ from raw XML')
        text = paragraph_text(p)
        python_text = Paragraph(python_p, doc).text
        # python-docx omits explicit page breaks from Paragraph.text.
        check(text.replace('\n', '') == python_text.replace('\n', ''), 'Independent paragraph text differs')
        rows.append({'body_index': list(body).index(p) if p.getparent() is body else None,
                     'text': text, 'true_empty': true_empty(p), 'run_count': len(p.findall('w:r', NS)),
                     'style_id': sid, 'font': resolved, 'direct': raw_direct, 'python_docx_direct': py_direct,
                     'page_breaks': len(p.findall('.//w:br[@w:type="page"]', NS))})
    return {'paragraphs': rows, 'text_node_count': len(main.findall('.//w:t', NS)),
            'body_blocks': [ET.QName(n).localname for n in body if ET.QName(n).localname != 'sectPr']}


def fresh(defaults=None, normal=None, layers=(), mark=None):
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Twips(12240), Twips(15840)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Twips(1440)
        remove(section._sectPr, 'docGrid')
    for node in list(doc.styles.element.findall('.//w:contextualSpacing', NS)):
        node.getparent().remove(node)
    defaults_node = doc.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr', NS)
    for child in list(defaults_node):
        defaults_node.remove(child)
    font_xml(defaults_node, defaults if defaults is not None else font())
    base = doc.styles['Normal']
    remove(base.element, 'rPr')
    if normal:
        font_xml(add(base.element, 'rPr'), normal)
    base.paragraph_format.space_before = base.paragraph_format.space_after = Pt(0)
    base.paragraph_format.line_spacing = 1
    for j, values in enumerate(layers):
        style = doc.styles.add_style('EmptyProbe' + str(j), WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = base
        font_xml(add(style.element, 'rPr'), values)
        base = style
    for text in ('BeforeBlank', '', 'AfterBlank'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        if text:
            run = p.add_run(text)
            run.font.name, run.font.size = 'DejaVu Sans', Pt(12)
            run.bold = run.italic = False
        else:
            p.style = base
            if mark is not None:
                font_xml(add(p._p.get_or_add_pPr(), 'rPr'), mark)
    return doc


def sources(directory):
    plans = {
        'defaults-12': ({}, font()),
        'style-18': ({'normal': font('DejaVu Serif', 36)}, font('DejaVu Serif', 36)),
        'mark-overrides-style': ({'normal': font('DejaVu Serif', 36, True, True), 'mark': font('DejaVu Sans', 28)}, font('DejaVu Sans', 28)),
        'partial-chain': ({'defaults': font('DejaVu Serif', 22), 'normal': {'halfPoints': 26}, 'layers': [{'family': 'DejaVu Sans', 'bold': True}, {'italic': True}], 'mark': {'halfPoints': 34}}, font('DejaVu Sans', 34, True, True)),
        'style-true-toggles': ({'normal': {'bold': True, 'italic': True}, 'layers': [{'bold': True, 'italic': True}]}, font()),
        'style-false-noop': ({'normal': {'bold': True, 'italic': True}, 'layers': [{'bold': False, 'italic': False}]}, font(bold=True, italic=True)),
        'direct-false-reset': ({'normal': {'bold': True, 'italic': True}, 'mark': {'bold': False, 'italic': False}}, font()),
        'partial-mark-family': ({'normal': font('DejaVu Serif', 30, True, False), 'mark': {'family': 'DejaVu Sans'}}, font('DejaVu Sans', 30, True, False)),
        'half-point-size': ({'mark': {'halfPoints': 25}}, font(halfPoints=25)),
        'direct-bold-italic': ({'mark': {'bold': True, 'italic': True, 'halfPoints': 40}}, font(halfPoints=40, bold=True, italic=True)),
    }
    items = {}
    for name, (kwargs, expected) in plans.items():
        path = save(fresh(**kwargs), directory / (name + '-source.docx'))
        # This reference is independently authored with explicit direct marks,
        # not made by reading/re-encoding the product model.
        reference = save(fresh(mark=expected), directory / (name + '-explicit-reference.docx'))
        items[name] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned source-generated directed fixture',
                       'expected': {'1': expected}, 'render': True,
                       'reference': {'path': str(reference), 'sha256': digest(reference)}}
    for name, expected in [('source-minimum-size', font(halfPoints=2)), ('source-maximum-size', font(halfPoints=3276)),
                           ('source-31-utf16-name', font('a' * 31))]:
        path = save(fresh(mark=expected), directory / (name + '.docx'))
        items[name] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned valid boundary fixture, structural only',
                       'expected': {'1': expected}, 'render': False}
    themed = fresh()
    fonts = themed.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr/w:rFonts', NS)
    fonts.set(qn('w:asciiTheme'), 'minorHAnsi'); fonts.set(qn('w:hAnsiTheme'), 'minorHAnsi')
    staging = save(themed, directory / 'theme-precedence-staging.docx')
    path = directory / 'theme-precedence-source.docx'
    mutate_package(staging, path, 'word/theme/theme1.xml', lambda root: root.find('a:themeElements/a:fontScheme/a:minorFont/a:latin', NS).set('typeface', 'DejaVu Serif'))
    reference = save(fresh(mark=font('DejaVu Serif')), directory / 'theme-precedence-explicit-reference.docx')
    items['theme-precedence'] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned same-level theme-over-literal Latin control',
                               'expected': {'1': font('DejaVu Serif')}, 'render': True,
                               'reference': {'path': str(reference), 'sha256': digest(reference)}}
    bare = fresh()
    bare.paragraphs[1].add_run('')
    path = save(bare, directory / 'bare-empty-run-source.docx')
    items['bare-empty-run'] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned bare empty-run geometry-neutral source',
                               'expected': {'1': font()}, 'render': True}
    identity = fresh(mark=font('DejaVu Serif', 30))
    table = identity.add_table(rows=2, cols=2)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.text = f'Identity {ri},{ci}'
    path = save(identity, directory / 'body-table-identity-source.docx')
    items['body-table-identity'] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned distinct cell-text identity control',
                                  'expected': {'1': font('DejaVu Serif', 30)}, 'render': False}
    mixed = fresh(defaults={'halfPoints': 24})
    known = mixed.styles.add_style('KnownBlank', WD_STYLE_TYPE.PARAGRAPH)
    known.base_style = mixed.styles['Normal']
    font_xml(add(known.element, 'rPr'), font('DejaVu Serif', 30))
    mixed.paragraphs[1].style = known
    mixed.add_paragraph()
    path = save(mixed, directory / 'mixed-known-and-missing-font-source.docx')
    items['mixed-known-and-missing-font'] = {'path': str(path), 'sha256': digest(path), 'kind': 'owned independent per-blank font proof control',
                                           'expected': {'1': font('DejaVu Serif', 30)}, 'expected_unproved': [3], 'render': False}
    originals = [
        ('known-contextual-control', 'tmp/word-contextual-spacing/run-_a6int64/empty-neighbor-explicit-reference.docx', '34f60b81f669e29932c8ec04223d5f71bf828bc7d491e26429ef8dfb4a176ff5', {'1': font()}, None),
        ('original-table-alignment', 'tmp/table-corpus/table-alignment.docx', 'ee36729e0d050e529fea0b444006a194bea72b35cf8958ae58c6624c26156b41', {str(i): font('Aptos', 24) for i in (1, 3, 5, 7, 9, 11)}, 'table-alignment.docx'),
        ('original-NoHeadFoot', 'tmp/paragraph-indentation-source/poi/NoHeadFoot.docx', 'dc6a67b1c118658acadc87ba41200f0b90ead05ada95e7ce6fa72a97700430c5', {str(i): font('Calibri', 22) for i in (1, 3, 5, 6)}, 'NoHeadFoot.docx'),
    ]
    for name, relative, sha, expected, filename in originals:
        path = ROOT / relative
        check(path.is_file() and digest(path) == sha, name + ': original missing or changed')
        data = path.read_bytes()
        items[name] = {'path': str(path), 'sha256': sha, 'expected': expected, 'render': True,
                       'git_blob_sha1': hashlib.sha1(f'blob {len(data)}\0'.encode() + data).hexdigest(),
                       'kind': 'unchanged upstream Apache POI regression fixture, not a user production file' if filename else 'unchanged prior independently authored contextual-spacing control',
                       'source_url': f'https://github.com/apache/poi/blob/{REVISION}/test-data/document/{filename}' if filename else None}
    original = Path(items['known-contextual-control']['path'])
    counterpart = directory / 'contextual-no-run-counterpart.docx'
    def remove_empty_run(root):
        p = root.find('w:body', NS)[1]
        runs = p.findall('w:r', NS)
        check(len(runs) == 1 and len(runs[0]) == 0, 'Known control no longer has exactly one empty run')
        p.remove(runs[0])
    mutate_package(original, counterpart, 'word/document.xml', remove_empty_run)
    items['contextual-no-run-counterpart'] = {'path': str(counterpart), 'sha256': digest(counterpart), 'expected': {'1': font()}, 'render': True,
                                            'kind': 'explicit derivative of known control with only its empty run removed; not the unchanged original',
                                            'derived_from_sha256': digest(original)}
    for name, item in items.items():
        item['independent_source'] = snapshot(Path(item['path']))
        actual = {str(p['body_index']): p['font'] for p in item['independent_source']['paragraphs'] if p['body_index'] is not None and p['true_empty'] and p['font'] is not None}
        check(actual == item['expected'], name + ': fixed expectation differs from independent source cascade: ' + repr(actual))
        unproved = [p['body_index'] for p in item['independent_source']['paragraphs'] if p['body_index'] is not None and p['true_empty'] and p['font'] is None]
        check(unproved == item.get('expected_unproved', []), name + ': independent unresolved identities differ')
        if 'reference' in item:
            ref = snapshot(Path(item['reference']['path']))
            check(ref['paragraphs'][1]['font'] == item['expected']['1'], name + ': authored explicit reference differs')
            item['reference']['independent'] = ref
    return items


def warning_sources(directory):
    cases = {}
    def write(name, mutate, expected_warning=True):
        doc = fresh()
        mutate(doc)
        path = save(doc, directory / (name + '-source.docx'))
        cases[name] = {'path': str(path), 'sha256': digest(path), 'warning': expected_warning,
                       'kind': 'owned boundary fixture; unsupported proof normalizes with existing import policy unchanged'}
    def decorated_run(d):
        run = d.paragraphs[1].add_run('')
        run.font.name, run.font.size = 'DejaVu Serif', Pt(18)
    write('font-bearing-empty-run-not-target', decorated_run, False)
    write('empty-text-element-not-target', lambda d: add(d.paragraphs[1].add_run('')._r, 't'), False)
    write('attributed-empty-run-not-target', lambda d: d.paragraphs[1].add_run('')._r.set(qn('w:rsidR'), '00112233'), False)
    write('whitespace-is-not-empty', lambda d: d.paragraphs[1].add_run('  '), False)
    write('tab-is-not-empty', lambda d: d.paragraphs[1].add_run('\t'), False)
    write('line-break-is-not-empty', lambda d: d.paragraphs[1].add_run().add_break(), False)
    from docx.enum.text import WD_BREAK
    write('page-break-is-not-empty', lambda d: d.paragraphs[1].add_run().add_break(WD_BREAK.PAGE), False)
    def inherited_mark(d):
        font_xml(add(d.styles['Normal'].element.get_or_add_pPr(), 'rPr'), font(halfPoints=38))
    write('inherited-paragraph-mark', inherited_mark)
    def default_mark(d):
        font_xml(add(d.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr', NS), 'rPr'), font(halfPoints=38))
    write('default-paragraph-mark', default_mark)
    def character_mark(d):
        style = d.styles.add_style('EmptyMarkCharacter', WD_STYLE_TYPE.CHARACTER)
        style.font.size = Pt(18)
        add(add(d.paragraphs[1]._p.get_or_add_pPr(), 'rPr'), 'rStyle', val=style.style_id)
    write('character-style-paragraph-mark', character_mark)
    def slots(d):
        add(add(d.paragraphs[1]._p.get_or_add_pPr(), 'rPr'), 'rFonts', ascii='DejaVu Sans', hAnsi='DejaVu Serif')
    write('unresolved-latin-slot-disagreement', slots)
    def missing_font(d):
        remove(d.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr', NS), 'rFonts')
    write('missing-latin-font', missing_font)
    def theme_slot(d):
        add(add(d.paragraphs[1]._p.get_or_add_pPr(), 'rPr'), 'rFonts', asciiTheme='minorEastAsia', hAnsiTheme='minorEastAsia')
    write('unsupported-theme-slot', theme_slot)
    def list_identity(d):
        d.paragraphs[1].style = d.styles['List Bullet']
    write('list-inheritance-normalization', list_identity)
    def table_only(d):
        d.paragraphs[1].add_run('Middle')
        d.add_table(rows=1, cols=1)
    write('table-empty-not-body-target', table_only)
    def unproved_identity(d):
        d.paragraphs[0].style = d.styles['List Bullet']
    write('list-neighbor-breaks-body-proof', unproved_identity)
    return cases


def blocked_sources(directory):
    cases = {}
    for name, needle in [('hidden-text', 'Hidden text'), ('tracked-insertion', 'Tracked changes')]:
        doc = fresh()
        run = doc.paragraphs[0].runs[0]
        if name == 'hidden-text':
            add(run._r.get_or_add_rPr(), 'vanish')
        else:
            p = doc.paragraphs[0]._p
            p.remove(run._r)
            wrapper = add(p, 'ins', id='1', author='Independent fixture')
            wrapper.append(run._r)
        path = save(doc, directory / (name + '-blocked-source.docx'))
        cases[name] = {'path': str(path), 'sha256': digest(path), 'needle': needle,
                       'kind': 'owned existing-import-blocker fixture, not a new empty-font blocker'}
    return cases


def run_cli(cli, commands, *args, success=True):
    command = [str(cli), *map(str, args)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=60)
    commands.append({'command': command, 'returncode': result.returncode, 'stderr': result.stderr})
    check((result.returncode == 0) == success, 'Unexpected CLI result: ' + ' '.join(command) + '\n' + result.stderr)
    return json.loads(result.stdout) if success else result.stderr


def walk(node):
    yield node
    for child in node.get('content', []):
        yield from walk(child)


def strip_font(model):
    result = copy.deepcopy(model)
    for node in walk(result):
        node.get('attrs', {}).pop(ATTR, None)
        if node.get('attrs') == {}:
            del node['attrs']
    return result


def model_expected(model):
    return [n.get('attrs', {}).get(ATTR) for n in walk(model) if n['type'] in ('paragraph', 'heading')]


def verify_output(path, model):
    actual = snapshot(path)
    expected = model_expected(model)
    # Page-break model blocks serialize to a separate OOXML paragraph.
    paragraphs = [p for p in actual['paragraphs'] if not p['page_breaks']]
    check(len(paragraphs) == len(expected), 'Export paragraph identity count differs')
    for row, want in zip(paragraphs, expected):
        if want is not None:
            check(row['true_empty'] and row['run_count'] == 0 and row['text'] == '', 'Export fabricated content for an empty font')
            check(row['font'] == want, 'Export font differs from model')
            check(row['direct'] == dict(ascii=want['family'], hAnsi=want['family'], halfPoints=want['halfPoints'], bold=want['bold'], italic=want['italic']), 'Export requires complete direct pPr/rPr fields')
        else:
            check(row['direct'] is None, 'Export fabricated paragraph-mark properties without metadata')
    return actual


def export_model(cli, commands, directory, name, model, source=None):
    model_path, output = directory / (name + '.json'), directory / (name + '.docx')
    with model_path.open('x') as stream:
        json.dump(model, stream, ensure_ascii=False, indent=2)
    run_cli(cli, commands, 'create-docx', model_path, output)
    check(run_cli(cli, commands, 'read-docx-model', output) == model, name + ': full trusted model differs')
    independent = verify_output(output, model)
    targets = [output] + ([] if source is None else [source])
    before = {str(p): digest(p) for p in targets}
    for target in targets:
        error = run_cli(cli, commands, 'create-docx', model_path, target, success=False)
        check('already exists' in error, name + ': no-clobber error differs')
    check(all(digest(p) == value for p, value in before.items()), 'No-clobber modified source/output')
    return {'path': str(output), 'sha256': digest(output), 'model_path': str(model_path),
            'full_trusted_model_equal': True, 'independent': independent, 'no_clobber': True}


def rejection_checks(cli, commands, directory, evidence):
    good = {'type': 'doc', 'content': [{'type': 'paragraph', 'attrs': {ATTR: font()}}]}
    invalid = {}
    values = [None, False, 1, [], '', {}, dict(font(), unknown=1)]
    values += [{k: v for k, v in font().items() if k != missing} for missing in font()]
    for field, bad_values in [('family', ['', ' ', ' DejaVu Sans', 'DejaVu Sans ', '\ufeffDejaVu Sans', 'a\nb', 'a' * 32, '😀' * 16, '\uffff', 1, None]),
                              ('halfPoints', [0, 1, 3277, -1, 24.0, 24.5, True, '24', None]),
                              ('bold', [0, 1, 'true', None]), ('italic', [0, 1, 'false', None])]:
        values.extend(dict(font(), **{field: value}) for value in bad_values)
    for i, value in enumerate(values):
        model = copy.deepcopy(good)
        model['content'][0]['attrs'][ATTR] = value
        invalid['value-' + str(i)] = model
    for name, node in [('text', {'type': 'text', 'text': 'x'}), ('empty-text-node', {'type': 'text', 'text': ''}),
                       ('whitespace-text', {'type': 'text', 'text': ' '}), ('tab-text', {'type': 'text', 'text': '\t'}),
                       ('hard-break', {'type': 'hardBreak'})]:
        model = copy.deepcopy(good)
        model['content'][0]['content'] = [node]
        invalid['nonempty-' + name] = model
    heading = copy.deepcopy(good)
    heading['content'][0].update(type='heading', content=[{'type': 'text', 'text': 'x'}])
    heading['content'][0]['attrs']['level'] = 2
    invalid['nonempty-heading'] = heading
    paragraph = {'type': 'paragraph'}
    table = {'type': 'table', 'content': [{'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [paragraph]}]}]}
    listing = {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [paragraph]}]}
    placements = [('root', {'type': 'doc', 'content': [paragraph]}, []),
                  ('page-break', {'type': 'doc', 'content': [{'type': 'pageBreak'}]}, [0]),
                  ('text-node', {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'x'}]}]}, [0, 0]),
                  ('hard-break', {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'hardBreak'}]}]}, [0, 0]),
                  ('table', {'type': 'doc', 'content': [table]}, [0]),
                  ('table-row', {'type': 'doc', 'content': [table]}, [0, 0]),
                  ('table-cell', {'type': 'doc', 'content': [table]}, [0, 0, 0]),
                  ('list', {'type': 'doc', 'content': [listing]}, [0]),
                  ('list-item', {'type': 'doc', 'content': [listing]}, [0, 0])]
    for name, source, indexes in placements:
        model = copy.deepcopy(source)
        node = model
        for index in indexes:
            node = node['content'][index]
        node['attrs'] = {ATTR: font()}
        invalid['misplaced-' + name] = model
    for name, model in invalid.items():
        path, output = directory / ('invalid-' + name + '.json'), directory / ('invalid-' + name + '.docx')
        path.write_text(json.dumps(model))
        error = run_cli(cli, commands, 'create-docx', path, output, success=False)
        check(not output.exists(), name + ': rejected model created output')
        if name.startswith('misplaced-'):
            check(ATTR in error, name + ': wrong-placement refusal did not identify the font attribute')
        evidence[name] = {'stderr': error, 'no_output': True}


def authored_models(cli, commands, directory):
    results = {}
    for name, f in [('minimum', font(halfPoints=2)), ('maximum', font(halfPoints=3276)),
                    ('31-utf16', font('a' * 31)), ('astral-31-utf16', font('😀' * 15 + 'a'))]:
        model = {'type': 'doc', 'content': [{'type': 'paragraph', 'attrs': {ATTR: f}}]}
        results[name] = export_model(cli, commands, directory, 'authored-' + name, model)
    p = lambda kind='paragraph': {'type': kind, 'attrs': dict({ATTR: font('DejaVu Serif', 32, True, True)}, **({'level': 2} if kind == 'heading' else {}))}
    model = {'type': 'doc', 'content': [p(), p('heading'),
             {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [p()]}]},
             {'type': 'orderedList', 'content': [{'type': 'listItem', 'content': [p()]}]},
             {'type': 'table', 'content': [{'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [p()]}]}]}]}
    results['allowed-containers'] = export_model(cli, commands, directory, 'authored-allowed-containers', model)
    return results


def tamper_checks(cli, commands, directory, source):
    results = {}
    for tag, value in [('sz', '30'), ('b', '1'), ('i', '1'), ('rFonts', 'DejaVu Serif')]:
        def mutate(root, tag=tag, value=value):
            n = root.find('.//w:pPr/w:rPr/w:' + tag, NS)
            n.set(qn('w:ascii' if tag == 'rFonts' else 'w:val'), value)
        path = directory / ('tampered-' + tag + '.docx')
        mutate_package(source, path, 'word/document.xml', mutate)
        error = run_cli(cli, commands, 'read-docx-model', path, success=False)
        results[tag] = {'stderr': error, 'rejected': True}
        with zipfile.ZipFile(path) as package:
            parts = {name: package.read(name) for name in package.namelist()}
        stored = json.loads(parts['rust-wps/model-v2.json'])
        stored['hashes']['word/document.xml'] = hashlib.sha256(parts['word/document.xml']).hexdigest()
        parts['rust-wps/model-v2.json'] = json.dumps(stored, ensure_ascii=False, separators=(',', ':')).encode()
        forged = directory / ('tampered-' + tag + '-repaired-hash.docx')
        with zipfile.ZipFile(forged, 'x', zipfile.ZIP_DEFLATED) as package:
            for name, data in parts.items():
                package.writestr(name, data)
        error = run_cli(cli, commands, 'read-docx-model', forged, success=False)
        results[tag + '-repaired-hash'] = {'stderr': error, 'rejected': True, 'scope': 'part hash repaired but stale model rejected by exact reencoding'}
    return results


def measure_pdf(path):
    document = fitz.open(path)
    spans, horizontal = [], []
    for page_index, page in enumerate(document):
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                for span in line['spans']:
                    if span['text'].strip():
                        spans.append({'text': span['text'], 'x': round(span['origin'][0], 5), 'y': round(span['origin'][1], 5),
                                      'font': span['font'], 'size': span['size'], 'page': page_index})
        ys = []
        for drawing in page.get_drawings():
            for item in drawing['items']:
                if item[0] == 'l' and abs(item[1].y - item[2].y) < .05 and abs(item[1].x - item[2].x) > 100:
                    ys.append(round((item[1].y + item[2].y) / 2, 5))
        distinct = []
        for y in sorted(ys):
            if not distinct or y - distinct[-1] > .1:
                distinct.append(y)
        horizontal.append(distinct)
    return {'path': str(path), 'sha256': digest(path), 'page_count': len(document),
            'page_sizes': [[p.rect.width, p.rect.height] for p in document], 'spans': spans, 'horizontal_borders': horizontal}


def render_all(soffice, directory, paths, commands):
    out, cache = directory / 'pdf', directory / 'lo-cache'
    out.mkdir(); cache.mkdir()
    paths = list(dict.fromkeys(Path(p).resolve() for p in paths))
    check(len({p.stem for p in paths}) == len(paths), 'PDF basenames collide')
    results = {}
    for index, path in enumerate(paths):
        profile = directory / f'lo-profile-{index:03d}'
        command = [soffice, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(path)]
        result = subprocess.run(command, capture_output=True, text=True, timeout=90, env=dict(os.environ, XDG_CACHE_HOME=str(cache)))
        commands.append({'command': command, 'returncode': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr})
        check(result.returncode == 0, 'LibreOffice conversion failed: ' + result.stderr)
        results[str(path)] = measure_pdf(out / (path.stem + '.pdf'))
    return results


def compare_baselines(a, b, label):
    check(a['page_count'] == b['page_count'] and a['page_sizes'] == b['page_sizes'], label + ': PDF page geometry differs')
    check([s['text'] for s in a['spans']] == [s['text'] for s in b['spans']], label + ': PDF text differs')
    deltas = []
    for left, right in zip(a['spans'], b['spans']):
        check((left['font'], left['size'], left['page']) == (right['font'], right['size'], right['page']), label + ': neighboring consumer font/page differs')
        deltas.append(abs(left['y'] - right['y']))
    maximum = max(deltas, default=0)
    check(maximum <= TOLERANCE, f'{label}: baseline delta {maximum}pt exceeds {TOLERANCE}')
    return {'max_absolute_baseline_delta_pt': maximum, 'tolerance_pt': TOLERANCE, 'scope': 'controlled Latin text baseline y only; x/style/full-document fidelity not claimed'}


def recorded_span_deltas(source, output):
    """Record all coordinates, but do not turn original styling drift into a pass."""
    same_text = [s['text'] for s in source['spans']] == [s['text'] for s in output['spans']]
    return {'text_sequence_equal': same_text, 'source_page_count': source['page_count'], 'output_page_count': output['page_count'],
            'deltas': [{'text': a['text'], 'page_delta': b['page'] - a['page'], 'x_delta_pt': round(b['x'] - a['x'], 5),
                        'y_delta_pt': round(b['y'] - a['y'], 5), 'source_font': a['font'], 'output_font': b['font']}
                       for a, b in zip(source['spans'], output['spans'])] if same_text else [],
            'scope': 'record-only diagnostic; no full-document, horizontal, styling or font-availability fidelity claim'}


def table_gap_result(source, new, old):
    measurements = {}
    for name, item in [('source', source), ('new', new), ('no_attribute_old_cli', old)]:
        ys = item['horizontal_borders'][0]
        check(len(ys) == 24, name + ': expected six tables with four horizontal borders each')
        measurements[name] = {'table_top_y': ys[::4],
                              'intertable_gaps_pt': [round(ys[i + 4] - ys[i + 3], 5) for i in range(0, 20, 4)]}
    for left, right in zip(measurements['source']['intertable_gaps_pt'], measurements['new']['intertable_gaps_pt']):
        check(abs(left - right) <= TOLERANCE, 'Original table gap baseline not restored')
    measurements['new_table_top_minus_source_pt'] = [round(b - a, 5) for a, b in zip(measurements['source']['table_top_y'], measurements['new']['table_top_y'])]
    measurements['old_table_top_minus_source_pt'] = [round(b - a, 5) for a, b in zip(measurements['source']['table_top_y'], measurements['no_attribute_old_cli']['table_top_y'])]
    measurements['scope'] = 'blank intertable vertical gaps only; original widths, colors, x positions, and automatic geometry remain normalized'
    return measurements


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli', type=Path)
    parser.add_argument('--baseline-cli', type=Path, default=ROOT / 'tmp/table-vertical-final2-cli')
    parser.add_argument('--root', type=Path, default=ROOT / 'tmp/empty-font-independent')
    parser.add_argument('--generate-only', action='store_true')
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--soffice', default=shutil.which('soffice'))
    args = parser.parse_args()
    check(args.generate_only or args.cli, 'Pass a frozen --cli or --generate-only')
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix='run-', dir=args.root)).resolve()
    report = directory / 'report.json'
    script_snapshot = directory / Path(__file__).name
    shutil.copyfile(__file__, script_snapshot)
    evidence = {'status': 'running', 'script_sha256': digest(__file__), 'python': platform.python_version(),
                'script_snapshot': str(script_snapshot),
                'versions': {p: version(p) for p in ('python-docx', 'lxml', 'PyMuPDF')}, 'commands': [],
                'sources': {}, 'cases': {}, 'warning_sources': {}, 'blocked_sources': {}, 'boundaries': {}, 'rejected_sources': {}, 'invalid_models': {}, 'tamper': {},
                'render_strategy': 'Exactly one new soffice process and unique profile per input; XDG_CACHE_HOME inside this evidence directory',
                'limitations': ['No browser/native print, Microsoft Word or WPS verification.',
                               'Latin blank baseline font only; no installed-font guarantee or full-document fidelity.',
                               'Original upstream fixtures are unchanged public regression documents, not user production documents.',
                               'Table/list source inheritance, character-style marks and inherited paragraph marks are warning normalization.',
                               'Authored empty metadata in list/table/heading models does not imply source-import preservation for those containers.']}
    try:
        evidence['sources'] = sources(directory)
        evidence['warning_sources'] = warning_sources(directory)
        evidence['blocked_sources'] = blocked_sources(directory)
        paths = [item['path'] for item in evidence['sources'].values() if item['render']]
        paths += [item['reference']['path'] for item in evidence['sources'].values() if 'reference' in item]
        if not args.generate_only:
            cli, old_cli = args.cli.resolve(), args.baseline_cli.resolve()
            check(digest(old_cli) == OLD_SHA, 'Frozen baseline CLI identity changed')
            evidence['cli'] = {'path': str(cli), 'sha256': digest(cli)}
            evidence['baseline_cli'] = {'path': str(old_cli), 'sha256': digest(old_cli)}
            for name, item in evidence['sources'].items():
                source = Path(item['path'])
                imported = run_cli(cli, evidence['commands'], 'import-docx', source)
                model = imported['document']
                actual = {str(i): n.get('attrs', {}).get(ATTR) for i, n in enumerate(model['content']) if ATTR in n.get('attrs', {})}
                check(actual == item['expected'], name + ': source identity or empty font expectation differs: ' + repr(actual))
                check(all(not n.get('content') for n in walk(model) if ATTR in n.get('attrs', {})), name + ': nonempty node received metadata')
                if item['expected']:
                    check(PRESERVED in json.dumps(imported['compatibility']), name + ': fixed font normalization warning absent')
                if item.get('expected_unproved'):
                    check(UNPROVED in json.dumps(imported['compatibility']), name + ': partial-resolution warning absent')
                old_imported = run_cli(old_cli, evidence['commands'], 'import-docx', source)
                check(strip_font(model) == old_imported['document'], name + ': unrelated import model changed')
                new = export_model(cli, evidence['commands'], directory, name + '-new', model, source)
                reset = strip_font(model)
                reset_result = export_model(cli, evidence['commands'], directory, name + '-reset', reset, source)
                old_output = directory / (name + '-old-no-attribute.docx')
                run_cli(old_cli, evidence['commands'], 'create-docx', reset_result['model_path'], old_output)
                check(digest(old_output) == reset_result['sha256'], name + ': entire no-attribute DOCX bytes changed')
                check(run_cli(old_cli, evidence['commands'], 'read-docx-model', old_output) == reset, name + ': old trusted model differs')
                with zipfile.ZipFile(new['path']) as a, zipfile.ZipFile(reset_result['path']) as b:
                    check(set(a.namelist()) == set(b.namelist()), name + ': empty font changed package part inventory')
                    check(all(a.read(part) == b.read(part) for part in a.namelist() if part not in ('word/document.xml', 'rust-wps/model-v2.json')), name + ': empty font changed a global package part')
                    check(a.read('word/styles.xml') == b.read('word/styles.xml'), name + ': global styles changed for empty font')
                    direct, reset_xml = ET.fromstring(a.read('word/document.xml')), ET.fromstring(b.read('word/document.xml'))
                    for ppr in direct.findall('.//w:pPr', NS):
                        remove(ppr, 'rPr')
                        if len(ppr) == 0 and not ppr.attrib and not ppr.text:
                            ppr.getparent().remove(ppr)
                    for ppr in reset_xml.findall('.//w:pPr', NS):
                        if len(ppr) == 0 and not ppr.attrib and not ppr.text:
                            ppr.getparent().remove(ppr)
                    check(ET.tostring(direct, method='c14n') == ET.tostring(reset_xml, method='c14n'), name + ': document XML changed outside paragraph-mark properties')
                evidence['cases'][name] = {'compatibility': imported['compatibility'], 'old_compatibility': old_imported['compatibility'],
                                          'only_new_model_field_is_empty_font': True, 'new': new, 'reset': reset_result,
                                          'old_no_attribute': {'path': str(old_output), 'sha256': digest(old_output), 'entire_package_bytes_equal': True},
                                          'global_styles_unchanged': True, 'only_document_and_embedded_model_parts_changed': True,
                                          'only_paragraph_mark_properties_changed_in_document_xml': True}
                if item['render']:
                    paths += [new['path']]
                if name in ('defaults-12', 'known-contextual-control', 'contextual-no-run-counterpart', 'original-table-alignment', 'original-NoHeadFoot'):
                    paths += [str(old_output)]
            for name, item in evidence['warning_sources'].items():
                imported = run_cli(cli, evidence['commands'], 'import-docx', item['path'])
                old = run_cli(old_cli, evidence['commands'], 'import-docx', item['path'])
                check(not any(ATTR in n.get('attrs', {}) for n in walk(imported['document'])), name + ': unproved font attached')
                check(imported['document'] == old['document'], name + ': existing model behavior changed')
                if item['warning']:
                    check(UNPROVED in json.dumps(imported['compatibility']), name + ': normalization warning missing')
                evidence['boundaries'][name] = {'compatibility': imported['compatibility'], 'old_model_equal': True, 'no_empty_font_attached': True}
            for name, item in evidence['blocked_sources'].items():
                new_error = run_cli(cli, evidence['commands'], 'import-docx', item['path'], success=False)
                old_error = run_cli(old_cli, evidence['commands'], 'import-docx', item['path'], success=False)
                check(item['needle'].lower() in new_error.lower() and item['needle'].lower() in old_error.lower(), name + ': existing blocker diagnosis changed')
                evidence['rejected_sources'][name] = {'new_stderr': new_error, 'old_stderr': old_error, 'existing_blocker_preserved': True}
            evidence['authored_models'] = authored_models(cli, evidence['commands'], directory)
            rejection_checks(cli, evidence['commands'], directory, evidence['invalid_models'])
            evidence['tamper'] = tamper_checks(cli, evidence['commands'], directory, Path(evidence['cases']['defaults-12']['new']['path']))
        if args.render:
            check(args.soffice, 'LibreOffice unavailable')
            evidence['soffice_version'] = subprocess.run([args.soffice, '--version'], capture_output=True, text=True).stdout.strip()
            report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
            evidence['renders'] = render_all(args.soffice, directory, paths, evidence['commands'])
            comparisons = {}
            for name, item in evidence['sources'].items():
                if not item['render']:
                    continue
                source_pdf = evidence['renders'][str(Path(item['path']).resolve())]
                if 'reference' in item:
                    reference_pdf = evidence['renders'][str(Path(item['reference']['path']).resolve())]
                    comparisons[name + '-source-vs-independent-reference'] = compare_baselines(source_pdf, reference_pdf, name + ' source/reference')
                if not args.generate_only and (name in ('known-contextual-control', 'bare-empty-run', 'contextual-no-run-counterpart') or 'reference' in item):
                    new_pdf = evidence['renders'][evidence['cases'][name]['new']['path']]
                    comparisons[name + '-source-vs-new'] = compare_baselines(source_pdf, new_pdf, name + ' source/new')
            if not args.generate_only:
                name = 'original-table-alignment'; item = evidence['sources'][name]; case = evidence['cases'][name]
                comparisons['unchanged-original-table-gaps'] = table_gap_result(evidence['renders'][item['path']], evidence['renders'][case['new']['path']], evidence['renders'][case['old_no_attribute']['path']])
                for name in ('defaults-12', 'known-contextual-control', 'contextual-no-run-counterpart', 'original-table-alignment', 'original-NoHeadFoot'):
                    item, case = evidence['sources'][name], evidence['cases'][name]
                    source_pdf = evidence['renders'][item['path']]
                    comparisons[name + '-old-record-only'] = recorded_span_deltas(source_pdf, evidence['renders'][case['old_no_attribute']['path']])
                    comparisons[name + '-new-record-only'] = recorded_span_deltas(source_pdf, evidence['renders'][case['new']['path']])
            evidence['comparisons'] = comparisons
        for item in list(evidence['sources'].values()) + list(evidence['warning_sources'].values()) + list(evidence['blocked_sources'].values()):
            check(digest(item['path']) == item['sha256'], 'Original/source mutated during checks')
        evidence['all_original_and_generated_sources_unchanged'] = True
        evidence['status'] = 'generated' if args.generate_only else 'passed'
    except Exception as error:
        evidence['status'] = 'failed'
        evidence['error'] = str(error)
        raise
    finally:
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
        print(report)


if __name__ == '__main__':
    main()
