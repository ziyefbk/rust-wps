"""Independent bounded table-position regression.

Source fixtures are directed synthetic python-docx documents, never original
user files. ORIGINAL POI fixtures are separately hash-pinned in the companion
manifest; alignment is an autofit normalization positive, indentation a refusal.
LibreOffice PDF probes validate controlled external-consumer border positions,
not the app's rendering, pagination, Windows/WebView2, Microsoft Word or WPS.
Explicit-page overflow is separately refused; historical absent-page position
limits and edits are still exercised without silently changing table values.

  python scripts/check-word-table-position.py --generate-only
  python scripts/check-word-table-position.py --cli target/debug/rust-wps --render --corpus-dir tmp/table-corpus
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import platform
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from pypdf import PdfReader

_spec = importlib.util.spec_from_file_location("width_checks", Path(__file__).with_name("check-word-table-widths.py"))
w = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(w)
ALIGN = "tableAlign"
INDENT = "tableIndentTwips"
WIDTH = "columnWidthsTwips"
INHERITED_WARNING = "Inherited table indentation normalized; automatic table geometry not preserved"


def snapshot(path):
    result = w.independent_snapshot(path)
    with zipfile.ZipFile(path) as z:
        root = w.ET.fromstring(z.read("word/document.xml"))
        result["page"] = [{w.ET.QName(k).localname: v for k, v in n.attrib.items()}
                          for n in root.findall(".//w:sectPr/w:pgSz", w.NS)]
        result["margins"] = [{w.ET.QName(k).localname: v for k, v in n.attrib.items()}
                             for n in root.findall(".//w:sectPr/w:pgMar", w.NS)]
        result["table_cell_margins"] = [[{"edge": w.ET.QName(n).localname, **w.declaration(n)}
                                         for n in t.findall("w:tblPr/w:tblCellMar/*", w.NS)]
                                        for t in root.findall(".//w:tbl", w.NS)]
        result["settings_xml"] = z.read("word/settings.xml").decode() if "word/settings.xml" in z.namelist() else None
    result["python_docx_alignment"] = []
    result["body_paragraph_texts"] = [p.text for p in Document(path).paragraphs]
    for table in Document(path).tables:
        try:
            value = table.alignment
            result["python_docx_alignment"].append(None if value is None else {"value": int(value), "name": str(value)})
        except ValueError:
            result["python_docx_alignment"].append("unsupported-by-python-docx-enum; raw XML checked")
    return result


def target(xml, widths, align=None, indent=None):
    return {**(w.fixed_expected(widths, xml["texts"], xml["header_rows"]) if widths is not None else
            {"widths": None, "texts": xml["texts"], "header_rows": xml["header_rows"]}),
            "align": align, "indent": indent}


def verify_positions(snap, targets, label, source=False):
    w.verify_independent(snap, targets, label, fixed_required=not source)
    for i, (xml, expected) in enumerate(zip(snap["ooxml"], targets)):
        actual_align = None if xml["alignment"] is None else xml["alignment"].get("val")
        if source:
            actual_align = {"start": "left", "end": "right"}.get(actual_align, actual_align)
        w.check(actual_align == expected["align"], f"{label} table {i}: alignment differs")
        indent = expected.get("source_indent", expected["indent"]) if source else expected["indent"]
        expected_indent = None if indent is None else {"type": "dxa", "w": str(indent)}
        w.check(xml["indent"] == expected_indent, f"{label} table {i}: absent/explicit indent differs")
        if not source:
            py_align = snap["python_docx_alignment"][i]
            expected_value = {None: None, "left": 0, "center": 1, "right": 2}[expected["align"]]
            w.check((py_align is None) if expected_value is None else py_align["value"] == expected_value,
                    f"{label} table {i}: python-docx alignment differs")
        if not source and (expected["align"] is not None or expected["indent"] is not None):
            margins = {m["edge"]: {k: v for k, v in m.items() if k != "edge"} for m in snap["table_cell_margins"][i]}
            w.check(margins.get("left") == margins.get("right") == {"type": "dxa", "w": "108"},
                    f"{label} table {i}: positioned export lacks explicit 108twip horizontal cell margins")


def verify_model(model, targets, label):
    tables = w.verify_model(model, targets, label)
    for i, (table, expected) in enumerate(zip(w.model_tables(model), targets)):
        attrs = table.get("attrs", {})
        for key, value in [(ALIGN, expected["align"]), (INDENT, expected["indent"])]:
            w.check((key not in attrs) if value is None else attrs.get(key) == value,
                    f"{label} table {i}: {key} missing, invented or changed")
    return tables


def make_fixtures(directory):
    sources = {}

    def save(name, doc, expected=None, reject=None, render=False, page_fit=False, legacy_no_page=False):
        path = directory / (name + ".docx")
        w.check(not path.exists(), "Fixture exists: " + str(path))
        doc.save(path)
        sources[name] = {"path": str(path), "sha256": w.digest(path), "bytes": path.stat().st_size,
                         "fixture_kind": "directed synthetic python-docx fixture, not an original upstream/user document",
                         "expected_action": "reject" if reject else "roundtrip", "unsupported_form": reject,
                         "source_independent": snapshot(path), "expected": expected, "render_probe": render,
                         "page_fit_refusal": page_fit, "legacy_absent_page": legacy_no_page}
        if page_fit:
            tables = sources[name]["source_independent"]["ooxml"]
            targets = [{"widths": [int(g["w"]) for g in table["grid"]] if table["layout"] == {"type": "fixed"} else None,
                        "indent": None if table["indent"] is None else int(table["indent"]["w"])} for table in tables]
            sources[name]["page_fit_evidence"] = w.external_page_fit(path, targets)
            w.check(not sources[name]["page_fit_evidence"]["fits"], name + ": page-fit rejection fixture actually fits")
        if expected is not None:
            verify_positions(sources[name]["source_independent"], expected, name + " source", source=True)

    def fixture(name, jc=None, indent=None, typ="dxa", layout="fixed", reject=None, mutate=None, render=False, effective_indent=None):
        doc = w.fresh_document()
        table = w.add_table(doc, [1800, 4200], rows=1, probe=True, layout=layout)
        if jc is not None:
            w.child(table._tbl.tblPr, "jc", val=jc)
        if indent is not None:
            w.child(table._tbl.tblPr, "tblInd", w=indent, **({"type": typ} if typ is not None else {}))
        if mutate:
            mutate(doc, table)
        # Expected text/width declarations are fixed by fixture construction.
        texts = [[" ".join(w.WRAP_WORDS), "RightProbe"]]
        expected = [{"widths": [1800, 4200] if layout == "fixed" else None, "texts": texts,
                     "header_rows": [False], "align": {"start": "left", "end": "right"}.get(jc, jc), "indent": indent}]
        if effective_indent is not None:
            expected[0].update(indent=effective_indent, source_indent=indent)
        if name in ("auto-left-indent", "fixed-max-combined", "auto-max-combined"):
            save("reject-page-fit-" + name, doc, reject="known table extent exceeds explicit A4 content width", page_fit=True)
            w.without_page_declaration(doc)
            save(name, doc, expected, render=render, legacy_no_page=True)
        else:
            save(name, doc, None if reject else expected, reject, render)

    for jc in [None, "left", "center", "right", "start", "end"]:
        fixture("fixed-align-" + str(jc), jc, render=jc in [None, "left", "center", "right"])
    for jc in [None, "left", "start"]:
        for indent in [0, 720]:
            fixture(f"fixed-{jc}-indent-{indent}", jc, indent, render=jc == "left")
    for jc in ["center", "right"]:
        fixture(f"fixed-{jc}-explicit-zero", jc, 0, render=True)
    for jc in ["left", "center", "right", "start", "end"]:
        fixture("auto-align-" + jc, jc, layout=None)
    fixture("auto-left-indent", "left", 720, layout=None)
    fixture("fixed-max-combined", "left", 25680)
    fixture("auto-max-combined", "left", 22680, layout=None)
    for name, value, typ in [("negative", -1, "dxa"), ("fractional", "1.5", "dxa"), ("missing-type", 720, None),
                             ("nil", 1234, "nil"), ("pct", 5678, "pct"), ("auto", 9012, "auto"),
                             ("over-limit", 31681, "dxa")]:
        fixture("reject-indent-" + name, "left", value, typ, reject=name)
    for jc in ["center", "right", "end"]:
        fixture("reject-nonleft-indent-" + jc, jc, 720, reject="nonzero indentation with non-left alignment")
    fixture("reject-fixed-combined-bound", "left", 25681, reject="width+indent exceeds31680")
    fixture("reject-auto-combined-bound", "left", 22681, layout=None, reject="9000+indent exceeds31680")
    for jc in ["both", "invalid"]:
        fixture("reject-align-" + jc, jc, reject="unsupported alignment")
    for value in [0, 720]:
        def modern(doc, table):
            for n in doc.settings.element.findall(".//w:compatSetting", w.NS):
                if n.get(w.qn("w:name")) == "compatibilityMode":
                    n.set(w.qn("w:val"), "15")
        fixture(f"reject-modern-indent-{value}", "left", value, reject="modern compatibility border-edge semantics", mutate=modern)
    def margin(value, cell=False):
        def apply(doc, table):
            p = w.child(table.cell(0, 0)._tc.get_or_add_tcPr() if cell else table._tbl.tblPr,
                        "tcMar" if cell else "tblCellMar")
            for edge in ["left", "right"]:
                w.child(p, edge, w=value, type="dxa")
        return apply
    for amount in [0, 216]:
        fixture(f"reject-custom-margin-{amount}", "left", 720, reject="custom horizontal margin", mutate=margin(amount))
    fixture("reject-cell-margin", "left", 720, reject="cell-level horizontal margin override", mutate=margin(216, True))
    def malformed_margin(kind):
        def apply(doc, table):
            props = table._tbl.tblPr
            parent = props
            if kind == "wrapped-container": parent = w.child(props, "wrapper")
            mar = w.child(parent, "tblCellMar")
            if kind == "foreign-container": mar.tag = "{urn:position-probe}tblCellMar"
            if kind == "container-attribute": mar.set(w.qn("w:extra"), "1")
            side_parent = mar
            if kind == "wrapped-side":
                side_parent = w.ET.SubElement(mar, "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
            left = w.child(side_parent, "left", w=108, type="dxa")
            w.child(mar, "right", w=108, type="dxa")
            if kind == "foreign-side": left.tag = "{urn:position-probe}left"
            if kind == "duplicate-side": w.child(mar, "left", w=108, type="dxa")
            if kind == "duplicate-container":
                other = w.child(props, "tblCellMar"); w.child(other, "left", w=108, type="dxa")
            if kind == "ambiguous-side": w.child(mar, "start", w=108, type="dxa")
            if kind == "side-attribute": left.set(w.qn("w:extra"), "1")
            if kind == "side-missing-type": del left.attrib[w.qn("w:type")]
            if kind == "side-child": w.child(left, "hidden")
            if kind == "unknown-side": w.child(mar, "hidden")
        return apply
    for kind in ["wrapped-container", "wrapped-side", "foreign-container", "foreign-side", "duplicate-side",
                 "duplicate-container", "ambiguous-side", "container-attribute", "side-attribute", "side-missing-type",
                 "side-child", "unknown-side"]:
        fixture("reject-margin-" + kind, "left", 720, reject="ambiguous or malformed horizontal margin: " + kind,
                mutate=malformed_margin(kind))

    fixture("reject-bidi", "start", 720, reject="bidi table", mutate=lambda d, t: w.child(t._tbl.tblPr, "bidiVisual"))
    fixture("reject-duplicate-align", "left", reject="duplicate alignment", mutate=lambda d, t: w.child(t._tbl.tblPr, "jc", val="right"))
    fixture("reject-duplicate-indent", "left", 0, reject="duplicate indent", mutate=lambda d, t: w.child(t._tbl.tblPr, "tblInd", w=0, type="dxa"))
    fixture("reject-row-align", "left", reject="row-level alignment", mutate=lambda d, t: w.child(t.rows[0]._tr.get_or_add_trPr(), "jc", val="right"))
    fixture("reject-row-property-exception", "left", reject="row table-property exception", mutate=lambda d, t: w.child(w.child(t.rows[0]._tr, "tblPrEx"), "jc", val="right"))
    def mode(value):
        def apply(doc, table):
            for n in doc.settings.element.findall(".//w:compatSetting", w.NS):
                if n.get(w.qn("w:name")) == "compatibilityMode":
                    if value is None: n.getparent().remove(n)
                    else: n.set(w.qn("w:val"), str(value))
        return apply
    for value in [None, 11, 12]:
        fixture(f"fixed-compatible-mode-{value}", "left", 720, mutate=mode(value), render=True)
    fixture("reject-unknown-mode", "left", 720, reject="unsupported compatibility mode", mutate=mode(13))
    def selected_style(style_id="TableGrid", compatibility=14, custom=False):
        def apply(doc, table):
            if custom:
                from docx.enum.style import WD_STYLE_TYPE
                style = doc.styles.add_style(style_id, WD_STYLE_TYPE.TABLE)
                style.base_style = doc.styles["Normal Table"]
            w.child(table._tbl.tblPr, "tblStyle", val=style_id)
            mode(compatibility)(doc, table)
        return apply
    for style_id, custom in [("TableGrid", False), ("TableNormal", False), ("PositionBaseZero", True)]:
        fixture("fixed-inherited-zero-" + style_id, "left", mutate=selected_style(style_id, custom=custom), effective_indent=0, render=True)
    fixture("fixed-inherited-zero-center", "center", mutate=selected_style(), effective_indent=0, render=True)
    fixture("fixed-inherited-zero-right", "right", mutate=selected_style(), effective_indent=0, render=True)
    fixture("reject-modern-inherited-zero", "left", reject="modern fixed inherited zero indentation", mutate=selected_style(compatibility=15))
    fixture("auto-modern-inherited-zero", "left", layout=None, mutate=selected_style(compatibility=15))
    fixture("reject-inherited-custom-margin", "left", reject="legacy inherited zero with custom horizontal margin",
            mutate=lambda d,t: (selected_style()(d,t), margin(216)(d,t)))
    # Header and multiple-table evidence is separate from one-table PDF probes.
    doc = w.fresh_document()
    t = w.add_table(doc, [1200, 4800], rows=3, header=True); w.child(t._tbl.tblPr, "jc", val="center")
    doc.add_paragraph("BetweenTables")
    t = w.add_table(doc, [1000, 2000, 3000], rows=2); w.child(t._tbl.tblPr, "jc", val="left"); w.child(t._tbl.tblPr, "tblInd", w=360, type="dxa")
    path = directory / "snapshot-only.docx"; doc.save(path); raw = w.xml_snapshot(path); path.unlink()
    save("multiple-header-position", doc, [target(raw[0], [1200, 4800], "center"), target(raw[1], [1000, 2000, 3000], "left", 360)])
    manifest = directory / "expected-fixtures.json"
    manifest.write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")
    return sources, manifest


def exported(cli, evidence, directory, name, model, targets):
    model_path, path = directory / (name + ".json"), directory / (name + ".docx")
    w.write_model(model_path, model)
    w.run_cli(cli, evidence["commands"], "create-docx", model_path, path)
    result = {"path": str(path), "model_path": str(model_path), "sha256": w.digest(path), "independent": snapshot(path)}
    verify_positions(result["independent"], targets, name)
    reopened = w.run_cli(cli, evidence["commands"], "read-docx-model", path)
    verify_model(reopened, targets, name + " trusted reopen")
    w.check(reopened == model, name + ": full trusted model equality failed")
    result["trusted_full_model_equal"] = True
    result["external_page_fit"] = w.external_page_fit(path, targets)
    if not result["external_page_fit"]["fits"]:
        w.check("pageGeometry" not in model.get("attrs", {}), "Explicit-page model wrote overflowing table")
        error = w.run_cli(cli, evidence["commands"], "import-docx", path, success=False)
        w.check(w.PAGE_FIT in error, name + ": wrong external page-fit refusal")
        result.update(external_reimport_refusal=error, external_reimport_expected_action="refuse emitted explicit page overflow; trusted legacy model remains exact")
    else:
        external = w.run_cli(cli, evidence["commands"], "import-docx", path)["document"]
        verify_model(external, targets, name + " external reimport")
        result["external_position_and_width_equal"] = True
    error = w.run_cli(cli, evidence["commands"], "create-docx", model_path, path, success=False)
    w.check("Output already exists" in error and w.digest(path) == result["sha256"], "Output overwritten: " + name)
    result["overwrite_refused"] = True
    return result


def rendered(soffice, directory, path, expected):
    widths = expected["widths"]
    result = w.render_pdf(soffice, directory, path, widths)
    snap = snapshot(path)
    page, margins = snap["page"][-1], snap["margins"][-1]
    page_width, left, right = int(page["w"]) / 20, int(margins["left"]) / 20, int(margins["right"]) / 20
    result.update(page_width_pt=float(PdfReader(result["path"]).pages[0].mediabox.width),
                  document_page_width_pt=page_width, left_margin_pt=left, right_margin_pt=right)
    w.check(abs(result["page_width_pt"] - page_width) < 0.1, "PDF page width mismatch")
    width = sum(widths) / 20
    align, indent = expected["align"], expected["indent"]
    x = left + (page_width - left - right - width) / 2 if align == "center" else page_width - right - width if align == "right" else left + (indent / 20 - 5.4 if indent is not None else 0)
    result["expected_left_border_pt"] = x
    result["left_border_error_pt"] = result["vertical_border_x_pt"][0] - x
    w.check(abs(result["left_border_error_pt"]) < 0.8, f"Border position mismatch: {path}: {result['left_border_error_pt']}")
    return result


def bad_models(cli, evidence, directory, base):
    cases = []
    for key, values in [(ALIGN, [None, "", "start", "end", "both", 0, True, [], {}]),
                        (INDENT, [None, -1, 0.5, "720", True, [], {}, 31681])]:
        for i, value in enumerate(values):
            model = copy.deepcopy(base); next(w.model_tables(model)).setdefault("attrs", {})[key] = value
            cases.append((key + "-" + str(i), model))
    for name, attrs in [("nonleft-center", {ALIGN: "center", INDENT: 1}), ("nonleft-right", {ALIGN: "right", INDENT: 1}),
                        ("combined-fixed", {WIDTH: [1800, 4200], INDENT: 25681}),
                        ("combined-auto", {INDENT: 22681})]:
        model = copy.deepcopy(base); next(w.model_tables(model))["attrs"] = attrs; cases.append((name, model))
    for key, value in [(ALIGN, "left"), (INDENT, 0)]:
        for location in ["doc", "row", "cell", "paragraph"]:
            model = copy.deepcopy(base); table = next(w.model_tables(model)); row = table["content"][0]; cell = row["content"][0]
            node = {"doc": model, "row": row, "cell": cell, "paragraph": cell["content"][0]}[location]
            node.setdefault("attrs", {})[key] = value; cases.append((key + "-location-" + location, model))
    results = []
    for name, model in cases:
        jp, dp = directory / ("bad-" + name + ".json"), directory / ("bad-" + name + ".docx")
        w.write_model(jp, model); error = w.run_cli(cli, evidence["commands"], "create-docx", jp, dp, success=False)
        w.check(not dp.exists(), "Invalid model wrote output: " + name)
        results.append({"name": name, "error": error, "output_absent": True})
    return results


def corpus_check(root, cli, evidence, directory, safe_model_path):
    manifest_path = Path(__file__).with_name("table-position-corpus.json")
    manifest = json.loads(manifest_path.read_text())
    results = []
    for expected in manifest["fixtures"]:
        path = root / expected["name"]; data = path.read_bytes()
        git_hash = hashlib.sha1(f"blob {len(data)}\0".encode() + data).hexdigest()
        w.check(len(data) == expected["bytes"] and w.digest(path) == expected["sha256"] and git_hash == expected["git_blob_sha"], "Original corpus provenance mismatch")
        snap = snapshot(path)
        w.check(len(snap["ooxml"]) == expected["table_count"], "Original table count changed")
        result = {**expected, "path": str(path), "source_independent": snap, "source_unchanged": True,
                  "fixture_kind": "original unmodified pinned Apache POI regression fixture, not a user production document"}
        if expected["expected_action"] == "autofit-position-roundtrip":
            targets = []
            for i, (xml, align) in enumerate(zip(snap["ooxml"], expected["source_alignments"])):
                w.check(xml["alignment"] == (None if align is None else {"val": align}), "Original alignment declarations changed")
                w.check(xml["layout"] is None and xml["total"] == {"w": "0", "type": "auto"}, "Original was not autofit")
                w.check(xml["indent"] is None and xml["texts"] == expected["cell_texts"], "Original table content/indent changed")
                w.check(xml["grid"] == [{"w": str(n)} for n in expected["source_grid_twips"]], "Original grid changed")
                targets.append(target(xml, None, {"start": "left", "end": "right"}.get(align, align)))
            if cli:
                result["audit"] = w.run_cli(cli, evidence["commands"], "audit-docx", path)
                w.check(w.WARNING in result["audit"]["risks"] and INHERITED_WARNING in result["audit"]["risks"] and w.RISK not in result["audit"]["risks"], "Original auto position import missing correct warning")
                model = w.run_cli(cli, evidence["commands"], "import-docx", path)["document"]
                verify_model(model, targets, "Original autofit imported")
                result["roundtrip"] = exported(cli, evidence, directory, "original-auto-alignment", model, targets)
                w.check(result["roundtrip"]["independent"]["body_paragraph_texts"] == snap["body_paragraph_texts"],
                        "Original corpus body paragraph text changed")
                result["body_paragraph_text_preserved"] = True
                result["numeric_width_or_full_layout_preservation_claimed"] = False
        else:
            pinned = json.loads(Path(__file__).with_name("table-width-corpus.json").read_text())["fixtures"][0]
            for xml, declaration in zip(snap["ooxml"], pinned["tables"]):
                w.check(xml["layout"] == pinned["layout"] and xml["alignment"] == pinned["alignment"] and xml["indent"] == declaration["indent"], "Original refusal geometry changed")
            if cli:
                result["audit"] = w.run_cli(cli, evidence["commands"], "audit-docx", path)
                w.check(w.RISK in result["audit"]["risks"], "Original indent geometry risk missing")
                result["refusal"] = w.run_cli(cli, evidence["commands"], "import-docx", path, success=False)
                w.check(w.RISK in result["refusal"], "Original indent refused for unrelated reason")
        if cli:
            error = w.run_cli(cli, evidence["commands"], "create-docx", safe_model_path, path, success=False)
            w.check("Output already exists" in error, "Original source overwrite not refused")
            result["source_overwrite_refused"] = True
        w.check(w.digest(path) == expected["sha256"], "Original source changed")
        results.append(result)
    return {"manifest_path": str(manifest_path), "manifest_sha256": w.digest(manifest_path), "repository": manifest["repository"],
            "revision": manifest["revision"], "fixtures": results}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli", type=Path, default=Path("target/debug/rust-wps"))
    parser.add_argument("--root", type=Path, default=Path("tmp/word-table-position"))
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--corpus-dir", type=Path)
    args = parser.parse_args()
    cli = None if args.generate_only else args.cli.resolve()
    soffice = shutil.which("soffice")
    w.check(cli is None or cli.is_file(), "CLI missing")
    w.check(not args.render or soffice, "LibreOffice missing")
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.root)).resolve()
    report = directory / "report.json"
    evidence = {"status": "running", "script_sha256": w.digest(Path(__file__)), "width_helper_sha256": w.digest(Path(w.__file__)),
                "python_version": platform.python_version(), "python_docx_version": version("python-docx"), "pypdf_version": version("pypdf"),
                "cli": str(cli), "cli_sha256": None if cli is None else w.digest(cli), "commands": [], "sources": {},
                "semantic_sources": [
                    {"title": "Microsoft OOXML TableIndentation", "url": "https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.wordprocessing.tableindentation?view=openxml-3.0.1", "finding": "ISO leading-edge indentation; inherited when absent; pct/auto ignored; standard ignores indent unless left-aligned"},
                    {"title": "Microsoft Word deviation: table indentation", "url": "https://learn.microsoft.com/en-us/openspecs/office_standards/ms-oe376/a129dffd-ef75-4faa-ab1f-47c3fd4cff26", "finding": "Word does not ignore indentation for non-left alignment, unlike ISO wording"},
                    {"title": "Microsoft OOXML TableJustification", "url": "https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.wordprocessing.tablejustification?view=openxml-3.0.1", "finding": "Alignment relative to text margins; bidiVisual reverses interpretation; absent inherits styles"},
                    {"title": "Microsoft table measurement notes", "url": "https://learn.microsoft.com/en-us/openspecs/office_standards/ms-oi29500/d9b4a6f6-8b07-43c5-87f2-34c1238911a4", "finding": "Word's nonnegative width rule excludes tblInd; negative indentation is real OOXML behavior but excluded from this bounded feature"},
                    {"title": "Microsoft table width unit values", "url": "https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.wordprocessing.tablewidthunitvalues?view=openxml-3.0.1", "finding": "nil means no width; dxa means twentieths of a point"}
                ],
                "limitations": "Synthetic fixed-table source fixtures use known legacy compatibility modes and 108twip horizontal cell margins. "
                "Model indent uses text-edge semantics; absent and explicit zero are distinct. Mode15/custom-margin indent sources are refusal fixtures. "
                "Original autofit alignment corpus validates canonical alignment and explicit width normalization, not numerical original-width/position fidelity. "
                "PDF probes establish this installed LibreOffice consumer only, not app/native/Word/WPS rendering or pagination."}
    try:
        sources, manifest = make_fixtures(directory)
        evidence.update(sources=sources, expected_manifest=str(manifest), expected_manifest_sha256=w.digest(manifest))
        if args.render:
            evidence["libreoffice_version"] = subprocess.run([soffice, "--version"], capture_output=True, text=True, check=True).stdout.strip()
        safe_model = safe_model_path = None
        for name, item in sources.items():
            path = Path(item["path"])
            if cli:
                item["audit"] = w.run_cli(cli, evidence["commands"], "audit-docx", path)
                if item["expected_action"] == "reject":
                    w.check(item.get("page_fit_refusal") or w.RISK in item["audit"]["risks"], name + ": precise geometry risk missing")
                    item["refusal"] = w.run_cli(cli, evidence["commands"], "import-docx", path, success=False)
                    diagnostic = w.PAGE_FIT if item.get("page_fit_refusal") else w.RISK
                    w.check(diagnostic in item["refusal"], name + ": refused for unrelated reason")
                else:
                    w.check(w.RISK not in item["audit"]["risks"], name + ": supported source refused")
                    if item["expected"][0]["widths"] is None:
                        w.check(w.WARNING in item["audit"]["risks"], name + ": width normalization warning missing")
                    if name == "auto-modern-inherited-zero":
                        w.check(INHERITED_WARNING in item["audit"]["risks"], "Modern auto inherited-position warning missing")
                    model = w.run_cli(cli, evidence["commands"], "import-docx", path)["document"]
                    if item.get("legacy_absent_page"):
                        w.check(w.PAGE_WARNING in item["audit"]["risks"] and "pageGeometry" not in model.get("attrs", {}), name + ": missing-page normalization not explicit")
                    verify_model(model, item["expected"], name + " import")
                    item["unchanged_export"] = exported(cli, evidence, directory, name + "-unchanged", model, item["expected"])
                    if safe_model is None:
                        safe_model, safe_model_path = model, item["unchanged_export"]["model_path"]
                    edited, targets = copy.deepcopy(model), copy.deepcopy(item["expected"])
                    for table, expected in zip(w.model_tables(edited), targets):
                        align, indent = ("right", 0) if expected["align"] == "left" else ("left", 360)
                        table.setdefault("attrs", {}).update({ALIGN: align, INDENT: indent})
                        expected.update(align=align, indent=indent)
                    # Only position attrs changed; text, paragraph/style attrs, widths and structure remain identical.
                    original_copy, edited_copy = copy.deepcopy(model), copy.deepcopy(edited)
                    for document in [original_copy, edited_copy]:
                        for table in w.model_tables(document):
                            for key in [ALIGN, INDENT]: table.get("attrs", {}).pop(key, None)
                    w.check(original_copy == edited_copy, "Position edit changed non-position content")
                    item["edited_expected"] = targets
                    if w.model_exceeds_page(edited):
                        # The exact historical position edit is now unsafe for an explicit page.
                        # Prove atomic refusal, then exercise the same old edit in legacy mode.
                        jp, dp = directory / (name + "-edited-overflow.json"), directory / (name + "-edited-overflow.docx")
                        w.write_model(jp, edited)
                        error = w.run_cli(cli, evidence["commands"], "create-docx", jp, dp, success=False)
                        w.check(w.PAGE_FIT in error and not dp.exists(), name + ": explicit-page edit did not refuse precisely without output")
                        item["explicit_page_edit_refusal"] = {"error": error, "output_absent": True, "model_path": str(jp)}
                        edited = copy.deepcopy(edited)
                        edited["attrs"].pop("pageGeometry")
                        item["edited_export_scope"] = "Historical position edit in explicit legacy absent-page model; same edit on preserved source page refused"
                    item["edited_export"] = exported(cli, evidence, directory, name + "-edited", edited, targets)
                    if args.render and item["render_probe"]:
                        item["unchanged_pdf"] = rendered(soffice, directory, Path(item["unchanged_export"]["path"]), item["expected"][0])
                        item["edited_pdf"] = rendered(soffice, directory, Path(item["edited_export"]["path"]), targets[0])
            if args.render and item["render_probe"]:
                item["source_pdf"] = rendered(soffice, directory, path, item["expected"][0])
                if "unchanged_pdf" in item:
                    a, b = item["source_pdf"], item["unchanged_pdf"]
                    w.check(all(abs(x-y) < 0.6 for x,y in zip(a["vertical_border_x_pt"], b["vertical_border_x_pt"])), name + ": source/export absolute borders differ")
                    w.check(a["wrap_line_count"] == b["wrap_line_count"], name + ": source/export wrapping differs")
                    item["source_export_geometry_and_wrapping_match"] = True
            if args.render and name in ["reject-modern-indent-0", "reject-modern-indent-720", "reject-modern-inherited-zero"]:
                # Consumer calibration explains the refusal; it is never a supported import/roundtrip.
                pdf = w.render_pdf(soffice, directory, path, [1800, 4200])
                modern_indent = 720 if name == "reject-modern-indent-720" else 0
                expected_x = 72 + modern_indent / 20
                w.check(abs(pdf["vertical_border_x_pt"][0] - expected_x) < 0.8, "Modern compatibility calibration changed")
                pdf.update(scope="Rejected modern source calibration only; not supported model geometry", expected_left_border_pt=expected_x)
                item["rejected_source_semantics_pdf"] = pdf
            w.check(w.digest(path) == item["sha256"], "Source changed: " + name)
            item["source_unchanged"] = True
        if args.render:
            stages = ["source_pdf", "unchanged_pdf"] if cli else ["source_pdf"]
            evidence["pdf_assertions"] = {}
            for stage in stages:
                absent = sources["fixed-align-None"][stage]["vertical_border_x_pt"][0]
                zero = sources["fixed-left-indent-0"][stage]["vertical_border_x_pt"][0]
                positive = sources["fixed-left-indent-720"][stage]["vertical_border_x_pt"][0]
                w.check(4.5 < absent-zero < 6.3 and abs(positive-zero-36) < 0.3, "Explicit-zero/indent displacement relation changed")
                evidence["pdf_assertions"][stage] = {"absent_border_pt": absent, "explicit_zero_border_pt": zero,
                    "indent720_border_pt": positive, "indent720_minus_zero_pt": positive-zero,
                    "absent_minus_zero_pt": absent-zero, "absolute_offsets_checked_against_page_and_margins": True}
            modern = sources["reject-modern-indent-0"]["rejected_source_semantics_pdf"]["vertical_border_x_pt"][0]
            legacy = sources["fixed-left-indent-0"]["source_pdf"]["vertical_border_x_pt"][0]
            w.check(modern-legacy > 4.5, "Modern/legacy semantic split no longer demonstrated")
            evidence["pdf_assertions"]["modern_vs_legacy_zero_shift_pt"] = modern-legacy
        if cli:
            # Restore default must truly remove attrs, not write explicit zero/left.
            restored = copy.deepcopy(safe_model)
            for table in w.model_tables(restored):
                table.setdefault("attrs", {}).update({ALIGN: "right", INDENT: 0})
                for key in [ALIGN, INDENT]: table["attrs"].pop(key)
            w.check(restored == safe_model, "Restore-default changed unrelated model content")
            evidence["restored_default"] = exported(cli, evidence, directory, "restored-default", restored, sources["fixed-align-None"]["expected"])
            w.check(evidence["restored_default"]["sha256"] == sources["fixed-align-None"]["unchanged_export"]["sha256"],
                    "Same default model did not restore exact deterministic DOCX bytes")
            evidence["invalid_model_refusals"] = bad_models(cli, evidence, directory, safe_model)
            for item in sources.values():
                path = Path(item["path"])
                error = w.run_cli(cli, evidence["commands"], "create-docx", safe_model_path, path, success=False)
                w.check("Output already exists" in error and w.digest(path) == item["sha256"], "Source no-clobber failed")
                item["source_overwrite_refused"] = True
        if args.corpus_dir:
            evidence["original_corpus"] = corpus_check(args.corpus_dir.resolve(), cli, evidence, directory, safe_model_path)
        evidence["status"] = "fixtures-verified" if cli is None else "passed"
        evidence["summary"] = {"positive_sources": sum(i["expected_action"] == "roundtrip" for i in sources.values()),
                               "refusal_sources": sum(i["expected_action"] == "reject" for i in sources.values()),
                               "invalid_models_refused": len(evidence.get("invalid_model_refusals", [])),
                               "explicit_page_fit_refusals": sum(bool(i.get("page_fit_refusal")) for i in sources.values()),
                               "legacy_absent_page_sources": sum(bool(i.get("legacy_absent_page")) for i in sources.values()),
                               "explicit_page_edit_refusals": sum("explicit_page_edit_refusal" in i for i in sources.values()),
                               "rendered_sources": sum("source_pdf" in i for i in sources.values()),
                               "rejected_sources_rendered_for_calibration": sum("rejected_source_semantics_pdf" in i for i in sources.values()),
                               "rendered_exports": sum("unchanged_pdf" in i for i in sources.values()) + sum("edited_pdf" in i for i in sources.values()),
                               "original_corpus_files": len(evidence.get("original_corpus", {}).get("fixtures", []))}
    except Exception as error:
        evidence.update(status="failed", error=f"{type(error).__name__}: {error}")
        raise
    finally:
        evidence["source_hashes_unchanged_on_exit"] = all(w.digest(Path(i["path"])) == i["sha256"] for i in evidence["sources"].values())
        evidence["script_unchanged_during_run"] = evidence["script_sha256"] == w.digest(Path(__file__))
        evidence["width_helper_unchanged_during_run"] = evidence["width_helper_sha256"] == w.digest(Path(w.__file__))
        evidence["cli_unchanged_during_run"] = cli is None or evidence["cli_sha256"] == w.digest(cli)
        report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
        print(report)


if __name__ == "__main__":
    main()
