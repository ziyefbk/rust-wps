"""Independent DOCX hidden-text regression; never edits input documents.

Produces targeted python-docx fixtures (not real-world corpus). Optionally
renders them using an installed LibreOffice and checks PDF text with pypdf.
The baseline mode proves the old importer drops hidden source text. Each run
uses a new directory and records source hashes, CLI outcomes and reader data.
"""
import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfReader

RISK = "Hidden text cannot be preserved in an editable copy"
VISIBLE = "VisibleProbe"
HIDDEN = "HiddenProbe"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_fixtures(directory):
    result = {}
    for case in ["direct", "inherited", "false-only"]:
        doc = Document()
        p = doc.add_paragraph()
        p.add_run(VISIBLE + " ")
        run = p.add_run(HIDDEN)
        if case == "inherited":
            style = doc.styles.add_style("HiddenCharacter", WD_STYLE_TYPE.CHARACTER)
            style.font.hidden = True
            run.style = style
        else:
            run.font.hidden = case == "direct"
        path = directory / f"{case}.docx"
        doc.save(path)
        read = Document(path)
        assert read.paragraphs[0].text == VISIBLE + " " + HIDDEN
        source_run = read.paragraphs[0].runs[1]
        if case == "inherited":
            assert source_run.font.hidden is None and source_run.style.font.hidden is True
        else:
            assert source_run.font.hidden is (case == "direct")
        result[case] = {"path": str(path), "sha256": digest(path),
                        "independent_text": read.paragraphs[0].text,
                        "direct_hidden": source_run.font.hidden,
                        "style_hidden": source_run.style.font.hidden}
    return result


def run_cli(cli, *args, success=True):
    outcome = subprocess.run([str(cli), *map(str, args)], text=True,
                             capture_output=True, timeout=60, encoding="utf-8")
    assert (outcome.returncode == 0) is success, outcome.stdout + outcome.stderr
    return json.loads(outcome.stdout) if success else outcome.stderr


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli", type=Path)
    parser.add_argument("--root", type=Path, default=Path("tmp/hidden-text"))
    parser.add_argument("--expect-baseline-loss", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--corpus", type=Path, help="Locally downloaded pinned Tika DOCX; see hidden-text-corpus.json")
    args = parser.parse_args()
    if args.expect_baseline_loss and not args.cli:
        parser.error("--expect-baseline-loss requires --cli")
    args.root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.root)).resolve()
    evidence = {"fixture_kind": "generated targeted python-docx packages, not real-world corpus",
                "sources": make_fixtures(directory), "baseline_loss": args.expect_baseline_loss}
    cli = args.cli.resolve() if args.cli else None
    evidence["cli_checked"] = cli is not None
    for case, source in evidence["sources"].items():
        path = Path(source["path"])
        if cli:
            audit = run_cli(cli, "audit-docx", path)
            if case != "false-only" and not args.expect_baseline_loss:
                assert RISK in audit["risks"]
                source["import_error"] = run_cli(cli, "import-docx", path, success=False)
                assert RISK in source["import_error"]
            else:
                imported = run_cli(cli, "import-docx", path)
                serialized = json.dumps(imported["document"])
                assert VISIBLE in serialized
                assert (HIDDEN in serialized) is (case == "false-only")
                model = directory / f"{case}.json"
                model.write_text(json.dumps(imported["document"]), encoding="utf-8")
                output = directory / f"{case}-export.docx"
                run_cli(cli, "create-docx", model, output)
                independent = "\n".join(p.text for p in Document(output).paragraphs)
                assert VISIBLE in independent
                assert (HIDDEN in independent) is (case == "false-only")
                source["export_independent_text"] = independent
                output_hash = digest(output)
                run_cli(cli, "create-docx", model, output, success=False)
                run_cli(cli, "create-docx", model, path, success=False)
                assert digest(output) == output_hash
                source["overwrite_refused"] = True
        if args.render:
            assert shutil.which("soffice"), "LibreOffice is required for --render"
            out = directory / "rendered"
            out.mkdir(exist_ok=True)
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                            str(out), str(path)], check=True, capture_output=True, timeout=90)
            pdf = out / f"{case}.pdf"
            pages = PdfReader(pdf).pages
            rendered = "\n".join(page.extract_text() or "" for page in pages)
            assert VISIBLE in rendered, rendered
            assert (HIDDEN in rendered) is (case == "false-only"), rendered
            source["rendered_pdf"] = str(pdf)
            source["rendered_pages"] = len(pages)
            source["rendered_text"] = rendered
        assert digest(path) == source["sha256"], "Source changed"
    if args.corpus:
        manifest = json.loads(Path(__file__).with_name("hidden-text-corpus.json").read_text(encoding="utf-8"))
        path = args.corpus.resolve()
        assert path.stat().st_size == manifest["bytes"]
        assert digest(path) == manifest["sha256"]
        doc = Document(path)
        hidden = [run.text for p in doc.paragraphs for run in p.runs if run.font.hidden]
        assert hidden == manifest["hidden_runs"], hidden
        corpus = {"path": str(path), "sha256": digest(path), "hidden_runs": hidden,
                  "fixture_kind": "Apache Tika upstream regression fixture; also contains revisions/comments"}
        if cli and not args.expect_baseline_loss:
            assert RISK in run_cli(cli, "audit-docx", path)["risks"]
            # Other unsupported features may be the first refusal. This file
            # proves real-source detection, not isolation of the new guard.
            corpus["import_error"] = run_cli(cli, "import-docx", path, success=False)
        if args.render:
            out = directory / "rendered"
            out.mkdir(exist_ok=True)
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                            str(out), str(path)], check=True, capture_output=True, timeout=90)
            pdf = out / (path.stem + ".pdf")
            pages = PdfReader(pdf).pages
            rendered = "\n".join(page.extract_text() or "" for page in pages)
            assert "Lorem ipsum dolor sit amet" in rendered
            assert all(text.strip() not in rendered for text in hidden)
            corpus.update(rendered_pdf=str(pdf), rendered_pages=len(pages), rendered_text=rendered)
        assert digest(path) == manifest["sha256"]
        evidence["upstream_corpus"] = corpus
    evidence["limitations"] = (
        "python-docx verifies source content and declared hidden flags; LibreOffice/PDF text "
        "only verifies these targeted visible/hidden probes. Not a Word/WPS layout oracle, "
        "native WebView2 test, complete style resolution, or hidden-text editing support."
    )
    report = directory / "report.json"
    report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(report)


if __name__ == "__main__":
    main()
