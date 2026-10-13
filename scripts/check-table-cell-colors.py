"""Independent direct-cell-fill and local-auto-color acceptance.

Fixtures are authored with python-docx/OOXML, not the product writer. The POI
original is independently hash-pinned and never modified. Consumer comparisons
use one fresh LibreOffice process/profile per file. No browser/native/Word/WPS
claim. Every run keeps its own sources, report, command log and script snapshot.
"""
import argparse
import copy
import hashlib
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt
from lxml import etree as ET
import fitz

ROOT = Path(__file__).resolve().parent.parent
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = {'w': W, 'a': A, 'r': R}
Q = lambda name: '{' + W + '}' + name
FILL, AUTO = 'cellFillRgb', 'cellTextColorMode'
NORMALIZED = 'Some table cell fills are unproved and remain normalized'
AUTO_NORMALIZED = 'Some filled-cell automatic text colors are unproved; profile or explicit text colors remain in use'
THEME_FIXED = 'Direct table cell theme fills resolved to fixed RGB; dynamic theme linkage is not preserved'
AUTO_PRESERVED = 'Filled-cell automatic text color is preserved locally; consumers choose contrast and editor preview is an approximation'
BINDING_NORMALIZED = 'Table cell color identity could not be bound to converted content; original color normalization retained'
ORIGINAL_SHA = 'ee36729e0d050e529fea0b444006a194bea72b35cf8958ae58c6624c26156b41'
OLD_SHA = 'ed77eea6ce7d6f26eede7ee33af3655f42e1eeebbc9060ac086bb57c6a6d9430'
ORIGINAL_URL = 'https://github.com/apache/poi/blob/ae62bb5116b9aee19ebd5834e3a82066132c9f7f/test-data/document/table-alignment.docx'


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def child(parent, tag, **attrs):
    item = ET.SubElement(parent, Q(tag))
    for key, value in attrs.items():
        item.set(Q(key), str(value))
    return item


def one(parent, tag):
    item = parent.find('w:' + tag, NS)
    if item is None:
        item = ET.Element(Q(tag)); parent.insert(0, item)
    return item


def remove(parent, tag):
    for item in parent.findall('w:' + tag, NS):
        parent.remove(item)


def set_color(pr, rgb):
    remove(pr, 'color'); child(pr, 'color', val=rgb)


def tree_xml(tree):
    return ET.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)


def parts(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def write_parts(path, content):
    with zipfile.ZipFile(path, 'x', zipfile.ZIP_DEFLATED) as package:
        for name, data in content.items():
            package.writestr(name, data)


def mutate(source, target, callback):
    content = parts(source)
    main, styles, theme, rels = [ET.fromstring(content[name]) for name in
        ['word/document.xml', 'word/styles.xml', 'word/theme/theme1.xml', 'word/_rels/document.xml.rels']]
    callback(main, styles, theme, rels, content)
    for name, tree in [('word/document.xml', main), ('word/styles.xml', styles), ('word/theme/theme1.xml', theme), ('word/_rels/document.xml.rels', rels)]:
        content[name] = tree_xml(tree)
    write_parts(target, content)


def walk(model):
    yield model
    for node in model.get('content', []):
        yield from walk(node)


def cells(model):
    return [n for n in walk(model) if n['type'] in ('tableCell', 'tableHeader')]


def fields(model):
    return [{key: n['attrs'][key] for key in (FILL, AUTO) if key in n.get('attrs', {})} for n in cells(model)]


def stripped(model):
    model = copy.deepcopy(model)
    for n in walk(model):
        for key in (FILL, AUTO):
            n.get('attrs', {}).pop(key, None)
        if n.get('attrs') == {}:
            n.pop('attrs')
    return model


def fresh(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    normal = doc.styles['Normal']
    normal.font.name = 'Liberation Serif'; normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    default = doc.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr', NS)
    remove(default, 'color')
    table = doc.add_table(rows=2, cols=3)
    table.style = 'Table Grid'; table.autofit = False
    width = table._tbl.tblPr.find('w:tblW', NS); width.set(Q('w'), '9000'); width.set(Q('type'), 'dxa')
    for col in table.columns: col.width = Inches(3000 / 1440)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.width = Inches(3000 / 1440)
            cell.paragraphs[0].add_run(f'R{r}C{c} Auto')
            if r == 0: child(cell._tc.get_or_add_tcPr(), 'shd', val='clear', color='auto', fill='E97132')
    doc.add_paragraph('AFTER TABLE')
    check(not path.exists(), 'Refusing to overwrite source')
    doc.save(path)


def fixture_sources(directory):
    sources = {}; base = directory / 'literal-auto-source.docx'; fresh(base)
    fill_auto = [{FILL: 'E97132', AUTO: 'auto'} for _ in range(3)] + [{} for _ in range(3)]
    fill_rgb = [{FILL: 'E97132'} for _ in range(3)] + [{} for _ in range(3)]
    def add(name, callback=None, expected=None, warnings=(), render=False, foreground=None):
        path = base if name == 'literal-auto' else directory / (name + '-source.docx')
        if callback: mutate(base, path, callback)
        sources[name] = {'path': str(path), 'sha256': digest(path), 'kind': 'independently authored directed python-docx/OOXML source',
                         'expected': copy.deepcopy(fill_auto if expected is None else expected), 'warnings': list(warnings), 'render': render}
        if foreground: sources[name]['foreground'] = foreground
        return path
    add('literal-auto', render=True)
    def defaults(m,s,t,r,p,value): set_color(s.find('w:docDefaults/w:rPrDefault/w:rPr',NS),value)
    add('default-black', lambda *x: defaults(*x,'000000'), fill_rgb, render=True, foreground='000000')
    add('default-green', lambda *x: defaults(*x,'008000'), fill_rgb, render=True, foreground='008000')
    add('default-auto', lambda *x: defaults(*x,'auto'), render=True)
    def mixed(m,s,t,r,p):
        for para in m.findall('.//w:tbl/w:tr[1]/w:tc/w:p',NS):
            for text,c in [(' Black','000000'),(' Green','008000')]:
                run=child(para,'r');set_color(child(run,'rPr'),c)
                value=child(run,'t');value.set('{http://www.w3.org/XML/1998/namespace}space','preserve');value.text=text
    add('mixed-auto-explicit', mixed, render=True)
    def style(m,s,t,r,p,kind,base_value,leaf_value,direct=None):
        defaults(m,s,t,r,p,'112233')
        b=child(s,'style',type=kind,styleId='ColorBase');child(b,'name',val='Color base');set_color(child(b,'rPr'),base_value)
        f=child(s,'style',type=kind,styleId='ColorLeaf');child(f,'name',val='Color leaf');child(f,'basedOn',val='ColorBase');set_color(child(f,'rPr'),leaf_value)
        for para in m.findall('.//w:tbl/w:tr[1]/w:tc/w:p',NS):
            if kind=='paragraph':child(one(para,'pPr'),'pStyle',val='ColorLeaf')
            for run in para.findall('w:r',NS):
                if kind=='character':child(one(run,'rPr'),'rStyle',val='ColorLeaf')
                if direct:set_color(one(run,'rPr'),direct)
    add('paragraph-chain-auto',lambda *x: style(*x,'paragraph','008000','auto'),render=True)
    add('paragraph-chain-green',lambda *x: style(*x,'paragraph','auto','008000'),fill_rgb,render=True,foreground='008000')
    add('character-chain-auto',lambda *x: style(*x,'character','008000','auto'),render=True)
    add('character-chain-green',lambda *x: style(*x,'character','auto','008000'),fill_rgb,render=True,foreground='008000')
    add('direct-auto-over-style',lambda *x: style(*x,'paragraph','000000','008000','auto'),render=True)
    add('direct-black-over-auto',lambda *x: style(*x,'paragraph','008000','auto','000000'),fill_rgb,render=True,foreground='000000')
    def themed(m,s,t,r,p,mismatch=False):
        t.find('a:themeElements/a:clrScheme/a:accent2/a:srgbClr',NS).set('val','E97132' if not mismatch else '008000')
        for sh in m.findall('.//w:tcPr/w:shd',NS):sh.set(Q('themeFill'),'accent2')
    add('theme-match',themed,warnings=[THEME_FIXED],render=True)
    def renamed_theme(m,s,t,r,p):
        themed(m,s,t,r,p)
        name='word/meta/actual-colors.xml';p[name]=tree_xml(t)
        rel=next(n for n in r if n.get('Type','').endswith('/theme'));rel.set('Target','meta/actual-colors.xml')
        content_types=ET.fromstring(p['[Content_Types].xml'])
        tag='{http://schemas.openxmlformats.org/package/2006/content-types}Override'
        ET.SubElement(content_types,tag,PartName='/'+name,ContentType='application/vnd.openxmlformats-officedocument.theme+xml')
        p['[Content_Types].xml']=tree_xml(content_types)
    add('theme-renamed',renamed_theme,warnings=[THEME_FIXED],render=True)
    def unrelated_theme(m,s,t,r,p):
        themed(m,s,t,r,p)
        for rel in list(r):
            if rel.get('Type','').endswith('/theme'):r.remove(rel)
    add('theme-not-related',unrelated_theme,[{}]*6,[NORMALIZED])
    add('theme-mismatch',lambda *x:themed(*x,mismatch=True),[{}]*6,[NORMALIZED])
    def mutate_shd(m,s,t,r,p,**attrs):
        for sh in m.findall('.//w:tcPr/w:shd',NS):
            for key,val in attrs.items():
                if val is None:sh.attrib.pop(Q(key),None)
                else:sh.set(Q(key),val)
    for name, attrs in [('pattern-pct20',dict(val='pct20')),('pattern-stripe',dict(val='diagStripe')),('pattern-solid',dict(val='solid')),('pattern-unknown',dict(val='futurePattern')),
                        ('missing-pattern',dict(val=None)),('fill-auto',dict(fill='auto')),('missing-fill',dict(fill=None)),('theme-tint',dict(themeFill='accent2',themeFillTint='80')),
                        ('theme-shade',dict(themeFill='accent2',themeFillShade='BF')),('malformed-rgb',dict(fill='12345'))]:
        add(name,lambda *x,a=attrs:mutate_shd(*x,**a),[{}]*6,[NORMALIZED])
    def style_shading(m,s,t,r,p,conditional=False,unused=False):
        st=child(s,'style',type='table',styleId='ColorTable');child(st,'name',val='Color table')
        if not unused:
            table=m.find('.//w:tbl',NS);pr=one(table,'tblPr');remove(pr,'tblStyle');child(pr,'tblStyle',val='ColorTable')
        target=child(st,'tblStylePr',type='firstRow') if conditional else st
        child(child(target,'tcPr'),'shd',val='clear',fill='123456')
    add('inherited-shading',style_shading,[{}]*6,[NORMALIZED])
    add('conditional-shading',lambda *x:style_shading(*x,conditional=True),[{}]*6,[NORMALIZED])
    add('unused-conditional-shading',lambda *x:style_shading(*x,conditional=True,unused=True),render=False)
    def table_shading(m,s,t,r,p): child(one(m.find('.//w:tbl',NS),'tblPr'),'shd',val='clear',fill='123456')
    add('table-shading',table_shading,[{}]*6,[NORMALIZED])
    def overlays(m,s,t,r,p,where,tag):
        for para in m.findall('.//w:tbl/w:tr[1]/w:tc/w:p',NS):
            target=one(para,'pPr') if where=='paragraph' else one(para.find('w:r',NS),'rPr')
            child(target,tag,**(dict(val='clear',fill='FFFFFF') if tag=='shd' else dict(val='yellow')))
    add('paragraph-shading',lambda *x:overlays(*x,'paragraph','shd'),[{}]*6,[NORMALIZED])
    add('run-shading',lambda *x:overlays(*x,'run','shd'),[{}]*6,[NORMALIZED])
    add('run-highlight',lambda *x:overlays(*x,'run','highlight'),fill_rgb,[AUTO_NORMALIZED])
    def style_background(m,s,t,r,p,tag):
        st=child(s,'style',type='paragraph',styleId='PaintedPara');child(st,'name',val='Painted paragraph');child(st,'basedOn',val='Normal')
        child(child(st,'rPr'),tag,**(dict(val='clear',fill='FFFFFF') if tag=='shd' else dict(val='yellow')))
        for para in m.findall('.//w:tbl/w:tr[1]/w:tc/w:p',NS):child(one(para,'pPr'),'pStyle',val='PaintedPara')
    add('style-run-shading',lambda *x:style_background(*x,'shd'),fill_rgb,[AUTO_NORMALIZED])
    add('style-run-highlight',lambda *x:style_background(*x,'highlight'),fill_rgb,[AUTO_NORMALIZED])
    def table_color(m,s,t,r,p,conditional=False):
        st=s.xpath('./w:style[@w:styleId="TableGrid"]',namespaces=NS)[0]
        target=child(st,'tblStylePr',type='firstRow') if conditional else st
        set_color(child(target,'rPr'),'008000')
    add('table-style-color',table_color,fill_rgb,[AUTO_NORMALIZED])
    add('conditional-table-color',lambda *x:table_color(*x,conditional=True),fill_rgb,[AUTO_NORMALIZED])
    def default_character(m,s,t,r,p):
        st=s.xpath('./w:style[@w:type="character" and @w:default="1"]',namespaces=NS)[0];set_color(one(st,'rPr'),'008000')
    add('default-character-color',default_character,fill_rgb,[AUTO_NORMALIZED])
    def duplicate_color(m,s,t,r,p):
        for run in m.findall('.//w:tbl/w:tr[1]/w:tc/w:p/w:r',NS):
            pr=one(run,'rPr');child(pr,'color',val='000000');child(pr,'color',val='auto')
    add('duplicate-run-color',duplicate_color,fill_rgb,[AUTO_NORMALIZED])
    def wrapped_color(m,s,t,r,p):
        for run in m.findall('.//w:tbl/w:tr[1]/w:tc/w:p/w:r',NS):child(child(one(run,'rPr'),'custom'),'color',val='auto')
    add('wrapped-run-color',wrapped_color,fill_rgb,[AUTO_NORMALIZED])
    partial=copy.deepcopy(fill_auto);partial[0]={}
    add('partial-direct-pattern',lambda m,s,t,r,p:m.find('.//w:tcPr/w:shd',NS).set(Q('val'),'pct20'),partial,[NORMALIZED])
    add('clear-pattern-color-is-not-text',lambda *x:mutate_shd(*x,color='008000'),render=True)
    def numbered_cell(m,s,t,r,p):child(one(m.find('.//w:tbl/w:tr/w:tc/w:p',NS),'pPr'),'pStyle',val='ListNumber')
    add('numbered-cell-binding-fallback',numbered_cell,[{}]*6,[BINDING_NORMALIZED])
    # Authored references deliberately operate on source declarations, never on
    # application output or values learned from an import response.
    def edited(m,s,t,r,p):m.find('.//w:tcPr/w:shd',NS).set(Q('fill'),'156082')
    reference=directory/'literal-auto-edited-reference.docx';mutate(base,reference,edited)
    sources['literal-auto']['reference']={'path':str(reference),'sha256':digest(reference)}
    references={}
    for variant in ('reset-fill','reset-mode','reset-all'):
        def reset_reference(m,s,t,r,p,variant=variant):
            cell=m.find('.//w:tbl/w:tr/w:tc',NS)
            if variant in ('reset-fill','reset-all'):remove(one(cell,'tcPr'),'shd')
            else:cell.find('w:tcPr/w:shd',NS).set(Q('fill'),'156082')
            if variant in ('reset-mode','reset-all'):
                for run in cell.findall('w:p/w:r',NS):set_color(one(run,'rPr'),'000000')
        reference=directory/('literal-auto-'+variant+'-reference.docx');mutate(base,reference,reset_reference)
        references[variant]={'path':str(reference),'sha256':digest(reference)}
    sources['literal-auto']['reset_references']=references
    return sources


def run(cli, evidence, *args, success=True):
    command=[str(cli),*[str(x) for x in args]]
    result=subprocess.run(command,capture_output=True,text=True,timeout=60)
    evidence['commands'].append({'command':command,'returncode':result.returncode,'stderr':result.stderr,'stdout':result.stdout})
    check((result.returncode==0)==success,'Unexpected CLI outcome: '+' '.join(command)+'\n'+result.stderr)
    return json.loads(result.stdout) if success else result.stderr


def independent(path):
    doc=Document(path);content=parts(path);main=ET.fromstring(content['word/document.xml'])
    xmlcells=main.findall('.//w:tbl/w:tr/w:tc',NS)
    typed=[c for t in doc.tables for r in t.rows for c in r.cells]
    check(len(xmlcells)==len(typed),'Typed/XML cell counts differ')
    out=[]
    for cell,typed_cell in zip(xmlcells,typed):
        sh=cell.find('w:tcPr/w:shd',NS)
        runs=[]
        for p in cell.findall('w:p',NS):
            for r in p.findall('w:r',NS):
                text=''.join(r.xpath('.//w:t/text()',namespaces=NS));color=r.find('w:rPr/w:color',NS)
                runs.append({'text':text,'color':None if color is None else color.get(Q('val'))})
        out.append({'paragraphs':[p.text for p in typed_cell.paragraphs], 'fill':None if sh is None else {ET.QName(k).localname:v for k,v in sh.attrib.items()},'runs':runs})
    return out


def verify_export(path,model):
    actual=independent(path);expected=cells(model);check(len(actual)==len(expected),'Export cell count changed')
    for a,c in zip(actual,expected):
        attrs=c.get('attrs',{});want=attrs.get(FILL)
        if want is None and c['type']=='tableHeader':want='E8EEF7'
        check((a['fill'] or {}).get('fill')==want,'Export fill differs from model')
        if FILL in attrs:check(a['fill']==dict(val='clear',color='auto',fill=want),'Export is not explicit clear RGB')
        wanted=[]
        for p in c.get('content',[]):
            for n in walk(p):
                if n['type']!='text':continue
                marks=n.get('marks',[]);rgb=next((m['attrs']['rgb'] for m in marks if m['type']=='textColor'),None)
                if rgb is None and any(m['type']=='link' for m in marks):rgb='0563C1'
                if rgb is None and attrs.get(AUTO)=='auto':rgb='auto'
                wanted.append({'text':n.get('text',''),'color':rgb})
        check(a['runs']==wanted,'Export run colors or exact text differ: '+repr((a['runs'],wanted)))
    return actual


def export(cli,evidence,directory,name,model,source=None):
    model_path=directory/(name+'.json');out=directory/(name+'.docx')
    with model_path.open('x') as f:json.dump(model,f,ensure_ascii=False,indent=2)
    run(cli,evidence,'create-docx',model_path,out)
    check(run(cli,evidence,'read-docx-model',out)==model,'Trusted model differs')
    snap=verify_export(out,model)
    for target in [out]+([Path(source)] if source else []):
        before=digest(target);error=run(cli,evidence,'create-docx',model_path,target,success=False)
        check('already exists' in error and digest(target)==before,'No-clobber failed')
    return {'path':str(out),'sha256':digest(out),'model_path':str(model_path),'trusted_model_equal':True,'no_clobber':True,'independent':snap}


def invalid_models(cli,evidence,directory,base):
    result={}
    for field,values in [(FILL,[None,'',False,0,[],{},'12345','1234567','abcdef','#AABBCC','GG0000','auto']), (AUTO,[None,'',False,True,0,[],{},'AUTO','black',' auto'])]:
        for i,value in enumerate(values):
            model=copy.deepcopy(base);cells(model)[0].setdefault('attrs',{})[field]=value
            name=f'invalid-{field}-{i}';mp=directory/(name+'.json');out=directory/(name+'.docx');mp.write_text(json.dumps(model))
            error=run(cli,evidence,'create-docx',mp,out,success=False);check(not out.exists(),'Invalid model created output')
            result[name]={'error':error,'output_absent':True}
    for name,index in [('doc',None),('table',0),('row',1),('paragraph',3),('text',4)]:
        for field,value in [(FILL,'123456'),(AUTO,'auto')]:
            model=copy.deepcopy(base);target=model if index is None else list(walk(model))[index+1]
            target.setdefault('attrs',{})[field]=value
            tag='invalid-placement-'+name+'-'+field;mp=directory/(tag+'.json');out=directory/(tag+'.docx');mp.write_text(json.dumps(model))
            error=run(cli,evidence,'create-docx',mp,out,success=False);check(not out.exists(),'Misplaced model created output')
            result[tag]={'error':error,'output_absent':True}
    return result


def tamper(cli,evidence,directory,source):
    result={}
    for name,change in [('fill',lambda m:m.find('.//w:tcPr/w:shd',NS).set(Q('fill'),'ABCDEF')),
                        ('auto',lambda m:m.find('.//w:rPr/w:color',NS).set(Q('val'),'000000'))]:
        original=parts(source);main=ET.fromstring(original['word/document.xml']);change(main);original['word/document.xml']=tree_xml(main)
        for repaired in (False,True):
            modified=copy.deepcopy(original)
            if repaired:
                embedded=json.loads(modified['rust-wps/model-v2.json']);embedded['hashes']['word/document.xml']=hashlib.sha256(modified['word/document.xml']).hexdigest()
                modified['rust-wps/model-v2.json']=json.dumps(embedded,separators=(',',':')).encode()
            tag='tamper-'+name+('-repaired' if repaired else '');path=directory/(tag+'.docx');write_parts(path,modified)
            result[tag]={'error':run(cli,evidence,'read-docx-model',path,success=False),'rejected':True,'sha256':digest(path)}
    return result


def render(soffice,directory,paths,evidence):
    out=directory/'pdf';out.mkdir();result={}
    for i,path in enumerate(dict.fromkeys(str(Path(p).resolve()) for p in paths)):
        path=Path(path);profile=directory/f'lo-profile-{i:03d}'
        command=[soffice,f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','pdf','--outdir',str(out),str(path)]
        ran=subprocess.run(command,capture_output=True,text=True,timeout=90,env=dict(os.environ,XDG_CACHE_HOME=str(directory/f'cache-{i:03d}')))
        evidence['commands'].append({'command':command,'returncode':ran.returncode,'stdout':ran.stdout,'stderr':ran.stderr})
        check(ran.returncode==0,'Independent renderer failed: '+ran.stderr)
        pdf=out/(path.stem+'.pdf');doc=fitz.open(pdf);spans=[];fills=[]
        for pi,page in enumerate(doc):
            for b in page.get_text('dict')['blocks']:
                for line in b.get('lines',[]):
                    for s in line['spans']:spans.append(dict(page=pi,text=s['text'],color=f'{s["color"]:06X}',origin=list(s['origin']),font=s['font'],size=s['size']))
            for item in page.get_drawings():
                if item.get('fill') is not None:fills.append(dict(page=pi,rgb=''.join(f'{round(v*255):02X}' for v in item['fill']),rect=list(item['rect'])))
        result[str(path)]={'path':str(pdf),'sha256':digest(pdf),'pages':len(doc),'page_sizes':[[p.rect.width,p.rect.height] for p in doc],'spans':spans,'fills':fills}
    return result


def compare(a,b,strict_geometry=True):
    check(a['page_sizes']==b['page_sizes'],'Consumer pages differ')
    check([(s['text'],s['color']) for s in a['spans']]==[(s['text'],s['color']) for s in b['spans']],'Consumer text/color differs')
    dx=max((abs(x['origin'][0]-y['origin'][0]) for x,y in zip(a['spans'],b['spans'])),default=0)
    dy=max((abs(x['origin'][1]-y['origin'][1]) for x,y in zip(a['spans'],b['spans'])),default=0)
    check(dy<=.3,'Consumer vertical baseline differs')
    if strict_geometry:check(dx<=.6,'Consumer horizontal baseline differs')
    check(Counter(x['rgb'] for x in a['fills'])==Counter(x['rgb'] for x in b['fills']),'Consumer fill colors/counts differ')
    return {'text_and_color_equal':True,'fill_rgb_counts_equal':True,'max_x_delta_pt':dx,'max_y_delta_pt':dy,'horizontal_fidelity_claimed':strict_geometry}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cli',type=Path);parser.add_argument('--cli-sha256')
    parser.add_argument('--baseline-cli',type=Path,default=ROOT/'tmp/empty-font-final-cli')
    parser.add_argument('--root',type=Path,default=ROOT/'tmp/table-cell-colors-independent')
    parser.add_argument('--original',type=Path,default=ROOT/'tmp/table-corpus/table-alignment.docx')
    parser.add_argument('--render',action='store_true');parser.add_argument('--generate-only',action='store_true')
    parser.add_argument('--ui-driver',type=Path,default=ROOT/'tests/ui/edit-table-cell-colors.mjs')
    args=parser.parse_args();check(args.generate_only or args.cli,'Pass --cli or --generate-only')
    args.root.mkdir(parents=True,exist_ok=True);directory=Path(tempfile.mkdtemp(prefix='run-',dir=args.root)).resolve()
    shutil.copyfile(__file__,directory/Path(__file__).name)
    evidence={'status':'running','script_sha256':digest(__file__),'python':platform.python_version(),'commands':[],'sources':{},'cases':{},'limitations':['No browser/native print or Microsoft Word/WPS verification','Original horizontal widths remain normalized','Controlled sources and model edits are deliberately authored validation cases','Each PDF uses a fresh process and unique profile','Automatic foreground remains consumer-selected; no hardcoded white RGB']}
    report=directory/'report.json';paths=[]
    try:
        evidence['sources']=fixture_sources(directory)
        check(digest(args.original)==ORIGINAL_SHA,'Pinned original differs')
        expected=[]
        for _ in range(6):expected.extend([{FILL:'E97132',AUTO:'auto'} for _ in range(3)]+[{} for _ in range(6)])
        evidence['sources']['original-table-alignment']={'path':str(args.original.resolve()),'sha256':ORIGINAL_SHA,'bytes':15519,'git_blob_sha1':'76fd85c338a1515c2ba2736953ba8424a9bc9367','source_url':ORIGINAL_URL,'kind':'untouched pinned upstream Apache POI regression fixture; not a user production file','expected':expected,'warnings':[THEME_FIXED,AUTO_PRESERVED],'render':True}
        if not args.generate_only:
            cli=args.cli.resolve();check(cli.exists(),'Frozen CLI missing')
            if args.cli_sha256:check(digest(cli)==args.cli_sha256,'Frozen CLI identity differs')
            check(digest(args.baseline_cli)==OLD_SHA,'Historical CLI identity differs')
            evidence['cli']={'path':str(cli),'sha256':digest(cli)};evidence['baseline_cli']={'path':str(args.baseline_cli),'sha256':OLD_SHA}
            for name,info in evidence['sources'].items():
                source=Path(info['path']);imported=run(cli,evidence,'import-docx',source);model=imported['document']
                check(fields(model)==info['expected'],name+': unexpected preserved fields '+repr(fields(model)))
                if info.get('foreground'):
                    for cell in cells(model)[:3]:
                        colors=[next((m['attrs']['rgb'] for m in n.get('marks',[]) if m['type']=='textColor'),None) for n in walk(cell) if n['type']=='text']
                        check(colors and all(c==info['foreground'] for c in colors),name+': source-declared foreground differs')
                for warning in info['warnings']:check(warning in imported['compatibility']['risks'],name+': missing warning '+warning)
                if any(AUTO in f for f in info['expected']):check(AUTO_PRESERVED in imported['compatibility']['risks'],name+': missing automatic disclosure')
                old=run(args.baseline_cli,evidence,'import-docx',source)
                check(stripped(model)==old['document'],name+': unrelated imported model change')
                exported=export(cli,evidence,directory,name+'-export',model,source)
                reset=stripped(model);reset_export=export(cli,evidence,directory,name+'-reset',reset)
                old_path=directory/(name+'-historical.docx');run(args.baseline_cli,evidence,'create-docx',reset_export['model_path'],old_path)
                check(digest(old_path)==reset_export['sha256'],name+': historical whole-package bytes changed')
                current_parts,reset_parts=parts(exported['path']),parts(reset_export['path'])
                check(set(current_parts)==set(reset_parts),name+': package inventory changed')
                check(all(current_parts[p]==reset_parts[p] for p in current_parts if p not in ('word/document.xml','rust-wps/model-v2.json')),name+': global package parts changed')
                evidence['cases'][name]={'compatibility':imported['compatibility'],'new':exported,'reset':reset_export,'historical_sha256':digest(old_path),'historical_bytes_equal':True,'global_parts_unchanged':True,'only_new_model_fields_changed':True}
                if info['render']:paths.extend([source,Path(exported['path'])])
                report.write_text(json.dumps(evidence,ensure_ascii=False,indent=2))
            base=json.loads(Path(evidence['cases']['literal-auto']['new']['model_path']).read_text())
            edited=copy.deepcopy(base);cells(edited)[0]['attrs'][FILL]='156082'
            edit=export(cli,evidence,directory,'literal-auto-edited-export',edited)
            evidence['model_edit']={'method':'explicit JSON first-cell-only edit','new':edit,'reference':evidence['sources']['literal-auto']['reference']}
            check(stripped(edited)==stripped(base),'Edit changed unrelated model')
            paths.extend([Path(edit['path']),Path(evidence['sources']['literal-auto']['reference']['path'])])
            if args.ui_driver.exists():
                driver_dir=directory/'ui-driver';command=['node',str(args.ui_driver.resolve()),evidence['cases']['literal-auto']['new']['model_path'],str(driver_dir),'0','0',json.dumps({FILL:'156082',AUTO:'auto'})]
                result=subprocess.run(command,cwd=ROOT,capture_output=True,text=True,timeout=60)
                evidence['commands'].append({'command':command,'returncode':result.returncode,'stdout':result.stdout,'stderr':result.stderr})
                check(result.returncode==0,'Actual UI candidate driver failed: '+result.stderr)
                evidence['ui_driver']={'path':str(args.ui_driver.resolve()),'sha256':digest(args.ui_driver),'metadata':json.loads(result.stdout),'exports':{}}
                driver_before=json.loads((driver_dir/'before.json').read_text())
                for variant in ('before','edited','undo','redo','reset-fill','reset-mode','reset-all'):
                    value=json.loads((driver_dir/(variant+'.json')).read_text())
                    if variant in ('before','undo'):check(value==driver_before,'UI undo changed before document')
                    evidence['ui_driver']['exports'][variant]=export(cli,evidence,directory,'ui-'+variant+'-export',value)
                    if variant in ('edited','reset-fill','reset-mode','reset-all'):
                        ref=evidence['sources']['literal-auto']['reference'] if variant=='edited' else evidence['sources']['literal-auto']['reset_references'][variant]
                        paths.extend([Path(ref['path']),Path(evidence['ui_driver']['exports'][variant]['path'])])
                check(json.loads((driver_dir/'edited.json').read_text())==json.loads((driver_dir/'redo.json').read_text()),'UI redo changed edited document')
                protected={str(p):digest(p) for p in driver_dir.glob('*.json')}
                retry=subprocess.run(command,cwd=ROOT,capture_output=True,text=True,timeout=60)
                evidence['commands'].append({'command':command,'returncode':retry.returncode,'stdout':retry.stdout,'stderr':retry.stderr})
                check(retry.returncode!=0 and 'EEXIST' in retry.stderr and all(digest(p)==h for p,h in protected.items()),'UI evidence helper overwrote files')
                evidence['ui_driver']['no_clobber']=True
                # Repeat the actual edit on the genuine source-derived model.
                # Its independent reference preserves original horizontal
                # geometry, so that difference remains separately measured.
                original_driver=directory/'ui-original-driver'
                original_command=['node',str(args.ui_driver.resolve()),evidence['cases']['original-table-alignment']['new']['model_path'],str(original_driver),'0','0',json.dumps({FILL:'156082',AUTO:'auto'})]
                result=subprocess.run(original_command,cwd=ROOT,capture_output=True,text=True,timeout=60)
                evidence['commands'].append({'command':original_command,'returncode':result.returncode,'stdout':result.stdout,'stderr':result.stderr})
                check(result.returncode==0,'Original UI candidate driver failed: '+result.stderr)
                evidence['original_ui_driver']={'metadata':json.loads(result.stdout),'exports':{}}
                for variant in ('before','edited','undo','redo','reset-fill','reset-mode','reset-all'):
                    value=json.loads((original_driver/(variant+'.json')).read_text())
                    evidence['original_ui_driver']['exports'][variant]=export(cli,evidence,directory,'ui-original-'+variant+'-export',value)
                def original_edit(m,s,t,r,p):
                    sh=m.find('.//w:tcPr/w:shd',NS);sh.set(Q('fill'),'156082');sh.attrib.pop(Q('themeFill'),None)
                original_ref=directory/'original-edited-independent-reference.docx';mutate(args.original,original_ref,original_edit)
                evidence['original_ui_driver']['reference']={'path':str(original_ref),'sha256':digest(original_ref),'kind':'deliberate independent derivative of untouched original; first cell fill only'}
                paths.extend([original_ref,Path(evidence['original_ui_driver']['exports']['edited']['path'])])
            evidence['invalid_models']=invalid_models(cli,evidence,directory,stripped(base))
            evidence['tamper']=tamper(cli,evidence,directory,Path(evidence['cases']['literal-auto']['new']['path']))
            if args.render:
                soffice=shutil.which('soffice');check(soffice,'No independent renderer')
                evidence['soffice_version']=subprocess.run([soffice,'--version'],capture_output=True,text=True).stdout.strip()
                evidence['renders']=render(soffice,directory,paths,evidence)
                comparisons={}
                for name,info in evidence['sources'].items():
                    if info['render']:comparisons[name]=compare(evidence['renders'][info['path']],evidence['renders'][evidence['cases'][name]['new']['path']],name!='original-table-alignment')
                comparisons['edited']=compare(evidence['renders'][evidence['model_edit']['reference']['path']],evidence['renders'][edit['path']])
                if 'ui_driver' in evidence:
                    for variant in ('edited','reset-fill','reset-mode','reset-all'):
                        ref=evidence['sources']['literal-auto']['reference'] if variant=='edited' else evidence['sources']['literal-auto']['reset_references'][variant]
                        comparisons['ui-'+variant]=compare(evidence['renders'][ref['path']],evidence['renders'][evidence['ui_driver']['exports'][variant]['path']])
                    comparisons['ui-original-edited']=compare(evidence['renders'][evidence['original_ui_driver']['reference']['path']],evidence['renders'][evidence['original_ui_driver']['exports']['edited']['path']],False)
                evidence['comparisons']=comparisons
            check(digest(cli)==evidence['cli']['sha256'],'Frozen CLI changed during run')
        evidence['source_hashes_unchanged']=all(digest(s['path'])==s['sha256'] for s in evidence['sources'].values())
        check(evidence['source_hashes_unchanged'],'Original/source file changed')
        evidence['status']='passed'
    except Exception as error:
        evidence['status']='failed';evidence['error']=str(error);raise
    finally:
        report.write_text(json.dumps(evidence,ensure_ascii=False,indent=2));print(report,flush=True)


if __name__=='__main__':main()
