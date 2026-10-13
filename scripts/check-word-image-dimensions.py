"""Independent bounded inline-image-size regression; no input is overwritten.

Synthetic PNG/JPEG fixtures are controlled python-docx documents. The separate
original fixture manifest pins an unchanged upstream document; no binaries are
committed. Both extent declarations, image MIME/bytes, trusted model equality,
resizing/reset, refusals and source/output no-clobber are checked independently.
LibreOffice PDF is an external consumer, not app/browser/Word/WPS validation.

  python scripts/check-word-image-dimensions.py --generate-only
  python scripts/check-word-image-dimensions.py --cli target/debug/rust-wps --render --corpus-dir tmp/image-corpus
"""
import argparse
import base64
import copy
import hashlib
import importlib.util
import io
import json
import math
import os
import platform
import posixpath
import shutil
import subprocess
import tempfile
import zipfile
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Emu
from lxml import etree as ET
from PIL import Image, ImageDraw
from pypdf import PdfReader

_spec = importlib.util.spec_from_file_location("width_checks", Path(__file__).with_name("check-word-table-widths.py"))
w = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(w)
NS = {**w.NS, "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
      "ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
WIDTH, HEIGHT = "displayWidthEmu", "displayHeightEmu"
MAX_WIDTH, MAX_HEIGHT = 5715000, 7620000
SOURCE_RISK = "Unsupported image geometry cannot be preserved in an editable copy"
PLACEMENT_WARNING = "Image-only paragraph spacing and alignment preserved; other page and paragraph layout may be normalized"
IMAGE_PARAGRAPH = {"spaceBefore": 0, "spaceAfter": 0, "lineSpacing": 240, "textAlign": "left",
                   "paragraphFont": {"family": "DejaVu Sans", "halfPoints": 24, "bold": False, "italic": False},
                   "runFont": {"family": "DejaVu Sans", "halfPoints": 24, "bold": False, "italic": False}}
check, digest, run_cli = w.check, w.digest, w.run_cli


def element(prefix, local, **attrs):
    n = ET.Element("{" + NS[prefix] + "}" + local)
    for key, value in attrs.items():
        n.set(key, str(value))
    return n


def images(node):
    if node["type"] == "image":
        yield node
    for child in node.get("content", []):
        yield from images(child)


def snapshot(path):
    """Read actual XML relations/content types independently from the app."""
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        rels = {r.get("Id"): dict(r.attrib) for r in ET.fromstring(z.read("word/_rels/document.xml.rels"))}
        types = ET.fromstring(z.read("[Content_Types].xml"))
        defaults = {n.get("Extension"): n.get("ContentType") for n in types.findall("ct:Default", NS)}
        overrides = {n.get("PartName"): n.get("ContentType") for n in types.findall("ct:Override", NS)}
        result = {"images": [], "anchor_count": len(root.findall(".//wp:anchor", NS)),
                  "body_paragraph_texts": ["".join(n.itertext()) for n in root.findall("w:body/w:p", NS)]}
        for inline in root.findall(".//wp:inline", NS):
            extent = inline.find("wp:extent", NS)
            inner = inline.find("a:graphic/a:graphicData/pic:pic/pic:spPr/a:xfrm/a:ext", NS)
            blip = inline.find(".//a:blip", NS)
            item = {"wp_extent": None if extent is None else dict(extent.attrib),
                    "a_extent": None if inner is None else dict(inner.attrib)}
            if blip is not None:
                rel = rels.get(blip.get("{" + NS["r"] + "}embed"))
                item["relationship"] = rel
                if rel and rel.get("TargetMode") != "External":
                    target = rel["Target"]
                    part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("word/" + target)
                    if part in z.namelist():
                        data = z.read(part)
                        item.update(media_part=part, bytes=len(data), sha256=hashlib.sha256(data).hexdigest(),
                                    mime=overrides.get("/" + part, defaults.get(part.rsplit(".", 1)[-1])))
                        try:
                            decoded = Image.open(io.BytesIO(data))
                            item.update(pixel_width=decoded.width, pixel_height=decoded.height)
                        except Exception:
                            pass
            result["images"].append(item)
    try:
        result["python_docx_dimensions"] = [[int(s.width), int(s.height)] for s in Document(path).inline_shapes]
    except Exception as exc:
        result["python_docx_error"] = str(exc)
    return result


def verify_snapshot(actual, expected, label):
    check(actual["anchor_count"] == 0, label + ": unexpected floating image")
    check(len(actual["images"]) == len(expected), label + ": image count differs")
    for i, (got, want) in enumerate(zip(actual["images"], expected)):
        pair = {"cx": str(want["width_emu"]), "cy": str(want["height_emu"])}
        check(got["wp_extent"] == got["a_extent"] == pair, f"{label} image{i}: inconsistent extent: {got}")
        for key in ["sha256", "mime", "bytes", "pixel_width", "pixel_height"]:
            check(got.get(key) == want.get(key), f"{label} image{i}: {key} changed")
    check(actual.get("python_docx_dimensions") == [[x["width_emu"], x["height_emu"]] for x in expected],
          label + ": python-docx dimensions disagree")


def verify_model(model, expected, label, dimensions=True):
    actual = list(images(model))
    check(len(actual) == len(expected), label + ": model image count differs")
    for i, (got, want) in enumerate(zip(actual, expected)):
        attrs = got["attrs"]
        if dimensions:
            check([attrs.get(WIDTH), attrs.get(HEIGHT)] == [want["width_emu"], want["height_emu"]],
                  f"{label} image{i}: model dimensions changed")
        else:
            check(WIDTH not in attrs and HEIGHT not in attrs, label + ": reset did not remove pair")
        head, encoded = attrs["src"].split(",", 1)
        data = base64.b64decode(encoded, validate=True)
        check(head == "data:" + want["mime"] + ";base64", label + ": model MIME changed")
        check(len(data) == want["bytes"] and hashlib.sha256(data).hexdigest() == want["sha256"],
              label + ": embedded source image bytes changed")


def expected_image(data, mime, width, height):
    im = Image.open(io.BytesIO(data))
    return {"width_emu": width, "height_emu": height, "mime": mime, "bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(), "pixel_width": im.width, "pixel_height": im.height}


def unsafe_metadata_blobs():
    im = Image.new("RGB", (120, 80), "#2673ae")
    result = {}
    for orientation in [1, 6]:
        exif = Image.Exif();exif[274] = orientation
        out = io.BytesIO();im.save(out, "JPEG", exif=exif)
        result["jpeg-exif" + str(orientation)] = (out.getvalue(), "image/jpeg")
    out = io.BytesIO();im.save(out, "PNG", exif=exif)
    result["png-exif"] = (out.getvalue(), "image/png")
    out = io.BytesIO();im.save(out, "PNG", save_all=True,
                              append_images=[Image.new("RGB", (120, 80), "#ad3326")], duration=200, loop=0)
    result["png-animated"] = (out.getvalue(), "image/png")
    return result


def make_fixtures(directory):
    raster = Image.new("RGB", (120, 80), "#176eab")
    draw = ImageDraw.Draw(raster)
    draw.rectangle((0, 0, 119, 79), outline="black", width=3)
    draw.rectangle((8, 8, 45, 35), fill="#e84b45")
    data = {}
    for fmt, mime in [("PNG", "image/png"), ("JPEG", "image/jpeg")]:
        out = io.BytesIO()
        raster.save(out, fmt, **({"quality": 91} if fmt == "JPEG" else {}))
        data[fmt] = (out.getvalue(), mime)
    data.update(unsafe_metadata_blobs())
    sources = {}

    def create(name, specs=None, mutation=None, reject=None, render=False, package_mutation=None):
        specs = specs or [("PNG", 1270000, 635000)]
        doc = w.fresh_document()
        if render:
            doc.add_paragraph("BeforeImage")
        expected = []
        for fmt, width, height in specs:
            blob, mime = data[fmt]
            shape = doc.add_paragraph().add_run().add_picture(io.BytesIO(blob), width=Emu(width), height=Emu(height))
            # python-docx omits these; LibreOffice supplies a 9 pt left gap for
            # absent distL. Explicit zeros isolate the displayed rectangle.
            for side in ["distT", "distB", "distL", "distR"]:
                shape._inline.set(side, "0")
            expected.append(expected_image(blob, mime, width, height))
        if render:
            doc.add_paragraph("AfterImage")
        if mutation:
            mutation(doc)
        path = directory / (name + ".docx")
        check(not path.exists(), "Fixture destination already exists")
        doc.save(path)
        if package_mutation:
            with zipfile.ZipFile(path) as z:
                parts = {n: z.read(n) for n in z.namelist()}
            package_mutation(parts)
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
                for n, b in parts.items(): z.writestr(n, b)
        item = {"path": str(path), "sha256": digest(path), "bytes": path.stat().st_size,
                "fixture_kind": "Directed synthetic python-docx/Pillow fixture, not an original upstream or user document",
                "expected_action": "reject" if reject else "roundtrip", "unsupported_form": reject,
                "expected": None if reject else expected, "render_probe": render,
                "source_independent": snapshot(path)}
        if not reject:
            verify_snapshot(item["source_independent"], expected, name + " source")
        sources[name] = item

    for fmt in ["PNG", "JPEG"]:
        create(fmt.lower() + "-landscape", [(fmt, 1270000, 635000)], render=True)
        create(fmt.lower() + "-stretched", [(fmt, 762000, 1524000)], render=True)
        create(fmt.lower() + "-nonpixel", [(fmt, 1000001, 700003)], render=True)
    create("min-positive", [("PNG", 1, 1)])
    create("max-bounds", [("JPEG", MAX_WIDTH, MAX_HEIGHT)])
    create("repeated-bytes", [("PNG", 914400, 609600), ("PNG", 1828800, 304800)], render=True)
    create("mixed-types", [("PNG", 914400, 609600), ("JPEG", 1371600, 914400)], render=True)
    def preserved_empty_paragraphs(doc):
        first, second = doc.paragraphs
        first._p.addprevious(OxmlElement("w:p"))
        first._p.addnext(OxmlElement("w:p"))
        second._p.addnext(OxmlElement("w:p"))
    create("empty-paragraph-order", [("PNG", 914400, 609600), ("PNG", 1828800, 304800)],
           mutation=preserved_empty_paragraphs)
    sources["empty-paragraph-order"]["expected_root_order"] = ["paragraph", "image", "paragraph", "image", "paragraph"]

    def inline(doc): return doc.element.find(".//wp:inline", NS)
    def first(doc, query): return inline(doc).find(query, NS)
    def append_zero(doc):
        inline(doc).insert(1, element("wp", "effectExtent", l=0, t=0, r=0, b=0))
        first(doc, ".//pic:blipFill").insert(1, element("a", "srcRect", l=0, t=0, r=0, b=0))
    create("explicit-zero-crop-effects", mutation=append_zero)
    create("empty-effects-list", mutation=lambda d: first(d, ".//pic:spPr").append(element("a", "effectLst")))

    def both_extents(doc, key, value):
        for path in ["wp:extent", ".//pic:spPr/a:xfrm/a:ext"]:
            n = first(doc, path)
            if value is None: del n.attrib[key]
            else: n.set(key, str(value))
    for key, values in [("cx", [0, -1, "1.5", None, MAX_WIDTH + 1]), ("cy", [0, -1, "1.5", None, MAX_HEIGHT + 1])]:
        for value in values:
            create(f"reject-{key}-{value}", mutation=lambda d, k=key, v=value: both_extents(d, k, v), reject="invalid or missing positive bounded integer extent")
    create("reject-mismatched-inner", mutation=lambda d: first(d, ".//a:xfrm/a:ext").set("cx", "1270001"), reject="inner and outer extent mismatch")
    create("reject-missing-extent", mutation=lambda d: inline(d).remove(first(d, "wp:extent")), reject="missing wp extent")
    create("reject-duplicate-extent", mutation=lambda d: inline(d).append(copy.deepcopy(first(d, "wp:extent"))), reject="duplicate wp extent")
    create("reject-duplicate-inner", mutation=lambda d: first(d, ".//a:xfrm").append(copy.deepcopy(first(d, ".//a:xfrm/a:ext"))), reject="duplicate inner extent")
    def wrapped(doc):
        n = first(doc, "wp:extent"); inline(doc).remove(n)
        p = element("wp", "unknown");p.append(n);inline(doc).append(p)
    create("reject-wrapped-extent", mutation=wrapped, reject="extent hidden by unknown wrapper")
    create("reject-foreign-extent", mutation=lambda d: setattr(first(d, "wp:extent"), "tag", "{urn:untrusted}extent"), reject="foreign namespace extent")
    create("reject-extent-extra-attribute", mutation=lambda d: first(d, "wp:extent").set("mystery", "1"), reject="unknown extent attribute")
    for attr, value in [("rot", 5400000), ("flipH", 1), ("flipV", 1)]:
        create("reject-transform-" + attr, mutation=lambda d, a=attr, v=value: first(d, ".//a:xfrm").set(a, str(v)), reject="rotation or flip")
    create("reject-nonzero-offset", mutation=lambda d: first(d, ".//a:xfrm/a:off").set("x", "12700"), reject="nonzero picture offset")
    create("reject-crop", mutation=lambda d: first(d, ".//pic:blipFill").insert(1, element("a", "srcRect", l=1000)), reject="nonzero crop")
    create("reject-effect-extent", mutation=lambda d: inline(d).insert(1, element("wp", "effectExtent", l=0, t=0, r=12700, b=0)), reject="nonzero effects")
    def shadow_effect(doc):
        n = element("a", "effectLst");n.append(element("a", "outerShdw", blurRad=12700))
        first(doc, ".//pic:spPr").append(n)
    create("reject-effects-list", mutation=shadow_effect, reject="unsupported nonempty effects declaration")
    create("reject-round-rect", mutation=lambda d: first(d, ".//a:prstGeom").set("prst", "roundRect"), reject="nonrectangular shape")
    create("reject-inline-distance", mutation=lambda d: inline(d).set("distL", "12700"), reject="nonzero inline distance")
    create("reject-floating-anchor", mutation=lambda d: setattr(inline(d), "tag", "{" + NS["wp"] + "}anchor"), reject="floating image")
    create("reject-unknown-extension", mutation=lambda d: first(d, ".//a:blip").append(element("a", "extLst")), reject="unknown extension")
    create("reject-mixed-text", mutation=lambda d: d.paragraphs[0].add_run("SiblingText"), reject="image and text in same paragraph")
    create("reject-mixed-tab", mutation=lambda d: d.paragraphs[0].add_run().add_tab(), reject="image and tab in same paragraph")
    create("reject-mixed-break", mutation=lambda d: d.paragraphs[0].add_run().add_break(), reject="image and break in same paragraph")
    create("reject-multiple-pictures", mutation=lambda d: d.paragraphs[0].add_run().add_picture(io.BytesIO(data["PNG"][0]), width=Emu(914400)), reject="multiple pictures in same paragraph")
    def list_context(doc):
        p = doc.paragraphs[0]._p.get_or_add_pPr();num = OxmlElement("w:numPr")
        for tag, val in [("ilvl", 0), ("numId", 1)]:
            x = OxmlElement("w:" + tag);x.set(w.qn("w:val"), str(val));num.append(x)
        p.append(num)
    create("reject-list-image", mutation=list_context, reject="image in list paragraph")
    def cell_context(doc):
        p = doc.paragraphs[0]._p;doc.add_table(rows=1, cols=1).cell(0, 0)._tc.append(p)
    create("reject-table-image", mutation=cell_context, reject="image in table cell")
    def hyperlink_context(doc):
        p = doc.paragraphs[0]._p;run = p.find("w:r", NS);p.remove(run)
        h = OxmlElement("w:hyperlink");h.set(w.qn("w:anchor"), "bookmark");h.append(run);p.append(h)
    create("reject-hyperlink-image", mutation=hyperlink_context, reject="image in hyperlink")
    def unsupported_text_layout(doc, tag, attrs, inherited):
        if tag == "textDirection":
            parent = (doc.styles["Normal"].element if inherited else doc.paragraphs[0]._p).get_or_add_pPr()
        else:
            parent = (doc.styles["Normal"].element if inherited else doc.paragraphs[0].runs[0]._r).get_or_add_rPr()
        node = OxmlElement("w:" + tag)
        for key, value in attrs.items():node.set(w.qn("w:" + key), value)
        parent.append(node)
    for tag, attrs in [("textDirection", {"val": "tbRl"}), ("eastAsianLayout", {"combine": "1"}),
                       ("w", {"val": "200"}), ("vertAlign", {"val": "superscript"})]:
        for inherited in [False, True]:
            create("reject-" + ("inherited-" if inherited else "direct-") + tag,
                   mutation=lambda d, t=tag, a=attrs, i=inherited: unsupported_text_layout(d, t, a, i),
                   reject="direct or inherited vertical/combined/scaled/superscript image-run context")
    def external(parts):
        n = ET.fromstring(parts["word/_rels/document.xml.rels"])
        for rel in n:
            if rel.get("Type", "").endswith("/image"):
                rel.set("Target", "https://example.invalid/probe.png");rel.set("TargetMode", "External")
        parts["word/_rels/document.xml.rels"] = ET.tostring(n)
    create("reject-external-relationship", package_mutation=external, reject="external image relationship (never fetched)")
    def wrong_mime(parts):
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b'ContentType="image/png"', b'ContentType="image/jpeg"')
    create("reject-mime-conflict", package_mutation=wrong_mime, reject="PNG media with JPEG declared MIME")
    def jpeg_in_png(parts):
        for part in list(parts):
            if part.startswith("word/media/") and part.endswith(".png"):
                parts[part] = data["JPEG"][0]
    create("reject-bitmap-format-conflict", package_mutation=jpeg_in_png, reject="JPEG bytes with PNG declared type and extension")
    def missing_relationship(parts):
        root = ET.fromstring(parts["word/_rels/document.xml.rels"])
        for rel in list(root):
            if rel.get("Type", "").endswith("/image"):root.remove(rel)
        parts["word/_rels/document.xml.rels"] = ET.tostring(root)
    create("reject-missing-relationship", package_mutation=missing_relationship, reject="missing embedded image relationship")
    for name in unsafe_metadata_blobs():
        create("reject-" + name, [(name, 1270000, 635000)], reject="EXIF/APP1 or animated image consumer ambiguity")
    manifest = directory / "synthetic-manifest.json"
    manifest.write_text(json.dumps(sources, indent=2), encoding="utf-8")
    return sources


def export_verify(cli, evidence, directory, name, model, expected, dimensions=True):
    model_path = directory / (name + ".json")
    w.write_model(model_path, model)
    output = directory / (name + ".docx")
    run_cli(cli, evidence["commands"], "create-docx", model_path, output)
    independent = snapshot(output)
    verify_snapshot(independent, expected, name + " output")
    reopened = run_cli(cli, evidence["commands"], "read-docx-model", output)
    check(reopened == model, name + ": trusted full model equality failed")
    verify_model(reopened, expected, name + " trusted model", dimensions)
    external = run_cli(cli, evidence["commands"], "import-docx", output)["document"]
    verify_model(external, expected, name + " external reimport")
    check([n["attrs"].get("imageParagraph") for n in images(external)] ==
          [n["attrs"].get("imageParagraph") for n in images(model)], name + ": external image paragraph metadata changed")
    before = digest(output)
    error = run_cli(cli, evidence["commands"], "create-docx", model_path, output, success=False)
    check("Output already exists" in error and digest(output) == before, name + ": no-clobber failed")
    return {"model_path": str(model_path), "path": str(output), "sha256": before, "independent": independent,
            "full_model_equal_after_authored_reopen": True, "external_reimport_dimensions_bytes_equal": True,
            "existing_output_overwrite_refused": True}


def render_pdf(soffice, directory, path, expected, left_bound=None):
    out = directory / "libreoffice-pdf";out.mkdir(exist_ok=True)
    profile = directory / "lo-profile"
    proc = subprocess.run([str(soffice), f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "pdf",
                           "--outdir", str(out), str(path)], capture_output=True, text=True, timeout=120,
                          env=dict(os.environ, XDG_CACHE_HOME=str(directory / "lo-cache")))
    pdf = out / (path.stem + ".pdf")
    check(proc.returncode == 0 and pdf.is_file(), proc.stdout + proc.stderr)
    reader = PdfReader(pdf)
    found, texts = [], []
    for page_index, page in enumerate(reader.pages):
        resources = page["/Resources"].get_object()
        xobjects = resources.get("/XObject", {}).get_object() if "/XObject" in resources else {}
        def operand(op, args, cm, tm):
            if op == b"Do" and args[0] in xobjects:
                obj = xobjects[args[0]].get_object()
                if obj.get("/Subtype") == "/Image":
                    found.append({"page": page_index + 1, "x_pt": float(cm[4]), "y_pt": float(cm[5]),
                                  "width_pt": math.hypot(cm[0], cm[1]), "height_pt": math.hypot(cm[2], cm[3]),
                                  "page_width_pt": float(page.mediabox.width), "page_height_pt": float(page.mediabox.height),
                                  "pdf_pixel_width": int(obj["/Width"]), "pdf_pixel_height": int(obj["/Height"])})
        def visit_text(text, cm, tm, font, size):
            if text.strip():texts.append({"page": page_index + 1, "text": text.strip(), "x_pt": tm[4], "y_pt": tm[5], "font_size": size})
        page.extract_text(visitor_operand_before=operand, visitor_text=visit_text)
    check(len(found) == len(expected), f"{path.name}: PDF image count differs ({len(found)}/{len(expected)})")
    for got, want in zip(found, expected):
        check(abs(got["width_pt"] - want["width_emu"] / 12700) <= 0.15 and
              abs(got["height_pt"] - want["height_emu"] / 12700) <= 0.15,
              f"{path.name}: PDF physical image size differs: {got}")
        check(got["x_pt"] >= -0.1 and got["y_pt"] >= -0.1 and
              got["x_pt"] + got["width_pt"] <= got["page_width_pt"] + 0.1 and
              got["y_pt"] + got["height_pt"] <= got["page_height_pt"] + 0.1,
              path.name + ": PDF image clipped by page bounds")
        if left_bound is not None:
            check(abs(got["x_pt"] - left_bound) <= 0.75, path.name + ": PDF left placement differs")
    return {"path": str(pdf), "sha256": digest(pdf), "page_count": len(reader.pages), "images": found, "texts": texts,
            "dimension_tolerance_pt": 0.15, "absolute_vertical_position_claimed_preserved": False}


def invalid_models(cli, evidence, directory, model):
    cases = {"both-null": (None, None), "width-null": (None, 635000), "height-null": (1270000, None),
             "zero-width": (0, 635000), "zero-height": (1270000, 0), "negative-width": (-1, 635000),
             "negative-height": (1270000, -1), "string-width": ("1270000", 635000),
             "string-height": (1270000, "635000"), "bool-width": (True, 635000), "bool-height": (1270000, False),
             "fractional-width": (1270000.5, 635000), "fractional-height": (1270000, 635000.5),
             "large-width": (MAX_WIDTH + 1, 635000), "large-height": (1270000, MAX_HEIGHT + 1),
             "array-width": ([1270000], 635000), "object-height": (1270000, {"height": 635000})}
    result = []
    for name, pair in cases.items():
        bad = copy.deepcopy(model);attrs = next(images(bad))["attrs"]
        attrs[WIDTH], attrs[HEIGHT] = pair
        result.append((name, bad))
    for key in [WIDTH, HEIGHT]:
        bad = copy.deepcopy(model);del next(images(bad))["attrs"][key]
        result.append(("partial-" + key, bad))
    for kind in ["paragraph", "blockquote"]:
        image = copy.deepcopy(next(images(model)))
        result.append(("nested-" + kind, {"type": "doc", "content": [{"type": kind, "content": [image]}]}))
    image = copy.deepcopy(next(images(model)))
    result.append(("nested-list", {"type": "doc", "content": [{"type": "bulletList", "content": [
        {"type": "listItem", "content": [{"type": "paragraph"}, image]}]}]}))
    result.append(("nested-table", {"type": "doc", "content": [{"type": "table", "content": [
        {"type": "tableRow", "content": [{"type": "tableCell", "content": [{"type": "paragraph"}, image]}]}]}]}))
    for name, (blob, mime) in unsafe_metadata_blobs().items():
        bad = copy.deepcopy(model)
        next(images(bad))["attrs"]["src"] = "data:" + mime + ";base64," + base64.b64encode(blob).decode()
        result.append((name, bad))
    for name, bad in result:
        p = directory / ("invalid-" + name + ".json");w.write_model(p, bad)
        out = p.with_suffix(".docx")
        error = run_cli(cli, evidence["commands"], "create-docx", p, out, success=False)
        check(not out.exists(), name + ": invalid model published output")
        evidence["invalid_models"][name] = error


def run_supported(name, item, cli, evidence, directory, soffice):
    path = Path(item["path"]);expected = item["expected"]
    imported = run_cli(cli, evidence["commands"], "import-docx", path)
    model = imported["document"]
    item["compatibility"] = imported["compatibility"]
    check(PLACEMENT_WARNING in item["compatibility"]["risks"], name + ": missing bounded paragraph/layout warning")
    verify_model(model, expected, name + " imported")
    wanted_paragraph = item.get("expected_image_paragraph", IMAGE_PARAGRAPH)
    check([n["attrs"].get("imageParagraph") for n in images(model)] == [wanted_paragraph] * len(expected),
          name + ": exact image paragraph metadata differs")
    if item.get("expected_root_order"):
        check([n["type"] for n in model["content"]] == item["expected_root_order"], name + ": genuine empty paragraph order changed")
        check(all(not n.get("content") for n in model["content"] if n["type"] == "paragraph"), name + ": expected empty paragraph has content")
    item["unchanged_export"] = export_verify(cli, evidence, directory, name + "-unchanged", model, expected)
    edited, edited_expected = copy.deepcopy(model), copy.deepcopy(expected)
    for index, (node, want) in enumerate(zip(images(edited), edited_expected)):
        width, height = 1524000 + index * 12700, 1016000 + index * 12700
        node["attrs"].update({WIDTH: width, HEIGHT: height})
        want.update(width_emu=width, height_emu=height)
    check(edited != model, name + ": edit probe did not change model")
    item["edited_export"] = export_verify(cli, evidence, directory, name + "-edited", edited, edited_expected)
    baseline, reset = copy.deepcopy(model), copy.deepcopy(edited)
    for m in [baseline, reset]:
        for node in images(m):
            node["attrs"].pop(WIDTH);node["attrs"].pop(HEIGHT)
    check(baseline == reset, name + ": reset changed unrelated model content")
    reset_expected = copy.deepcopy(expected)
    for want in reset_expected:
        scale = min(1, 600 / want["pixel_width"], 800 / want["pixel_height"])
        want.update(width_emu=round(want["pixel_width"] * scale * 9525), height_emu=round(want["pixel_height"] * scale * 9525))
    item["reset_export"] = export_verify(cli, evidence, directory, name + "-reset", reset, reset_expected, dimensions=False)
    baseline_export = export_verify(cli, evidence, directory, name + "-baseline", baseline, reset_expected, dimensions=False)
    check(item["reset_export"]["sha256"] == baseline_export["sha256"], name + ": reset deterministic baseline bytes differ")
    item["reset_baseline_model_and_docx_bytes_equal"] = True
    error = run_cli(cli, evidence["commands"], "create-docx", item["unchanged_export"]["model_path"], path, success=False)
    check("Output already exists" in error and digest(path) == item["sha256"], name + ": source protection failed")
    item["source_overwrite_refused"] = True
    if soffice and item.get("render_probe"):
        for key, docx, target in [("source_render", path, expected),
                                  ("unchanged_render", Path(item["unchanged_export"]["path"]), expected),
                                  ("edited_render", Path(item["edited_export"]["path"]), edited_expected)]:
            item[key] = render_pdf(soffice, directory, docx, target, 72)
        deltas = []
        for source, output in zip(item["source_render"]["images"], item["unchanged_render"]["images"]):
            delta = {key: output[key] - source[key] for key in ["x_pt", "y_pt", "width_pt", "height_pt"]}
            check(all(abs(value) <= 0.30 for value in delta.values()), name + ": controlled source/export image placement differs")
            deltas.append(delta)
        source_marker = next(t for t in item["source_render"]["texts"] if t["text"] == "AfterImage")
        output_marker = next(t for t in item["unchanged_render"]["texts"] if t["text"] == "AfterImage")
        baseline_delta = output_marker["y_pt"] - source_marker["y_pt"]
        check(abs(baseline_delta) <= 0.30, name + ": controlled following baseline differs")
        item["controlled_placement_comparison"] = {"images": deltas, "following_baseline_delta_pt": baseline_delta, "tolerance_pt": 0.30}
    return model


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli", type=Path, default=Path("target/debug/rust-wps"))
    parser.add_argument("--out", type=Path, default=Path("tmp/word-image-dimensions"))
    parser.add_argument("--corpus-dir", type=Path)
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--soffice")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.out)).resolve()
    cli = args.cli.resolve()
    soffice = (args.soffice or shutil.which("soffice")) if args.render else None
    check(not args.render or soffice, "LibreOffice is required for --render")
    paths = [Path(__file__).resolve(), Path(__file__).with_name("image-dimensions-corpus.json").resolve(),
             Path(__file__).with_name("check-word-table-widths.py").resolve()]
    if not args.generate_only: paths.append(cli)
    initial = {str(p): digest(p) for p in paths}
    evidence = {"status": "running", "python": platform.python_version(),
                "dependencies": {n: version(n) for n in ["python-docx", "lxml", "Pillow", "pypdf"]},
                "commands": [], "invalid_models": {}, "original_corpus": {}, "tool_hashes": initial,
                "authored_reopen_command": "read-docx-model", "external_reimport_command": "import-docx",
                "limitations": "Dimension/bytes scope only. LibreOffice external consumer does not verify app CSS, browser focus, native printing, Word/WPS, full source page or paragraph layout."}
    try:
        evidence["sources"] = make_fixtures(directory)
        if soffice:
            evidence["libreoffice_version"] = subprocess.run([soffice, "--version"], capture_output=True, text=True, check=True, timeout=30).stdout.strip()
        if not args.generate_only:
            seed = None
            for name, item in evidence["sources"].items():
                if item["expected_action"] == "reject":
                    item["audit"] = run_cli(cli, evidence["commands"], "audit-docx", item["path"])
                    check(SOURCE_RISK in item["audit"]["risks"], name + ": missing explicit unsupported image diagnostic")
                    item["import_error"] = run_cli(cli, evidence["commands"], "import-docx", item["path"], success=False)
                else:
                    model = run_supported(name, item, cli, evidence, directory, soffice)
                    if seed is None: seed = model
                check(digest(Path(item["path"])) == item["sha256"], name + ": source changed")
            invalid_models(cli, evidence, directory, seed)
            if args.corpus_dir:
                manifest = json.loads(Path(__file__).with_name("image-dimensions-corpus.json").read_text())
                for pinned in manifest["fixtures"]:
                    path = (args.corpus_dir / pinned["name"]).resolve();blob = path.read_bytes()
                    check(len(blob) == pinned["bytes"] and digest(path) == pinned["sha256"], "Original corpus hash mismatch")
                    check(hashlib.sha1(b"blob " + str(len(blob)).encode() + b"\0" + blob).hexdigest() == pinned["git_blob_sha"], "Original corpus Git blob mismatch")
                    item = {**pinned, "path": str(path), "expected": pinned["images"], "source_independent": snapshot(path), "render_probe": False,
                            "expected_image_paragraph": {"spaceBefore": 0, "spaceAfter": 200, "lineSpacing": 276, "textAlign": "left",
                                "paragraphFont": {"family": "Calibri", "halfPoints": 22, "bold": False, "italic": False},
                                "runFont": {"family": "Calibri", "halfPoints": 22, "bold": False, "italic": False}}}
                    evidence["original_corpus"][pinned["name"]] = item
                    verify_snapshot(item["source_independent"], item["expected"], pinned["name"] + " original")
                    run_supported("original-" + path.stem, item, cli, evidence, directory, None)
                    if soffice:
                        # Record original source/output absolute positions without claiming their
                        # unmodelled paragraph and page styling have been preserved.
                        for key, docx, expected in [("source_render", path, item["expected"]),
                                                   ("unchanged_render", Path(item["unchanged_export"]["path"]), item["expected"])]:
                            item[key] = render_pdf(soffice, directory, docx, expected)
                    check(digest(path) == pinned["sha256"], "Original upstream bytes changed")
        check(all(digest(Path(p)) == sha for p, sha in initial.items()), "CLI or validation inputs changed during run")
        evidence["summary"] = {"supported_synthetic": sum(x["expected_action"] == "roundtrip" for x in evidence["sources"].values()),
                               "refused_synthetic": sum(x["expected_action"] == "reject" for x in evidence["sources"].values()),
                               "invalid_models": len(evidence["invalid_models"]), "original_positive": len(evidence["original_corpus"]),
                               "pdfs": len(list((directory / "libreoffice-pdf").glob("*.pdf")))}
        evidence["status"] = "generated" if args.generate_only else "passed"
    except Exception as exc:
        evidence["status"] = "failed";evidence["error"] = str(exc)
        raise
    finally:
        (directory / "report.json").write_text(json.dumps(evidence, indent=2, ensure_ascii=False), encoding="utf-8")
        print(directory / "report.json", flush=True)
    print(json.dumps(evidence["summary"]))


if __name__ == "__main__":
    main()
