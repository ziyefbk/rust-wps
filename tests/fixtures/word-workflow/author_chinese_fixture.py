"""Independent python-docx authored source, never derived from product model/importer."""
from pathlib import Path
import datetime
import hashlib
import json
import zipfile
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'independent-chinese-minutes.docx'
TEXTS = [
    '项目会议纪要 MinutesSource',
    '会议时间：2026年10月16日。参会人员：林舟、陈悦。讨论事项：文档编辑、保存与恢复。',
    '行动安排：林舟整理测试结果，陈悦核对原始文件。请在周五前完成复核，并保留所有原件。',
]

assert not OUTPUT.exists(), f'Non-overwrite fixture policy: {OUTPUT}'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25.4)
document.core_properties.title = '独立编写的中文会议纪要测试源文件'
document.core_properties.subject = 'Authored python-docx fixture, not a genuine third-party document'
document.core_properties.author = 'Independent workflow verification'
document.core_properties.created = datetime.datetime(2026, 10, 16, tzinfo=datetime.timezone.utc)
document.core_properties.modified = datetime.datetime(2026, 10, 16, tzinfo=datetime.timezone.utc)

for index, text in enumerate(TEXTS):
    paragraph = document.add_paragraph(style='Heading 1' if index == 0 else 'Normal')
    formatting = paragraph.paragraph_format
    formatting.space_before = Pt(0)
    formatting.space_after = Pt(6)
    formatting.line_spacing = 1.0
    formatting.keep_with_next = False
    formatting.keep_together = False
    formatting.page_break_before = False
    formatting.widow_control = True
    run = paragraph.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(18 if index == 0 else 11)
    run.font.bold = index == 0
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for slot, value in [('ascii', 'Arial'), ('hAnsi', 'Arial'), ('eastAsia', 'Microsoft YaHei'), ('cs', 'Arial')]:
        fonts.set(qn('w:' + slot), value)

document.save(OUTPUT)
check = Document(OUTPUT)
assert [p.text for p in check.paragraphs] == TEXTS
assert [p.style.style_id for p in check.paragraphs] == ['Heading1', 'Normal', 'Normal']
assert len(check.tables) == 0 and len(check.sections) == 1
with zipfile.ZipFile(OUTPUT) as archive:
    assert not any(name.startswith('rust-wps/') for name in archive.namelist())
    assert not any(name.startswith(('word/header', 'word/footer')) for name in archive.namelist())

expected = {
    'classification': 'Independent authored fixture; not a genuine third-party document',
    'generator': 'python-docx; no product CLI, model, importer, or serialization used',
    'path': OUTPUT.name,
    'sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    'paragraphs': TEXTS,
    'paragraph_styles': ['Heading1', 'Normal', 'Normal'],
    'paragraph_count': 3,
    'table_count': 0,
    'section_count': 1,
    'page_mm': [210, 297],
    'margins_mm': {'top': 25.4, 'bottom': 25.4, 'left': 25.4, 'right': 25.4},
    'direct_spacing': {'before_pt': 0, 'after_pt': 6, 'line_multiple': 1},
}
(ROOT / 'independent-chinese-minutes.expected.json').write_text(json.dumps(expected, indent=2, ensure_ascii=False) + '\n')
print(json.dumps(expected, indent=2, ensure_ascii=False))
